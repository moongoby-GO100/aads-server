"""Tenant-scoped OBYS business registry and durable ledger uploads."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

from fastapi import HTTPException
from openpyxl import load_workbook

from app.core.obys_db import obys_db_url

MAX_BYTES = 10 * 1024 * 1024
MAX_UNPACKED_BYTES = 50 * 1024 * 1024
UPLOAD_ROOT = Path(os.getenv("OBYS_UPLOAD_ROOT", "app/data/yeoljeong_finance/uploads/ledgers"))
CATEGORY_EXTENSIONS = {
    "sales": {".csv", ".xlsx", ".pdf", ".jpg", ".jpeg", ".png"},
    "purchase": {".csv", ".xlsx", ".pdf", ".jpg", ".jpeg", ".png"},
    "transaction": {".csv", ".xlsx"},
    "card": {".csv", ".xlsx", ".pdf"},
}
LEDGER_CATEGORIES = frozenset({"sales", "purchase", "transaction", "card"})
GENERIC_LEDGER_CATEGORIES = frozenset({"sales", "purchase", "transaction"})
MIME_BY_EXTENSION = {
    ".csv": {"text/csv", "application/csv", "text/plain", "application/vnd.ms-excel"},
    ".xlsx": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
    ".pdf": {"application/pdf"},
    ".jpg": {"image/jpeg"},
    ".jpeg": {"image/jpeg"},
    ".png": {"image/png"},
}
SIGNATURES = {
    ".pdf": b"%PDF-",
    ".png": b"\x89PNG\r\n\x1a\n",
    ".jpg": b"\xff\xd8\xff",
    ".jpeg": b"\xff\xd8\xff",
    ".xlsx": b"PK\x03\x04",
}


def _tenant(user: dict[str, Any]) -> UUID:
    try:
        return UUID(str(user.get("tenant_id") or ""))
    except (ValueError, AttributeError) as exc:
        raise HTTPException(status_code=403, detail="유효한 테넌트가 필요합니다") from exc


def _actor(user: dict[str, Any]) -> str:
    return str(user.get("email") or user.get("user_id") or user.get("sub") or "unknown")[:320]


def tenant_session_for_user(user: dict[str, Any]) -> dict[str, Any]:
    """Build a session from the active tenant membership, never legacy HR files."""
    tenant_id = str(_tenant(user))
    membership = user.get("current_membership") or {}
    membership_tenant = str(membership.get("tenant_id") or "").strip()
    status = str(membership.get("status") or "").strip().lower()
    role = str(membership.get("role") or "").strip().lower()
    if membership_tenant != tenant_id or status != "active" or role not in {"owner", "admin", "member", "viewer"}:
        raise HTTPException(status_code=403, detail="현재 테넌트의 활성 멤버십이 필요합니다")
    can_manage = role in {"owner", "admin"}
    labels = {"owner": "소유자", "admin": "관리자", "member": "구성원", "viewer": "조회전용"}
    return {
        "user": {
            "id": str(user.get("user_id") or ""),
            "email": str(user.get("email") or ""),
            "name": str(user.get("name") or ""),
            "tenant_id": tenant_id,
            "is_admin": can_manage,
        },
        "tenant": user.get("current_tenant") or {"id": tenant_id},
        "permissions": {
            "role": role,
            "role_label": labels[role],
            "can_view": True,
            "can_edit_local_data": role != "viewer",
            "can_manage_settings": can_manage,
            "can_manage_automation": can_manage,
            "can_import_settlements": role != "viewer",
            "can_manage_onboarding": can_manage,
            "can_upload_own_documents": role != "viewer",
        },
    }


def _require_write(user: dict[str, Any]) -> None:
    membership = user.get("current_membership") or {}
    same_tenant = str(membership.get("tenant_id") or "") == str(user.get("tenant_id") or "")
    role = str(membership.get("role") or "").strip().lower()
    status = str(membership.get("status") or "").strip().lower()
    if not same_tenant or status != "active" or role not in {"owner", "admin", "member"}:
        raise HTTPException(status_code=403, detail="이 원장을 등록·수정할 권한이 없습니다")


def _safe_name(filename: str) -> str:
    if not filename or Path(filename).name != filename or "\x00" in filename or "/" in filename or "\\" in filename:
        raise HTTPException(status_code=400, detail="안전하지 않은 파일명입니다")
    cleaned = re.sub(r"[^0-9A-Za-z._가-힣 -]", "_", filename).strip(" .")
    if not cleaned:
        raise HTTPException(status_code=400, detail="파일명이 비어 있습니다")
    return cleaned[:240]


def _validate_file(category: str, filename: str, content_type: str, data: bytes) -> tuple[str, str]:
    if category not in CATEGORY_EXTENSIONS:
        raise HTTPException(status_code=400, detail="지원하지 않는 업로드 구분입니다")
    name = _safe_name(filename)
    ext = Path(name).suffix.lower()
    if ext not in CATEGORY_EXTENSIONS[category]:
        raise HTTPException(status_code=415, detail="지원하지 않는 파일 형식입니다")
    if not data:
        raise HTTPException(status_code=422, detail="빈 원장은 업로드할 수 없습니다")
    if len(data) > MAX_BYTES:
        raise HTTPException(status_code=413, detail="파일은 10MB 이하여야 합니다")
    mime = (content_type or "").split(";", 1)[0].strip().lower()
    if mime not in MIME_BY_EXTENSION[ext]:
        raise HTTPException(status_code=415, detail="확장자와 MIME 형식이 일치하지 않습니다")
    signature = SIGNATURES.get(ext)
    if signature and not data.startswith(signature):
        raise HTTPException(status_code=415, detail="파일 서명이 확장자와 일치하지 않습니다")
    if ext == ".xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                members = archive.infolist()
                unpacked = sum(item.file_size for item in members)
                packed = sum(max(item.compress_size, 1) for item in members)
                if len(members) > 2000 or unpacked > MAX_UNPACKED_BYTES or unpacked / max(packed, 1) > 100:
                    raise HTTPException(status_code=413, detail="Excel 압축 해제 한도를 초과했습니다")
        except zipfile.BadZipFile as exc:
            raise HTTPException(status_code=422, detail="손상된 Excel 파일입니다") from exc
    return name, ext


def _raw_rows(ext: str, data: bytes) -> list[dict[str, Any]]:
    try:
        if ext == ".csv":
            text = None
            for encoding in ("utf-8-sig", "cp949"):
                try:
                    text = data.decode(encoding)
                    break
                except UnicodeDecodeError:
                    pass
            if text is None:
                raise HTTPException(status_code=422, detail="CSV 문자 인코딩을 읽을 수 없습니다")
            reader = csv.DictReader(io.StringIO(text))
            if not reader.fieldnames or not any(str(item or "").strip() for item in reader.fieldnames):
                raise HTTPException(status_code=422, detail="CSV 헤더가 없습니다")
            rows = list(reader)
        else:
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            values = book.active.iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(values, ())]
            if not headers or not any(headers):
                raise HTTPException(status_code=422, detail="Excel 헤더가 없습니다")
            rows = [dict(zip(headers, row)) for row in values if any(value not in (None, "") for value in row)]
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=422, detail="원장 파일을 해석할 수 없습니다") from exc
    if not rows:
        raise HTTPException(status_code=422, detail="빈 원장은 업로드할 수 없습니다")
    return rows


def _value(row: dict[str, Any], *names: str) -> Any:
    def normalize(value: Any) -> str:
        return re.sub(r"[\s_()-]", "", str(value)).lower()

    normalized = {normalize(key): value for key, value in row.items()}
    return next((normalized[normalize(name)] for name in names if normalize(name) in normalized), "")


def _canonical(category: str, rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], int]:
    accepted: list[dict[str, Any]] = []
    rejected = 0
    for row in rows:
        raw_date = _value(row, "date", "거래일자", "매출일자", "일자")
        raw_amount = _value(row, "amount", "금액", "매출액", "결제금액", "입금액", "출금액")
        if category == "transaction" and _value(row, "amount", "금액") in (None, ""):
            try:
                incoming = Decimal(re.sub(r"[^0-9.-]", "", str(_value(row, "입금액") or 0)))
                outgoing = Decimal(re.sub(r"[^0-9.-]", "", str(_value(row, "출금액") or 0)))
                if incoming < 0 or outgoing < 0 or (incoming and outgoing):
                    raise ValueError("ambiguous bank direction")
                raw_amount = incoming - outgoing
            except (ValueError, InvalidOperation):
                rejected += 1
                continue
        try:
            occurred = raw_date.date() if isinstance(raw_date, datetime) else date.fromisoformat(str(raw_date)[:10].replace(".", "-").replace("/", "-"))
            amount = Decimal(re.sub(r"[^0-9.-]", "", str(raw_amount)))
        except (ValueError, InvalidOperation):
            rejected += 1
            continue
        counterparty = str(_value(row, "counterparty", "거래처", "매입처", "가맹점", "상대방") or "").strip()
        description = str(_value(row, "description", "적요", "항목", "품목", "메모") or "").strip()
        stable = "|".join((category, occurred.isoformat(), str(amount), counterparty, description))
        accepted.append({"occurred_on": occurred, "amount": amount, "counterparty": counterparty, "description": description, "source_hash": hashlib.sha256(stable.encode()).hexdigest(), "payload": row})
    if not accepted:
        raise HTTPException(status_code=422, detail="반영할 수 있는 원장 행이 없습니다")
    return accepted, rejected


def preview_upload(*, category: str, filename: str, content_type: str, data: bytes) -> dict[str, Any]:
    """Validate and parse an upload without filesystem or database writes."""
    original, ext = _validate_file(category, filename, content_type, data)
    if ext not in {".csv", ".xlsx"}:
        raise HTTPException(status_code=415, detail="원장 미리보기는 .xlsx 또는 .csv만 지원합니다")
    raw = _raw_rows(ext, data)
    accepted, rejected = _canonical(category, raw)
    accepted_hashes = {row["source_hash"] for row in accepted}
    duplicate_rows = len(accepted) - len(accepted_hashes)
    columns = [str(key) for key in raw[0].keys()]
    return {
        "filename": original,
        "category": category,
        "byte_size": len(data),
        "sha256": hashlib.sha256(data).hexdigest(),
        "columns": columns,
        "mapping": {
            "occurred_on": next((name for name in columns if _value({name: name}, "date", "거래일자", "매출일자", "일자")), ""),
            "amount": next((name for name in columns if _value({name: name}, "amount", "금액", "매출액", "결제금액", "입금액", "출금액")), ""),
        },
        "preview_rows": [
            {"occurred_on": row["occurred_on"].isoformat(), "amount": str(row["amount"]),
             "counterparty": row["counterparty"], "description": row["description"]}
            for row in accepted[:20]
        ],
        "accepted_rows": len(accepted),
        "duplicate_rows": duplicate_rows,
        "rejected_rows": rejected,
        "requires_confirmation": True,
        "source_hashes": list(accepted_hashes),
    }


async def _connect():
    import asyncpg

    return await asyncpg.connect(obys_db_url(), timeout=5)


async def _require_business(conn: Any, tenant_id: UUID, business_id: str) -> None:
    owns = await conn.fetchval(
        "SELECT EXISTS(SELECT 1 FROM yeoljeong_businesses WHERE id=$1 AND tenant_id=$2 AND deleted_at IS NULL)",
        business_id,
        tenant_id,
    )
    if not owns:
        raise HTTPException(status_code=404, detail="현재 테넌트의 사업자를 찾을 수 없습니다")


async def list_businesses(*, user: dict[str, Any]) -> list[dict[str, Any]]:
    tenant_id = _tenant(user)
    conn = await _connect()
    try:
        rows = await conn.fetch("SELECT id,entity_type,name,registration_no,representative,tax_type,opened_at,address,memo,created_at,updated_at FROM yeoljeong_businesses WHERE tenant_id=$1 AND deleted_at IS NULL ORDER BY sort_order,id", tenant_id)
        return [dict(row) for row in rows]
    finally:
        await conn.close()


async def create_business(*, user: dict[str, Any], payload: dict[str, Any]) -> dict[str, Any]:
    tenant_id = _tenant(user)
    business_id = str(payload.get("id") or f"biz-{uuid4().hex[:16]}").strip()
    name = str(payload.get("name") or "").strip()
    if not name or not re.fullmatch(r"[A-Za-z0-9_-]{3,64}", business_id):
        raise HTTPException(status_code=400, detail="사업자 ID와 상호를 확인하십시오")
    conn = await _connect()
    try:
        try:
            row = await conn.fetchrow("""INSERT INTO yeoljeong_businesses (id,tenant_id,entity_type,name,registration_no,representative,tax_type,opened_at,address,memo,sort_order,updated_by) VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10,(SELECT COALESCE(MAX(sort_order),0)+1 FROM yeoljeong_businesses WHERE tenant_id=$2),$11) RETURNING *""", business_id, tenant_id, str(payload.get("entity_type") or "individual"), name, str(payload.get("registration_no") or ""), str(payload.get("representative") or ""), str(payload.get("tax_type") or ""), str(payload.get("opened_at") or ""), str(payload.get("address") or ""), str(payload.get("memo") or ""), _actor(user))
        except Exception as exc:
            if exc.__class__.__name__ == "UniqueViolationError":
                raise HTTPException(status_code=409, detail="이미 사용 중인 사업자 ID입니다") from exc
            raise
        return dict(row)
    finally:
        await conn.close()


async def update_business(*, user: dict[str, Any], business_id: str, payload: dict[str, Any]) -> dict[str, Any]:
    tenant_id = _tenant(user)
    allowed = {"name", "registration_no", "representative", "tax_type", "opened_at", "address", "memo", "entity_type"}
    updates = {key: str(value) for key, value in payload.items() if key in allowed and value is not None}
    if not updates:
        raise HTTPException(status_code=400, detail="변경할 항목이 없습니다")
    columns = list(updates)
    assignments = ",".join(f"{column}=${index}" for index, column in enumerate(columns, 1))
    conn = await _connect()
    try:
        row = await conn.fetchrow(f"UPDATE yeoljeong_businesses SET {assignments},updated_by=${len(columns)+1},updated_at=NOW() WHERE id=${len(columns)+2} AND tenant_id=${len(columns)+3} AND deleted_at IS NULL RETURNING *", *[updates[column] for column in columns], _actor(user), business_id, tenant_id)
        if not row:
            raise HTTPException(status_code=404, detail="사업자를 찾을 수 없습니다")
        return dict(row)
    finally:
        await conn.close()


async def create_upload(*, user: dict[str, Any], business_id: str, category: str, filename: str, content_type: str, data: bytes) -> dict[str, Any]:
    _require_write(user)
    if category not in LEDGER_CATEGORIES:
        raise HTTPException(status_code=400, detail="지원하지 않는 원장 업로드 구분입니다")
    tenant_id = _tenant(user)
    original, ext = _validate_file(category, filename, content_type, data)
    parsed, rejected = ([], 0) if ext in {".pdf", ".jpg", ".jpeg", ".png"} else _canonical(category, _raw_rows(ext, data))
    digest, upload_id = hashlib.sha256(data).hexdigest(), uuid4()
    stored_name = f"{upload_id.hex}{ext}"
    directory = UPLOAD_ROOT / str(tenant_id) / hashlib.sha256(business_id.encode()).hexdigest()[:32] / category
    target = directory / stored_name
    conn = await _connect()
    try:
        await _require_business(conn, tenant_id, business_id)
        existing = await conn.fetchrow("SELECT * FROM yeoljeong_uploads WHERE tenant_id=$1 AND business_id=$2 AND category=$3 AND sha256=$4 AND deleted_at IS NULL", tenant_id, business_id, category, digest)
        if existing:
            result = dict(existing)
            result["status"] = "duplicate"
            return result
        directory.mkdir(parents=True, exist_ok=True)
        target.write_bytes(data)
        status = "pending_review" if not parsed else "imported"
        async with conn.transaction():
            await conn.execute("""INSERT INTO yeoljeong_uploads (id,tenant_id,business_id,category,original_filename,stored_filename,content_type,byte_size,sha256,status,rejected_rows,created_by) VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10,$11,$12)""", upload_id, tenant_id, business_id, category, original, stored_name, content_type, len(data), digest, status, rejected, _actor(user))
            imported = duplicates = 0
            for row in parsed:
                command = await conn.execute("""INSERT INTO yeoljeong_uploaded_ledger_rows (id,tenant_id,business_id,upload_id,category,source_hash,occurred_on,counterparty,description,amount,payload) VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10,$11::jsonb) ON CONFLICT (tenant_id,business_id,category,source_hash) DO NOTHING""", uuid4(), tenant_id, business_id, upload_id, category, row["source_hash"], row["occurred_on"], row["counterparty"], row["description"], row["amount"], json.dumps(row["payload"], ensure_ascii=False, default=str))
                imported += command.endswith(" 1")
                duplicates += command.endswith(" 0")
            await conn.execute("UPDATE yeoljeong_uploads SET imported_rows=$2,duplicate_rows=$3,status=$4 WHERE id=$1", upload_id, imported, duplicates, "duplicate" if parsed and not imported else status)
        return {"id": str(upload_id), "tenant_id": str(tenant_id), "business_id": business_id, "category": category, "original_filename": original, "status": "duplicate" if parsed and not imported else status, "imported_rows": imported, "duplicate_rows": duplicates, "rejected_rows": rejected}
    except Exception:
        if target.exists():
            target.unlink()
        raise
    finally:
        await conn.close()


async def list_uploads(*, user: dict[str, Any], business_id: str, category: str | None = None) -> list[dict[str, Any]]:
    tenant_id = _tenant(user)
    conn = await _connect()
    try:
        await _require_business(conn, tenant_id, business_id)
        rows = await conn.fetch("SELECT id,business_id,category,original_filename,content_type,byte_size,sha256,status,imported_rows,duplicate_rows,rejected_rows,error_message,created_by,created_at FROM yeoljeong_uploads WHERE tenant_id=$1 AND business_id=$2 AND deleted_at IS NULL AND ($3::text IS NULL OR category=$3) ORDER BY created_at DESC LIMIT 100", tenant_id, business_id, category)
        return [dict(row) for row in rows]
    finally:
        await conn.close()


async def list_ledger_rows(*, user: dict[str, Any], business_id: str, category: str, limit: int | None = 500, date_from: str | None = None, date_to: str | None = None) -> list[dict[str, Any]]:
    tenant_id = _tenant(user)
    if category not in LEDGER_CATEGORIES:
        raise HTTPException(status_code=400, detail="지원하지 않는 원장 구분입니다")
    conn = await _connect()
    try:
        await _require_business(conn, tenant_id, business_id)
        rows = await conn.fetch("SELECT id,business_id,upload_id,category,occurred_on,counterparty,description,amount,payload,created_at FROM yeoljeong_uploaded_ledger_rows WHERE tenant_id=$1 AND business_id=$2 AND category=$3 AND ($5::date IS NULL OR occurred_on >= $5) AND ($6::date IS NULL OR occurred_on <= $6) ORDER BY occurred_on DESC NULLS LAST,created_at DESC LIMIT $4", tenant_id, business_id, category, None if limit is None else min(max(limit, 1), 1000), _iso_date(date_from) if date_from else None, _iso_date(date_to) if date_to else None)
        return [dict(row) for row in rows]
    finally:
        await conn.close()


async def get_ledger_row(*, user: dict[str, Any], business_id: str, category: str, entry_id: UUID) -> dict[str, Any] | None:
    tenant_id = _tenant(user)
    conn = await _connect()
    try:
        await _require_business(conn, tenant_id, business_id)
        row = await conn.fetchrow("SELECT * FROM yeoljeong_uploaded_ledger_rows WHERE id=$1 AND tenant_id=$2 AND business_id=$3 AND category=$4", entry_id, tenant_id, business_id, category)
        return dict(row) if row else None
    finally:
        await conn.close()


async def download_path(*, user: dict[str, Any], upload_id: UUID) -> tuple[Path, str, str]:
    tenant_id = _tenant(user)
    conn = await _connect()
    try:
        row = await conn.fetchrow("SELECT business_id,category,stored_filename,original_filename,content_type FROM yeoljeong_uploads WHERE id=$1 AND tenant_id=$2 AND deleted_at IS NULL", upload_id, tenant_id)
        if not row:
            raise HTTPException(status_code=404, detail="업로드 파일을 찾을 수 없습니다")
        await _require_business(conn, tenant_id, row["business_id"])
        path = UPLOAD_ROOT / str(tenant_id) / hashlib.sha256(row["business_id"].encode()).hexdigest()[:32] / row["category"] / row["stored_filename"]
        if not path.is_file() or UPLOAD_ROOT.resolve() not in path.resolve().parents:
            raise HTTPException(status_code=404, detail="업로드 파일을 찾을 수 없습니다")
        return path, row["original_filename"], row["content_type"]
    finally:
        await conn.close()


async def delete_upload(*, user: dict[str, Any], upload_id: UUID) -> bool:
    _require_write(user)
    tenant_id = _tenant(user)
    conn = await _connect()
    try:
        row = await conn.fetchrow("SELECT business_id FROM yeoljeong_uploads WHERE id=$1 AND tenant_id=$2 AND deleted_at IS NULL", upload_id, tenant_id)
        if not row:
            raise HTTPException(status_code=404, detail="업로드 파일을 찾을 수 없습니다")
        await _require_business(conn, tenant_id, row["business_id"])
        # Canonical rows are derived from this upload. Remove them atomically so
        # a same-SHA retry can import fresh rows instead of hitting stale hashes.
        async with conn.transaction():
            await conn.execute(
                "DELETE FROM yeoljeong_uploaded_ledger_rows WHERE upload_id=$1 AND tenant_id=$2",
                upload_id,
                tenant_id,
            )
            await conn.execute("UPDATE yeoljeong_uploads SET deleted_at=NOW() WHERE id=$1 AND tenant_id=$2", upload_id, tenant_id)
        return True
    finally:
        await conn.close()


async def create_card_upload(
    *, user: dict[str, Any], business_id: str, filename: str,
    content_type: str, data: bytes,
) -> dict[str, Any]:
    """Persist card evidence without mixing it into the bank/general ledger."""
    _require_write(user)
    tenant_id = _tenant(user)
    original, ext = _validate_file("card", filename, content_type, data)
    digest, upload_id = hashlib.sha256(data).hexdigest(), uuid4()
    stored_name = f"{upload_id.hex}{ext}"
    directory = UPLOAD_ROOT / str(tenant_id) / hashlib.sha256(business_id.encode()).hexdigest()[:32] / "card"
    target = directory / stored_name
    conn = await _connect()
    try:
        await _require_business(conn, tenant_id, business_id)
        existing = await conn.fetchrow(
            """SELECT * FROM yeoljeong_card_uploads
                 WHERE tenant_id=$1 AND business_id=$2 AND sha256=$3 AND deleted_at IS NULL""",
            tenant_id, business_id, digest,
        )
        if existing:
            result = dict(existing)
            result["status"] = "duplicate"
            return result
        directory.mkdir(parents=True, exist_ok=True)
        target.write_bytes(data)
        row = await conn.fetchrow(
            """INSERT INTO yeoljeong_card_uploads
                   (id,tenant_id,business_id,original_filename,stored_filename,content_type,
                    byte_size,sha256,status,created_by)
                 VALUES ($1,$2,$3,$4,$5,$6,$7,$8,'pending_review',$9)
                 RETURNING id,business_id,original_filename,content_type,byte_size,sha256,
                           status,imported_rows,rejected_rows,error_message,created_by,created_at""",
            upload_id, tenant_id, business_id, original, stored_name, content_type,
            len(data), digest, _actor(user),
        )
        return dict(row)
    except Exception:
        if target.exists():
            target.unlink()
        raise
    finally:
        await conn.close()


async def list_card_uploads(*, user: dict[str, Any], business_id: str) -> list[dict[str, Any]]:
    tenant_id = _tenant(user)
    conn = await _connect()
    try:
        await _require_business(conn, tenant_id, business_id)
        rows = await conn.fetch(
            """SELECT id,business_id,original_filename,content_type,byte_size,sha256,status,
                      imported_rows,rejected_rows,error_message,created_by,created_at
                 FROM yeoljeong_card_uploads
                WHERE tenant_id=$1 AND business_id=$2 AND deleted_at IS NULL
                ORDER BY created_at DESC LIMIT 100""",
            tenant_id, business_id,
        )
        return [dict(row) for row in rows]
    finally:
        await conn.close()


async def download_card_upload(*, user: dict[str, Any], upload_id: UUID) -> tuple[Path, str, str]:
    tenant_id = _tenant(user)
    conn = await _connect()
    try:
        row = await conn.fetchrow(
            """SELECT business_id,stored_filename,original_filename,content_type
                 FROM yeoljeong_card_uploads
                WHERE id=$1 AND tenant_id=$2 AND deleted_at IS NULL""",
            upload_id, tenant_id,
        )
        if not row:
            raise HTTPException(status_code=404, detail="카드 업로드 파일을 찾을 수 없습니다")
        await _require_business(conn, tenant_id, row["business_id"])
        path = UPLOAD_ROOT / str(tenant_id) / hashlib.sha256(row["business_id"].encode()).hexdigest()[:32] / "card" / row["stored_filename"]
        if not path.is_file() or UPLOAD_ROOT.resolve() not in path.resolve().parents:
            raise HTTPException(status_code=404, detail="카드 업로드 파일을 찾을 수 없습니다")
        return path, row["original_filename"], row["content_type"]
    finally:
        await conn.close()


async def delete_card_upload(*, user: dict[str, Any], upload_id: UUID) -> bool:
    _require_write(user)
    tenant_id = _tenant(user)
    conn = await _connect()
    try:
        command = await conn.execute(
            """UPDATE yeoljeong_card_uploads SET deleted_at=NOW()
                 WHERE id=$1 AND tenant_id=$2 AND deleted_at IS NULL""",
            upload_id, tenant_id,
        )
        if command.endswith(" 0"):
            raise HTTPException(status_code=404, detail="카드 업로드 파일을 찾을 수 없습니다")
        return True
    finally:
        await conn.close()


def _ledger_category(category: str) -> str:
    value = str(category or "").strip().lower()
    if value not in {"sales", "purchase"}:
        raise HTTPException(status_code=400, detail="원장 구분은 sales 또는 purchase여야 합니다")
    return value


def _money_parts(payload: dict[str, Any]) -> tuple[Decimal, Decimal, Decimal]:
    try:
        supply = Decimal(str(payload.get("supply_amount", 0)))
        tax = Decimal(str(payload.get("tax_amount", 0)))
        total = Decimal(str(payload.get("total_amount", 0)))
    except InvalidOperation as exc:
        raise HTTPException(status_code=422, detail="금액 형식이 올바르지 않습니다") from exc
    maximum = Decimal("9999999999999999.99")
    if any(not value.is_finite() or value < 0 or value > maximum for value in (supply, tax, total)):
        raise HTTPException(status_code=422, detail="금액은 0 이상 허용 범위 이하여야 합니다")
    if supply + tax != total:
        raise HTTPException(status_code=422, detail="합계는 공급가와 세액의 합이어야 합니다")
    return supply, tax, total


def _iso_date(value: Any, field: str = "occurred_on") -> date:
    try:
        return date.fromisoformat(str(value or "")[:10])
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=f"{field} 날짜 형식이 올바르지 않습니다") from exc


def _iso_datetime(value: Any) -> datetime:
    try:
        parsed = datetime.fromisoformat(str(value or "").replace("Z", "+00:00"))
    except ValueError as exc:
        raise HTTPException(status_code=422, detail="occurred_at 일시 형식이 올바르지 않습니다") from exc
    if parsed.tzinfo is None:
        raise HTTPException(status_code=422, detail="occurred_at에는 시간대가 필요합니다")
    return parsed


async def list_manual_entries(*, user: dict[str, Any], business_id: str, category: str, date_from: str | None = None, date_to: str | None = None, limit: int | None = 1000) -> list[dict[str, Any]]:
    tenant_id, category = _tenant(user), _ledger_category(category)
    start = _iso_date(date_from, "date_from") if date_from else None
    end = _iso_date(date_to, "date_to") if date_to else None
    conn = await _connect()
    try:
        await _require_business(conn, tenant_id, business_id)
        rows = await conn.fetch("""SELECT id,business_id,category,occurred_on,counterparty,description,supply_amount,tax_amount,total_amount,source,created_at,updated_at FROM yeoljeong_manual_ledger_entries WHERE tenant_id=$1 AND business_id=$2 AND category=$3 AND deleted_at IS NULL AND ($4::date IS NULL OR occurred_on >= $4) AND ($5::date IS NULL OR occurred_on <= $5) ORDER BY occurred_on DESC,created_at DESC LIMIT $6""", tenant_id, business_id, category, start, end, limit)
        return [dict(row) for row in rows]
    finally:
        await conn.close()


async def create_manual_entry(*, user: dict[str, Any], category: str, payload: dict[str, Any]) -> dict[str, Any]:
    _require_write(user)
    tenant_id, category = _tenant(user), _ledger_category(category)
    business_id, occurred = str(payload.get("business_id") or "").strip(), _iso_date(payload.get("occurred_on"))
    supply, tax, total = _money_parts(payload)
    conn = await _connect()
    try:
        await _require_business(conn, tenant_id, business_id)
        row = await conn.fetchrow("""INSERT INTO yeoljeong_manual_ledger_entries (id,tenant_id,business_id,category,occurred_on,counterparty,description,supply_amount,tax_amount,total_amount,source,created_by,updated_by) VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10,'manual',$11,$11) RETURNING *""", uuid4(), tenant_id, business_id, category, occurred, str(payload.get("counterparty") or "").strip(), str(payload.get("description") or "").strip(), supply, tax, total, _actor(user))
        return dict(row)
    finally:
        await conn.close()


async def get_manual_entry(*, user: dict[str, Any], category: str, entry_id: UUID) -> dict[str, Any]:
    tenant_id, category = _tenant(user), _ledger_category(category)
    conn = await _connect()
    try:
        row = await conn.fetchrow("SELECT * FROM yeoljeong_manual_ledger_entries WHERE id=$1 AND tenant_id=$2 AND category=$3 AND deleted_at IS NULL", entry_id, tenant_id, category)
        if not row:
            raise HTTPException(status_code=404, detail="원장 항목을 찾을 수 없습니다")
        await _require_business(conn, tenant_id, row["business_id"])
        return dict(row)
    finally:
        await conn.close()


async def update_manual_entry(*, user: dict[str, Any], category: str, entry_id: UUID, payload: dict[str, Any]) -> dict[str, Any]:
    _require_write(user)
    current = await get_manual_entry(user=user, category=category, entry_id=entry_id)
    merged = {**current, **payload}
    supply, tax, total = _money_parts(merged)
    occurred = _iso_date(merged.get("occurred_on"))
    tenant_id = _tenant(user)
    conn = await _connect()
    try:
        row = await conn.fetchrow("""UPDATE yeoljeong_manual_ledger_entries SET occurred_on=$4,counterparty=$5,description=$6,supply_amount=$7,tax_amount=$8,total_amount=$9,updated_by=$10,updated_at=NOW() WHERE id=$1 AND tenant_id=$2 AND category=$3 AND source='manual' AND deleted_at IS NULL RETURNING *""", entry_id, tenant_id, _ledger_category(category), occurred, str(merged.get("counterparty") or "").strip(), str(merged.get("description") or "").strip(), supply, tax, total, _actor(user))
        if not row:
            raise HTTPException(status_code=404, detail="수정 가능한 수기 원장을 찾을 수 없습니다")
        return dict(row)
    finally:
        await conn.close()


async def delete_manual_entry(*, user: dict[str, Any], category: str, entry_id: UUID) -> None:
    _require_write(user)
    tenant_id = _tenant(user)
    conn = await _connect()
    try:
        command = await conn.execute("UPDATE yeoljeong_manual_ledger_entries SET deleted_at=NOW(),updated_by=$4,updated_at=NOW() WHERE id=$1 AND tenant_id=$2 AND category=$3 AND source='manual' AND deleted_at IS NULL", entry_id, tenant_id, _ledger_category(category), _actor(user))
        if command.endswith(" 0"):
            raise HTTPException(status_code=404, detail="삭제 가능한 수기 원장을 찾을 수 없습니다")
    finally:
        await conn.close()


def _public_card(row: Any) -> dict[str, Any]:
    value = dict(row)
    value["card_number_masked"] = f"**** **** **** {value.pop('card_last4')}"
    return value


async def list_card_transactions(*, user: dict[str, Any], business_id: str, date_from: str | None = None, date_to: str | None = None, limit: int | None = 1000) -> list[dict[str, Any]]:
    tenant_id = _tenant(user)
    start = _iso_date(date_from, "date_from") if date_from else None
    end = _iso_date(date_to, "date_to") if date_to else None
    conn = await _connect()
    try:
        await _require_business(conn, tenant_id, business_id)
        rows = await conn.fetch("""SELECT id,business_id,occurred_at,merchant,description,supply_amount,tax_amount,total_amount,card_last4,source,created_at,updated_at FROM yeoljeong_card_transactions WHERE tenant_id=$1 AND business_id=$2 AND deleted_at IS NULL AND ($3::date IS NULL OR (occurred_at AT TIME ZONE 'Asia/Seoul')::date >= $3) AND ($4::date IS NULL OR (occurred_at AT TIME ZONE 'Asia/Seoul')::date <= $4) ORDER BY occurred_at DESC LIMIT $5""", tenant_id, business_id, start, end, limit)
        return [_public_card(row) for row in rows]
    finally:
        await conn.close()


async def create_card_transaction(*, user: dict[str, Any], payload: dict[str, Any]) -> dict[str, Any]:
    _require_write(user)
    tenant_id, business_id = _tenant(user), str(payload.get("business_id") or "").strip()
    supply, tax, total = _money_parts(payload)
    last4 = str(payload.get("card_last4") or "")
    if not re.fullmatch(r"\d{4}", last4):
        raise HTTPException(status_code=422, detail="카드번호는 last4 네 자리만 입력하십시오")
    conn = await _connect()
    try:
        await _require_business(conn, tenant_id, business_id)
        row = await conn.fetchrow("""INSERT INTO yeoljeong_card_transactions (id,tenant_id,business_id,occurred_at,merchant,description,supply_amount,tax_amount,total_amount,card_last4,source,created_by,updated_by) VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10,'manual',$11,$11) RETURNING *""", uuid4(), tenant_id, business_id, _iso_datetime(payload.get("occurred_at")), str(payload.get("merchant") or "").strip(), str(payload.get("description") or "").strip(), supply, tax, total, last4, _actor(user))
        return _public_card(row)
    finally:
        await conn.close()


async def get_card_transaction(*, user: dict[str, Any], transaction_id: UUID) -> dict[str, Any]:
    tenant_id = _tenant(user)
    conn = await _connect()
    try:
        row = await conn.fetchrow("SELECT * FROM yeoljeong_card_transactions WHERE id=$1 AND tenant_id=$2 AND deleted_at IS NULL", transaction_id, tenant_id)
        if not row:
            raise HTTPException(status_code=404, detail="카드 거래를 찾을 수 없습니다")
        await _require_business(conn, tenant_id, row["business_id"])
        return _public_card(row)
    finally:
        await conn.close()


async def update_card_transaction(*, user: dict[str, Any], transaction_id: UUID, payload: dict[str, Any]) -> dict[str, Any]:
    _require_write(user)
    current = await get_card_transaction(user=user, transaction_id=transaction_id)
    merged = {**current, **payload}
    supply, tax, total = _money_parts(merged)
    last4 = str(payload.get("card_last4") or str(current.get("card_number_masked") or "")[-4:])
    tenant_id = _tenant(user)
    conn = await _connect()
    try:
        row = await conn.fetchrow("""UPDATE yeoljeong_card_transactions SET occurred_at=$3,merchant=$4,description=$5,supply_amount=$6,tax_amount=$7,total_amount=$8,card_last4=$9,updated_by=$10,updated_at=NOW() WHERE id=$1 AND tenant_id=$2 AND source='manual' AND deleted_at IS NULL RETURNING *""", transaction_id, tenant_id, _iso_datetime(merged.get("occurred_at")), str(merged.get("merchant") or "").strip(), str(merged.get("description") or "").strip(), supply, tax, total, last4, _actor(user))
        if not row:
            raise HTTPException(status_code=404, detail="수정 가능한 카드 거래를 찾을 수 없습니다")
        return _public_card(row)
    finally:
        await conn.close()


async def delete_card_transaction(*, user: dict[str, Any], transaction_id: UUID) -> None:
    _require_write(user)
    tenant_id = _tenant(user)
    conn = await _connect()
    try:
        command = await conn.execute("UPDATE yeoljeong_card_transactions SET deleted_at=NOW(),updated_by=$3,updated_at=NOW() WHERE id=$1 AND tenant_id=$2 AND source='manual' AND deleted_at IS NULL", transaction_id, tenant_id, _actor(user))
        if command.endswith(" 0"):
            raise HTTPException(status_code=404, detail="삭제 가능한 카드 거래를 찾을 수 없습니다")
    finally:
        await conn.close()


async def list_bank_transactions(*, user: dict[str, Any], business_id: str, date_from: str | None = None, date_to: str | None = None, limit: int | None = 1000) -> list[dict[str, Any]]:
    tenant_id = _tenant(user)
    start = _iso_date(date_from, "date_from") if date_from else None
    end = _iso_date(date_to, "date_to") if date_to else None
    conn = await _connect()
    try:
        await _require_business(conn, tenant_id, business_id)
        rows = await conn.fetch(
            """SELECT id,business_id,occurred_at,direction,amount,balance,counterparty,memo,category,account_label,source,created_at,updated_at
                 FROM yeoljeong_manual_bank_transactions
                WHERE tenant_id=$1 AND business_id=$2 AND deleted_at IS NULL
                  AND ($3::date IS NULL OR (occurred_at AT TIME ZONE 'Asia/Seoul')::date >= $3)
                  AND ($4::date IS NULL OR (occurred_at AT TIME ZONE 'Asia/Seoul')::date <= $4)
                ORDER BY occurred_at DESC LIMIT $5""",
            tenant_id, business_id, start, end, limit,
        )
        return [dict(row) for row in rows]
    finally:
        await conn.close()


async def create_bank_transaction(*, user: dict[str, Any], payload: dict[str, Any]) -> dict[str, Any]:
    _require_write(user)
    tenant_id = _tenant(user)
    business_id = str(payload.get("business_id") or "").strip()
    conn = await _connect()
    try:
        await _require_business(conn, tenant_id, business_id)
        row = await conn.fetchrow(
            """INSERT INTO yeoljeong_manual_bank_transactions
                    (id,tenant_id,business_id,occurred_at,direction,amount,balance,counterparty,memo,category,account_label,source,created_by,updated_by)
                VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10,$11,'manual',$12,$12) RETURNING *""",
            uuid4(), tenant_id, business_id, _iso_datetime(payload.get("occurred_at")),
            str(payload.get("direction") or ""), int(payload.get("amount") or 0), payload.get("balance"),
            str(payload.get("counterparty") or "").strip(), str(payload.get("memo") or "").strip(),
            str(payload.get("category") or "").strip(), str(payload.get("account_label") or "").strip(), _actor(user),
        )
        return dict(row)
    finally:
        await conn.close()


async def get_bank_transaction(*, user: dict[str, Any], transaction_id: UUID) -> dict[str, Any]:
    tenant_id = _tenant(user)
    conn = await _connect()
    try:
        row = await conn.fetchrow(
            "SELECT * FROM yeoljeong_manual_bank_transactions WHERE id=$1 AND tenant_id=$2 AND deleted_at IS NULL",
            transaction_id, tenant_id,
        )
        if not row:
            raise HTTPException(status_code=404, detail="은행 거래를 찾을 수 없습니다")
        await _require_business(conn, tenant_id, row["business_id"])
        return dict(row)
    finally:
        await conn.close()


async def update_bank_transaction(*, user: dict[str, Any], transaction_id: UUID, payload: dict[str, Any]) -> dict[str, Any]:
    _require_write(user)
    current = await get_bank_transaction(user=user, transaction_id=transaction_id)
    merged = {**current, **payload}
    tenant_id = _tenant(user)
    conn = await _connect()
    try:
        row = await conn.fetchrow(
            """UPDATE yeoljeong_manual_bank_transactions
                  SET occurred_at=$3,direction=$4,amount=$5,balance=$6,counterparty=$7,memo=$8,category=$9,account_label=$10,updated_by=$11,updated_at=NOW()
                WHERE id=$1 AND tenant_id=$2 AND source='manual' AND deleted_at IS NULL RETURNING *""",
            transaction_id, tenant_id, _iso_datetime(merged.get("occurred_at")), str(merged.get("direction") or ""),
            int(merged.get("amount") or 0), merged.get("balance"), str(merged.get("counterparty") or "").strip(),
            str(merged.get("memo") or "").strip(), str(merged.get("category") or "").strip(),
            str(merged.get("account_label") or "").strip(), _actor(user),
        )
        if not row:
            raise HTTPException(status_code=404, detail="수정 가능한 은행 거래를 찾을 수 없습니다")
        return dict(row)
    finally:
        await conn.close()


async def delete_bank_transaction(*, user: dict[str, Any], transaction_id: UUID) -> None:
    _require_write(user)
    tenant_id = _tenant(user)
    conn = await _connect()
    try:
        command = await conn.execute(
            "UPDATE yeoljeong_manual_bank_transactions SET deleted_at=NOW(),updated_by=$3,updated_at=NOW() WHERE id=$1 AND tenant_id=$2 AND source='manual' AND deleted_at IS NULL",
            transaction_id, tenant_id, _actor(user),
        )
        if command.endswith(" 0"):
            raise HTTPException(status_code=404, detail="삭제 가능한 은행 거래를 찾을 수 없습니다")
    finally:
        await conn.close()


JOURNAL_SOURCES = {
    "uploaded_ledger_row": ("yeoljeong_uploaded_ledger_rows", "occurred_on", "amount", "description", ""),
    "manual_ledger_entry": ("yeoljeong_manual_ledger_entries", "occurred_on", "total_amount", "description", "AND deleted_at IS NULL"),
    "card_transaction": ("yeoljeong_card_transactions", "occurred_at", "total_amount", "description", "AND deleted_at IS NULL"),
    "bank_transaction": ("yeoljeong_manual_bank_transactions", "occurred_at", "amount", "memo", "AND deleted_at IS NULL"),
}
JOURNAL_STATUSES = frozenset({"draft", "needs_review", "approved", "posted", "reversed"})
JOURNAL_WRITE_FROZEN_CODE = "journal_write_moved_to_acct"
JOURNAL_WRITE_FROZEN_MESSAGE = "전표 정본은 ACCT(회계원장)입니다. 오비서에서는 전표를 생성·수정할 수 없습니다."


def journal_writes_are_frozen() -> bool:
    """Default-deny OBYS journal writes; false is the explicit rollback switch."""
    return os.getenv("OBYS_JOURNAL_WRITE_FROZEN", "true").strip().lower() not in {
        "0", "false", "no", "off",
    }


def require_journal_write_enabled() -> None:
    if journal_writes_are_frozen():
        raise HTTPException(
            status_code=409,
            detail={"code": JOURNAL_WRITE_FROZEN_CODE, "message": JOURNAL_WRITE_FROZEN_MESSAGE},
        )


def _require_approve(user: dict[str, Any]) -> None:
    membership = user.get("current_membership") or {}
    if (str(membership.get("tenant_id") or "") != str(user.get("tenant_id") or "")
            or str(membership.get("status") or "").lower() != "active"
            or str(membership.get("role") or "").lower() not in {"owner", "admin"}):
        raise HTTPException(status_code=403, detail="전표를 확정하거나 취소할 권한이 없습니다")


async def create_journal(*, user: dict[str, Any], business_id: str, source_type: str, source_id: UUID) -> dict[str, Any]:
    """Create one conservative, balanced draft per source (idempotent)."""
    require_journal_write_enabled()
    _require_write(user)
    if source_type not in JOURNAL_SOURCES:
        raise HTTPException(status_code=400, detail="지원하지 않는 전표 원본입니다")
    tenant_id = _tenant(user)
    table, date_column, amount_column, description_column, active_predicate = JOURNAL_SOURCES[source_type]
    conn = await _connect()
    try:
        await _require_business(conn, tenant_id, business_id)
        async with conn.transaction():
            # Serialize the check/create pair as well as retaining the unique
            # index: concurrent retries return the original voucher, not 409.
            await conn.execute(
                "SELECT pg_advisory_xact_lock(hashtextextended($1, 0))",
                f"{tenant_id}:{business_id}:{source_type}:{source_id}",
            )
            existing = await conn.fetchrow(
                "SELECT * FROM yeoljeong_journal_vouchers WHERE tenant_id=$1 AND business_id=$2 AND source_type=$3 AND source_id=$4",
                tenant_id, business_id, source_type, source_id,
            )
            if existing:
                result = dict(existing)
                result["idempotent"] = True
                return result
            source = await conn.fetchrow(
                f"SELECT {date_column} AS transaction_date,{amount_column} AS total_amount,{description_column} AS description "
                f"FROM {table} WHERE id=$1 AND tenant_id=$2 AND business_id=$3 {active_predicate}",
                source_id, tenant_id, business_id,
            )
            if not source:
                raise HTTPException(status_code=404, detail="현재 사업자의 전표 원본을 찾을 수 없습니다")
            total = Decimal(str(source["total_amount"] or 0))
            if total <= 0:
                raise HTTPException(status_code=422, detail="0원 이하 거래는 전표로 만들 수 없습니다")
            voucher_id = uuid4()
            voucher_no = f"OBYS-{datetime.now().strftime('%Y%m%d')}-{voucher_id.hex[:8].upper()}"
            transaction_date = source["transaction_date"]
            if isinstance(transaction_date, datetime):
                transaction_date = transaction_date.date()
            row = await conn.fetchrow(
                """INSERT INTO yeoljeong_journal_vouchers
                    (id,tenant_id,business_id,voucher_no,transaction_date,description,status,source_type,source_id,
                     supply_amount,tax_amount,total_amount,evidence_source,created_by)
                   VALUES ($1,$2,$3,$4,$5,$6,'needs_review',$7,$8,$9,0,$9,$7,$10) RETURNING *""",
                voucher_id, tenant_id, business_id, voucher_no, transaction_date,
                str(source["description"] or "")[:500], source_type, source_id, total, _actor(user),
            )
            # Conservative suspense accounts deliberately require accountant review.
            await conn.executemany(
                """INSERT INTO yeoljeong_journal_lines
                    (id,voucher_id,line_no,side,account_code,account_name,amount,tax_code,memo)
                   VALUES ($1,$2,$3,$4,$5,$6,$7,'', $8)""",
                [(uuid4(), voucher_id, 1, "debit", "9998", "차변 계정 검토", total, "자동분개 검토 필요"),
                 (uuid4(), voucher_id, 2, "credit", "9999", "대변 계정 검토", total, "자동분개 검토 필요")],
            )
        return dict(row)
    finally:
        await conn.close()


async def list_journals(*, user: dict[str, Any], business_id: str, status: str | None = None) -> list[dict[str, Any]]:
    tenant_id = _tenant(user)
    if status and status not in JOURNAL_STATUSES:
        raise HTTPException(status_code=400, detail="지원하지 않는 전표 상태입니다")
    conn = await _connect()
    try:
        await _require_business(conn, tenant_id, business_id)
        rows = await conn.fetch(
            """SELECT v.*,COALESCE(jsonb_agg(jsonb_build_object(
                       'id',l.id,'line_no',l.line_no,'side',l.side,'account_code',l.account_code,
                       'account_name',l.account_name,'amount',l.amount,'tax_code',l.tax_code,'memo',l.memo)
                       ORDER BY l.line_no) FILTER (WHERE l.id IS NOT NULL),'[]') AS lines
                 FROM yeoljeong_journal_vouchers v LEFT JOIN yeoljeong_journal_lines l ON l.voucher_id=v.id
                WHERE v.tenant_id=$1 AND v.business_id=$2 AND ($3::text IS NULL OR v.status=$3)
                GROUP BY v.id ORDER BY v.transaction_date DESC,v.created_at DESC LIMIT 500""",
            tenant_id, business_id, status,
        )
        return [dict(row) for row in rows]
    finally:
        await conn.close()


async def update_journal(*, user: dict[str, Any], voucher_id: UUID, payload: dict[str, Any]) -> dict[str, Any]:
    require_journal_write_enabled()
    _require_write(user)
    tenant_id = _tenant(user)
    conn = await _connect()
    try:
        current = await conn.fetchrow("SELECT * FROM yeoljeong_journal_vouchers WHERE id=$1 AND tenant_id=$2", voucher_id, tenant_id)
        if not current:
            raise HTTPException(status_code=404, detail="전표를 찾을 수 없습니다")
        await _require_business(conn, tenant_id, current["business_id"])
        if current["status"] not in {"draft", "needs_review"}:
            raise HTTPException(status_code=409, detail="검토 중인 전표만 보정할 수 있습니다")
        lines = payload.get("lines") or []
        if len(lines) < 2:
            raise HTTPException(status_code=422, detail="전표는 차변·대변 2개 이상의 분개가 필요합니다")
        debit = sum((Decimal(str(line.get("amount") or 0)) for line in lines if line.get("side") == "debit"), Decimal("0"))
        credit = sum((Decimal(str(line.get("amount") or 0)) for line in lines if line.get("side") == "credit"), Decimal("0"))
        if debit <= 0 or debit != credit or debit != current["total_amount"]:
            raise HTTPException(status_code=422, detail="차변·대변과 거래 합계가 일치해야 합니다")
        if any(line.get("side") not in {"debit", "credit"} or not str(line.get("account_code") or "").strip() or not str(line.get("account_name") or "").strip() for line in lines):
            raise HTTPException(status_code=422, detail="계정과목과 차대 구분을 확인하십시오")
        async with conn.transaction():
            await conn.execute("DELETE FROM yeoljeong_journal_lines WHERE voucher_id=$1", voucher_id)
            await conn.executemany(
                """INSERT INTO yeoljeong_journal_lines
                    (id,voucher_id,line_no,side,account_code,account_name,amount,tax_code,memo)
                   VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9)""",
                [(uuid4(), voucher_id, index, line["side"], str(line["account_code"])[:40],
                  str(line["account_name"])[:120], Decimal(str(line["amount"])),
                  str(line.get("tax_code") or "")[:40], str(line.get("memo") or "")[:500])
                 for index, line in enumerate(lines, 1)],
            )
            row = await conn.fetchrow(
                "UPDATE yeoljeong_journal_vouchers SET description=$3,status='draft',updated_at=NOW() WHERE id=$1 AND tenant_id=$2 RETURNING *",
                voucher_id, tenant_id, str(payload.get("description") or current["description"])[:500],
            )
        return dict(row)
    finally:
        await conn.close()


async def transition_journal(*, user: dict[str, Any], voucher_id: UUID, action: str) -> dict[str, Any]:
    require_journal_write_enabled()
    _require_approve(user)
    tenant_id = _tenant(user)
    transitions = {"approve": ({"draft", "needs_review"}, "approved"), "post": ({"approved"}, "posted")}
    if action not in transitions:
        raise HTTPException(status_code=400, detail="지원하지 않는 전표 상태 전환입니다")
    conn = await _connect()
    try:
        row = await conn.fetchrow("SELECT * FROM yeoljeong_journal_vouchers WHERE id=$1 AND tenant_id=$2", voucher_id, tenant_id)
        if not row:
            raise HTTPException(status_code=404, detail="전표를 찾을 수 없습니다")
        await _require_business(conn, tenant_id, row["business_id"])
        allowed, target = transitions[action]
        if row["status"] not in allowed:
            raise HTTPException(status_code=409, detail="현재 상태에서는 요청한 전표 전환을 할 수 없습니다")
        balance = await conn.fetchrow("SELECT COALESCE(SUM(amount) FILTER (WHERE side='debit'),0) debit,COALESCE(SUM(amount) FILTER (WHERE side='credit'),0) credit,COUNT(*) count FROM yeoljeong_journal_lines WHERE voucher_id=$1", voucher_id)
        if balance["count"] < 2 or balance["debit"] != balance["credit"] or balance["debit"] != row["total_amount"]:
            raise HTTPException(status_code=422, detail="균형 및 합계 검증을 통과하지 못했습니다")
        return dict(await conn.fetchrow(
            """UPDATE yeoljeong_journal_vouchers SET status=$3,approved_by=CASE WHEN $3='approved' THEN $4 ELSE approved_by END,
                       approved_at=CASE WHEN $3='approved' THEN NOW() ELSE approved_at END,
                       posted_at=CASE WHEN $3='posted' THEN NOW() ELSE posted_at END,updated_at=NOW()
                 WHERE id=$1 AND tenant_id=$2 RETURNING *""", voucher_id, tenant_id, target, _actor(user)))
    finally:
        await conn.close()


async def reverse_journal(*, user: dict[str, Any], voucher_id: UUID) -> dict[str, Any]:
    require_journal_write_enabled()
    _require_approve(user)
    tenant_id = _tenant(user)
    conn = await _connect()
    try:
        original = await conn.fetchrow("SELECT * FROM yeoljeong_journal_vouchers WHERE id=$1 AND tenant_id=$2", voucher_id, tenant_id)
        if not original:
            raise HTTPException(status_code=404, detail="전표를 찾을 수 없습니다")
        await _require_business(conn, tenant_id, original["business_id"])
        async with conn.transaction():
            original = await conn.fetchrow(
                "SELECT * FROM yeoljeong_journal_vouchers WHERE id=$1 AND tenant_id=$2 FOR UPDATE",
                voucher_id, tenant_id,
            )
            if original["status"] != "posted":
                existing = await conn.fetchrow(
                    "SELECT * FROM yeoljeong_journal_vouchers WHERE reversal_of_id=$1 AND tenant_id=$2",
                    voucher_id, tenant_id,
                )
                if existing:
                    result = dict(existing)
                    result["idempotent"] = True
                    return result
                raise HTTPException(status_code=409, detail="확정된 전표만 역분개할 수 있습니다")
            reversal_id = uuid4()
            number = f"RV-{original['voucher_no']}-{reversal_id.hex[:6].upper()}"
            original_lines = await conn.fetch(
                "SELECT line_no,side,account_code,account_name,amount,tax_code,memo FROM yeoljeong_journal_lines WHERE voucher_id=$1 ORDER BY line_no",
                voucher_id,
            )
            reversal = await conn.fetchrow(
                """INSERT INTO yeoljeong_journal_vouchers
                    (id,tenant_id,business_id,voucher_no,transaction_date,description,status,source_type,source_id,
                     supply_amount,tax_amount,total_amount,evidence_source,export_status,created_by,reversal_of_id,posted_at)
                   VALUES ($1,$2,$3,$4,CURRENT_DATE,$5,'posted',$6,$7,$8,$9,$10,$11,'not_exported',$12,$13,NOW()) RETURNING *""",
                reversal_id, tenant_id, original["business_id"], number, f"역분개: {original['description']}",
                original["source_type"], uuid4(), original["supply_amount"], original["tax_amount"],
                original["total_amount"], original["evidence_source"], _actor(user), voucher_id,
            )
            await conn.executemany(
                """INSERT INTO yeoljeong_journal_lines (id,voucher_id,line_no,side,account_code,account_name,amount,tax_code,memo)
                   VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9)""",
                [(uuid4(), reversal_id, line["line_no"], "credit" if line["side"] == "debit" else "debit",
                  line["account_code"], line["account_name"], line["amount"], line["tax_code"],
                  f"역분개: {line['memo']}") for line in original_lines],
            )
            await conn.execute("UPDATE yeoljeong_journal_vouchers SET status='reversed',reversed_at=NOW(),updated_at=NOW() WHERE id=$1", voucher_id)
        return dict(reversal)
    finally:
        await conn.close()
