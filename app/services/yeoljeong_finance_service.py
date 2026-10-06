"""JSON-backed store for the Yeoljeong store assistant app.

The app is a static SPA, so this service keeps the first operational HR flow
small and explicit: employee join requests, onboarding documents, contracts,
payroll statements, and delivery account status.
"""
from __future__ import annotations

import asyncio
import base64
import binascii
import csv
import fcntl
import hashlib
import io
import json
import logging
import math
import mimetypes
import os
import re
import secrets
import shutil
import threading
import time
import urllib.error
import urllib.request
import zipfile
from contextlib import AsyncExitStack
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from urllib.parse import urlparse
from uuid import UUID, uuid4

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from app.services.auth_challenge_orchestrator import approved_operator_input, classify_portal_state, make_resume_token
from app.services.browser_collection_audit import SITE_STAGE_LOG_SCHEMA, append_site_stage_log

KST = timezone(timedelta(hours=9))
DATA_DIR = Path(os.getenv("YEOLJEONG_FINANCE_DATA_DIR", "app/data/yeoljeong_finance"))
UPLOAD_DIR = DATA_DIR / "uploads" / "onboarding"
EVIDENCE_UPLOAD_DIR = DATA_DIR / "uploads" / "evidence"
MAX_UPLOAD_BYTES = 15 * 1024 * 1024
MAX_SIGNATURE_BYTES = 256 * 1024
DELIVERY_SYNC_STALE_AFTER = timedelta(minutes=15)
DELIVERY_BACKFILL_STALE_AFTER = timedelta(
    minutes=max(20, int(os.getenv("YEOLJEONG_DELIVERY_BACKFILL_STALE_MINUTES", "60")))
)
DELIVERY_BACKFILL_MAX_ATTEMPTS = max(1, int(os.getenv("YEOLJEONG_DELIVERY_BACKFILL_MAX_ATTEMPTS", "3")))
DELIVERY_CHALLENGE_TIMEOUT = timedelta(minutes=20)
DELIVERY_CHALLENGE_MAX_ATTEMPTS = 3
DELIVERY_OPERATOR_ACTION_COOLDOWN_MINUTES = max(
    1,
    int(os.getenv("YEOLJEONG_DELIVERY_OPERATOR_ACTION_COOLDOWN_MINUTES", "45")),
)
DELIVERY_OPERATOR_ACTION_COOLDOWN_CODES = {
    "BAEMIN_SECURITY_BLOCKED",
    "COUPANGEATS_SECURITY_BLOCKED",
    "DDANGYO_NUMERIC_CAPTCHA_REQUIRED",
    "MISSING_CREDENTIALS",
    "PC_AGENT_LOGIN_REQUIRED",
    "PC_AGENT_SESSION_REQUIRED",
    "PORTAL_AUTH_CHALLENGE",
    "PORTAL_BLOCKED",
}
BAEMIN_SECURITY_BLOCK_COOLDOWN_MINUTES = max(
    1,
    int(os.getenv("YEOLJEONG_BAEMIN_SECURITY_BLOCK_COOLDOWN_MINUTES", "45")),
)
BAEMIN_SECURITY_BLOCK_CODES = {"BAEMIN_SECURITY_BLOCKED", "PORTAL_BLOCKED", "SECURITY_BLOCKED"}
CONTRACT_SIGNATURE_CONSENT_VERSION = "yeoljeong-contract-sign-v1"
try:
    CONTRACT_SIGN_LINK_TTL_DAYS = max(1, int(os.getenv("OBYS_CONTRACT_SIGN_LINK_TTL_DAYS", "14")))
except ValueError:
    CONTRACT_SIGN_LINK_TTL_DAYS = 14
_CONTRACT_SIGN_LOCK = threading.Lock()

DOCUMENT_TYPES: list[dict[str, str]] = [
    {
        "type": "resident_register",
        "label": "주민등록등본",
        "requirement": "필수",
        "notice": "주민등록번호 뒷자리와 가족정보는 마스킹본을 원칙으로 합니다.",
    },
    {"type": "id_card", "label": "신분증", "requirement": "필수", "notice": "주민등록번호 뒷자리는 마스킹합니다."},
    {"type": "bankbook", "label": "통장사본", "requirement": "필수", "notice": "급여 입금 계좌 확인용입니다."},
    {"type": "health_certificate", "label": "보건증", "requirement": "필수", "notice": "식품위생 업종 제출 서류입니다."},
    {"type": "employment_contract_info", "label": "계약 정보 확인서", "requirement": "선택", "notice": "근무조건 확인용입니다."},
    {"type": "foreign_registration", "label": "외국인등록증", "requirement": "외국인 조건부", "notice": "체류자격과 취업 가능 여부를 확인합니다."},
    {"type": "visa_status_certificate", "label": "체류자격 증빙", "requirement": "외국인 조건부", "notice": "비자 유형별 취업 가능 범위를 확인합니다."},
    {"type": "work_permission_confirmation", "label": "취업활동 가능 확인", "requirement": "외국인 조건부", "notice": "고용허가/시간제취업 허가 등을 확인합니다."},
    {"type": "foreign_employment_report", "label": "외국인 고용 신고/변동 확인", "requirement": "외국인 조건부", "notice": "신고 대상 여부와 처리일을 기록합니다."},
    {"type": "other", "label": "기타 서류", "requirement": "선택", "notice": "수집 사유를 메모에 남깁니다."},
]

CONNECTOR_LABELS = {
    "baemin": "배민셀프서비스",
    "coupangeats": "쿠팡이츠",
    "yogiyo": "요기요",
    "ddangyo": "땡겨요",
    "matepos": "메이트포스",
    "shinhan_business": "신한은행 기업",
    "ibk_business": "기업은행 기업",
    "coupang_supplier": "쿠팡 매입처",
    "marketbom": "마켓봄",
    "newtong": "뉴통",
    "baljugo": "발주고",
    "supplier_custom_portal": "기타 매입처 주문프로그램",
    "supplier_statement_upload": "거래내역서/영수증 업로드",
    "hometax": "홈택스",
    "tax_invoice_upload": "계산서/증빙 업로드",
    "card_pg": "카드사/PG 매출",
    "utility_bills": "공과금 고지서",
    "accountant_tax_agent": "세무대리인/회계프로그램",
}
PLATFORM_LABELS = {
    key: CONNECTOR_LABELS[key]
    for key in ("baemin", "coupangeats", "yogiyo", "ddangyo")
}
DELIVERY_AGENT_VAULT_ORIGINS: dict[str, tuple[str, ...]] = {
    "baemin": ("https://biz-member.baemin.com", "https://self.baemin.com"),
    "coupangeats": (
        "https://store.coupangeats.com",
        "https://xauth.coupang.com",
        "https://login.coupang.com",
    ),
    "yogiyo": ("https://ceo.yogiyo.co.kr",),
    "ddangyo": ("https://boss.ddangyo.com",),
}
DELIVERY_UPLOAD_COLLECTION_MODES = {"portal-csv", "csv-upload", "statement-upload", "upload_queue", "manual"}
DELIVERY_COLLECTION_STATUSES = {"queued", "running", "succeeded", "partial", "action_required", "failed"}
DELIVERY_ACTION_REQUIRED_STATUSES = {
    "blocked",
    "credential_required",
    "credentials_missing",
    "connector_not_configured",
    "portal_action_required",
    "upload_required",
}
FINANCIAL_TRANSACTION_SERVICES = {
    "shinhan_business",
    "ibk_business",
    "card_pg",
}
TRANSACTION_SOURCE_BY_SERVICE = {
    "shinhan_business": "bank",
    "ibk_business": "bank",
    "card_pg": "card",
}
BANK_QUICK_SERVICE_CONFIG = {
    "shinhan_business": {
        "label": "신한은행 간편서비스",
        "login_url": "https://bank.shinhan.com/rib/easy/index.jsp#210000000000",
        "enrollment": "기업뱅킹에서 간편조회 허용 계좌 등록 후 간편서비스 계좌조회로 거래내역을 확인합니다.",
    },
    "ibk_business": {
        "label": "IBK기업은행 빠른서비스",
        "login_url": "https://kiup.ibk.co.kr/uib/jsp/guest/qcs/qcs10/qcs1010/PQCS101000_i.jsp",
        "enrollment": "기업뱅킹의 빠른조회서비스 신청/해제에서 대상 계좌를 등록한 뒤 빠른조회로 거래내역을 확인합니다.",
    },
}
BANK_AGENT_VAULT_ORIGINS: dict[str, tuple[str, ...]] = {
    # Keep the legacy corporate origin because CEO Password Manager imports
    # were saved before the collector moved to Shinhan EasyView.
    "shinhan_business": (
        "https://bank.shinhan.com",
        "https://bizbank.shinhan.com",
    ),
    "ibk_business": (
        "https://kiup.ibk.co.kr",
        "https://mybank.ibk.co.kr",
    ),
}


def _normalize_bank_quick_login_url(service: str, login_url: Any) -> str:
    """Keep Shinhan quick-service runs on the corporate simple-account page."""
    configured = BANK_QUICK_SERVICE_CONFIG.get(service, {})
    fallback = str(configured.get("login_url") or "").strip()
    raw = str(login_url or "").strip()
    if service == "shinhan_business":
        if not raw or "bizbank.shinhan.com" in raw or "bank.shinhan.com/rib/easy/index.jsp" in raw:
            return fallback
    if service == "ibk_business":
        if not raw or "mybank.ibk.co.kr" in raw or "PQCS102000" in raw:
            return fallback
    return raw or fallback

DEFAULT_CATEGORY_RULES: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("식자재", ("쌀", "백미", "고춧가루", "소스", "김치", "대파", "양파", "고기", "식자재")),
    ("배달앱", ("배민", "배달의민족", "요기요", "쿠팡이츠", "정산")),
    ("임차료", ("월세", "임대료", "관리비")),
    ("인건비", ("급여", "4대보험", "고용", "알바", "직원")),
    ("공과금", ("전기", "가스", "수도", "통신", "인터넷")),
    ("카드수수료", ("카드수수료", "수수료")),
    ("비품", ("비품", "소모품", "주방", "용기", "포장")),
)

SETTINGS_TABLES = ("yeoljeong_businesses", "yeoljeong_branches", "yeoljeong_settings")
HR_LEDGER_TABLES = (
    "yeoljeong_employee_join_requests",
    "yeoljeong_onboarding_documents",
    "yeoljeong_contracts",
    "yeoljeong_payroll_statements",
)
DELIVERY_LEDGER_TABLES = (
    "yeoljeong_platform_accounts",
    "yeoljeong_delivery_sales",
    "yeoljeong_delivery_settlements",
    "yeoljeong_delivery_reviews",
    "yeoljeong_delivery_ads",
    "yeoljeong_delivery_collection_status",
)
DELIVERY_RECORD_TYPES = ("sales", "settlements", "reviews", "ads")
DELIVERY_DATA_LEDGER_NAMES = {
    "delivery_sales",
    "delivery_settlements",
    "delivery_reviews",
    "delivery_ads",
}


def _delivery_empty_counts() -> dict[str, int]:
    return {kind: 0 for kind in DELIVERY_RECORD_TYPES}


def _delivery_empty_record_lists() -> dict[str, list[dict[str, Any]]]:
    return {kind: [] for kind in DELIVERY_RECORD_TYPES}


def _delivery_record_kind_from_ledger_name(name: str) -> str:
    if name.startswith("delivery_"):
        return name.removeprefix("delivery_")
    return name


def _delivery_int_value(value: Any) -> int:
    if isinstance(value, bool):
        return 0
    if isinstance(value, (int, float)):
        return int(value)
    text = re.sub(r"[^0-9.-]", "", str(value or ""))
    if not text or text in {"-", ".", "-."}:
        return 0
    try:
        return int(round(float(text)))
    except ValueError:
        return 0


def _delivery_record_has_text(record: dict[str, Any], *fields: str) -> bool:
    return any(str(record.get(field) or "").strip() for field in fields)


def _delivery_record_is_meaningful(kind: str, record: dict[str, Any]) -> bool:
    """Reject placeholder rows parsed from portal chrome or empty tables."""
    normalized_kind = _delivery_record_kind_from_ledger_name(str(kind or ""))
    if normalized_kind == "sales":
        return _delivery_int_value(record.get("gross_amount")) > 0 and _delivery_record_has_text(
            record,
            "order_id",
            "order_no",
            "source_id",
        )
    if normalized_kind == "settlements":
        amount = max(
            _delivery_int_value(record.get("settlement_amount")),
            _delivery_int_value(record.get("sales_amount")),
        )
        return amount > 0 and _delivery_record_has_text(record, "settlement_id", "order_id", "order_no", "source_id")
    if normalized_kind == "reviews":
        return bool(str(record.get("review_text") or "").strip()) or _delivery_int_value(record.get("rating")) > 0
    if normalized_kind == "ads":
        return _delivery_record_has_text(record, "campaign_name", "ad_product") or max(
            _delivery_int_value(record.get("cost_amount")),
            _delivery_int_value(record.get("impressions")),
            _delivery_int_value(record.get("clicks")),
            _delivery_int_value(record.get("orders")),
            _delivery_int_value(record.get("sales_amount")),
        ) > 0
    return True


def _delivery_filter_meaningful_records(
    records: Any,
) -> tuple[dict[str, list[dict[str, Any]]], dict[str, int]]:
    filtered = _delivery_empty_record_lists()
    rejected = _delivery_empty_counts()
    if not isinstance(records, dict):
        return filtered, rejected
    for kind in DELIVERY_RECORD_TYPES:
        rows = records.get(kind) or []
        if not isinstance(rows, list):
            continue
        for row in rows:
            if not isinstance(row, dict):
                rejected[kind] += 1
                continue
            if _delivery_record_is_meaningful(kind, row):
                filtered[kind].append(row)
            else:
                rejected[kind] += 1
    return filtered, rejected
JSON_LEDGER_FILES = (
    "employee_join_requests",
    "onboarding_documents",
    "contracts",
    "payroll_statements",
)
DB_LEDGER_TABLE_BY_NAME = {
    "employee_join_requests": "yeoljeong_employee_join_requests",
    "onboarding_documents": "yeoljeong_onboarding_documents",
    "contracts": "yeoljeong_contracts",
    "payroll_statements": "yeoljeong_payroll_statements",
    "platform_accounts": "yeoljeong_platform_accounts",
    "delivery_sales": "yeoljeong_delivery_sales",
    "delivery_settlements": "yeoljeong_delivery_settlements",
    "delivery_reviews": "yeoljeong_delivery_reviews",
    "delivery_ads": "yeoljeong_delivery_ads",
    "delivery_collection_status": "yeoljeong_delivery_collection_status",
}
GENERIC_DB_LEDGER_NAMES = {
    "platform_accounts",
    "delivery_sales",
    "delivery_settlements",
    "delivery_reviews",
    "delivery_ads",
    "delivery_collection_status",
}

CONTRACT_TEMPLATE_META = {
    "freelancer": {
        "document_kind": "freelancer_service_contract",
        "template_version": "majangbiseo-freelancer-2026-07-identity-table-v3",
        "print_title": "3.3% 프리랜서 용역계약서",
    },
    "confidentiality": {
        "document_kind": "confidentiality_agreement",
        "template_version": "majangbiseo-confidentiality-2026-07-identity-table-v3",
        "print_title": "보안 및 개인정보 보호 서약서",
    },
    "default": {
        "document_kind": "standard_employment_contract",
        "template_version": "majangbiseo-employment-2026-07-identity-table-v3",
        "print_title": "표준근로계약서",
    },
}

EMPLOYMENT_CONTRACT_TYPES = {"part_time", "regular", "manager"}
VALID_CONTRACT_TYPES = EMPLOYMENT_CONTRACT_TYPES | {"freelancer", "confidentiality"}
CONTRACT_SNAPSHOT_EXCLUDED_FIELDS = {
    "sign_token",
    "sign_token_hash",
    "signed_snapshot",
    "signed_snapshot_sha256",
    "signature_data_uri",
    "updated_at",
    # 서명 뒤에 붙는 보관·교부 메타. 봉인 스냅샷에 섞이면 해시가 흔들린다.
    "signed_pdf_path",
    "signed_pdf_sha256",
    "signed_pdf_bytes",
    "signed_pdf_generated_at",
    "signed_pdf_error",
    "delivered_at",
    "delivery_channel",
}

CANONICAL_BUSINESSES: list[dict[str, Any]] = [
    {
        "id": "biz-junghwa",
        "entityType": "corporation",
        "name": "열정국밥 중화점",
        "registrationNo": "710-86-04499",
        "representative": "오윤희",
        "taxType": "일반과세",
        "openedAt": "2026-07-01",
        "address": "서울특별시 중랑구 봉화산로27길 8, 1층(중화동)",
        "memo": "법인사업자 / 법인등록번호 110111-0961922 / 주류판매신고번호 146-5-11334",
    },
    {
        "id": "biz-sungshin",
        "entityType": "individual",
        "name": "열정국밥 성신여대점",
        "registrationNo": "",
        "representative": "",
        "taxType": "일반과세",
        "openedAt": "",
        "address": "",
        "memo": "개인사업자 2",
    },
    {
        "id": "biz-eonni-naengmyeon",
        "entityType": "individual",
        "name": "언니냉면",
        "registrationNo": "",
        "representative": "",
        "taxType": "일반과세",
        "openedAt": "",
        "address": "",
        "memo": "계정표 기준 자동수집 사업자",
    },
    {
        "id": "biz-mia",
        "entityType": "individual",
        "name": "열정국밥_미아점",
        "registrationNo": "874-21-02160",
        "representative": "최미미",
        "taxType": "일반과세",
        "openedAt": "2025-04-01",
        "address": "서울특별시 강북구 도봉로76길 42, 1층 점포일부(좌측)",
        "memo": "개인사업자 3 / 주류판매신고번호 210-5-62608",
    },
]

CANONICAL_BRANCHES: list[dict[str, Any]] = [
    {"id": "branch-junghwa", "name": "중화점", "businessId": "biz-junghwa", "status": "active", "phone": "", "address": "서울특별시 중랑구 봉화산로27길 8, 1층(중화동)"},
    {"id": "branch-sungshin", "name": "성신여대점", "businessId": "biz-sungshin", "status": "active", "phone": "", "address": ""},
    {"id": "branch-gangbuk-mia", "name": "열정국밥_미아점", "businessId": "biz-mia", "status": "active", "phone": "", "address": "서울특별시 강북구 도봉로76길 42, 1층 점포일부(좌측)"},
]

# 사업자등록증 기준 미비 여부를 판정하는 항목(설정 JSON 키). 자리표시자는 값이 아니라
# 상태이므로 needs_registration_info / missing_registration_fields 로만 내보낸다.
BUSINESS_REGISTRATION_FIELDS = ("registrationNo", "representative", "openedAt", "address")
CANONICAL_BUSINESS_IDS = {item["id"] for item in CANONICAL_BUSINESSES}

# 사업자 등록 4항목의 자리표시자. 값이 아니라 "아직 없음" 이므로 설정 동기화(save_settings_persisted)
# 에서는 빈 값과 같게 취급한다 — 들어오는 쪽도, DB 에 이미 들어 있는 쪽도.
BUSINESS_PLACEHOLDER_VALUES = ("기초등록 필요", "미등록")
# 설정 동기화가 "비어 있을 때만" 채우는 컬럼. 관리자 직접 수정(obys_upload_service.update_business)은 별개 경로.
BUSINESS_FILL_ONLY_COLUMNS = ("registration_no", "representative", "opened_at", "address")


def _business_registration_value(value: Any) -> str:
    text = str(value or "").strip()
    return "" if text in BUSINESS_PLACEHOLDER_VALUES else text


def _build_business_settings_upsert_sql() -> str:
    placeholders = ", ".join("'" + value.replace("'", "''") + "'" for value in BUSINESS_PLACEHOLDER_VALUES)
    fill_only = ",\n".join(
        f"""                               {column} = CASE
                                   WHEN COALESCE(TRIM(yeoljeong_businesses.{column}), '') = ''
                                     OR TRIM(yeoljeong_businesses.{column}) IN ({placeholders})
                                   THEN EXCLUDED.{column}
                                   ELSE yeoljeong_businesses.{column}
                               END"""
        for column in BUSINESS_FILL_ONLY_COLUMNS
    )
    return f"""
                        INSERT INTO yeoljeong_businesses
                            (id, entity_type, name, registration_no, representative, tax_type,
                             opened_at, address, memo, sort_order, updated_by)
                        VALUES ($1, $2, $3, $4, $5, $6, $7, $8, $9, $10, $11)
                        ON CONFLICT (id) DO UPDATE
                           SET entity_type = EXCLUDED.entity_type,
                               name = EXCLUDED.name,
{fill_only},
                               tax_type = EXCLUDED.tax_type,
                               memo = EXCLUDED.memo,
                               sort_order = EXCLUDED.sort_order,
                               updated_by = EXCLUDED.updated_by,
                               updated_at = NOW(),
                               deleted_at = NULL
                        """


BUSINESS_SETTINGS_UPSERT_SQL = _build_business_settings_upsert_sql()
CANONICAL_BRANCH_NAMES = {item["name"] for item in CANONICAL_BRANCHES}
MIA_BUSINESS_ID = "biz-mia"
MIA_BRANCH_NAME = "열정국밥_미아점"
BRANCH_ALIASES = {
    "열정국밥 강북미아점": MIA_BRANCH_NAME,
    "강북미아점": MIA_BRANCH_NAME,
    "미아점": MIA_BRANCH_NAME,
}
BUSINESS_BY_BRANCH = {
    **{item["name"]: item["businessId"] for item in CANONICAL_BRANCHES},
    "성신여대역점": "biz-eonni-naengmyeon",
}


def _now() -> str:
    return datetime.now(KST).isoformat(timespec="seconds")


def _delivery_baemin_security_cooldown_until(reference: datetime | None = None) -> str:
    base = reference or datetime.now(KST)
    return (base + timedelta(minutes=BAEMIN_SECURITY_BLOCK_COOLDOWN_MINUTES)).isoformat(timespec="seconds")


def _delivery_operator_action_cooldown_until(reference: datetime | None = None) -> str:
    base = reference or datetime.now(KST)
    return (base + timedelta(minutes=DELIVERY_OPERATOR_ACTION_COOLDOWN_MINUTES)).isoformat(timespec="seconds")


def _delivery_result_requires_operator_cooldown(public_status: str, public_error_code: str) -> bool:
    return (
        str(public_status or "").strip() == "action_required"
        and str(public_error_code or "").strip().upper() in DELIVERY_OPERATOR_ACTION_COOLDOWN_CODES
    )


def _ensure_dirs() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    _evidence_upload_dir().mkdir(parents=True, exist_ok=True)


def _evidence_upload_dir() -> Path:
    return DATA_DIR / "uploads" / "evidence"


def _path(name: str) -> Path:
    _ensure_dirs()
    return DATA_DIR / f"{name}.json"


def _read_file_rows(name: str) -> list[dict[str, Any]]:
    path = _path(name)
    if not path.exists():
        return []
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        raise HTTPException(status_code=500, detail=f"{name} 저장소 JSON이 손상되었습니다")
    if isinstance(data, dict):
        data = data.get(name) or data.get("items") or data.get("records") or []
    return data if isinstance(data, list) else []


def _write_file_rows(name: str, rows: list[dict[str, Any]]) -> None:
    path = _path(name)
    tmp = path.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(path)


def _read_json_object(name: str) -> dict[str, Any]:
    path = _path(name)
    if not path.exists():
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        raise HTTPException(status_code=500, detail=f"{name} 저장소 JSON이 손상되었습니다")
    return data if isinstance(data, dict) else {}


def _write_json_object(name: str, data: dict[str, Any]) -> None:
    path = _path(name)
    tmp = path.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(path)


logger = logging.getLogger(__name__)


def _db_url() -> str:
    """오비서 업무 DSN 만 본다. AADS ``DATABASE_URL`` 로 폴백하지 않는다.

    폴백을 두면 환경변수 하나가 빠졌을 때 오비서가 조용히 aads DB 의
    ``yeoljeong_*`` 를 읽고 쓰게 되고, 두 DB 가 갈라진 뒤에야 드러난다.
    같은 이유로 ``app/core/obys_db.py`` 는 2026-09-19 에 이미 폴백을 걷어냈는데
    이 경로만 남아 있었다. 진아서버 이전에서는 이 경로가 그대로 교차 쓰기가
    되므로 여기서 끊는다.
    """
    url = (
        os.getenv("YEOLJEONG_FINANCE_DATABASE_URL")
        or os.getenv("OBYS_DATABASE_URL")
        or ""
    ).strip()
    if not url:
        logger.error(
            "yeoljeong_finance: OBYS_DATABASE_URL 미설정 — aads DB 로 폴백하지 않고 "
            "DB 경로를 비활성화한다"
        )
        return ""
    return url.replace("postgresql://", "postgres://")


def _db_available() -> bool:
    return bool(_db_url())


def _run_db(coro: Any) -> Any:
    if not _db_available():
        close = getattr(coro, "close", None)
        if callable(close):
            close()
        return None
    try:
        asyncio.get_running_loop()
    except RuntimeError:
        pass
    else:
        close = getattr(coro, "close", None)
        if callable(close):
            close()
        return None
    try:
        return asyncio.run(coro)
    except Exception:
        close = getattr(coro, "close", None)
        if callable(close):
            close()
        return None


def _run_async(coro: Any) -> Any:
    try:
        asyncio.get_running_loop()
    except RuntimeError:
        pass
    else:
        close = getattr(coro, "close", None)
        if callable(close):
            close()
        return None
    try:
        return asyncio.run(coro)
    except Exception:
        close = getattr(coro, "close", None)
        if callable(close):
            close()
        return None


async def _db_table_exists(table: str) -> bool:
    import asyncpg

    conn = await asyncpg.connect(_db_url(), timeout=5)
    try:
        return bool(await conn.fetchval("SELECT to_regclass($1) IS NOT NULL", f"public.{table}"))
    finally:
        await conn.close()


def _table_ready(table: str) -> bool:
    return bool(_run_db(_db_table_exists(table)))


def _iso(value: Any) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat(timespec="seconds")
        except TypeError:
            return value.isoformat()
    return str(value)


def _pg_ts(value: Any) -> datetime | None:
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).astimezone(KST)
    except ValueError:
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%Y/%m/%d %H:%M:%S", "%Y/%m/%d"):
            try:
                parsed = datetime.strptime(text, fmt)
                return parsed.replace(tzinfo=KST)
            except ValueError:
                continue
    return None


def _payload_dict(value: Any) -> dict[str, Any]:
    parsed = _jsonb_object(value)
    return parsed if isinstance(parsed, dict) else {}


HR_TENANT_LEDGER_NAMES = {
    "employee_join_requests",
    "onboarding_documents",
    "contracts",
    "payroll_statements",
}
PAYROLL_AMOUNT_LIMIT = 9_000_000_000_000_000


def _tenant_id(user: dict[str, Any] | None) -> str:
    """Return only a verified active membership tenant; never fall back to user.tenant_id."""
    membership = (user or {}).get("current_membership")
    user_tenant = str((user or {}).get("tenant_id") or "").strip()
    if not isinstance(membership, dict):
        raise HTTPException(status_code=403, detail="활성 테넌트 멤버십이 필요합니다")
    membership_tenant = str(membership.get("tenant_id") or "").strip()
    if (
        str(membership.get("status") or "").strip().lower() != "active"
        or not membership_tenant
        or not user_tenant
        or membership_tenant != user_tenant
    ):
        raise HTTPException(status_code=403, detail="활성 테넌트 멤버십이 필요합니다")
    try:
        UUID(membership_tenant)
    except (TypeError, ValueError, AttributeError):
        raise HTTPException(status_code=403, detail="유효한 테넌트 멤버십이 필요합니다") from None
    return membership_tenant


def _record_tenant_uuid(record: dict[str, Any]) -> UUID | None:
    value = str(record.get("tenant_id") or "").strip()
    try:
        return UUID(value) if value else None
    except (TypeError, ValueError, AttributeError):
        return None


def _payroll_integer(value: Any, *, field: str, default: int = 0) -> int:
    """Parse an exact payroll integer without accepting bool, float, exponent, or truncation."""
    if value is None or value == "":
        return default
    if isinstance(value, bool) or isinstance(value, float):
        raise ValueError(f"{field}: exact integer required")
    if isinstance(value, int):
        parsed = value
    elif isinstance(value, str) and re.fullmatch(r"[+-]?\d+", value.strip()):
        parsed = int(value.strip())
    else:
        raise ValueError(f"{field}: exact integer required")
    if abs(parsed) > PAYROLL_AMOUNT_LIMIT:
        raise ValueError(f"{field}: out of range")
    return parsed


def _payroll_integer_for_api(value: Any, *, field: str, default: int = 0) -> int:
    try:
        return _payroll_integer(value, field=field, default=default)
    except ValueError:
        raise HTTPException(status_code=400, detail=f"{field}은 범위 내 정수로 입력하십시오") from None


def _db_row_to_record(name: str, row: Any) -> dict[str, Any]:
    item = dict(row)
    if name in GENERIC_DB_LEDGER_NAMES:
        payload = _payload_dict(item.get("payload"))
        return {
            **payload,
            "id": str(payload.get("id") or item.get("row_id") or ""),
            "business_id": payload.get("business_id") or item.get("business_id") or "",
            "branch": payload.get("branch") or item.get("branch") or "",
            "created_at": payload.get("created_at") or _iso(item.get("created_at")),
            "updated_at": payload.get("updated_at") or _iso(item.get("updated_at")),
        }
    if name == "employee_join_requests":
        payload = _payload_dict(item.get("request_payload"))
        email = str(payload.get("email") or item.get("employee_email") or "").strip().lower()
        return {
            **payload,
            "id": str(payload.get("id") or item.get("id") or ""),
            "email": email,
            "email_masked": payload.get("email_masked") or item.get("employee_email_masked") or _mask_email(email),
            "name": payload.get("name") or item.get("employee_name") or "",
            "phone": payload.get("phone") or item.get("phone") or "",
            "business_id": item.get("business_id") or "",
            "branch": payload.get("branch") or item.get("branch") or "",
            "role": payload.get("role") or item.get("role") or "employee",
            "status": payload.get("status") or item.get("status") or "pending",
            "review_memo": payload.get("review_memo") or item.get("review_memo") or "",
            "requested_by": payload.get("requested_by") or item.get("requested_by") or "",
            "reviewed_by": payload.get("reviewed_by") or item.get("reviewed_by") or "",
            "requested_at": payload.get("requested_at") or _iso(item.get("requested_at")),
            "reviewed_at": payload.get("reviewed_at") or _iso(item.get("reviewed_at")),
            "created_at": payload.get("created_at") or _iso(item.get("created_at")),
            "updated_at": payload.get("updated_at") or _iso(item.get("updated_at")),
            "tenant_id": str(item.get("tenant_id") or ""),
        }
    if name == "onboarding_documents":
        payload = _payload_dict(item.get("metadata"))
        email = str(payload.get("employee_email") or item.get("employee_email") or "").strip().lower()
        return {
            **payload,
            "id": str(payload.get("id") or item.get("id") or ""),
            "employee_request_id": payload.get("employee_request_id") or item.get("employee_request_id") or "",
            "employee_email": email,
            "employee_email_masked": payload.get("employee_email_masked") or item.get("employee_email_masked") or _mask_email(email),
            "employee_name": payload.get("employee_name") or item.get("employee_name") or "",
            "business_id": item.get("business_id") or "",
            "branch": payload.get("branch") or item.get("branch") or "",
            "document_type": payload.get("document_type") or item.get("document_type") or "",
            "document_label": payload.get("document_label") or item.get("document_label") or "",
            "requirement": payload.get("requirement") or item.get("requirement") or "",
            "status": payload.get("status") or item.get("status") or "uploaded",
            "original_filename": payload.get("original_filename") or item.get("original_filename") or "",
            "stored_filename": payload.get("stored_filename") or item.get("stored_filename") or "",
            "content_type": payload.get("content_type") or item.get("content_type") or "",
            "size_bytes": payload.get("size_bytes") if payload.get("size_bytes") is not None else int(item.get("size_bytes") or 0),
            "issue_date": payload.get("issue_date") or item.get("issue_date") or "",
            "memo": payload.get("memo") or item.get("memo") or "",
            "review_memo": payload.get("review_memo") or item.get("review_memo") or "",
            "uploaded_by": payload.get("uploaded_by") or item.get("uploaded_by") or "",
            "reviewed_by": payload.get("reviewed_by") or item.get("reviewed_by") or "",
            "uploaded_at": payload.get("uploaded_at") or _iso(item.get("uploaded_at")),
            "reviewed_at": payload.get("reviewed_at") or _iso(item.get("reviewed_at")),
            "created_at": payload.get("created_at") or _iso(item.get("created_at")),
            "updated_at": payload.get("updated_at") or _iso(item.get("updated_at")),
            "tenant_id": str(item.get("tenant_id") or ""),
            # 무결성·이력 필드는 물리 컬럼이 원본이다. metadata 에 같은 키가 있어도 컬럼이 이긴다.
            "expires_at": _iso(item.get("expires_at")) or str(payload.get("expires_at") or ""),
            "sha256": str(item.get("sha256") or payload.get("sha256") or ""),
            "stored_path": str(item.get("stored_path") or payload.get("stored_path") or ""),
            "superseded_by": str(item.get("superseded_by") or payload.get("superseded_by") or ""),
        }
    if name == "contracts":
        payload = _payload_dict(item.get("contract_payload"))
        email = str(payload.get("employee_email") or item.get("employee_email") or "").strip().lower()
        return {
            **payload,
            "id": str(payload.get("id") or item.get("id") or ""),
            "employee_email": email,
            "employee_email_masked": payload.get("employee_email_masked") or item.get("employee_email_masked") or _mask_email(email),
            "employee_name": payload.get("employee_name") or item.get("employee_name") or "",
            "business_id": item.get("business_id") or "",
            "branch": payload.get("branch") or item.get("branch") or "",
            "contract_type": payload.get("contract_type") or item.get("contract_type") or "part_time",
            "document_kind": payload.get("document_kind") or item.get("document_kind") or "standard_employment_contract",
            "template_version": payload.get("template_version") or item.get("template_version") or "",
            "print_title": payload.get("print_title") or item.get("print_title") or "",
            "status": payload.get("status") or item.get("status") or "draft",
            "requested_at": payload.get("requested_at") or _iso(item.get("requested_at")),
            "signed_at": payload.get("signed_at") or _iso(item.get("signed_at")),
            "created_at": payload.get("created_at") or _iso(item.get("created_at")),
            "updated_at": payload.get("updated_at") or _iso(item.get("updated_at")),
            "tenant_id": str(item.get("tenant_id") or ""),
        }
    if name == "payroll_statements":
        payload = _payload_dict(item.get("statement_payload"))
        email = str(payload.get("employee_email") or item.get("employee_email") or "").strip().lower()
        amounts: dict[str, int | None] = {}
        invalid_amount_fields: list[str] = []
        for field in ("gross_pay", "tax_withholding", "insurance_deduction", "other_deduction", "net_pay"):
            try:
                amounts[field] = _payroll_integer(item.get(field), field=field)
            except ValueError:
                amounts[field] = None
                invalid_amount_fields.append(field)
        return {
            **payload,
            "id": str(payload.get("id") or item.get("id") or ""),
            "employee_email": email,
            "employee_email_masked": payload.get("employee_email_masked") or item.get("employee_email_masked") or _mask_email(email),
            "employee_name": payload.get("employee_name") or item.get("employee_name") or "",
            "tenant_id": str(item.get("tenant_id") or ""),
            "business_id": item.get("business_id") or "",
            "branch": payload.get("branch") or item.get("branch") or "",
            "payroll_month": item.get("payroll_month") or "",
            **amounts,
            "payroll_validation_errors": invalid_amount_fields,
            "status": item.get("status") or "draft",
            "created_at": payload.get("created_at") or _iso(item.get("created_at")),
            "updated_at": payload.get("updated_at") or _iso(item.get("updated_at")),
        }
    return item


async def _db_fetch_ledger(name: str, tenant_id: str | None = None) -> list[dict[str, Any]] | None:
    import asyncpg

    table = DB_LEDGER_TABLE_BY_NAME.get(name)
    if not table:
        return None
    conn = await asyncpg.connect(_db_url(), timeout=5)
    try:
        ready = await conn.fetchval("SELECT to_regclass($1) IS NOT NULL", f"public.{table}")
        if not ready:
            return None
        key = "row_id" if name in GENERIC_DB_LEDGER_NAMES else "id"
        if name in HR_TENANT_LEDGER_NAMES:
            if not tenant_id:
                return []
            rows = await conn.fetch(
                f"SELECT * FROM {table} h WHERE h.deleted_at IS NULL AND h.tenant_id = $1::uuid "
                "AND EXISTS (SELECT 1 FROM yeoljeong_business_tenant_mapping m "
                "WHERE m.business_id = h.business_id AND m.tenant_id = h.tenant_id) "
                f"ORDER BY updated_at DESC NULLS LAST, created_at DESC NULLS LAST, {key} DESC",
                UUID(tenant_id),
            )
        else:
            rows = await conn.fetch(f"SELECT * FROM {table} WHERE deleted_at IS NULL ORDER BY updated_at DESC NULLS LAST, created_at DESC NULLS LAST, {key} DESC")
        return [_db_row_to_record(name, row) for row in rows]
    finally:
        await conn.close()


def _db_payload_record(name: str, record: dict[str, Any]) -> dict[str, Any]:
    if name == "platform_accounts":
        return {key: value for key, value in record.items() if key not in _ACCOUNT_SECRET_FIELDS}
    return record


def _attach_local_account_secrets(
    db_rows: list[dict[str, Any]], file_rows: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    """Restore account secrets from the protected file without copying them to DB."""
    file_by_id = {str(row.get("id") or ""): row for row in file_rows if row.get("id")}
    merged_rows: list[dict[str, Any]] = []
    for db_row in db_rows:
        merged = {**db_row}
        local = file_by_id.get(str(db_row.get("id") or ""), {})
        for field in _ACCOUNT_SECRET_FIELDS:
            if local.get(field):
                merged[field] = local[field]
        merged_rows.append(merged)
    return merged_rows


async def _db_upsert_ledger(name: str, record: dict[str, Any]) -> bool:
    import asyncpg

    table = DB_LEDGER_TABLE_BY_NAME.get(name)
    if not table:
        return False
    if name in DELIVERY_DATA_LEDGER_NAMES and not _delivery_record_is_meaningful(name, record):
        return False
    conn = await asyncpg.connect(_db_url(), timeout=5)
    try:
        ready = await conn.fetchval("SELECT to_regclass($1) IS NOT NULL", f"public.{table}")
        if not ready:
            return False
        payload = json.dumps(_db_payload_record(name, record), ensure_ascii=False)
        record_id = str(record.get("id") or uuid4())
        now = _now()
        if name in GENERIC_DB_LEDGER_NAMES:
            await conn.execute(
                f"""
                INSERT INTO {table} (row_id, business_id, branch, created_at, updated_at, payload)
                VALUES ($1, $2, $3, $4::timestamptz, $5::timestamptz, $6::jsonb)
                ON CONFLICT (row_id) DO UPDATE SET
                    business_id = EXCLUDED.business_id,
                    branch = EXCLUDED.branch,
                    updated_at = EXCLUDED.updated_at,
                    deleted_at = NULL,
                    payload = EXCLUDED.payload
                """,
                record_id,
                str(record.get("business_id") or ""),
                str(record.get("branch") or ""),
                _pg_ts(record.get("created_at") or now),
                _pg_ts(record.get("updated_at") or now),
                payload,
            )
            return True
        if name == "employee_join_requests":
            result = await conn.execute(
                """
                INSERT INTO yeoljeong_employee_join_requests
                    (id, employee_email, employee_email_masked, employee_name, phone, business_id, branch, role, status,
                     tenant_id, request_payload, review_memo, requested_by, reviewed_by, requested_at, reviewed_at, updated_at, deleted_at)
                VALUES ($1, $2, $3, $4, $5, $6, $7, $8, $9, $10::uuid, $11::jsonb, $12, $13, $14, $15::timestamptz, $16::timestamptz, $17::timestamptz, NULL)
                ON CONFLICT (id) DO UPDATE SET
                    employee_email = EXCLUDED.employee_email,
                    employee_email_masked = EXCLUDED.employee_email_masked,
                    employee_name = EXCLUDED.employee_name,
                    phone = EXCLUDED.phone,
                    business_id = EXCLUDED.business_id,
                    branch = EXCLUDED.branch,
                    role = EXCLUDED.role,
                    status = EXCLUDED.status,
                    request_payload = EXCLUDED.request_payload,
                    review_memo = EXCLUDED.review_memo,
                    requested_by = EXCLUDED.requested_by,
                    reviewed_by = EXCLUDED.reviewed_by,
                    requested_at = EXCLUDED.requested_at,
                    reviewed_at = EXCLUDED.reviewed_at,
                    updated_at = EXCLUDED.updated_at,
                    deleted_at = NULL
                WHERE yeoljeong_employee_join_requests.tenant_id = EXCLUDED.tenant_id
                """,
                record_id,
                str(record.get("email") or record.get("employee_email") or "").strip().lower(),
                str(record.get("email_masked") or record.get("employee_email_masked") or ""),
                str(record.get("name") or record.get("employee_name") or ""),
                str(record.get("phone") or ""),
                str(record.get("business_id") or ""),
                str(record.get("branch") or ""),
                str(record.get("role") or "employee"),
                str(record.get("status") or "pending"),
                _record_tenant_uuid(record),
                payload,
                str(record.get("review_memo") or ""),
                str(record.get("requested_by") or ""),
                str(record.get("reviewed_by") or ""),
                _pg_ts(record.get("requested_at") or record.get("created_at") or now),
                _pg_ts(record.get("reviewed_at")),
                _pg_ts(record.get("updated_at") or now),
            )
            return result != "INSERT 0 0"
        if name == "onboarding_documents":
            result = await conn.execute(
                """
                INSERT INTO yeoljeong_onboarding_documents
                    (id, employee_request_id, employee_email, employee_email_masked, employee_name, business_id, branch,
                     document_type, document_label, requirement, status, original_filename, stored_filename, content_type,
                     size_bytes, issue_date, memo, review_memo, uploaded_by, reviewed_by, uploaded_at, reviewed_at, updated_at, tenant_id, metadata,
                     expires_at, sha256, stored_path, superseded_by, deleted_at)
                VALUES ($1, $2, $3, $4, $5, $6, $7, $8, $9, $10, $11, $12, $13, $14, $15, $16, $17, $18, $19, $20,
                        $21::timestamptz, $22::timestamptz, $23::timestamptz, $24::uuid, $25::jsonb,
                        $26::date, $27, $28, $29, NULL)
                ON CONFLICT (id) DO UPDATE SET
                    employee_request_id = EXCLUDED.employee_request_id,
                    employee_email = EXCLUDED.employee_email,
                    employee_email_masked = EXCLUDED.employee_email_masked,
                    employee_name = EXCLUDED.employee_name,
                    business_id = EXCLUDED.business_id,
                    branch = EXCLUDED.branch,
                    document_type = EXCLUDED.document_type,
                    document_label = EXCLUDED.document_label,
                    requirement = EXCLUDED.requirement,
                    status = EXCLUDED.status,
                    original_filename = EXCLUDED.original_filename,
                    stored_filename = EXCLUDED.stored_filename,
                    content_type = EXCLUDED.content_type,
                    size_bytes = EXCLUDED.size_bytes,
                    issue_date = EXCLUDED.issue_date,
                    memo = EXCLUDED.memo,
                    review_memo = EXCLUDED.review_memo,
                    uploaded_by = EXCLUDED.uploaded_by,
                    reviewed_by = EXCLUDED.reviewed_by,
                    uploaded_at = EXCLUDED.uploaded_at,
                    reviewed_at = EXCLUDED.reviewed_at,
                    updated_at = EXCLUDED.updated_at,
                    metadata = EXCLUDED.metadata,
                    expires_at = EXCLUDED.expires_at,
                    sha256 = EXCLUDED.sha256,
                    stored_path = EXCLUDED.stored_path,
                    superseded_by = EXCLUDED.superseded_by,
                    deleted_at = NULL
                WHERE yeoljeong_onboarding_documents.tenant_id = EXCLUDED.tenant_id
                """,
                record_id,
                str(record.get("employee_request_id") or ""),
                str(record.get("employee_email") or "").strip().lower(),
                str(record.get("employee_email_masked") or ""),
                str(record.get("employee_name") or ""),
                str(record.get("business_id") or ""),
                str(record.get("branch") or ""),
                str(record.get("document_type") or ""),
                str(record.get("document_label") or ""),
                str(record.get("requirement") or ""),
                str(record.get("status") or "uploaded"),
                str(record.get("original_filename") or ""),
                str(record.get("stored_filename") or ""),
                str(record.get("content_type") or ""),
                int(record.get("size_bytes") or 0),
                str(record.get("issue_date") or ""),
                str(record.get("memo") or ""),
                str(record.get("review_memo") or ""),
                str(record.get("uploaded_by") or ""),
                str(record.get("reviewed_by") or ""),
                _pg_ts(record.get("uploaded_at") or record.get("created_at") or now),
                _pg_ts(record.get("reviewed_at")),
                _pg_ts(record.get("updated_at") or now),
                _record_tenant_uuid(record),
                payload,
                _pg_date(record.get("expires_at")),
                str(record.get("sha256") or ""),
                str(record.get("stored_path") or ""),
                str(record.get("superseded_by") or ""),
            )
            return result != "INSERT 0 0"
        if name == "contracts":
            token = str(record.get("sign_token") or "")
            token_hash = str(record.get("sign_token_hash") or "")
            result = await conn.execute(
                """
                INSERT INTO yeoljeong_contracts
                    (id, employee_email, employee_email_masked, employee_name, business_id, branch, contract_type,
                     document_kind, template_version, print_title, status, sign_token_hash, requested_at, signed_at,
                     created_by, requested_by, signer_email, signer_name, tenant_id, contract_payload, updated_at, deleted_at)
                VALUES ($1, $2, $3, $4, $5, $6, $7, $8, $9, $10, $11, $12, $13::timestamptz, $14::timestamptz,
                        $15, $16, $17, $18, $19::uuid, $20::jsonb, $21::timestamptz, NULL)
                ON CONFLICT (id) DO UPDATE SET
                    employee_email = EXCLUDED.employee_email,
                    employee_email_masked = EXCLUDED.employee_email_masked,
                    employee_name = EXCLUDED.employee_name,
                    business_id = EXCLUDED.business_id,
                    branch = EXCLUDED.branch,
                    contract_type = EXCLUDED.contract_type,
                    document_kind = EXCLUDED.document_kind,
                    template_version = EXCLUDED.template_version,
                    print_title = EXCLUDED.print_title,
                    status = EXCLUDED.status,
                    sign_token_hash = EXCLUDED.sign_token_hash,
                    requested_at = EXCLUDED.requested_at,
                    signed_at = EXCLUDED.signed_at,
                    created_by = EXCLUDED.created_by,
                    requested_by = EXCLUDED.requested_by,
                    signer_email = EXCLUDED.signer_email,
                    signer_name = EXCLUDED.signer_name,
                    contract_payload = EXCLUDED.contract_payload,
                    updated_at = EXCLUDED.updated_at,
                    deleted_at = NULL
                WHERE yeoljeong_contracts.tenant_id = EXCLUDED.tenant_id
                """,
                record_id,
                str(record.get("employee_email") or "").strip().lower(),
                str(record.get("employee_email_masked") or ""),
                str(record.get("employee_name") or ""),
                str(record.get("business_id") or ""),
                str(record.get("branch") or ""),
                str(record.get("contract_type") or "part_time"),
                str(record.get("document_kind") or "standard_employment_contract"),
                str(record.get("template_version") or ""),
                str(record.get("print_title") or ""),
                str(record.get("status") or "draft"),
                hashlib.sha256(token.encode("utf-8")).hexdigest() if token else token_hash,
                _pg_ts(record.get("requested_at")),
                _pg_ts(record.get("signed_at")),
                str(record.get("created_by") or ""),
                str(record.get("requested_by") or ""),
                str(record.get("signer_email") or ""),
                str(record.get("signer_name") or ""),
                _record_tenant_uuid(record),
                payload,
                _pg_ts(record.get("updated_at") or now),
            )
            return result != "INSERT 0 0"
        if name == "payroll_statements":
            result = await conn.execute(
                """
                INSERT INTO yeoljeong_payroll_statements
                    (id, employee_email, employee_email_masked, employee_name, business_id, branch, payroll_month,
                     gross_pay, tax_withholding, insurance_deduction, other_deduction, net_pay, status, created_by,
                     confirmed_by, confirmed_at, tenant_id, statement_payload, updated_at, deleted_at)
                VALUES ($1, $2, $3, $4, $5, $6, $7, $8, $9, $10, $11, $12, $13, $14, $15, $16::timestamptz,
                        $17::uuid, $18::jsonb, $19::timestamptz, NULL)
                ON CONFLICT (id) DO UPDATE SET
                    employee_email = EXCLUDED.employee_email,
                    employee_email_masked = EXCLUDED.employee_email_masked,
                    employee_name = EXCLUDED.employee_name,
                    business_id = EXCLUDED.business_id,
                    branch = EXCLUDED.branch,
                    payroll_month = EXCLUDED.payroll_month,
                    gross_pay = EXCLUDED.gross_pay,
                    tax_withholding = EXCLUDED.tax_withholding,
                    insurance_deduction = EXCLUDED.insurance_deduction,
                    other_deduction = EXCLUDED.other_deduction,
                    net_pay = EXCLUDED.net_pay,
                    status = EXCLUDED.status,
                    created_by = EXCLUDED.created_by,
                    confirmed_by = EXCLUDED.confirmed_by,
                    confirmed_at = EXCLUDED.confirmed_at,
                    statement_payload = EXCLUDED.statement_payload,
                    updated_at = EXCLUDED.updated_at,
                    deleted_at = NULL
                WHERE yeoljeong_payroll_statements.tenant_id = EXCLUDED.tenant_id
                """,
                record_id,
                str(record.get("employee_email") or "").strip().lower(),
                str(record.get("employee_email_masked") or ""),
                str(record.get("employee_name") or ""),
                str(record.get("business_id") or ""),
                str(record.get("branch") or ""),
                str(record.get("payroll_month") or ""),
                int(record.get("gross_pay") or 0),
                int(record.get("tax_withholding") or 0),
                int(record.get("insurance_deduction") or 0),
                int(record.get("other_deduction") or 0),
                int(record.get("net_pay") or 0),
                str(record.get("status") or "draft"),
                str(record.get("created_by") or ""),
                str(record.get("confirmed_by") or ""),
                _pg_ts(record.get("confirmed_at")),
                _record_tenant_uuid(record),
                payload,
                _pg_ts(record.get("updated_at") or now),
            )
            return result != "INSERT 0 0"
        return False
    finally:
        await conn.close()


async def _db_delete_ledger(name: str, row_id: str) -> bool:
    import asyncpg

    table = DB_LEDGER_TABLE_BY_NAME.get(name)
    if not table:
        return False
    conn = await asyncpg.connect(_db_url(), timeout=5)
    try:
        ready = await conn.fetchval("SELECT to_regclass($1) IS NOT NULL", f"public.{table}")
        if not ready:
            return False
        key = "row_id" if name in GENERIC_DB_LEDGER_NAMES else "id"
        await conn.execute(f"UPDATE {table} SET deleted_at = NOW() WHERE {key} = $1", str(row_id))
        return True
    finally:
        await conn.close()


async def _db_delete_hr_ledger(name: str, row_id: str, tenant_id: str) -> bool:
    import asyncpg

    table = DB_LEDGER_TABLE_BY_NAME.get(name)
    if name not in HR_TENANT_LEDGER_NAMES or not table:
        return False
    conn = await asyncpg.connect(_db_url(), timeout=5)
    try:
        result = await conn.execute(
            f"UPDATE {table} SET deleted_at = NOW() WHERE id = $1 AND tenant_id = $2::uuid AND deleted_at IS NULL",
            str(row_id), UUID(tenant_id),
        )
        return result.endswith(" 1")
    finally:
        await conn.close()


def _read(name: str) -> list[dict[str, Any]]:
    file_rows = _read_file_rows(name)
    if name not in DB_LEDGER_TABLE_BY_NAME:
        return file_rows
    db_rows = _run_db(_db_fetch_ledger(name))
    if isinstance(db_rows, list):
        if name == "platform_accounts":
            db_rows = _attach_local_account_secrets(db_rows, file_rows)
        if db_rows:
            db_ids = {str(row.get("id") or "") for row in db_rows}
            missing_rows = [row for row in file_rows if str(row.get("id") or "") and str(row.get("id") or "") not in db_ids]
            if missing_rows:
                for row in missing_rows:
                    _run_db(_db_upsert_ledger(name, row))
                merged = _run_db(_db_fetch_ledger(name))
                if isinstance(merged, list) and merged:
                    if name == "platform_accounts":
                        merged = _attach_local_account_secrets(merged, file_rows)
                    return merged
            return db_rows
        if file_rows:
            for row in file_rows:
                _run_db(_db_upsert_ledger(name, row))
            seeded = _run_db(_db_fetch_ledger(name))
            if isinstance(seeded, list) and seeded:
                if name == "platform_accounts":
                    seeded = _attach_local_account_secrets(seeded, file_rows)
                return seeded
    return file_rows


def _read_hr(name: str, user: dict[str, Any] | None) -> list[dict[str, Any]]:
    """Read HR rows with the tenant predicate in SQL; legacy NULL rows stay invisible."""
    if name not in HR_TENANT_LEDGER_NAMES:
        raise ValueError(f"not an HR tenant ledger: {name}")
    tenant_id = _tenant_id(user)
    if _db_available():
        rows = _run_db(_db_fetch_ledger(name, tenant_id))
        return rows if isinstance(rows, list) else []
    return [
        row
        for row in _read_file_rows(name)
        if str(row.get("tenant_id") or "").strip() == tenant_id and not row.get("deleted_at")
    ]


async def _db_business_tenant_matches(business_id: str, tenant_id: str) -> bool:
    import asyncpg

    conn = await asyncpg.connect(_db_url(), timeout=5)
    try:
        return bool(await conn.fetchval(
            "SELECT EXISTS (SELECT 1 FROM yeoljeong_business_tenant_mapping WHERE business_id = $1 AND tenant_id = $2::uuid)",
            business_id, UUID(tenant_id),
        ))
    finally:
        await conn.close()


def _require_business_for_tenant(business_id: Any, user: dict[str, Any] | None) -> str:
    tenant_id = _tenant_id(user)
    normalized = str(business_id or "").strip()
    if not normalized:
        raise HTTPException(status_code=400, detail="테넌트에 귀속된 사업자가 필요합니다")
    if _db_available() and not _run_db(_db_business_tenant_matches(normalized, tenant_id)):
        raise HTTPException(status_code=403, detail="테넌트에 귀속되지 않은 사업자입니다")
    return normalized


async def _db_business_tenant_id(business_id: str) -> str | None:
    import asyncpg

    conn = await asyncpg.connect(_db_url(), timeout=5)
    try:
        value = await conn.fetchval(
            "SELECT tenant_id::text FROM yeoljeong_business_tenant_mapping WHERE business_id = $1",
            business_id,
        )
        return str(value) if value else None
    finally:
        await conn.close()


async def _db_business_invite_info(business_id: str) -> dict[str, Any] | None:
    import asyncpg

    conn = await asyncpg.connect(_db_url(), timeout=5)
    try:
        name = await conn.fetchval(
            "SELECT name FROM yeoljeong_businesses WHERE id = $1 AND deleted_at IS NULL",
            business_id,
        )
        if name is None:
            return None
        rows = await conn.fetch(
            "SELECT name FROM yeoljeong_branches WHERE business_id = $1 AND deleted_at IS NULL ORDER BY sort_order, id",
            business_id,
        )
        return {"name": str(name), "branches": [str(row["name"]) for row in rows]}
    finally:
        await conn.close()


async def _db_business_ids_by_branch_name(branch: str) -> list[str]:
    import asyncpg

    conn = await asyncpg.connect(_db_url(), timeout=5)
    try:
        rows = await conn.fetch(
            """SELECT DISTINCT b.business_id
                 FROM yeoljeong_branches b
                 JOIN yeoljeong_businesses biz ON biz.id = b.business_id AND biz.deleted_at IS NULL
                WHERE b.name = $1 AND b.deleted_at IS NULL
                ORDER BY b.business_id""",
            branch,
        )
        return [str(row["business_id"]) for row in rows]
    finally:
        await conn.close()


async def _db_hr_record_tenant(name: str, row_id: str) -> str | None:
    import asyncpg

    table = DB_LEDGER_TABLE_BY_NAME.get(name)
    if name not in HR_TENANT_LEDGER_NAMES or not table:
        return None
    conn = await asyncpg.connect(_db_url(), timeout=5)
    try:
        value = await conn.fetchval(
            f"SELECT tenant_id::text FROM {table} WHERE id = $1 AND deleted_at IS NULL",
            str(row_id),
        )
        return str(value) if value else None
    finally:
        await conn.close()


async def _db_fetch_join_requests_by_email(email: str) -> list[dict[str, Any]] | None:
    import asyncpg

    table = DB_LEDGER_TABLE_BY_NAME["employee_join_requests"]
    conn = await asyncpg.connect(_db_url(), timeout=5)
    try:
        ready = await conn.fetchval("SELECT to_regclass($1) IS NOT NULL", f"public.{table}")
        if not ready:
            return None
        rows = await conn.fetch(
            f"SELECT * FROM {table} h WHERE h.deleted_at IS NULL AND h.tenant_id IS NOT NULL "
            "AND lower(h.employee_email) = $1 "
            "AND EXISTS (SELECT 1 FROM yeoljeong_business_tenant_mapping m "
            "WHERE m.business_id = h.business_id AND m.tenant_id = h.tenant_id) "
            "ORDER BY updated_at DESC NULLS LAST, created_at DESC NULLS LAST, id DESC",
            email,
        )
        return [_db_row_to_record("employee_join_requests", row) for row in rows]
    finally:
        await conn.close()


def _employer_scope_user(user: dict[str, Any], employer_tenant_id: str) -> dict[str, Any]:
    """가입요청 한 건을 고용주 테넌트에 읽고 쓰기 위한 범위 한정용 사용자 사본.

    호출자는 그 테넌트의 멤버가 아니다 — 이 사본은 upsert_join_request 안에서만 쓰고,
    권한 판정(_is_admin 등)에는 절대 넘기지 않는다.
    """
    scoped = dict(user)
    scoped["tenant_id"] = employer_tenant_id
    scoped["current_membership"] = {"tenant_id": employer_tenant_id, "status": "active", "role": "applicant"}
    return scoped


def _join_request_scope(business_id: str, email: str, user: dict[str, Any]) -> dict[str, Any]:
    """대상 사업자의 고용주 테넌트로 가입요청을 귀속시킬 범위를 정한다.

    호출자 JWT 테넌트가 아니라 yeoljeong_business_tenant_mapping 의 테넌트가 기준이다.
    다른 테넌트의 사업자에는 본인 이메일 요청만 낼 수 있다(남을 대신해 남의 테넌트에 쓰지 못한다).
    DB 가 없는 파일 모드에는 매핑이 없어 종전처럼 호출자 테넌트를 쓴다(_require_business_for_tenant 와 같은 규칙).
    """
    caller_tenant = _tenant_id(user)
    if not _db_available():
        return user
    employer = _run_db(_db_business_tenant_id(business_id))
    if not employer:
        raise HTTPException(status_code=400, detail="등록되지 않은 사업자입니다")
    employer = str(employer).strip()
    if employer == caller_tenant:
        return user
    if not email or email != _email(user):
        raise HTTPException(status_code=403, detail="다른 사업자에는 본인 가입요청만 등록할 수 있습니다")
    return _employer_scope_user(user, employer)


def _join_request_owner_tenant(request_id: str) -> str:
    if _db_available():
        return str(_run_db(_db_hr_record_tenant("employee_join_requests", request_id)) or "").strip()
    record = _find(_read_file_rows("employee_join_requests"), request_id)
    return str((record or {}).get("tenant_id") or "").strip()


def _require_review_tenant(request_id: str, user: dict[str, Any]) -> str:
    """가입요청의 귀속 테넌트와 호출자 테넌트가 같을 때만 승인·반려를 허용한다(다르면 403)."""
    tenant_id = _tenant_id(user)
    owner = _join_request_owner_tenant(request_id)
    if owner and owner != tenant_id:
        raise HTTPException(status_code=403, detail="다른 테넌트의 가입요청은 처리할 수 없습니다")
    return tenant_id


def _read_join_requests_by_email(email: str) -> list[dict[str, Any]]:
    """직원 본인 이메일의 가입요청을 고용주 테넌트와 무관하게 읽는다(본인 레코드만)."""
    if _db_available():
        rows = _run_db(_db_fetch_join_requests_by_email(email))
        return rows if isinstance(rows, list) else []
    return [
        row
        for row in _read_file_rows("employee_join_requests")
        if str(row.get("tenant_id") or "").strip() and str(row.get("email") or "").strip().lower() == email
    ]


def _require_hr_record(record: dict[str, Any] | None, user: dict[str, Any] | None, *, detail: str) -> dict[str, Any]:
    tenant_id = _tenant_id(user)
    if not record or str(record.get("tenant_id") or "").strip() != tenant_id:
        # Do not reveal whether the identifier exists in another tenant.
        raise HTTPException(status_code=404, detail=detail)
    return record


def _owned_hr_record(record: dict[str, Any], user: dict[str, Any]) -> dict[str, Any]:
    owned = dict(record)
    owned["tenant_id"] = _tenant_id(user)
    owned["business_id"] = _require_business_for_tenant(owned.get("business_id"), user)
    return owned


def _write(name: str, rows: list[dict[str, Any]]) -> None:
    _write_file_rows(name, rows)
    if name not in DB_LEDGER_TABLE_BY_NAME:
        return
    for row in rows:
        _run_db(_db_upsert_ledger(name, row))


def _write_hr_record(name: str, record: dict[str, Any], user: dict[str, Any] | None) -> None:
    tenant_id = _tenant_id(user)
    if str(record.get("tenant_id") or "").strip() != tenant_id:
        raise HTTPException(status_code=403, detail="다른 테넌트의 데이터는 저장할 수 없습니다")
    if _db_available():
        if not _run_db(_db_upsert_ledger(name, record)):
            raise HTTPException(status_code=404, detail="대상을 찾을 수 없습니다")
        return
    rows = _read_file_rows(name)
    existing = _find(rows, str(record.get("id") or ""))
    if existing and str(existing.get("tenant_id") or "").strip() != tenant_id:
        raise HTTPException(status_code=404, detail="대상을 찾을 수 없습니다")
    if existing:
        existing.clear()
        existing.update(record)
    else:
        rows.insert(0, record)
    _write_file_rows(name, rows)


def _delete_hr_record(name: str, row_id: str, user: dict[str, Any] | None) -> None:
    tenant_id = _tenant_id(user)
    if _db_available():
        if not _run_db(_db_delete_hr_ledger(name, row_id, tenant_id)):
            raise HTTPException(status_code=404, detail="대상을 찾을 수 없습니다")
        return
    rows = _read_file_rows(name)
    matched = next(
        (
            row
            for row in rows
            if str(row.get("id")) == str(row_id)
            and str(row.get("tenant_id") or "") == tenant_id
            and not row.get("deleted_at")
        ),
        None,
    )
    if not matched:
        raise HTTPException(status_code=404, detail="대상을 찾을 수 없습니다")
    if name == "onboarding_documents":
        # DB 모드(deleted_at)와 같게 행을 남긴다. 입사서류는 노무 분쟁의 근거다.
        now = _now()
        matched["deleted_at"] = now
        matched["updated_at"] = now
        _write_file_rows(name, rows)
        return
    _write_file_rows(name, [row for row in rows if row is not matched])


def _delete(name: str, row_id: str) -> None:
    _write_file_rows(name, [row for row in _read_file_rows(name) if str(row.get("id") or "") != str(row_id)])
    if name in DB_LEDGER_TABLE_BY_NAME:
        _run_db(_db_delete_ledger(name, row_id))


def _merge_by_id(current_items: Any, default_items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    current = {str(item.get("id")): item for item in current_items if isinstance(item, dict)} if isinstance(current_items, list) else {}
    merged: list[dict[str, Any]] = []
    seen: set[str] = set()
    for default_item in default_items:
        item = {**default_item, **current.get(default_item["id"], {})}
        item["id"] = default_item["id"]
        merged.append(item)
        seen.add(default_item["id"])
    if isinstance(current_items, list):
        for current_item in current_items:
            if not isinstance(current_item, dict):
                continue
            current_id = str(current_item.get("id") or "").strip()
            if not current_id or current_id in seen:
                continue
            merged.append({**current_item, "id": current_id})
            seen.add(current_id)
    return merged


def _canonicalize_ui_settings(settings: dict[str, Any]) -> dict[str, Any]:
    from app.services.obys_upload_service import registration_info_gaps

    businesses = _merge_by_id(settings.get("businesses"), CANONICAL_BUSINESSES)
    canonical_names = {item["id"]: item["name"] for item in CANONICAL_BUSINESSES}
    for item in businesses:
        item["entityType"] = item.get("entityType") or "individual"
        # 저장된 상호가 이긴다. canonical 상호는 저장값이 비었을 때만 쓴다(id 는 그대로).
        item["name"] = str(item.get("name") or "").strip() or canonical_names.get(item["id"], "")
        item["status"] = item.get("status") or "active"
        missing = registration_info_gaps(item, BUSINESS_REGISTRATION_FIELDS)
        item["needs_registration_info"] = bool(missing)
        item["missing_registration_fields"] = missing

    branches = _merge_by_id(settings.get("branches"), CANONICAL_BRANCHES)
    canonical_branch_names = {item["id"]: item["name"] for item in CANONICAL_BRANCHES}
    canonical_branch_businesses = {item["id"]: item["businessId"] for item in CANONICAL_BRANCHES}
    all_business_ids = {str(item.get("id") or "") for item in businesses}
    normalized_branches: list[dict[str, Any]] = []
    for item in branches:
        if item["id"] in canonical_branch_names:
            item["name"] = str(item.get("name") or "").strip() or canonical_branch_names[item["id"]]
            item["businessId"] = canonical_branch_businesses[item["id"]]
        else:
            item["name"] = str(item.get("name") or "").strip()
            business_id = str(item.get("businessId") or item.get("business_id") or "").strip()
            if not item["name"] or business_id not in all_business_ids:
                continue
            item["businessId"] = business_id
        item["status"] = item.get("status") or "active"
        normalized_branches.append(item)
    branches = normalized_branches
    all_branch_names = {str(item.get("name") or "") for item in branches}

    def normalize_business_ref(item: dict[str, Any]) -> dict[str, Any]:
        next_item = {**item}
        business_id = str(next_item.get("businessId") or next_item.get("business_id") or "").strip()
        if business_id not in all_business_ids:
            business_id = MIA_BUSINESS_ID
        next_item["businessId"] = business_id
        if "business_id" in next_item:
            next_item["business_id"] = business_id
        branch = str(next_item.get("branch") or "").strip()
        if branch and branch not in all_branch_names:
            next_item["branch"] = MIA_BRANCH_NAME
        service = str(next_item.get("service") or "").strip()
        if service in CONNECTOR_LABELS:
            next_item["service"] = service
            next_item["label"] = str(next_item.get("label") or CONNECTOR_LABELS[service]).strip()
            next_item["category"] = str(next_item.get("category") or "").strip()
            next_item["collectionMode"] = str(
                next_item.get("collectionMode") or next_item.get("collection_mode") or "browser-automation"
            ).strip()
            next_item["dataScope"] = str(next_item.get("dataScope") or next_item.get("data_scope") or "").strip()
            next_item["requiredProof"] = str(
                next_item.get("requiredProof") or next_item.get("required_proof") or ""
            ).strip()
            next_item["loginUrl"] = str(next_item.get("loginUrl") or next_item.get("login_url") or "").strip()
            next_item["username"] = str(next_item.get("username") or "").strip()
            next_item["businessRegistrationNoMasked"] = str(
                next_item.get("businessRegistrationNoMasked")
                or next_item.get("business_registration_no_masked")
                or ""
            ).strip()
            next_item["status"] = str(next_item.get("status") or "ready").strip()
        return next_item

    return {
        "businesses": businesses,
        "branches": branches,
        "accounts": [normalize_business_ref(item) for item in settings.get("accounts", []) if isinstance(item, dict)],
        "staff": [item for item in settings.get("staff", []) if isinstance(item, dict)],
        "integrations": [normalize_business_ref(item) for item in settings.get("integrations", []) if isinstance(item, dict)],
    }


def _jsonb_object(value: Any) -> dict[str, Any]:
    if isinstance(value, dict):
        return value
    if isinstance(value, str):
        try:
            parsed = json.loads(value)
        except json.JSONDecodeError:
            return {}
        return parsed if isinstance(parsed, dict) else {}
    return {}


def _encrypt_secret(value: str) -> str:
    try:
        from app.core.credential_vault import encrypt_value

        return encrypt_value(value)
    except Exception as exc:
        raise HTTPException(status_code=500, detail="계정 비밀번호 암호화에 실패했습니다") from exc


def _decrypt_secret(value: str) -> str:
    if not value:
        return ""
    try:
        from app.core.credential_vault import decrypt_value

        return decrypt_value(value)
    except Exception:
        return ""


def _normalize_origin(value: Any) -> str:
    text = str(value or "").strip().rstrip("/")
    if not text:
        return ""
    parsed = urlparse(text)
    if parsed.scheme and parsed.netloc:
        return f"{parsed.scheme.lower()}://{parsed.netloc.lower()}"
    if parsed.netloc:
        return f"https://{parsed.netloc.lower()}"
    return text.lower()


def _delivery_service_for_vault_origin(origin: Any) -> str:
    normalized = _normalize_origin(origin)
    for service, origins in DELIVERY_AGENT_VAULT_ORIGINS.items():
        if normalized in {_normalize_origin(item) for item in origins}:
            return service
    return ""


async def _db_fetch_delivery_agent_vault_credentials() -> list[dict[str, Any]] | None:
    import asyncpg

    origins = sorted(
        {_normalize_origin(origin) for service_origins in DELIVERY_AGENT_VAULT_ORIGINS.values() for origin in service_origins}
    )
    if not origins:
        return []
    conn = await asyncpg.connect(_db_url(), timeout=5)
    try:
        ready = await conn.fetchval("SELECT to_regclass($1) IS NOT NULL", "public.agent_vault_credentials")
        if not ready:
            return []
        rows = await conn.fetch(
            """
            SELECT id, work_key, origin, label, username_enc, password_enc, metadata
              FROM agent_vault_credentials
             WHERE is_active = TRUE
               AND origin = ANY($1::text[])
             ORDER BY work_key, origin, label
            """,
            origins,
        )
        return [dict(row) for row in rows]
    finally:
        await conn.close()


async def _db_fetch_bank_agent_vault_credentials(
    service: str,
    tenant_id: str,
) -> list[dict[str, Any]]:
    """Fetch encrypted bank login rows for one tenant and approved origins."""
    import asyncpg

    origins = sorted({_normalize_origin(value) for value in BANK_AGENT_VAULT_ORIGINS.get(service, ())})
    if not origins or not tenant_id:
        return []
    try:
        tenant_uuid = UUID(str(tenant_id))
    except (TypeError, ValueError, AttributeError):
        return []
    conn = await asyncpg.connect(_db_url(), timeout=5)
    try:
        ready = await conn.fetchval("SELECT to_regclass($1) IS NOT NULL", "public.agent_vault_credentials")
        if not ready:
            return []
        rows = await conn.fetch(
            """
            SELECT id, work_key, origin, label, username_enc, password_enc, metadata
              FROM agent_vault_credentials
             WHERE tenant_id = $1
               AND is_active = TRUE
               AND origin = ANY($2::text[])
             ORDER BY updated_at DESC, label
            """,
            tenant_uuid,
            origins,
        )
        return [dict(row) for row in rows]
    finally:
        await conn.close()


def _bank_login_from_agent_vault(
    *,
    service: str,
    tenant_id: str,
    expected_username: str,
    business_id: str,
    branch_names: set[str],
) -> dict[str, str]:
    """Resolve a Password Manager login without crossing tenant boundaries."""
    if not tenant_id or not _db_available():
        return {}
    rows = _run_db(_db_fetch_bank_agent_vault_credentials(service, tenant_id))
    if not isinstance(rows, list) or not rows:
        return {}
    expected = str(expected_username or "").strip().lower()

    def _rank(row: dict[str, Any]) -> tuple[int, int, int, str]:
        metadata = _jsonb_object(row.get("metadata"))
        username = _decrypt_secret(str(row.get("username_enc") or "")).strip().lower()
        row_business = str(metadata.get("business_id") or metadata.get("businessId") or "").strip()
        row_branch = str(metadata.get("branch") or "").strip()
        return (
            1 if expected and username == expected else 0,
            1 if business_id and row_business == business_id else 0,
            1 if branch_names and row_branch in branch_names else 0,
            str(row.get("id") or ""),
        )

    ranked = sorted(rows, key=_rank, reverse=True)
    best = ranked[0]
    best_rank = _rank(best)
    # A non-singleton Vault set needs at least one positive identity/scope match.
    if len(ranked) > 1 and not any(best_rank[:3]):
        return {}
    username = _decrypt_secret(str(best.get("username_enc") or "")).strip()
    password = _decrypt_secret(str(best.get("password_enc") or ""))
    if not username or not password:
        return {}
    return {
        "login_username": username,
        "login_password": password,
        "agent_vault_credential_id": str(best.get("id") or ""),
        "agent_vault_origin": _normalize_origin(best.get("origin")),
    }


def _hydrate_delivery_account_passwords_from_agent_vault(rows: list[dict[str, Any]]) -> int:
    """Copy matching Agent Vault password ciphertexts into delivery accounts.

    Platform accounts keep usernames in the Yeoljeong settings file, while
    Agent Vault stores imported browser credentials by origin. Matching by
    service origin and username lets automatic collection reuse already
    approved Vault credentials without exposing plaintext passwords.
    """
    targets = [
        row
        for row in rows
        if str(row.get("service") or "") in PLATFORM_LABELS
        and str(row.get("username") or "").strip()
        and not _has_secret_value(row, "password")
    ]
    if not targets or not _db_available():
        return 0

    vault_rows = _run_db(_db_fetch_delivery_agent_vault_credentials())
    if not isinstance(vault_rows, list) or not vault_rows:
        return 0

    by_service_scope: dict[tuple[str, str, str], dict[str, Any]] = {}
    by_service_username: dict[tuple[str, str], dict[str, Any]] = {}
    for vault_row in vault_rows:
        metadata = _jsonb_object(vault_row.get("metadata"))
        service = str(metadata.get("service") or "").strip() or _delivery_service_for_vault_origin(vault_row.get("origin"))
        encrypted_password = str(vault_row.get("password_enc") or "")
        if not service or not encrypted_password:
            continue
        business_id = str(metadata.get("business_id") or metadata.get("businessId") or "").strip()
        branch = BRANCH_ALIASES.get(
            str(metadata.get("branch") or "").strip(),
            str(metadata.get("branch") or "").strip(),
        )
        if business_id and branch:
            by_service_scope.setdefault((service, business_id, branch), vault_row)
        username = _decrypt_secret(str(vault_row.get("username_enc") or "")).strip()
        if not username:
            continue
        key = (service, username.lower())
        if key not in by_service_username:
            by_service_username[key] = vault_row

    if not by_service_scope and not by_service_username:
        return 0

    synced_at = _now()
    changed = 0
    for row in targets:
        service = str(row.get("service") or "").strip()
        username = str(row.get("username") or "").strip().lower()
        business_id = str(row.get("business_id") or "").strip()
        branch = BRANCH_ALIASES.get(str(row.get("branch") or "").strip(), str(row.get("branch") or "").strip())
        vault_row = by_service_scope.get((service, business_id, branch)) or by_service_username.get((service, username))
        if not vault_row:
            continue
        row["password_enc"] = str(vault_row.get("password_enc") or "")
        row["password_source"] = "agent_vault"
        row["agent_vault_credential_id"] = str(vault_row.get("id") or "")
        row["agent_vault_origin"] = _normalize_origin(vault_row.get("origin"))
        row["updated_at"] = synced_at
        if str(row.get("portal_status") or "") in {"action_required", "credential_required"}:
            row["portal_status"] = "credential_registered"
            row["portal_message"] = "Agent Vault 자격증명과 매칭되어 자동수집 비밀번호를 반영했습니다."
        changed += 1
    return changed


_ACCOUNT_SECRET_FIELD_MAP: dict[str, str] = {
    "password": "password_enc",
    "api_key": "api_key_enc",
    "client_secret": "client_secret_enc",
    "certificate_password": "certificate_password_enc",
    "account_no": "account_no_enc",
    "account_password": "account_password_enc",
    "business_registration_no": "business_registration_no_enc",
}


def _migrate_platform_account_secrets(rows: list[dict[str, Any]]) -> bool:
    changed = False
    for row in rows:
        for plaintext_field, encrypted_field in _ACCOUNT_SECRET_FIELD_MAP.items():
            plaintext = str(row.get(plaintext_field) or "")
            if not plaintext:
                continue
            if not row.get(encrypted_field):
                row[encrypted_field] = _encrypt_secret(plaintext)
            row.pop(plaintext_field, None)
            changed = True
    return changed


def _has_account_secret(row: dict[str, Any]) -> bool:
    return any(bool(row.get(field)) for field in _ACCOUNT_SECRET_FIELDS)


def _masked_digits(value: Any, *, visible_tail: int = 4) -> str:
    digits = re.sub(r"\D+", "", str(value or ""))
    if not digits:
        return ""
    if len(digits) <= visible_tail:
        return "*" * len(digits)
    return f"{'*' * max(3, len(digits) - visible_tail)}{digits[-visible_tail:]}"


def _has_secret_value(row: dict[str, Any], plaintext_field: str) -> bool:
    encrypted_field = _ACCOUNT_SECRET_FIELD_MAP.get(plaintext_field, "")
    return bool(encrypted_field and row.get(encrypted_field))


def _public_platform_account_status(row: dict[str, Any]) -> str:
    status = str(row.get("status") or "ready").strip()
    sync_status = str(row.get("last_sync_status") or row.get("portal_status") or "").strip()
    service = str(row.get("service") or "").strip()
    collection_mode = str(row.get("collection_mode") or row.get("collectionMode") or "").strip()
    has_browser_state = bool(
        str(row.get("storage_state_path") or row.get("browser_storage_state_path") or row.get("baemin_storage_state_path") or "").strip()
    )
    if sync_status in {"queued", "succeeded", "partial", "action_required", "failed"}:
        return sync_status
    if sync_status == "running":
        started_text = str(row.get("last_sync_at") or row.get("updated_at") or "").strip()
        try:
            started_at = datetime.fromisoformat(started_text.replace("Z", "+00:00"))
        except ValueError:
            started_at = None
        if started_at and datetime.now(started_at.tzinfo or KST) - started_at < timedelta(seconds=60):
            return "running"
        if service in PLATFORM_LABELS and collection_mode in DELIVERY_UPLOAD_COLLECTION_MODES and not has_browser_state:
            return "upload_required"
        return "credential_required" if service in PLATFORM_LABELS and not has_browser_state else "blocked"
    if service in PLATFORM_LABELS and not _has_secret_value(row, "password") and not has_browser_state:
        return "credential_required"
    if service in PLATFORM_LABELS and collection_mode in DELIVERY_UPLOAD_COLLECTION_MODES:
        return "upload_required"
    if service in BANK_QUICK_SERVICE_CONFIG and collection_mode == "bank-quick-service":
        for key in ("password", "account_no", "account_password", "business_registration_no"):
            if not _has_secret_value(row, key):
                return "credential_required"
    if service in FINANCIAL_TRANSACTION_SERVICES and collection_mode in {"bank-excel", "card-pg-report", "statement-upload"}:
        return "upload_required"
    return status


def _missing_connector_requirements(row: dict[str, Any]) -> list[str]:
    service = str(row.get("service") or "").strip()
    collection_mode = str(row.get("collection_mode") or row.get("collectionMode") or "").strip()
    missing: list[str] = []
    if service in PLATFORM_LABELS:
        if not str(row.get("username") or "").strip():
            missing.append("아이디")
        has_browser_state = bool(
            str(row.get("storage_state_path") or row.get("browser_storage_state_path") or row.get("baemin_storage_state_path") or "").strip()
        )
        if not _has_secret_value(row, "password") and not has_browser_state:
            missing.append("비밀번호 또는 PC Agent 로그인 세션")
        if collection_mode in DELIVERY_UPLOAD_COLLECTION_MODES and not has_browser_state:
            missing.append("브라우저 자동화 방식 선택 또는 포털 CSV/엑셀 업로드")
    if service in BANK_QUICK_SERVICE_CONFIG and collection_mode == "bank-quick-service":
        for label, key in (
            ("로그인 비밀번호", "password"),
            ("조회용 계좌번호", "account_no"),
            ("계좌비밀번호", "account_password"),
            ("사업자번호", "business_registration_no"),
        ):
            if not _has_secret_value(row, key):
                missing.append(label)
    return list(dict.fromkeys(missing))


def _mark_platform_account_sync_state(account: dict[str, Any], *, status: str, message: str, synced_at: str) -> None:
    account["last_sync_status"] = status
    account["portal_status"] = status
    account["portal_message"] = message
    account["last_sync_at"] = synced_at
    account["updated_at"] = synced_at
    account_id = str(account.get("id") or "")
    if not account_id:
        return
    rows = _read("platform_accounts")
    for row in rows:
        if str(row.get("id") or "") == account_id:
            row.update(
                {
                    "last_sync_status": status,
                    "portal_status": status,
                    "portal_message": message,
                    "last_sync_at": synced_at,
                    "updated_at": synced_at,
                }
            )
            _write("platform_accounts", rows)
            break


def _delivery_public_collection_status(status: Any) -> str:
    raw = str(status or "").strip().lower()
    if raw in DELIVERY_COLLECTION_STATUSES:
        return raw
    if raw in {"completed", "success"}:
        return "succeeded"
    if raw in {"no_records", "empty", "authenticated_no_rows"}:
        return "partial"
    if raw in DELIVERY_ACTION_REQUIRED_STATUSES:
        return "action_required"
    if raw in {"stale", "error"}:
        return "failed"
    return "failed" if raw else "failed"


def _delivery_public_error_code(status: str, error_code: Any) -> str:
    raw = str(error_code or "").strip()
    upper = raw.upper()
    if upper in {"", "NONE", "NULL"}:
        if status == "partial":
            return "AUTHENTICATED_NO_ROWS"
        return ""
    if upper == "PC_AGENT_SESSION_REQUIRED":
        return "PC_AGENT_SESSION_REQUIRED"
    if upper in {"ACCOUNT_NOT_REGISTERED", "CREDENTIAL_REQUIRED", "CREDENTIALS_MISSING"}:
        return "MISSING_CREDENTIALS"
    if upper == "BAEMIN_SECURITY_BLOCKED":
        return "BAEMIN_SECURITY_BLOCKED"
    if upper.endswith("_SECURITY_BLOCKED") or upper in {"SECURITY_BLOCKED", "PORTAL_BLOCKED"}:
        return "PORTAL_BLOCKED"
    if upper == "DDANGYO_NUMERIC_CAPTCHA_REQUIRED":
        return "DDANGYO_NUMERIC_CAPTCHA_REQUIRED"
    if upper in {"MFA_REQUIRED", "CAPTCHA_REQUIRED", "PORTAL_AUTH_CHALLENGE"}:
        return "PORTAL_AUTH_CHALLENGE"
    if upper in {"NO_RECORDS", "NO_ROWS", "AUTHENTICATED_NO_ROWS"}:
        return "AUTHENTICATED_NO_ROWS"
    return upper


# Fields that must never appear in API responses or logs.
_ACCOUNT_SECRET_FIELDS: frozenset[str] = frozenset(
    set(_ACCOUNT_SECRET_FIELD_MAP) | set(_ACCOUNT_SECRET_FIELD_MAP.values())
)


def _normalize_delivery_scope(business_id: Any, branch: Any) -> tuple[str, str]:
    normalized_branch = BRANCH_ALIASES.get(str(branch or "").strip(), str(branch or "").strip())
    normalized_business = str(business_id or "").strip()
    expected_business = BUSINESS_BY_BRANCH.get(normalized_branch)
    if normalized_business not in CANONICAL_BUSINESS_IDS:
        raise HTTPException(status_code=400, detail="등록되지 않은 사업자입니다")
    if not expected_business or expected_business != normalized_business:
        raise HTTPException(status_code=400, detail="사업자와 지점 연결이 일치하지 않습니다")
    return normalized_business, normalized_branch


def _normalize_connector_scope(service: str, business_id: Any, branch: Any) -> tuple[str, str]:
    if service in PLATFORM_LABELS:
        return _normalize_delivery_scope(business_id, branch)
    normalized_business = str(business_id or "").strip()
    normalized_branch = BRANCH_ALIASES.get(str(branch or "").strip(), str(branch or "").strip())
    if normalized_business not in CANONICAL_BUSINESS_IDS:
        raise HTTPException(status_code=400, detail="등록되지 않은 사업자입니다")
    if normalized_branch:
        expected_business = BUSINESS_BY_BRANCH.get(normalized_branch)
        if not expected_business or expected_business != normalized_business:
            raise HTTPException(status_code=400, detail="사업자와 지점 연결이 일치하지 않습니다")
    return normalized_business, normalized_branch


def _get_pool_or_none() -> Any | None:
    try:
        from app.core.db_pool import get_pool

        return get_pool()
    except Exception:
        return None


async def get_storage_status(user: dict[str, Any]) -> dict[str, Any]:
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="저장소 상태 확인 권한이 없습니다")
    ledger_files = sorted(set(JSON_LEDGER_FILES) | GENERIC_DB_LEDGER_NAMES)
    json_ledgers = []
    for name in ledger_files:
        path = _path(name)
        file_rows = _read_file_rows(name) if path.exists() else []
        json_ledgers.append(
            {
                "name": name,
                "path": str(path),
                "exists": path.exists(),
                "records": len(file_rows),
            }
        )

    pool = _get_pool_or_none()
    db_tables = {name: False for name in (*SETTINGS_TABLES, *HR_LEDGER_TABLES, *DELIVERY_LEDGER_TABLES)}
    if pool is not None:
        try:
            async with pool.acquire() as conn:
                rows = await conn.fetch(
                    """
                    SELECT table_name
                      FROM information_schema.tables
                     WHERE table_schema = 'public'
                       AND table_name = ANY($1::text[])
                    """,
                    list(db_tables),
                )
            existing = {row["table_name"] for row in rows}
            db_tables = {name: name in existing for name in db_tables}
        except Exception:
            pass
    elif _db_available():
        try:
            import asyncpg

            conn = await asyncpg.connect(_db_url(), timeout=5)
            try:
                rows = await conn.fetch(
                    """
                    SELECT table_name
                      FROM information_schema.tables
                     WHERE table_schema = 'public'
                       AND table_name = ANY($1::text[])
                    """,
                    list(db_tables),
                )
            finally:
                await conn.close()
            existing = {row["table_name"] for row in rows}
            db_tables = {name: name in existing for name in db_tables}
        except Exception:
            pass

    settings_db_ready = all(db_tables[name] for name in SETTINGS_TABLES)
    hr_db_ready = all(db_tables[name] for name in HR_LEDGER_TABLES)
    delivery_db_ready = all(db_tables[name] for name in DELIVERY_LEDGER_TABLES)
    return {
        "checked_at": _now(),
        "mode": "database+json-fallback" if any([settings_db_ready, hr_db_ready, delivery_db_ready]) else "json-only",
        "settings": {
            "source": "database" if settings_db_ready else "json",
            "tables": {name: db_tables[name] for name in SETTINGS_TABLES},
        },
        "hr_ledgers": {
            "source": "database" if hr_db_ready else "json",
            "tables": {name: db_tables[name] for name in HR_LEDGER_TABLES},
            "json_files": json_ledgers,
        },
        "delivery_ledgers": {
            "source": "database" if delivery_db_ready else "json",
            "tables": {name: db_tables[name] for name in DELIVERY_LEDGER_TABLES},
        },
        "migration": {
            "settings": "113_yeoljeong_finance_settings.sql",
            "hr_ledgers": "115_yeoljeong_finance_hr_ledgers.sql",
            "delivery_ledgers": "116_yeoljeong_finance_delivery_ledgers.sql",
            "hr_db_ready": hr_db_ready,
            "delivery_db_ready": delivery_db_ready,
            "note": "각 DB 테이블이 실제 적용되기 전까지 해당 원장은 기존 JSON 파일을 유지합니다.",
        },
    }


def _mask_email(email: str) -> str:
    email = str(email or "").strip().lower()
    if "@" not in email:
        return email
    local, domain = email.split("@", 1)
    if len(local) <= 2:
        masked = local[:1] + "*"
    else:
        masked = local[:2] + "*" * max(1, len(local) - 2)
    return f"{masked}@{domain}"


def _mask_phone(phone: str) -> str:
    digits = re.sub(r"\D+", "", str(phone or ""))
    if len(digits) < 7:
        return phone
    return f"{digits[:3]}-****-{digits[-4:]}"


def _safe_filename(filename: str) -> str:
    name = Path(filename or "upload.bin").name
    name = re.sub(r"[^A-Za-z0-9_.가-힣 -]+", "_", name).strip(" .")
    return name or "upload.bin"


def _find(rows: list[dict[str, Any]], item_id: str) -> dict[str, Any] | None:
    return next((row for row in rows if str(row.get("id")) == str(item_id)), None)


EMPLOYEE_ACCESS_ROLES: dict[str, dict[str, Any]] = {
    "employee": {
        "label": "직원",
        "can_edit_local_data": False,
        "can_manage_settings": False,
        "can_manage_automation": False,
        "can_import_settlements": False,
        "can_manage_onboarding": False,
    },
    "viewer": {
        "label": "조회 전용",
        "can_edit_local_data": False,
        "can_manage_settings": False,
        "can_manage_automation": False,
        "can_import_settlements": False,
        "can_manage_onboarding": False,
    },
    "member": {
        "label": "운영 입력자",
        "can_edit_local_data": True,
        "can_manage_settings": False,
        "can_manage_automation": False,
        "can_import_settlements": True,
        "can_manage_onboarding": False,
    },
    "admin": {
        "label": "운영 관리자",
        "can_edit_local_data": True,
        "can_manage_settings": True,
        "can_manage_automation": True,
        "can_import_settlements": True,
        "can_manage_onboarding": True,
    },
}


def _employee_access_role(role: Any) -> str:
    normalized = str(role or "employee").strip().lower()
    return normalized if normalized in EMPLOYEE_ACCESS_ROLES else "employee"


def _is_admin(user: dict[str, Any]) -> bool:
    email = _email(user)
    user_role = str(user.get("user_role") or "").strip().lower()
    privileged_principal = bool(user.get("is_internal_admin")) or user_role in {"ceo", "admin", "system"}
    if email and not privileged_principal:
        membership = user.get("current_membership")
        membership_valid = (
            isinstance(membership, dict)
            and str(membership.get("status") or "").strip().lower() == "active"
            and str(membership.get("tenant_id") or "").strip() == str(user.get("tenant_id") or "").strip()
        )
        employee_record = next(
            (
                row
                for row in (_read_hr("employee_join_requests", user) if membership_valid else [])
                if str(row.get("email") or "").strip().lower() == email
                and str(row.get("status") or "").strip().lower() == "approved"
            ),
            None,
        )
        if employee_record:
            return _employee_access_role(employee_record.get("role")) == "admin"
    tenant_role = str(user.get("tenant_role") or "").strip().lower()
    membership_role = str((user.get("current_membership") or {}).get("role") or "").strip().lower()
    return bool(
        user.get("is_admin")
        or user.get("is_internal_admin")
        or tenant_role in {"owner", "admin"}
        or membership_role in {"owner", "admin"}
        or user_role in {"ceo", "admin", "system"}
    )


def _email(user: dict[str, Any]) -> str:
    return str(user.get("email") or "").strip().lower()


def _filter_user(rows: list[dict[str, Any]], user: dict[str, Any], *email_keys: str) -> list[dict[str, Any]]:
    if _is_admin(user):
        return rows
    email = _email(user)
    return [row for row in rows if any(str(row.get(key) or "").strip().lower() == email for key in email_keys)]


def _document_meta(document_type: str) -> dict[str, str]:
    return next((item for item in DOCUMENT_TYPES if item["type"] == document_type), DOCUMENT_TYPES[-1])


def _required_document_types() -> list[dict[str, str]]:
    return [item for item in DOCUMENT_TYPES if item.get("requirement") == "필수"]


def list_document_types() -> list[dict[str, str]]:
    return DOCUMENT_TYPES


def session_for_user(user: dict[str, Any]) -> dict[str, Any]:
    _tenant_id(user)
    email = _email(user)
    joins = _read_hr("employee_join_requests", user)
    own = next((row for row in joins if str(row.get("email") or "").strip().lower() == email), None)
    user_role = str(user.get("user_role") or "").strip().lower()
    privileged_principal = bool(user.get("is_internal_admin")) or user_role in {"ceo", "admin", "system"}
    admin = _is_admin(user)
    if own and not privileged_principal:
        status = str(own.get("status") or "pending").strip().lower()
        if status == "approved":
            access_role = _employee_access_role(own.get("role"))
            role_config = EMPLOYEE_ACCESS_ROLES[access_role]
            permissions = {
                "role": access_role,
                "role_label": role_config["label"],
                "employee_request_status": status,
                "employee_request_id": own.get("id"),
                "can_view": True,
                "can_edit_local_data": role_config["can_edit_local_data"],
                "can_manage_settings": role_config["can_manage_settings"],
                "can_manage_automation": role_config["can_manage_automation"],
                "can_import_settlements": role_config["can_import_settlements"],
                "can_manage_onboarding": role_config["can_manage_onboarding"],
                "can_upload_own_documents": True,
            }
        else:
            permissions = {
                "role": f"employee_{status}",
                "role_label": "직원 가입 반려" if status == "rejected" else "직원 가입요청",
                "employee_request_status": status,
                "employee_request_id": own.get("id"),
                "can_view": True,
                "can_edit_local_data": False,
                "can_manage_settings": False,
                "can_manage_automation": False,
                "can_import_settlements": False,
                "can_manage_onboarding": False,
                "can_upload_own_documents": status != "rejected",
            }
    elif admin:
        permissions = {
            "role": "owner",
            "role_label": "총괄 운영관리자",
            "can_view": True,
            "can_edit_local_data": True,
            "can_manage_settings": True,
            "can_manage_automation": True,
            "can_import_settlements": True,
            "can_manage_onboarding": True,
            "can_upload_own_documents": True,
        }
    else:
        permissions = {
            "role": "member",
            "role_label": "운영관리자",
            "can_view": True,
            "can_edit_local_data": True,
            "can_manage_settings": False,
            "can_manage_automation": False,
            "can_import_settlements": False,
            "can_manage_onboarding": False,
            "can_upload_own_documents": True,
        }
    return {
        "user": {
            "id": user.get("user_id"),
            "email": email,
            "name": user.get("name") or "",
            "tenant_id": user.get("tenant_id") or "",
            "is_admin": admin,
        },
        "tenant": user.get("current_tenant") or {"id": user.get("tenant_id")},
        "permissions": permissions,
    }


INVITE_BUSINESS_ERROR = "초대할 수 없는 사업자입니다"


def _business_invite_info(business_id: str) -> dict[str, Any] | None:
    """사업자 상호와 지점명 목록. DB 모드는 DB 가 기준, 파일 모드는 기본 사업자·지점 목록이 기준이다."""
    if _db_available():
        info = _run_db(_db_business_invite_info(business_id))
        return info if isinstance(info, dict) else None
    business = next((item for item in CANONICAL_BUSINESSES if item["id"] == business_id), None)
    if not business:
        return None
    branches = [item["name"] for item in CANONICAL_BRANCHES if item["businessId"] == business_id]
    branches += [name for name, owner in BUSINESS_BY_BRANCH.items() if owner == business_id and name not in branches]
    return {"name": business["name"], "branches": branches}


def _invite_target_label(business_name: str, branch: str) -> str:
    return f"{business_name} · {branch}" if branch else business_name


def _clean_invite_targets(raw: Any) -> list[dict[str, str]]:
    cleaned: list[dict[str, str]] = []
    for item in raw if isinstance(raw, list) else []:
        if not isinstance(item, dict):
            continue
        business_id = str(item.get("business_id") or "").strip()
        branch = str(item.get("branch") or "").strip()
        branch = BRANCH_ALIASES.get(branch, branch)
        if {"business_id": business_id, "branch": branch} not in cleaned:
            cleaned.append({"business_id": business_id, "branch": branch})
    return cleaned


def _validated_invite_targets(payload: dict[str, Any], user: dict[str, Any]) -> list[dict[str, str]]:
    caller_tenant = _tenant_id(user)
    targets = _clean_invite_targets(payload.get("targets"))
    legacy = not targets
    if legacy:
        # 하위호환: targets 없이 branch 만 오면 종전처럼 지점명으로 사업자를 정한다.
        branch = str(payload.get("branch") or "").strip()
        branch = BRANCH_ALIASES.get(branch, branch)
        business_id = str(BUSINESS_BY_BRANCH.get(branch) or "")
        if not business_id:
            return []
        targets = [{"business_id": business_id, "branch": branch}]
    business_ids = [target["business_id"] for target in targets]
    if len(set(business_ids)) != len(business_ids):
        # 가입요청의 키는 (이메일, 사업자)라 같은 사업자의 지점 둘은 요청 하나로 합쳐진다.
        raise HTTPException(status_code=400, detail="한 사업자에서는 매장을 하나만 고를 수 있습니다")
    for target in targets:
        business_id, branch = target["business_id"], target["branch"]
        if not business_id:
            raise HTTPException(status_code=400, detail=INVITE_BUSINESS_ERROR)
        if _db_available():
            owner = _run_db(_db_business_tenant_id(business_id))
            if not owner or str(owner).strip() != caller_tenant:
                raise HTTPException(status_code=400, detail=INVITE_BUSINESS_ERROR)
        elif business_id not in CANONICAL_BUSINESS_IDS:
            raise HTTPException(status_code=400, detail=INVITE_BUSINESS_ERROR)
        info = _business_invite_info(business_id)
        if info is None:
            raise HTTPException(status_code=400, detail=INVITE_BUSINESS_ERROR)
        if branch and not legacy and branch not in info["branches"]:
            raise HTTPException(status_code=400, detail=f"{info['name']}의 지점이 아닙니다: {branch}")
    return targets


def _invite_view(invite: dict[str, Any], names: dict[str, str]) -> dict[str, Any]:
    """저장된 초대에 사업자 상호·매장 표시 문자열을 붙인다. 예전 초대(targets 없음)는 지점명에서 유추한다."""
    view = dict(invite)
    targets = _clean_invite_targets(invite.get("targets"))
    if not targets:
        branch = str(invite.get("branch") or "").strip()
        branch = BRANCH_ALIASES.get(branch, branch)
        business_id = str(invite.get("business_id") or BUSINESS_BY_BRANCH.get(branch) or "").strip()
        targets = [{"business_id": business_id, "branch": branch}] if (business_id or branch) else []
    labels: list[str] = []
    for target in targets:
        business_id = target["business_id"]
        if business_id and business_id not in names:
            info = _business_invite_info(business_id)
            names[business_id] = str(info["name"]) if info else ""
        business_name = names.get(business_id) or str(invite.get("business_name") or "") or business_id
        labels.append(_invite_target_label(business_name, target["branch"]) if business_name else target["branch"])
    first_id = targets[0]["business_id"] if targets else ""
    view["targets"] = targets
    view["business_id"] = first_id
    view["business_name"] = (names.get(first_id) or str(invite.get("business_name") or "")) if first_id else ""
    view["branch_labels"] = labels
    return view


def _invite_scope_tenant(user: dict[str, Any]) -> str | None:
    """초대 열람·취소 범위. None 은 내부관리자(전체 열람), 그 밖에는 호출자의 검증된 테넌트."""
    role = str(user.get("user_role") or "").strip().lower()
    if user.get("is_internal_admin") or role in {"ceo", "admin", "system"}:
        return None
    return _tenant_id(user)


def _invite_owner_tenants(invite: dict[str, Any]) -> set[str]:
    """초대 대상 사업자들의 소유 테넌트. _validated_invite_targets 와 같은 _db_business_tenant_id 판정을 쓴다."""
    targets = _clean_invite_targets(invite.get("targets"))
    if not targets:
        branch = str(invite.get("branch") or "").strip()
        branch = BRANCH_ALIASES.get(branch, branch)
        targets = [{"business_id": str(invite.get("business_id") or BUSINESS_BY_BRANCH.get(branch) or "").strip(), "branch": branch}]
    owners: set[str] = set()
    if not _db_available():
        return owners
    for target in targets:
        if target["business_id"]:
            owner = _run_db(_db_business_tenant_id(target["business_id"]))
            if owner and str(owner).strip():
                owners.add(str(owner).strip())
    return owners


def _invite_in_scope(invite: dict[str, Any], scope_tenant: str | None) -> bool:
    if scope_tenant is None:
        return True
    owners = _invite_owner_tenants(invite)
    # 테넌트를 특정할 수 없는 레거시 초대(매핑 없음·파일 모드)는 종전처럼 관리자에게 그대로 보인다.
    return not owners or scope_tenant in owners


def list_invites(user: dict[str, Any]) -> list[dict[str, Any]]:
    if not _is_admin(user):
        return []
    scope_tenant = _invite_scope_tenant(user)
    names: dict[str, str] = {}
    rows = sorted(_read("employee_invites"), key=lambda row: row.get("created_at", ""), reverse=True)
    return [_invite_view(row, names) for row in rows if _invite_in_scope(row, scope_tenant)]


def revoke_invite(invite_id: str, user: dict[str, Any]) -> dict[str, Any]:
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="직원 초대 취소 권한이 없습니다")
    rows = _read("employee_invites")
    invite = _find(rows, invite_id)
    if not invite:
        raise HTTPException(status_code=404, detail="초대를 찾을 수 없습니다")
    if not _invite_in_scope(invite, _invite_scope_tenant(user)):
        raise HTTPException(status_code=403, detail="다른 사업자의 초대는 취소할 수 없습니다")
    status = str(invite.get("status") or "").strip().lower()
    if status == "accepted":
        raise HTTPException(status_code=409, detail="이미 수락된 초대는 취소할 수 없습니다. 가입요청을 반려하십시오.")
    if status == "revoked":
        return _invite_view(invite, {})
    invite["status"] = "revoked"
    invite["revoked_at"] = _now()
    invite["revoked_by"] = _email(user)
    invite.pop("token", None)
    _write("employee_invites", rows)
    return _invite_view(invite, {})


def create_invite(payload: dict[str, Any], user: dict[str, Any]) -> dict[str, Any]:
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="직원 초대 권한이 없습니다")
    targets = _validated_invite_targets(payload, user)
    rows = _read("employee_invites")
    now = _now()
    token = secrets.token_urlsafe(24)
    branch = targets[0]["branch"] if targets else str(payload.get("branch") or "").strip()
    invite = {
        "id": str(uuid4()),
        "token": token,
        "phone": str(payload.get("phone") or "").strip(),
        "phone_masked": _mask_phone(str(payload.get("phone") or "")),
        "name": str(payload.get("name") or "").strip(),
        "branch": branch,
        "role": str(payload.get("role") or "member"),
        "status": "pending",
        "memo": str(payload.get("memo") or ""),
        "created_by": _email(user),
        "created_at": now,
        "expires_at": (datetime.now(KST) + timedelta(hours=int(payload.get("expires_in_hours") or 72))).isoformat(timespec="seconds"),
    }
    invite.update(_invite_view({**invite, "targets": targets}, {}))
    rows.insert(0, invite)
    _write("employee_invites", rows)
    return invite


def _invite_expired(invite: dict[str, Any]) -> bool:
    """expires_at 이 지난 초대인가. 값이 없거나 해석할 수 없는 예전 초대는 만료로 보지 않는다."""
    raw = str(invite.get("expires_at") or "").strip()
    if not raw:
        return False
    try:
        expires = datetime.fromisoformat(raw)
    except ValueError:
        return False
    if expires.tzinfo is None:
        expires = expires.replace(tzinfo=KST)
    return expires < datetime.now(KST)


def _find_invite(token: str) -> dict[str, Any]:
    token = str(token or "").strip()
    invite = next((row for row in _read("employee_invites") if token and row.get("token") == token), None)
    if not invite:
        raise HTTPException(status_code=404, detail="초대를 찾을 수 없습니다")
    if str(invite.get("status") or "").strip().lower() == "revoked":
        raise HTTPException(status_code=404, detail="취소된 초대입니다")
    if _invite_expired(invite):
        raise HTTPException(status_code=410, detail="만료된 초대입니다")
    return invite


def resolve_invite(token: str) -> dict[str, Any]:
    view = _invite_view(_find_invite(token), {})
    view.pop("token", None)
    return view


def _rollback_join_requests(done: list[tuple[dict[str, Any], dict[str, Any], dict[str, Any] | None]]) -> None:
    for scope_user, saved, previous in reversed(done):
        try:
            if previous:
                _write_hr_record("employee_join_requests", previous, scope_user)
            else:
                _delete_hr_record("employee_join_requests", str(saved.get("id") or ""), scope_user)
        except Exception:
            logger.exception("초대 수락 롤백 실패: %s", saved.get("id"))


def accept_invite(payload: dict[str, Any], user: dict[str, Any]) -> dict[str, Any]:
    invite = _find_invite(str(payload.get("token") or ""))
    if str(invite.get("status") or "").strip().lower() == "accepted" and str(invite.get("accepted_email") or "").strip().lower() != _email(user):
        raise HTTPException(status_code=409, detail="이미 다른 계정이 수락한 초대입니다")
    # 초대 토큰이 가리키는 매장이 기준이다 — 본문 branch 로 다른 사업자에 갈아타지 못한다.
    # upsert_join_request 가 그 사업자의 고용주 테넌트(매핑)에 요청을 귀속한다.
    stored_targets = _clean_invite_targets(invite.get("targets"))
    if stored_targets:
        targets = stored_targets
    else:
        targets = [{"business_id": "", "branch": str(invite.get("branch") or payload.get("branch") or "")}]
    requests_payloads = []
    for target in targets:
        item = {
            "name": payload.get("name") or invite.get("name") or "",
            "email": _email(user),
            "branch": target["branch"],
            "phone": payload.get("phone") or invite.get("phone") or "",
            "memo": payload.get("memo") or "전화번호 초대 링크로 회원가입",
            "invite_id": invite.get("id"),
        }
        if target["business_id"]:
            item["business_id"] = target["business_id"]
        requests_payloads.append(item)

    def _reject(exc: HTTPException, item: dict[str, Any]) -> HTTPException:
        where = item.get("branch") or item.get("business_id") or ""
        return HTTPException(status_code=400, detail=f"초대 수락 실패({where}): {exc.detail}" if where else str(exc.detail))

    # 1단계: 저장 없이 전 매장을 먼저 검증한다 — 한 곳이라도 안 되면 아무것도 만들지 않는다.
    for item in requests_payloads:
        try:
            _prepare_join_request(item, user)
        except HTTPException as exc:
            raise _reject(exc, item) from exc
    # 2단계: 저장. 중간에 실패하면 이미 만든 요청을 되돌린다.
    created: list[dict[str, Any]] = []
    done: list[tuple[dict[str, Any], dict[str, Any], dict[str, Any] | None]] = []
    for item in requests_payloads:
        try:
            _record, scope_user, previous = _prepare_join_request(item, user)
            saved = upsert_join_request(item, user)
        except Exception as exc:
            _rollback_join_requests(done)
            if isinstance(exc, HTTPException):
                raise _reject(exc, item) from exc
            raise
        created.append(saved)
        done.append((scope_user, saved, previous))
    rows = _read("employee_invites")
    target = _find(rows, invite["id"])
    if target:
        target["status"] = "accepted"
        target["accepted_at"] = _now()
        target["accepted_email"] = _email(user)
        _write("employee_invites", rows)
    # 단일 매장 호출부 호환: 첫 가입요청의 필드를 최상위에도 펼친다.
    return {**created[0], "requests": created, "request": created[0]}


def list_join_requests(user: dict[str, Any]) -> list[dict[str, Any]]:
    from app.core.obys_tenant import is_legacy_obys_tenant

    if is_legacy_obys_tenant(user):
        rows = _read_hr("employee_join_requests", user)
        return sorted(_filter_user(rows, user, "email"), key=lambda row: row.get("requested_at", ""), reverse=True)
    # 레거시 테넌트가 아니면 직원 본인이다 — 가입요청은 고용주 테넌트에 있으므로 본인 이메일 것만 읽는다.
    _tenant_id(user)
    email = _email(user)
    if not email:
        return []
    return sorted(_read_join_requests_by_email(email), key=lambda row: row.get("requested_at", ""), reverse=True)


JOIN_BRANCH_MISMATCH_ERROR = "직원의 사업자와 지점 연결이 일치하지 않습니다"
JOIN_NAME_REQUIRED_ERROR = "직원 이름(실명)을 입력해 주십시오"
JOIN_NAME_EMAIL_ERROR = "이름에 이메일 아이디를 넣지 말고 실명을 입력해 주십시오"


def _validate_join_business_branch(business_id: str, branch: str) -> None:
    """가입요청의 사업자·지점이 유효한지 판정한다. accept_invite 도 _prepare_join_request 로 여기를 거친다.

    DB 모드는 yeoljeong_businesses/branches 가 기준이라 새로 만든 사업자·지점도 통과한다.
    파일 모드(DB 없음)는 종전 코드 상수 판정을 그대로 쓴다. 테넌트 격리 판정은 여기가 아니라
    _join_request_scope·_owned_hr_record 가 맡는다.
    """
    if not (business_id and _db_available()):
        if branch and (business_id not in CANONICAL_BUSINESS_IDS or BUSINESS_BY_BRANCH.get(branch) != business_id):
            raise HTTPException(status_code=400, detail=JOIN_BRANCH_MISMATCH_ERROR)
        return
    info = _business_invite_info(business_id)
    if info is None:
        raise HTTPException(status_code=400, detail="등록되지 않은 사업자입니다")
    if not branch:
        return
    branches = list(info.get("branches") or [])
    if branch in branches:
        return
    # 지점이 하나도 없는 사업자는 사업자 자체가 매장이다 — 지점 칸에 사업자명을 적은 경우만 같은 매장으로 본다.
    if not branches and branch == str(info.get("name") or ""):
        return
    raise HTTPException(status_code=400, detail=JOIN_BRANCH_MISMATCH_ERROR)


def _resolve_join_business_by_branch(branch: str) -> str:
    """상수 BUSINESS_BY_BRANCH 에 없는 지점 이름을 DB 로 사업자에 해석한다. 파일 모드·미등록 지점은 빈 문자열."""
    if not (branch and _db_available()):
        return ""
    found = _run_db(_db_business_ids_by_branch_name(branch))
    business_ids = sorted({str(item).strip() for item in (found or []) if str(item).strip()})
    if len(business_ids) > 1:
        raise HTTPException(
            status_code=400,
            detail="같은 이름의 지점이 여러 사업자에 있습니다 — 사업자를 지정해 주십시오",
        )
    return business_ids[0] if business_ids else ""


def _prepare_join_request(
    payload: dict[str, Any], user: dict[str, Any]
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any] | None]:
    """가입요청 레코드를 검증·구성만 한다(저장 없음). (record, scope_user, 기존 레코드)를 돌려준다."""
    _tenant_id(user)
    email = str(payload.get("email") or _email(user)).strip().lower()
    if not email:
        raise HTTPException(status_code=400, detail="직원 이메일이 필요합니다")
    # 겸직: 가입요청의 식별 키는 (email, business_id) 다. 사업자가 다르면 별개 고용주라 별개 요청이다.
    payload_branch = str(payload.get("branch") or "").strip()
    payload_branch = BRANCH_ALIASES.get(payload_branch, payload_branch)
    payload_business_id = (
        str(payload.get("business_id") or "").strip()
        or str(BUSINESS_BY_BRANCH.get(payload_branch) or "")
        or _resolve_join_business_by_branch(payload_branch)
    )
    # 요청 레코드는 호출자(직원)의 새 테넌트가 아니라 대상 사업자의 고용주 테넌트에 귀속한다.
    scope_user = _join_request_scope(payload_business_id, email, user) if payload_business_id else user
    rows = _read_hr("employee_join_requests", scope_user)
    same_email_rows = [row for row in rows if str(row.get("email") or "").strip().lower() == email]
    existing = None
    if not payload_business_id:
        # payload 가 사업자를 특정하지 않으면 종전처럼 이메일만으로 매칭한다.
        existing = next(iter(same_email_rows), None)
    else:
        existing = next(
            (row for row in same_email_rows if str(row.get("business_id") or "").strip() == payload_business_id),
            None,
        )
        if existing is None:
            legacy_blank_rows = [row for row in same_email_rows if not str(row.get("business_id") or "").strip()]
            if len(legacy_blank_rows) == 1:
                # business_id 가 비어 있던 기존 1건에 채운다 — 새 행을 만들면 유령 요청이 남는다.
                existing = legacy_blank_rows[0]
            elif len(legacy_blank_rows) > 1:
                raise HTTPException(
                    status_code=400,
                    detail="같은 이메일의 사업자 미지정 가입요청이 2건 이상이라 자동으로 연결할 수 없습니다 — 수동 정리가 필요합니다",
                )
    now = _now()
    record = dict(existing) if existing else {"id": str(uuid4()), "requested_at": now}
    real_name = str(payload.get("name") or record.get("name") or "").strip()
    if not real_name:
        raise HTTPException(status_code=400, detail=JOIN_NAME_REQUIRED_ERROR)
    if "@" in real_name or real_name.lower() == email.split("@")[0]:
        raise HTTPException(status_code=400, detail=JOIN_NAME_EMAIL_ERROR)
    branch = BRANCH_ALIASES.get(
        str(payload.get("branch") or record.get("branch") or "").strip(),
        str(payload.get("branch") or record.get("branch") or "").strip(),
    )
    business_id = str(
        payload.get("business_id") or record.get("business_id") or BUSINESS_BY_BRANCH.get(branch) or payload_business_id
    ).strip()
    if not business_id:
        raise HTTPException(status_code=400, detail="회사(사업자)와 근무 점포를 선택해 주십시오")
    _validate_join_business_branch(business_id, branch)
    if business_id and business_id != payload_business_id:
        # 기존 요청에서 이어받은 사업자 — 읽은 테넌트와 다르면 다른 테넌트에 쓰지 않는다.
        if _tenant_id(_join_request_scope(business_id, email, user)) != _tenant_id(scope_user):
            raise HTTPException(status_code=403, detail="다른 테넌트의 데이터는 저장할 수 없습니다")
    record.update(
        {
            "name": real_name,
            "email": email,
            "email_masked": _mask_email(email),
            "phone": str(payload.get("phone") or record.get("phone") or "").strip(),
            "phone_masked": _mask_phone(str(payload.get("phone") or record.get("phone") or "")),
            "address": str(payload.get("address") or record.get("address") or "").strip(),
            "birth_date": str(payload.get("birth_date") or record.get("birth_date") or "").strip(),
            "nationality": str(payload.get("nationality") or record.get("nationality") or "대한민국").strip(),
            "business_id": business_id,
            "branch": branch,
            "memo": str(payload.get("memo") or record.get("memo") or ""),
            "invite_id": payload.get("invite_id") or record.get("invite_id") or "",
            "status": record.get("status") if existing else "pending",
            "updated_at": now,
        }
    )
    # 신원 고정: 본인 로그인 이메일로 낸 요청만 요청자 계정(JWT user_id)을 남긴다.
    # 승인 시 고용주 테넌트 멤버십은 이 값으로만 연결된다 — 본문 email 만으로는
    # 남의 계정을 연결하거나(승인) 회수하지(반려) 못한다.
    requester_email = _email(user)
    requester_user_id = str(user.get("user_id") or user.get("id") or "").strip()
    if requester_email and requester_email == email and requester_user_id:
        record["requester_user_id"] = requester_user_id
        record["requester_email"] = requester_email
        record["requested_by"] = requester_email
    elif not existing:
        record["registered_by"] = requester_email
    record = _owned_hr_record(record, scope_user)
    return record, scope_user, existing


def upsert_join_request(payload: dict[str, Any], user: dict[str, Any]) -> dict[str, Any]:
    record, scope_user, _existing = _prepare_join_request(payload, user)
    _write_hr_record("employee_join_requests", record, scope_user)
    return record


def review_join_request(request_id: str, action: str, memo: str, user: dict[str, Any]) -> dict[str, Any]:
    _require_review_tenant(request_id, user)
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="가입요청 승인 권한이 없습니다")
    record = _require_hr_record(_find(_read_hr("employee_join_requests", user), request_id), user, detail="가입요청을 찾을 수 없습니다")
    if action not in {"approved", "rejected"}:
        raise HTTPException(status_code=400, detail="action은 approved 또는 rejected여야 합니다")
    record["status"] = action
    record["review_memo"] = memo
    record["reviewed_by"] = _email(user)
    record["reviewed_at"] = _now()
    record["updated_at"] = record["reviewed_at"]
    _write_hr_record("employee_join_requests", record, user)
    return record


# ---------------------------------------------------------------------------
# 직원 승인 ↔ 고용주 테넌트 멤버십 (AADS-OBYS-EMPLOYEE-TENANT-MEMBERSHIP-20260930)
#
# 2026-09-30 진아서버 실측: 승인 직원 7명 중 고용주 테넌트(tenant_memberships)
# 소속은 0명이었다. 직원은 가입 때 자동 생성된 개인 워크스페이스로만 로그인해
# _read_hr 가 계약서를 못 찾았고, 그 워크스페이스에서는 owner 라 _is_admin 이
# 참이 되어 서명이 막혔다. 승인이 멤버십(role=member)을 만들고, 반려·퇴사가
# 그 승인이 만든 멤버십만 회수(status=removed)한다. 멤버십은 AADS 인증 DB,
# 가입요청·감사는 오비서 DB 라 한 트랜잭션이 아니다 — 승인 자체는 멤버십
# 실패와 무관하게 성공으로 남기고 결과를 응답과 감사에 담는다.
#
# 신원: 본인 로그인으로 제출한 요청은 upsert_join_request 가 남긴
# requester_user_id 로 연결한다(이메일이 바뀌었으면 identity_mismatch).
# 그 고정이 없는 요청 — 관리자 대리 등록·초대 흐름·예전 요청 — 은 요청
# 이메일로 saas_users 를 찾는다. 직원은 개인 워크스페이스 JWT 로 고용주
# 테넌트에 요청을 낼 수 없어(obys_tenant 게이트) 실제 요청 대부분이 이쪽이고,
# 이메일은 승인하는 owner/admin 이 보증한다. requester_* 가 남았는데 요청
# 이메일과 어긋나면(남의 이메일로 낸 흔적) 이메일로 찾지 않고 pending_identity.
# ---------------------------------------------------------------------------
MEMBERSHIP_AUDIT_LOG = "tenant_membership_audit"
MEMBERSHIP_AUDIT_LEDGER = "tenant_memberships"
_MEMBERSHIP_AUDIT_CLASSIFICATION = {
    "linked": "membership_linked",
    "unchanged": "membership_unchanged",
    "pending_account": "membership_pending",
    "pending_identity": "membership_pending",
    "identity_mismatch": "membership_skipped",
    "removed": "membership_removed",
    "kept_elevated": "membership_skipped",
    "not_member": "membership_skipped",
    "skipped": "membership_skipped",
    "error": "membership_failed",
}


def _join_request_account_id(record: dict[str, Any]) -> str:
    """본인이 제출한 가입요청이면 그 계정 id. 아니면 빈 문자열."""
    requester_id = str(record.get("requester_user_id") or "").strip()
    requester_email = str(record.get("requester_email") or "").strip().lower()
    email = str(record.get("email") or "").strip().lower()
    return requester_id if requester_id and requester_email and requester_email == email else ""


def _join_request_allows_email_lookup(record: dict[str, Any]) -> bool:
    """본인 제출 고정이 없는 요청(관리자 대리 등록·초대 흐름·예전 요청)은 승인자가 이메일을 보증한다.

    requester_* 가 남아 있는데 요청 이메일과 어긋나면 누군가 남의 이메일로 낸 흔적이다 —
    그때는 이메일로 계정을 찾지 않는다.
    """
    if _join_request_account_id(record):
        return False
    return not str(record.get("requester_user_id") or "").strip() and not str(record.get("requester_email") or "").strip()


def _membership_audit_row(
    link: dict[str, Any],
    *,
    action: str,
    source: str,
    record: dict[str, Any],
    actor_user_id: str,
    actor_email: str,
) -> dict[str, Any]:
    status = str(link.get("status") or "error")
    return {
        "id": str(uuid4()),
        "ledger_table": MEMBERSHIP_AUDIT_LEDGER,
        "classification": _MEMBERSHIP_AUDIT_CLASSIFICATION.get(status, "membership_failed"),
        "reason": f"{action}:{status}" + (f":{link['reason']}" if link.get("reason") else ""),
        "source": source,
        "business_id": str(record.get("business_id") or ""),
        "tenant_id": str(link.get("tenant_id") or record.get("tenant_id") or ""),
        "action": action,
        "actor_user_id": actor_user_id,
        "actor_email": actor_email,
        "employee_email": str(link.get("employee_email") or record.get("email") or "").strip().lower(),
        "employee_user_id": str(link.get("user_id") or ""),
        "join_request_id": str(record.get("id") or ""),
        "before_state": link.get("before"),
        "after_state": link.get("after"),
        "classified_at": _now(),
    }


async def _db_insert_membership_audit(row: dict[str, Any]) -> bool:
    import asyncpg

    conn = await asyncpg.connect(_db_url(), timeout=5)
    try:
        await conn.execute(
            """
            INSERT INTO yeoljeong_hr_tenant_attribution_audit
                (ledger_table, row_id, classification, reason, source, business_id, tenant_id,
                 action, actor_user_id, actor_email, employee_email, employee_user_id, join_request_id,
                 before_state, after_state, classified_at)
            VALUES ($1, $2, $3, $4, $5, NULLIF($6, ''), NULLIF($7, '')::uuid,
                    $8, NULLIF($9, ''), NULLIF($10, ''), $11, NULLIF($12, ''), NULLIF($13, ''),
                    $14::jsonb, $15::jsonb, $16::timestamptz)
            """,
            row["ledger_table"],
            row["id"],
            row["classification"],
            row["reason"],
            row["source"],
            row["business_id"],
            row["tenant_id"],
            row["action"],
            row["actor_user_id"],
            row["actor_email"],
            row["employee_email"],
            row["employee_user_id"],
            row["join_request_id"],
            json.dumps(row["before_state"], ensure_ascii=False, default=str) if row["before_state"] is not None else None,
            json.dumps(row["after_state"], ensure_ascii=False, default=str) if row["after_state"] is not None else None,
            _pg_ts(row["classified_at"]),
        )
        return True
    finally:
        await conn.close()


def _append_membership_audit_file(row: dict[str, Any]) -> None:
    """파일 원장 추가. 읽고-다시-쓰기 사이를 flock 으로 묶어 동시 승인에서 행이 사라지지 않게 한다."""
    safe_row = json.loads(json.dumps(row, ensure_ascii=False, default=str))
    lock_path = _path(MEMBERSHIP_AUDIT_LOG).with_suffix(".lock")
    with open(lock_path, "a+") as lock:
        fcntl.flock(lock.fileno(), fcntl.LOCK_EX)
        try:
            _write_file_rows(MEMBERSHIP_AUDIT_LOG, [safe_row] + _read_file_rows(MEMBERSHIP_AUDIT_LOG))
        finally:
            fcntl.flock(lock.fileno(), fcntl.LOCK_UN)


def _record_membership_audit_sync(row: dict[str, Any]) -> bool:
    """스레드에서 부른다. DB 가 있으면 감사 테이블, 없으면 파일 원장."""
    if _db_available():
        # _run_db 는 워커 스레드(이벤트 루프 없음)에서 asyncio.run 으로 돌리고 실패는 None 이다.
        if _run_db(_db_insert_membership_audit(row)):
            return True
        # 20260930 감사 마이그레이션은 자동 적용 HOLD 다. 코드가 먼저 나가 컬럼·CHECK 가
        # 없으면 INSERT 가 실패한다 — 그 사이 이벤트를 잃지 않도록 파일 원장에 남긴다.
        logger.error("membership audit DB insert failed, falling back to file: request=%s", row.get("join_request_id"))
        try:
            _append_membership_audit_file({**row, "db_insert_failed": True})
        except Exception as exc:  # noqa: BLE001
            logger.error("membership audit not recorded: request=%s err=%s", row.get("join_request_id"), exc)
            return False
        return True
    _append_membership_audit_file(row)
    return True


async def record_membership_audit(row: dict[str, Any]) -> bool:
    """감사 1행. 파일·DB I/O 는 이벤트 루프 밖(스레드)에서 한다 — 감사 실패가 승인을 되돌리지 않는다."""
    try:
        return await asyncio.to_thread(_record_membership_audit_sync, row)
    except Exception as exc:  # noqa: BLE001
        logger.error("membership audit not recorded: request=%s err=%s", row.get("join_request_id"), exc)
        return False


def _save_membership_link(request_id: str, link: dict[str, Any], user: dict[str, Any]) -> dict[str, Any] | None:
    """가입요청에 '이 승인이 만든 멤버십' 표시를 남긴다 — 반려가 회수할 수 있는 유일한 근거."""
    record = _find(_read_hr("employee_join_requests", user), request_id)
    if not record:
        return None
    status = link.get("status")
    marker = dict(record.get("membership_link") or {})
    if status == "linked" and link.get("owned_by_request"):
        marker = {
            "tenant_id": str(link.get("tenant_id") or ""),
            "user_id": str(link.get("user_id") or ""),
            "membership_id": str((link.get("after") or {}).get("membership_id") or ""),
            "linked_by": str(user.get("user_id") or user.get("id") or ""),
            "linked_at": _now(),
        }
    elif status == "removed":
        marker["revoked_at"] = _now()
    else:
        return record
    record["membership_link"] = marker
    _write_hr_record("employee_join_requests", record, user)
    return record


def _membership_revoke_target(record: dict[str, Any], previous_status: str, tenant_id: str) -> tuple[dict[str, Any], str]:
    """반려가 회수해도 되는 멤버십인가 — (표시, 거절 사유). 사유가 비어야 회수한다."""
    if str(previous_status or "").strip().lower() != "approved":
        return {}, "request_was_not_approved"
    marker = record.get("membership_link") or {}
    if not isinstance(marker, dict) or not marker.get("membership_id") or not marker.get("user_id"):
        return {}, "membership_not_linked_by_request"
    if marker.get("revoked_at"):
        return {}, "membership_already_revoked"
    if str(marker.get("tenant_id") or "") != tenant_id:
        return {}, "membership_linked_in_other_tenant"
    account_id = _join_request_account_id(record)
    if account_id and account_id != str(marker.get("user_id")):
        return {}, "membership_user_differs_from_requester"
    return marker, ""


async def sync_employee_tenant_membership(
    record: dict[str, Any],
    action: str,
    user: dict[str, Any],
    *,
    previous_status: str = "",
    source: str = "review_join_request",
    conn: Any = None,
) -> dict[str, Any]:
    """승인이면 고용주 테넌트 member 연결, 반려·퇴사면 이 요청이 만든 멤버십만 회수.

    실패해도 예외를 올리지 않는다(결과·감사에 error 로 남긴다).  ``conn`` 은
    employee_membership_lock 이 내준 연결 — 멤버십 변경을 그 락 트랜잭션 안에서 한다.
    """
    from app import auth as auth_module

    tenant_id = _tenant_id(user)
    employee_email = str(record.get("email") or record.get("employee_email") or "").strip().lower()
    actor_user_id = str(user.get("user_id") or user.get("id") or "")
    base = {"tenant_id": tenant_id, "employee_email": employee_email, "user_id": None, "before": None, "after": None}
    try:
        if action == "approved":
            allow_email_lookup = _join_request_allows_email_lookup(record)
            if allow_email_lookup:
                # 이메일 조회의 전제(요청 테넌트 == 가입요청 테넌트·사업자)를 여기서도 다시 확인한다.
                await asyncio.to_thread(_require_employee_membership_context, record, user)
            link = await auth_module.link_employee_tenant_membership(
                tenant_id=tenant_id,
                employee_user_id=_join_request_account_id(record) or None,
                employee_email=employee_email,
                invited_by=actor_user_id or None,
                allow_email_lookup=allow_email_lookup,
                context_tenant_id=tenant_id,
                conn=conn,
            )
        else:
            marker, reason = _membership_revoke_target(record, previous_status, tenant_id)
            if reason:
                link = {**base, "status": "skipped", "reason": reason}
            else:
                link = await auth_module.revoke_employee_tenant_membership(
                    tenant_id=tenant_id,
                    user_id=str(marker["user_id"]),
                    membership_id=str(marker["membership_id"]),
                    employee_email=employee_email,
                    conn=conn,
                )
                link.setdefault("employee_email", employee_email)
    except Exception as exc:  # noqa: BLE001
        logger.error("employee membership sync failed: request=%s action=%s err=%s", record.get("id"), action, exc)
        reason = "tenant_context_mismatch" if isinstance(exc, (auth_module.EmployeeTenantContextError, HTTPException)) else type(exc).__name__
        link = {**base, "status": "error", "reason": reason}
    audit = _membership_audit_row(
        link,
        action=action,
        source=source,
        record=record,
        actor_user_id=actor_user_id,
        actor_email=_email(user),
    )
    recorded = await record_membership_audit(audit)
    try:
        saved = await asyncio.to_thread(_save_membership_link, str(record.get("id") or ""), link, user)
    except Exception as exc:  # noqa: BLE001
        logger.error("membership link marker not saved: request=%s err=%s", record.get("id"), exc)
        saved = None
    if saved:
        record.update(saved)
    return {
        "status": link.get("status"),
        "reason": link.get("reason") or "",
        "tenant_id": tenant_id,
        "user_id": link.get("user_id"),
        "role": (link.get("after") or {}).get("role"),
        "membership_status": (link.get("after") or link.get("before") or {}).get("status"),
        "default_tenant": link.get("default_tenant"),
        "audit_recorded": recorded,
    }


def _require_employee_membership_context(record: dict[str, Any], user: dict[str, Any]) -> str:
    """이메일로 직원 계정을 찾기 전의 전제를 코드로 강제한다 — 부재·불일치는 예외.

    ① 요청에 검증된 테넌트 컨텍스트(JWT tenant + 활성 멤버십)가 있어야 하고(_tenant_id),
    ② 가입요청 레코드가 그 테넌트 것이어야 하며(_require_hr_record),
    ③ 가입요청의 사업자가 그 테넌트에 귀속돼 있어야 한다(_require_business_for_tenant).
    예전엔 "레거시 테넌트 JWT 컨텍스트가 이미 잡혀 있다" 를 주석으로만 전제했다.
    """
    tenant_id = _tenant_id(user)
    _require_hr_record(record, user, detail="가입요청을 찾을 수 없습니다")
    _require_business_for_tenant(record.get("business_id"), user)
    return tenant_id


def _precheck_join_review(request_id: str, action: str, user: dict[str, Any]) -> str:
    """검토 전에 락 키(직원 이메일)를 정하고, 승인이면 이메일 조회 전제를 먼저 확인한다.

    승인 저장 뒤에 컨텍스트 오류가 나면 '승인됐는데 멤버십 없음' 이 남는다 — 그래서
    저장 전에 막는다(403/404, 아무것도 저장되지 않음).
    """
    _require_review_tenant(request_id, user)
    record = _find(_read_hr("employee_join_requests", user), request_id)
    if record and action == "approved" and _join_request_allows_email_lookup(record):
        _require_employee_membership_context(record, user)
    return str((record or {}).get("email") or (record or {}).get("employee_email") or "").strip().lower()


def _review_join_request_with_previous(
    request_id: str, action: str, memo: str, user: dict[str, Any]
) -> tuple[dict[str, Any], str]:
    previous = _find(_read_hr("employee_join_requests", user), request_id) or {}
    previous_status = str(previous.get("status") or "")
    return review_join_request(request_id, action, memo, user), previous_status


async def review_join_request_with_membership(
    request_id: str,
    action: str,
    memo: str,
    user: dict[str, Any],
) -> dict[str, Any]:
    """가입요청 승인·반려 + 고용주 테넌트 멤버십 연결·회수.

    승인 판정(review_join_request)은 그대로 두고, 저장이 끝난 뒤에만 멤버십을
    바꾼다 — 권한·검증 실패로 승인이 안 됐는데 멤버십만 생기는 일이 없다.
    반려는 직전 상태가 approved 이고 그 승인이 남긴 membership_link 가 있을 때만 회수한다.

    직전 상태 읽기 → 저장 → 멤버십 연결·회수는 (고용주 테넌트, 직원 이메일) 단위
    pg_advisory_xact_lock 으로 한 줄로 세운다.  예) 승인 A 와 반려 B 가 동시에 돌면 B 가
    직전 상태를 pending 으로 읽어 회수를 건너뛰고, 그 사이 A 가 멤버십을 만들어
    '반려됐는데 멤버' 가 남는다.  락은 blue/green 두 컨테이너가 같이 보는 AADS 인증 DB 에
    건다(파일 flock 은 컨테이너마다 파일시스템이 달라 서로를 못 본다).
    인증 DB 에 닿지 못하면 락 없이 진행한다 — 그때 멤버십 변경도 같은 DB 라 실패하고
    error 로 감사에 남으며, 승인 자체는 되돌리지 않는다.
    """
    from app import auth as auth_module

    tenant_id = _tenant_id(user)
    employee_email = await asyncio.to_thread(_precheck_join_review, request_id, action, user)
    async with AsyncExitStack() as stack:
        conn = None
        try:
            conn = await stack.enter_async_context(auth_module.employee_membership_lock(tenant_id, employee_email))
        except Exception as exc:  # noqa: BLE001
            logger.error("employee membership lock unavailable: request=%s err=%s", request_id, exc)
        record, previous_status = await asyncio.to_thread(
            _review_join_request_with_previous, request_id, action, memo, user
        )
        membership = await sync_employee_tenant_membership(
            record, action, user, previous_status=previous_status, conn=conn
        )
    return {"request": record, "membership": membership}


def update_approved_employee_role(request_id: str, role: str, memo: str, user: dict[str, Any]) -> dict[str, Any]:
    _tenant_id(user)
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="직원 권한 변경 권한이 없습니다")
    access_role = _employee_access_role(role)
    if access_role != str(role or "").strip().lower():
        raise HTTPException(status_code=400, detail="지원하지 않는 직원 권한입니다")
    record = _require_hr_record(_find(_read_hr("employee_join_requests", user), request_id), user, detail="직원을 찾을 수 없습니다")
    if str(record.get("status") or "").strip().lower() != "approved":
        raise HTTPException(status_code=400, detail="승인 완료 직원만 권한을 변경할 수 있습니다")
    now = _now()
    record["role"] = access_role
    record["access_role_label"] = EMPLOYEE_ACCESS_ROLES[access_role]["label"]
    record["access_memo"] = str(memo or "")
    record["access_updated_by"] = _email(user)
    record["access_updated_at"] = now
    record["updated_at"] = now
    _write_hr_record("employee_join_requests", record, user)
    return record


def _record_business_id(record: dict[str, Any]) -> str:
    branch = BRANCH_ALIASES.get(str(record.get("branch") or "").strip(), str(record.get("branch") or "").strip())
    explicit = str(record.get("business_id") or record.get("businessId") or "").strip()
    return explicit if explicit in CANONICAL_BUSINESS_IDS else str(BUSINESS_BY_BRANCH.get(branch) or "")


def _row_in_business(row: dict[str, Any], business_id: str) -> bool:
    """business_id 가 비어 있으면 스코프하지 않는다(종전 동작). 행의 사업자를 못 정하는 레거시 행은 제외하지 않는다."""
    wanted = str(business_id or "").strip()
    if not wanted:
        return True
    row_business = _record_business_id(row)
    if not row_business:
        snapshot = row.get("signed_snapshot") if isinstance(row.get("signed_snapshot"), dict) else {}
        row_business = _record_business_id(snapshot) if snapshot else ""
    return not row_business or row_business == wanted


ONBOARDING_PROFILE_FIELDS = (
    "address",
    "birth_date",
    "nationality",
    "bank_name",
    "bank_account_holder",
    "bank_account_masked",
    "health_certificate_issue_date",
    "health_certificate_valid_until",
)


def _employee_onboarding_profile(
    documents: list[dict[str, Any]],
    *,
    employee_email: str,
    employee_request_id: str,
) -> dict[str, Any]:
    """Return privacy-minimised fields verified from an employee's uploaded documents."""
    email = str(employee_email or "").strip().lower()
    request_id = str(employee_request_id or "").strip()
    matched = [
        row
        for row in documents
        if str(row.get("status") or "").strip().lower() not in {"missing", "superseded"}
        and (
            (request_id and str(row.get("employee_request_id") or "").strip() == request_id)
            or (email and str(row.get("employee_email") or "").strip().lower() == email)
        )
    ]
    profile: dict[str, Any] = {}
    summaries: list[dict[str, str]] = []
    for row in matched:
        extracted = row.get("extracted_fields")
        if not isinstance(extracted, dict):
            extracted = {}
        for field in ONBOARDING_PROFILE_FIELDS:
            value = extracted.get(field)
            if value and not profile.get(field):
                profile[field] = value
        summaries.append(
            {
                "document_type": str(row.get("document_type") or ""),
                "document_label": str(row.get("document_label") or ""),
                "status": str(row.get("status") or "uploaded"),
                "issue_date": str(row.get("issue_date") or ""),
                "valid_until": str(
                    extracted.get("health_certificate_valid_until")
                    or row.get("valid_until")
                    or ""
                ),
            }
        )
    summaries.sort(key=lambda item: (item["document_label"], item["issue_date"]))
    profile["onboarding_documents"] = summaries
    summary_items: list[str] = []
    for item in summaries:
        valid_until = f"~{item['valid_until']}" if item["valid_until"] else ""
        summary_items.append(
            f"{item['document_label'] or item['document_type']}"
            f"({item['issue_date'] or '발급일 미입력'}{valid_until}, {item['status']})"
        )
    profile["onboarding_document_summary"] = ", ".join(summary_items)
    return profile


def list_approved_employees(user: dict[str, Any], business_id: str | None = None) -> list[dict[str, Any]]:
    _tenant_id(user)
    if not _is_admin(user):
        return []
    docs = _read_hr("onboarding_documents", user)
    contracts = _read_hr("contracts", user)
    payroll = _read_hr("payroll_statements", user)

    result = []
    for row in _read_hr("employee_join_requests", user):
        if row.get("status") != "approved":
            continue
        employee_business_id = _record_business_id(row)
        if business_id and employee_business_id != business_id:
            continue
        email = str(row.get("email") or "").strip().lower()
        employee = dict(row)
        employee["business_id"] = employee_business_id
        employee["branch"] = BRANCH_ALIASES.get(str(employee.get("branch") or ""), str(employee.get("branch") or ""))
        employee["email_masked"] = employee.get("email_masked") or _mask_email(email)
        document_profile = _employee_onboarding_profile(
            docs,
            employee_email=email,
            employee_request_id=str(employee.get("id") or ""),
        )
        for field in ONBOARDING_PROFILE_FIELDS:
            if not employee.get(field) and document_profile.get(field):
                employee[field] = document_profile[field]
        employee["onboarding_documents"] = document_profile["onboarding_documents"]
        employee["onboarding_document_summary"] = document_profile["onboarding_document_summary"]
        # 겸직 직원은 매장(사업자)마다 별개 고용이다 — 다른 매장 서류가 이 매장의 "있음"으로 섞이지 않게 집계한다.
        employee["onboarding_document_count"] = sum(
            1
            for item in docs
            if str(item.get("employee_email") or "").strip().lower() == email
            and str(item.get("status") or "").strip().lower() != "superseded"
            and _row_in_business(item, employee_business_id)
        )
        employee["contract_count"] = sum(
            1
            for item in contracts
            if str(item.get("employee_email") or "").strip().lower() == email
            and _row_in_business(item, employee_business_id)
        )
        employee["payroll_statement_count"] = sum(
            1
            for item in payroll
            if str(item.get("employee_email") or "").strip().lower() == email
            and _row_in_business(item, employee_business_id)
        )
        employee["needs_onboarding_documents"] = employee["onboarding_document_count"] == 0
        employee["needs_contract"] = employee["contract_count"] == 0
        employee["needs_payroll"] = employee["payroll_statement_count"] == 0
        derived = _derive_current_employment(
            contracts,
            employee_email=email,
            employee_request_id=str(employee.get("id") or ""),
            business_id=employee_business_id,
        )
        employee["current_employment"] = employee.get("current_employment") or None
        employee["current_employment_contract_id"] = employee.get("current_employment_contract_id") or ""
        employee["needs_employment_sync"] = not _employment_snapshot_is_current(employee, derived)
        result.append(employee)
    return sorted(result, key=lambda row: row.get("reviewed_at") or row.get("updated_at") or row.get("requested_at") or "", reverse=True)


ONBOARDING_DOCUMENT_EXPIRY_WARNING_DAYS = 30
# 보건증·외국인 체류 관련 서류는 만료일이 실무상 필수다. 외국인 고용 신고/변동 확인(foreign_employment_report)은
# 처리일을 적는 서류라 만료 개념이 없어 뺀다.
ONBOARDING_EXPIRY_REQUIRED_TYPES = frozenset(
    {"health_certificate", "foreign_registration", "visa_status_certificate", "work_permission_confirmation"}
)
ONBOARDING_SUPERSEDABLE_STATUSES = ("uploaded", "approved", "needs_fix", "rejected")
ONBOARDING_RESUBMIT_STATUSES = frozenset({"rejected", "needs_fix"})
_DOCUMENT_HEIC_BRANDS = frozenset({b"heic", b"heix", b"hevc", b"hevx", b"heim", b"heis", b"mif1", b"msf1"})
_DOCUMENT_HEAD_BYTES = 16


def _normalize_document_mime(content_type: Any, original_filename: str) -> str:
    guessed = mimetypes.guess_type(original_filename)[0] or "application/octet-stream"
    declared = str(content_type or "").strip().lower()
    if re.fullmatch(r"[a-z0-9.+-]+/[a-z0-9.+-]+", declared) and declared != "application/octet-stream":
        return declared
    return guessed


def _document_signature_ok(suffix: str, head: bytes) -> bool:
    """pdf·이미지는 선두 바이트가 포맷 시그니처여야 한다. 그 밖의 확장자는 확인하지 않는다."""
    if suffix == ".pdf":
        return head.startswith(b"%PDF-")
    if suffix in {".jpg", ".jpeg"}:
        return head.startswith(b"\xff\xd8\xff")
    if suffix == ".png":
        return head.startswith(b"\x89PNG")
    if suffix == ".webp":
        return head[:4] == b"RIFF" and head[8:12] == b"WEBP"
    if suffix in {".tif", ".tiff"}:
        return head[:4] in {b"II*\x00", b"MM\x00*", b"II+\x00", b"MM\x00+"}
    if suffix == ".heic":
        return head[4:8] == b"ftyp" and head[8:12] in _DOCUMENT_HEIC_BRANDS
    return True


def _document_expiry_fields(expires_at: Any, warning_days: int, today: date | None = None) -> dict[str, Any]:
    expiry_status, expiry_label, days_left = "", "", None
    expires = str(expires_at or "")[:10]
    if expires:
        try:
            days_left = (date.fromisoformat(expires) - (today or datetime.now(KST).date())).days
        except ValueError:
            days_left = None
        if days_left is not None and days_left < 0:
            expiry_status, expiry_label = "expired", "만료"
        elif days_left is not None and days_left <= warning_days:
            expiry_status, expiry_label = "expiring", "만료 임박"
    return {
        "expiry_status": expiry_status,
        "expiry_label": expiry_label,
        "days_until_expiry": days_left,
        "is_expired": expiry_status == "expired",
        "is_expiring": expiry_status == "expiring",
    }


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _onboarding_document_view(record: dict[str, Any], today: date | None = None) -> dict[str, Any]:
    row = dict(record)
    row.pop("stored_path", None)
    row["expires_at"] = str(row.get("expires_at") or "")[:10]
    row.update(_document_expiry_fields(row["expires_at"], ONBOARDING_DOCUMENT_EXPIRY_WARNING_DAYS, today))
    if str(row.get("status") or "").strip().lower() == "superseded":
        # 대체된 이전본의 만료는 더 이상 조치 대상이 아니다.
        row.update(expiry_status="", expiry_label="", is_expired=False, is_expiring=False)
    return row


def _onboarding_document_path(record: dict[str, Any]) -> Path | None:
    relative = str(record.get("stored_path") or record.get("stored_filename") or "")
    if not relative:
        return None
    root = UPLOAD_DIR.resolve()
    path = (root / relative).resolve()
    if root not in path.parents or not path.is_file():
        return None
    return path


def _admin_business_ids(user: dict[str, Any]) -> set[str] | None:
    """None 이면 테넌트의 모든 사업장. 승인된 직원 레코드로 관리자가 된 사용자는 그 사업장만 본다."""
    user_role = str(user.get("user_role") or "").strip().lower()
    email = _email(user)
    if bool(user.get("is_internal_admin")) or user_role in {"ceo", "admin", "system"} or not email:
        return None
    own = [
        row
        for row in _read_hr("employee_join_requests", user)
        if str(row.get("email") or "").strip().lower() == email
        and str(row.get("status") or "").strip().lower() == "approved"
    ]
    if not own:
        return None
    return {
        _record_business_id(row)
        for row in own
        if _employee_access_role(row.get("role")) == "admin" and _record_business_id(row)
    }


def _require_onboarding_business_scope(record: dict[str, Any], user: dict[str, Any]) -> None:
    """관리자가 자기 사업장 밖의 입사서류를 열람·검수·삭제하지 못하게 한다. 직원 본인 접근은 호출부가 막는다."""
    if not _is_admin(user):
        return
    allowed = _admin_business_ids(user)
    if allowed is not None and _record_business_id(record) not in allowed:
        raise HTTPException(status_code=403, detail="해당 사업장의 서류에 접근할 수 없습니다")


async def _db_supersede_onboarding_documents(record: dict[str, Any]) -> int:
    import asyncpg

    conn = await asyncpg.connect(_db_url(), timeout=5)
    try:
        result = await conn.execute(
            """
            UPDATE yeoljeong_onboarding_documents
               SET status = 'superseded', superseded_by = $1, updated_at = $2::timestamptz,
                   metadata = metadata || jsonb_build_object('status', 'superseded', 'superseded_by', $1::text, 'updated_at', $3::text)
             WHERE tenant_id = $4::uuid AND business_id = $5 AND lower(employee_email) = $6
               AND document_type = $7 AND id <> $1 AND deleted_at IS NULL
               AND status = ANY($8::text[])
               AND (uploaded_at, id) < ($9::timestamptz, $1)
            """,
            str(record["id"]),
            _pg_ts(record.get("updated_at")),
            str(record.get("updated_at") or ""),
            UUID(str(record["tenant_id"])),
            str(record.get("business_id") or ""),
            str(record.get("employee_email") or "").strip().lower(),
            str(record.get("document_type") or ""),
            list(ONBOARDING_SUPERSEDABLE_STATUSES),
            _pg_ts(record.get("uploaded_at")),
        )
        return int(result.rsplit(" ", 1)[-1])
    finally:
        await conn.close()


def _supersede_onboarding_file_rows(rows: list[dict[str, Any]], record: dict[str, Any]) -> list[str]:
    superseded: list[str] = []
    wanted_email = str(record.get("employee_email") or "").strip().lower()
    for row in rows:
        if (
            row.get("deleted_at")
            or str(row.get("id")) == str(record["id"])
            or str(row.get("tenant_id") or "") != str(record["tenant_id"])
            or _record_business_id(row) != str(record.get("business_id") or "")
            or str(row.get("employee_email") or "").strip().lower() != wanted_email
            or str(row.get("document_type") or "") != str(record.get("document_type") or "")
            or str(row.get("status") or "uploaded").strip().lower() not in ONBOARDING_SUPERSEDABLE_STATUSES
        ):
            continue
        row["status"] = "superseded"
        row["superseded_by"] = record["id"]
        row["updated_at"] = record["updated_at"]
        superseded.append(str(row.get("id")))
    return superseded


async def save_onboarding_document(
    *,
    employee_name: str,
    employee_email: str,
    branch: str,
    document_type: str,
    issue_date: str,
    memo: str,
    upload: UploadFile,
    user: dict[str, Any],
    expires_at: str = "",
) -> dict[str, Any]:
    tenant_id = _tenant_id(user)
    email = str(employee_email or _email(user)).strip().lower()
    if not email:
        raise HTTPException(status_code=400, detail="직원 이메일이 필요합니다")
    if not _is_admin(user) and email != _email(user):
        raise HTTPException(status_code=403, detail="본인 서류만 업로드할 수 있습니다")
    employee_rows = (
        await _db_fetch_ledger("employee_join_requests", tenant_id)
        if _db_available()
        else _read_hr("employee_join_requests", user)
    )
    employee = next(
        (
            row
            for row in (employee_rows or [])
            if str(row.get("email") or "").strip().lower() == email
            and str(row.get("status") or "pending").strip().lower() != "rejected"
        ),
        None,
    )
    employee_branch = str((employee or {}).get("branch") or "").strip()
    normalized_branch = BRANCH_ALIASES.get(employee_branch or branch.strip(), employee_branch or branch.strip())
    business_id = _record_business_id(employee or {"branch": normalized_branch})
    if not business_id or (_db_available() and not await _db_business_tenant_matches(business_id, tenant_id)):
        raise HTTPException(status_code=403, detail="테넌트에 귀속되지 않은 사업자입니다")
    meta = _document_meta(document_type)
    issued, expires = _business_document_dates(issue_date, expires_at)
    if str(document_type or "").strip() in ONBOARDING_EXPIRY_REQUIRED_TYPES and not expires:
        raise HTTPException(status_code=400, detail=f"{meta['label']}은(는) 만료일을 입력해야 합니다")
    original = _safe_filename(upload.filename or "document.bin")
    suffix = Path(original).suffix.lower()
    if suffix not in ONBOARDING_DOCUMENT_EXTENSIONS:
        raise HTTPException(status_code=400, detail="PDF·이미지·한글/워드 문서만 등록할 수 있습니다")
    mime = _normalize_document_mime(upload.content_type, original)
    doc_id = str(uuid4())
    stored_name = f"{doc_id}{suffix}"
    destination = UPLOAD_DIR / stored_name
    _ensure_dirs()
    tmp = destination.with_name(f"{stored_name}.{secrets.token_hex(4)}.tmp")
    digest = hashlib.sha256()
    size = 0
    head = b""
    try:
        with tmp.open("wb") as out:
            while True:
                chunk = await upload.read(1024 * 1024)
                if not chunk:
                    break
                size += len(chunk)
                if size > MAX_UPLOAD_BYTES:
                    raise HTTPException(status_code=413, detail="파일은 최대 15MB까지 업로드할 수 있습니다")
                if len(head) < _DOCUMENT_HEAD_BYTES:
                    head = (head + chunk)[:_DOCUMENT_HEAD_BYTES]
                digest.update(chunk)
                out.write(chunk)
        if size == 0:
            raise HTTPException(status_code=400, detail="빈 파일은 등록할 수 없습니다")
        if not _document_signature_ok(suffix, head):
            raise HTTPException(status_code=400, detail="파일 내용이 확장자와 일치하지 않습니다")
        os.chmod(tmp, 0o600)
        tmp.replace(destination)
    except BaseException:
        tmp.unlink(missing_ok=True)
        raise
    now = _now()
    record = {
        "id": doc_id,
        "employee_request_id": str((employee or {}).get("id") or ""),
        "employee_name": employee_name.strip() or str((employee or {}).get("name") or ""),
        "employee_email": email,
        "employee_email_masked": _mask_email(email),
        "business_id": business_id,
        "branch": normalized_branch,
        "document_type": document_type,
        "document_label": meta["label"],
        "requirement": meta["requirement"],
        "issue_date": issued,
        "expires_at": expires,
        "memo": memo,
        "original_filename": original,
        "stored_filename": stored_name,
        "stored_path": stored_name,
        "sha256": digest.hexdigest(),
        "superseded_by": "",
        "content_type": mime,
        "size_bytes": size,
        "status": "uploaded",
        "uploaded_by": _email(user),
        "uploaded_at": now,
        "updated_at": now,
        "tenant_id": tenant_id,
    }
    try:
        if _db_available():
            if not await _db_upsert_ledger("onboarding_documents", record):
                raise HTTPException(status_code=404, detail="입사서류를 저장할 수 없습니다")
            try:
                await _db_supersede_onboarding_documents(record)
            except Exception:
                # 새 서류는 이미 저장됐다. 이전본 정리만 실패한 것이므로 다음 업로드가 다시 정리한다.
                logger.exception("onboarding document supersede failed: document=%s", doc_id)
        else:
            rows = _read_file_rows("onboarding_documents")
            _supersede_onboarding_file_rows(rows, record)
            rows.insert(0, record)
            _write_file_rows("onboarding_documents", rows)
    except Exception:
        # 메타를 못 남긴 원본은 고아가 된다 — 지우고 실패를 그대로 올린다.
        destination.unlink(missing_ok=True)
        raise
    return _onboarding_document_view(record)


def _evidence_kind_label(kind: str) -> str:
    return {
        "bank_statement": "은행 거래내역",
        "supplier_statement": "매입처 거래내역서",
        "receipt_photo": "영수증 사진",
        "tax_invoice": "세금계산서",
        "utility_bill": "공과금 고지서",
        "card_pg_report": "카드/PG 리포트",
        "other": "기타 증빙",
    }.get(kind, kind or "기타 증빙")


def list_integration_evidence(user: dict[str, Any], business_id: str | None = None) -> list[dict[str, Any]]:
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="연동 증빙 조회 권한이 없습니다")
    rows = _read("integration_evidence")
    if business_id:
        rows = [row for row in rows if str(row.get("business_id") or "") == business_id]
    return sorted(rows, key=lambda row: str(row.get("uploaded_at") or row.get("created_at") or ""), reverse=True)


async def save_integration_evidence(
    *,
    service: str,
    business_id: str,
    branch: str,
    document_kind: str,
    vendor: str,
    amount: int,
    memo: str,
    upload: UploadFile,
    user: dict[str, Any],
) -> dict[str, Any]:
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="연동 증빙 업로드 권한이 없습니다")
    normalized_service = str(service or "").strip()
    if normalized_service not in CONNECTOR_LABELS:
        raise HTTPException(status_code=400, detail="지원하지 않는 연동 서비스입니다")
    normalized_business, normalized_branch = _normalize_connector_scope(normalized_service, business_id, branch)
    original = _safe_filename(upload.filename or "evidence.bin")
    suffix = Path(original).suffix.lower()
    evidence_id = str(uuid4())
    stored_name = f"{evidence_id}{suffix or '.bin'}"
    destination = _evidence_upload_dir() / stored_name
    _ensure_dirs()
    size = 0
    with destination.open("wb") as out:
        while True:
            chunk = await upload.read(1024 * 1024)
            if not chunk:
                break
            size += len(chunk)
            if size > MAX_UPLOAD_BYTES:
                destination.unlink(missing_ok=True)
                raise HTTPException(status_code=413, detail="파일은 최대 15MB까지 업로드할 수 있습니다")
            out.write(chunk)
    now = _now()
    content_type = upload.content_type or mimetypes.guess_type(original)[0] or "application/octet-stream"
    record = {
        "id": evidence_id,
        "service": normalized_service,
        "service_label": CONNECTOR_LABELS.get(normalized_service, normalized_service),
        "business_id": normalized_business,
        "branch": normalized_branch,
        "document_kind": document_kind or "other",
        "document_label": _evidence_kind_label(document_kind or "other"),
        "vendor": vendor.strip() or CONNECTOR_LABELS.get(normalized_service, normalized_service),
        "amount": int(amount or 0),
        "memo": memo,
        "original_filename": original,
        "stored_filename": stored_name,
        "content_type": content_type,
        "size_bytes": size,
        "status": "pending_review",
        "uploaded_by": _email(user),
        "uploaded_at": now,
        "created_at": now,
        "updated_at": now,
    }
    evidence_rows = _read("integration_evidence")
    evidence_rows.insert(0, record)
    _write("integration_evidence", evidence_rows)

    transaction = create_transaction(
        {
            "transaction_date": now[:10],
            "source_type": "integration_evidence",
            "source_file": original,
            "description": f"{record['service_label']} {record['document_label']} 업로드",
            "amount": record["amount"],
            "direction": "expense" if normalized_service not in PLATFORM_LABELS else "income",
            "category": _transaction_category(" ".join([record["service_label"], record["document_label"], record["vendor"], memo])),
            "approval_number": "",
            "order_number": "",
            "account_name": record["vendor"],
            "business_id": normalized_business,
            "branch": normalized_branch,
            "evidence_id": evidence_id,
            "status": "pending",
            "memo": memo or "업로드 증빙 확인 필요",
        }
    )
    return {"evidence": record, "transaction": transaction}


def get_integration_evidence(evidence_id: str, user: dict[str, Any]) -> tuple[dict[str, Any], Path]:
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="연동 증빙 다운로드 권한이 없습니다")
    document = _find(_read("integration_evidence"), evidence_id)
    if not document:
        raise HTTPException(status_code=404, detail="연동 증빙을 찾을 수 없습니다")
    stored = _safe_filename(str(document.get("stored_filename") or ""))
    path = _evidence_upload_dir() / stored
    if not path.exists():
        raise HTTPException(status_code=404, detail="연동 증빙 파일이 없습니다")
    return document, path


def _onboarding_missing_document_rows(
    *,
    existing_rows: list[dict[str, Any]],
    user: dict[str, Any],
    business_id: str | None = None,
) -> list[dict[str, Any]]:
    email_filter = "" if _is_admin(user) else _email(user)
    live_rows = [row for row in existing_rows if not row.get("deleted_at")]
    # 반려·보완요청·대체된 행은 "제출됨"이 아니다. 제출로 치는 것은 검수를 통과했거나 검수 대기 중인 행뿐이다.
    existing_keys = {
        (
            str(row.get("employee_email") or "").strip().lower(),
            str(row.get("document_type") or "").strip(),
        )
        for row in live_rows
        if str(row.get("status") or "uploaded").strip().lower() not in ONBOARDING_RESUBMIT_STATUSES | {"superseded"}
    }
    resubmit_rows: dict[tuple[str, str], dict[str, Any]] = {}
    for row in sorted(live_rows, key=lambda item: str(item.get("uploaded_at") or ""), reverse=True):
        if str(row.get("status") or "").strip().lower() in ONBOARDING_RESUBMIT_STATUSES:
            resubmit_rows.setdefault(
                (str(row.get("employee_email") or "").strip().lower(), str(row.get("document_type") or "").strip()),
                row,
            )
    rows: list[dict[str, Any]] = []
    admin_view = _is_admin(user)
    for employee in _read_hr("employee_join_requests", user):
        status = str(employee.get("status") or "pending").strip().lower()
        if status == "rejected":
            continue
        if admin_view and status != "approved":
            continue
        employee_email = str(employee.get("email") or "").strip().lower()
        if not employee_email:
            continue
        if email_filter and employee_email != email_filter:
            continue
        employee_business_id = _record_business_id(employee)
        if business_id and employee_business_id != business_id:
            continue
        employee_id = str(employee.get("id") or employee_email)
        for meta in _required_document_types():
            document_type = meta["type"]
            if (employee_email, document_type) in existing_keys:
                continue
            previous = resubmit_rows.get((employee_email, document_type))
            rows.append(
                {
                    "id": f"missing-{employee_id}-{document_type}",
                    "employee_request_id": employee_id,
                    "employee_name": employee.get("name") or "",
                    "employee_email": employee_email,
                    "employee_email_masked": employee.get("email_masked") or _mask_email(employee_email),
                    "business_id": employee_business_id,
                    "branch": BRANCH_ALIASES.get(str(employee.get("branch") or ""), str(employee.get("branch") or "")),
                    "document_type": document_type,
                    "document_label": meta["label"],
                    "requirement": meta["requirement"],
                    "status": "resubmit_required" if previous else "missing",
                    "status_label": "재제출 필요" if previous else "작성 필요",
                    "previous_document_id": str(previous.get("id") or "") if previous else "",
                    "previous_status": str(previous.get("status") or "") if previous else "",
                    "review_memo": str(previous.get("review_memo") or "") if previous else "",
                    "original_filename": "",
                    "size_bytes": 0,
                    "uploaded_at": "",
                    "updated_at": employee.get("reviewed_at") or employee.get("updated_at") or employee.get("requested_at") or "",
                    "is_placeholder": True,
                    "missing_document": True,
                    "employee_request_status": status,
                    **_document_expiry_fields("", ONBOARDING_DOCUMENT_EXPIRY_WARNING_DAYS),
                    "expires_at": "",
                }
            )
    return rows


def list_onboarding_documents(user: dict[str, Any], business_id: str | None = None) -> list[dict[str, Any]]:
    stored_rows = _read_hr("onboarding_documents", user)
    visible_rows = _filter_user(stored_rows, user, "employee_email", "uploaded_by")
    normalized_rows = []
    for item in visible_rows:
        row = dict(item)
        row["business_id"] = _record_business_id(row)
        row["branch"] = BRANCH_ALIASES.get(str(row.get("branch") or ""), str(row.get("branch") or ""))
        if business_id and _is_admin(user) and row["business_id"] != business_id:
            continue
        normalized_rows.append(_onboarding_document_view(row))
    rows = normalized_rows + _onboarding_missing_document_rows(
        existing_rows=stored_rows,
        user=user,
        business_id=business_id if _is_admin(user) else None,
    )
    return sorted(rows, key=lambda row: row.get("uploaded_at") or row.get("updated_at") or "", reverse=True)


def get_onboarding_document(document_id: str, user: dict[str, Any]) -> tuple[dict[str, Any], Path]:
    record = _require_hr_record(_find(_read_hr("onboarding_documents", user), document_id), user, detail="입사서류를 찾을 수 없습니다")
    _require_onboarding_business_scope(record, user)
    if not _is_admin(user) and str(record.get("employee_email") or "").strip().lower() != _email(user):
        raise HTTPException(status_code=403, detail="본인 서류만 열람할 수 있습니다")
    path = _onboarding_document_path(record)
    if path is None:
        raise HTTPException(status_code=404, detail="업로드 파일이 없습니다")
    expected_sha256 = str(record.get("sha256") or "")
    if expected_sha256 and _file_sha256(path) != expected_sha256:
        logger.error("onboarding document hash mismatch: document=%s", record.get("id"))
        raise HTTPException(status_code=409, detail="입사서류 원본이 등록 당시와 다릅니다")
    return record, path


def review_onboarding_document(document_id: str, status: str, memo: str, user: dict[str, Any]) -> dict[str, Any]:
    _tenant_id(user)
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="서류 검수 권한이 없습니다")
    record = _require_hr_record(_find(_read_hr("onboarding_documents", user), document_id), user, detail="입사서류를 찾을 수 없습니다")
    _require_onboarding_business_scope(record, user)
    if status not in {"approved", "rejected", "needs_fix", "uploaded"}:
        raise HTTPException(status_code=400, detail="올바르지 않은 서류 상태입니다")
    if str(record.get("status") or "").strip().lower() == "superseded":
        raise HTTPException(status_code=409, detail="재제출로 대체된 이전본은 검수할 수 없습니다")
    record["status"] = status
    record["review_memo"] = memo
    record["reviewed_by"] = _email(user)
    record["reviewed_at"] = _now()
    record["updated_at"] = record["reviewed_at"]
    _write_hr_record("onboarding_documents", record, user)
    return _onboarding_document_view(record)


def delete_onboarding_document(document_id: str, user: dict[str, Any]) -> None:
    record = _require_hr_record(_find(_read_hr("onboarding_documents", user), document_id), user, detail="입사서류를 찾을 수 없습니다")
    _require_onboarding_business_scope(record, user)
    if not _is_admin(user) and str(record.get("employee_email") or "").strip().lower() != _email(user):
        raise HTTPException(status_code=403, detail="본인 서류만 삭제할 수 있습니다")
    _delete_hr_record("onboarding_documents", document_id, user)


# --- 사업자 서류 (AADS-OBYS-BUSINESS-DOCUMENTS-20260930) ----------------------
# 사업자(법인/개인) 단위 서류. 원본은 디스크(OBYS_UPLOAD_ROOT)에, DB 에는 경로·해시만 둔다.
# 같은 종류를 다시 올리면 이전 건을 지우지 않고 superseded 로 남긴다 — 세무·노무 분쟁의 근거다.
# 서류 종류 목록은 여기가 원본이다. 화면은 목록 응답의 document_types 를 받아 쓴다.
BUSINESS_DOCUMENT_TYPES: list[dict[str, str]] = [
    {"type": "business_registration", "label": "사업자등록증", "requirement": "필수", "notice": "세무서 발급본. 정정 발급 시 새로 올립니다."},
    {"type": "business_permit", "label": "영업신고증", "requirement": "선택", "notice": "구청 위생과 발급본입니다."},
    {"type": "bankbook", "label": "통장사본", "requirement": "선택", "notice": "사업용 계좌 확인용입니다."},
    {"type": "lease_contract", "label": "임대차계약서", "requirement": "선택", "notice": "계약 만료일을 만료일에 적습니다."},
    {"type": "hygiene_training", "label": "위생교육수료증", "requirement": "선택", "notice": "영업자 위생교육 수료증입니다."},
    {"type": "fire_insurance", "label": "화재보험증서", "requirement": "선택", "notice": "보험 만기일을 만료일에 적습니다."},
    {"type": "corporate_registry", "label": "법인등기부등본", "requirement": "법인 조건부", "notice": "법인 사업자만 해당합니다."},
    {"type": "seal_certificate", "label": "인감증명서", "requirement": "선택", "notice": "발급일로부터 3개월 이내본을 권장합니다."},
    {"type": "representative_id", "label": "대표자신분증", "requirement": "선택", "notice": "주민등록번호 뒷자리는 마스킹합니다."},
    {"type": "tax_agent_delegation", "label": "세무대리인위임장", "requirement": "선택", "notice": "세무대리인 수임 동의 서류입니다."},
    {"type": "mail_order_report", "label": "통신판매업신고증", "requirement": "선택", "notice": "온라인 판매 시 해당합니다."},
    {"type": "other", "label": "기타", "requirement": "선택", "notice": "서류 이름을 메모에 남깁니다."},
]
BUSINESS_DOCUMENT_TYPE_CODES = {item["type"] for item in BUSINESS_DOCUMENT_TYPES}
BUSINESS_DOCUMENT_REQUIRED_TYPE = "business_registration"
BUSINESS_DOCUMENT_EXPIRY_WARNING_DAYS = 30
BUSINESS_DOCUMENT_EXTENSIONS = {".pdf", ".jpg", ".jpeg", ".png", ".webp", ".heic", ".tif", ".tiff", ".hwp", ".doc", ".docx"}
ONBOARDING_DOCUMENT_EXTENSIONS = set(BUSINESS_DOCUMENT_EXTENSIONS)
BUSINESS_DOCUMENT_MEMO_MAX = 1000
BUSINESS_DOCUMENT_TABLE = "yeoljeong_business_documents"


def list_business_document_types() -> list[dict[str, str]]:
    return BUSINESS_DOCUMENT_TYPES


def _business_document_meta(document_type: str) -> dict[str, str]:
    code = str(document_type or "").strip()
    meta = next((item for item in BUSINESS_DOCUMENT_TYPES if item["type"] == code), None)
    if not meta:
        raise HTTPException(status_code=400, detail="지원하지 않는 서류 종류입니다")
    return meta


def _business_document_date(value: Any, label: str) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    try:
        return date.fromisoformat(text[:10]).isoformat()
    except ValueError:
        raise HTTPException(status_code=400, detail=f"{label}은 YYYY-MM-DD 형식이어야 합니다") from None


def _business_document_dates(issue_date: Any, expires_at: Any) -> tuple[str, str]:
    issued = _business_document_date(issue_date, "발급일")
    expires = _business_document_date(expires_at, "만료일")
    if issued and expires and expires < issued:
        raise HTTPException(status_code=400, detail="만료일이 발급일보다 빠릅니다")
    return issued, expires


def _business_document_memo(value: Any) -> str:
    memo = str(value or "").strip()
    if len(memo) > BUSINESS_DOCUMENT_MEMO_MAX:
        raise HTTPException(status_code=400, detail=f"메모는 {BUSINESS_DOCUMENT_MEMO_MAX}자 이하여야 합니다")
    return memo


def _business_document_scope(business_id: Any, user: dict[str, Any]) -> tuple[str, str]:
    """관리자 + 테넌트 + 사업자 소속을 코드에서 강제한다. SQL WHERE 에도 같은 조건이 들어간다."""
    tenant_id = _tenant_id(user)
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="사업자 서류는 관리자만 다룰 수 있습니다")
    normalized = str(business_id or "").strip()
    if not normalized:
        raise HTTPException(status_code=400, detail="사업자를 선택하십시오")
    if _db_available():
        matched = _run_db(_db_business_tenant_matches(normalized, tenant_id))
        if matched is None:
            raise HTTPException(status_code=503, detail="사업자 소속을 확인할 수 없습니다")
        if not matched:
            raise HTTPException(status_code=403, detail="테넌트에 귀속되지 않은 사업자입니다")
        return tenant_id, normalized
    known = {str(item.get("id") or "") for item in get_settings(user)["settings"].get("businesses") or []}
    if normalized not in known:
        raise HTTPException(status_code=403, detail="테넌트에 귀속되지 않은 사업자입니다")
    return tenant_id, normalized


def _business_document_relpath(tenant_id: str, document_id: str, suffix: str) -> Path:
    # 계약서 PDF 와 같은 루트(OBYS_UPLOAD_ROOT)·같은 <tenant>/<종류>/ 배치, 파일명은 입사서류처럼 <id><ext>.
    return Path(str(UUID(tenant_id))) / "business_documents" / f"{document_id}{suffix}"


def _business_document_file(record: dict[str, Any]) -> Path | None:
    relative = str(record.get("stored_path") or "")
    if not relative:
        return None
    root = _contract_pdf_root().resolve()
    path = (root / relative).resolve()
    if root not in path.parents or not path.is_file():
        return None
    return path


def _business_document_view(record: dict[str, Any], today: date | None = None) -> dict[str, Any]:
    row = dict(record)
    row.pop("stored_path", None)
    meta = next((item for item in BUSINESS_DOCUMENT_TYPES if item["type"] == row.get("document_type")), None)
    row["document_label"] = (meta or {}).get("label") or row.get("document_label") or "기타"
    row["requirement"] = (meta or {}).get("requirement") or "선택"
    row["status_label"] = {"current": "최신", "superseded": "이전본", "deleted": "삭제"}.get(str(row.get("status") or ""), "")
    row.update(_document_expiry_fields(row.get("expires_at"), BUSINESS_DOCUMENT_EXPIRY_WARNING_DAYS, today))
    return row


def _business_document_restatus(rows: list[dict[str, Any]], tenant_id: str, business_id: str, document_types: set[str]) -> None:
    """(사업자, 종류)마다 삭제되지 않은 최신 1건만 current, 나머지는 superseded."""
    for document_type in document_types:
        group = [
            row for row in rows
            if str(row.get("tenant_id") or "") == tenant_id
            and str(row.get("business_id") or "") == business_id
            and str(row.get("document_type") or "") == document_type
            and not row.get("deleted_at")
        ]
        group.sort(key=lambda row: (str(row.get("uploaded_at") or ""), str(row.get("id") or "")), reverse=True)
        for index, row in enumerate(group):
            row["status"] = "current" if index == 0 else "superseded"


_BUSINESS_DOCUMENT_COLUMNS = (
    "id, tenant_id, business_id, document_type, document_label, original_filename, content_type, stored_path, "
    "sha256, byte_size, issue_date, expires_at, memo, status, uploaded_by, uploaded_at, updated_at, deleted_at"
)


def _db_business_document_record(row: Any) -> dict[str, Any]:
    item = dict(row)
    return {
        **item,
        "id": str(item.get("id") or ""),
        "tenant_id": str(item.get("tenant_id") or ""),
        "issue_date": item["issue_date"].isoformat() if item.get("issue_date") else "",
        "expires_at": item["expires_at"].isoformat() if item.get("expires_at") else "",
        "byte_size": int(item.get("byte_size") or 0),
        "uploaded_at": _iso(item.get("uploaded_at")),
        "updated_at": _iso(item.get("updated_at")),
        "deleted_at": _iso(item.get("deleted_at")) if item.get("deleted_at") else "",
    }


async def _db_business_documents_restatus(conn: Any, tenant_id: UUID, business_id: str, document_type: str) -> None:
    # 두 단계로 나눈다. 한 문장으로 current 를 옮기면 부분 유일 인덱스가 행 단위로 검사돼 중간에 충돌한다.
    await conn.execute(
        f"UPDATE {BUSINESS_DOCUMENT_TABLE} SET status = 'superseded' "
        "WHERE tenant_id = $1 AND business_id = $2 AND document_type = $3 AND deleted_at IS NULL AND status <> 'superseded'",
        tenant_id, business_id, document_type,
    )
    await conn.execute(
        f"UPDATE {BUSINESS_DOCUMENT_TABLE} SET status = 'current' WHERE id = ("
        f"SELECT l.id FROM {BUSINESS_DOCUMENT_TABLE} l "
        "WHERE l.tenant_id = $1 AND l.business_id = $2 AND l.document_type = $3 AND l.deleted_at IS NULL "
        "ORDER BY l.uploaded_at DESC, l.id DESC LIMIT 1)",
        tenant_id, business_id, document_type,
    )


async def _db_business_documents_fetch(tenant_id: str, business_id: str | None) -> list[dict[str, Any]] | None:
    import asyncpg

    conn = await asyncpg.connect(_db_url(), timeout=5)
    try:
        if not await conn.fetchval("SELECT to_regclass($1) IS NOT NULL", f"public.{BUSINESS_DOCUMENT_TABLE}"):
            return None
        rows = await conn.fetch(
            f"SELECT {_BUSINESS_DOCUMENT_COLUMNS} FROM {BUSINESS_DOCUMENT_TABLE} d "
            "WHERE d.tenant_id = $1 AND d.deleted_at IS NULL AND ($2::text IS NULL OR d.business_id = $2) "
            "AND EXISTS (SELECT 1 FROM yeoljeong_business_tenant_mapping m "
            "WHERE m.business_id = d.business_id AND m.tenant_id = d.tenant_id) "
            "ORDER BY d.uploaded_at DESC, d.id DESC",
            UUID(tenant_id), business_id,
        )
        return [_db_business_document_record(row) for row in rows]
    finally:
        await conn.close()


async def _db_business_document_get(tenant_id: str, document_id: str) -> dict[str, Any] | None | bool:
    """테넌트 안의 행을 돌려준다. 없으면 None, 다른 테넌트 행이면 False."""
    import asyncpg

    conn = await asyncpg.connect(_db_url(), timeout=5)
    try:
        row = await conn.fetchrow(
            f"SELECT {_BUSINESS_DOCUMENT_COLUMNS} FROM {BUSINESS_DOCUMENT_TABLE} WHERE id = $1 AND deleted_at IS NULL",
            UUID(document_id),
        )
        if not row:
            return None
        if str(row["tenant_id"]) != tenant_id:
            return False
        return _db_business_document_record(row)
    finally:
        await conn.close()


async def _db_business_document_insert(record: dict[str, Any]) -> bool:
    import asyncpg

    tenant_id = UUID(record["tenant_id"])
    conn = await asyncpg.connect(_db_url(), timeout=5)
    try:
        async with conn.transaction():
            # 같은 사업자에 대한 동시 업로드가 current 를 둘 만들지 않게 직렬화한다.
            await conn.execute("SELECT pg_advisory_xact_lock(hashtext($1))", f"{tenant_id}|{record['business_id']}")
            await conn.execute(
                f"INSERT INTO {BUSINESS_DOCUMENT_TABLE} ({_BUSINESS_DOCUMENT_COLUMNS}) VALUES "
                "($1, $2, $3, $4, $5, $6, $7, $8, $9, $10, $11::date, $12::date, $13, $14, $15, "
                "$16::timestamptz, $17::timestamptz, NULL)",
                UUID(record["id"]), tenant_id, record["business_id"], record["document_type"], record["document_label"],
                record["original_filename"], record["content_type"], record["stored_path"], record["sha256"],
                int(record["byte_size"]), _pg_date(record.get("issue_date")), _pg_date(record.get("expires_at")),
                record["memo"], "superseded", record["uploaded_by"],
                _pg_ts(record["uploaded_at"]), _pg_ts(record["updated_at"]),
            )
            await _db_business_documents_restatus(conn, tenant_id, record["business_id"], record["document_type"])
        return True
    finally:
        await conn.close()


async def _db_business_document_update(record: dict[str, Any], previous_type: str) -> bool:
    import asyncpg

    tenant_id = UUID(record["tenant_id"])
    conn = await asyncpg.connect(_db_url(), timeout=5)
    try:
        async with conn.transaction():
            await conn.execute("SELECT pg_advisory_xact_lock(hashtext($1))", f"{tenant_id}|{record['business_id']}")
            # 파일 관련 열(stored_path·sha256·byte_size·original_filename)은 여기서 바꾸지 않는다.
            result = await conn.execute(
                f"UPDATE {BUSINESS_DOCUMENT_TABLE} SET document_type = $4, document_label = $5, issue_date = $6::date, "
                "expires_at = $7::date, memo = $8, updated_at = $9::timestamptz, status = 'superseded' "
                "WHERE id = $1 AND tenant_id = $2 AND business_id = $3 AND deleted_at IS NULL",
                UUID(record["id"]), tenant_id, record["business_id"], record["document_type"], record["document_label"],
                _pg_date(record.get("issue_date")), _pg_date(record.get("expires_at")), record["memo"],
                _pg_ts(record["updated_at"]),
            )
            if not result.endswith(" 1"):
                return False
            # 위에서 잠시 superseded 로 내렸다. 두 종류 모두 최신 1건을 다시 current 로 세운다.
            for document_type in {previous_type, record["document_type"]}:
                await _db_business_documents_restatus(conn, tenant_id, record["business_id"], document_type)
        return True
    finally:
        await conn.close()


async def _db_business_document_soft_delete(record: dict[str, Any], deleted_at: str) -> bool:
    import asyncpg

    tenant_id = UUID(record["tenant_id"])
    conn = await asyncpg.connect(_db_url(), timeout=5)
    try:
        async with conn.transaction():
            await conn.execute("SELECT pg_advisory_xact_lock(hashtext($1))", f"{tenant_id}|{record['business_id']}")
            result = await conn.execute(
                f"UPDATE {BUSINESS_DOCUMENT_TABLE} SET deleted_at = $4::timestamptz, updated_at = $4::timestamptz, "
                "status = 'deleted' WHERE id = $1 AND tenant_id = $2 AND business_id = $3 AND deleted_at IS NULL",
                UUID(record["id"]), tenant_id, record["business_id"], _pg_ts(deleted_at),
            )
            if not result.endswith(" 1"):
                return False
            # 현행본을 지우면 바로 앞 리비전이 다시 현행이 된다.
            await _db_business_documents_restatus(conn, tenant_id, record["business_id"], record["document_type"])
        return True
    finally:
        await conn.close()


def _pg_date(value: Any) -> date | None:
    text = str(value or "").strip()
    return date.fromisoformat(text[:10]) if text else None


def _business_document_rows(tenant_id: str, business_id: str | None) -> list[dict[str, Any]]:
    if _db_available():
        rows = _run_db(_db_business_documents_fetch(tenant_id, business_id))
        if rows is None:
            raise HTTPException(status_code=503, detail="사업자 서류 저장소를 읽을 수 없습니다")
        return rows
    return [
        row for row in _read_file_rows("business_documents")
        if str(row.get("tenant_id") or "") == tenant_id
        and not row.get("deleted_at")
        and (business_id is None or str(row.get("business_id") or "") == business_id)
    ]


def _business_document_record(document_id: str, user: dict[str, Any]) -> dict[str, Any]:
    """문서 1건을 찾아 테넌트·관리자·사업자 소속을 다시 확인한다."""
    tenant_id = _tenant_id(user)
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="사업자 서류는 관리자만 다룰 수 있습니다")
    try:
        normalized_id = str(UUID(str(document_id)))
    except (TypeError, ValueError, AttributeError):
        raise HTTPException(status_code=404, detail="사업자 서류를 찾을 수 없습니다") from None
    if _db_available():
        found = _run_db(_db_business_document_get(tenant_id, normalized_id))
    else:
        row = next(
            (item for item in _read_file_rows("business_documents")
             if str(item.get("id") or "") == normalized_id and not item.get("deleted_at")),
            None,
        )
        found = None if row is None else (row if str(row.get("tenant_id") or "") == tenant_id else False)
    if found is False:
        raise HTTPException(status_code=403, detail="다른 테넌트의 사업자 서류입니다")
    if not found:
        raise HTTPException(status_code=404, detail="사업자 서류를 찾을 수 없습니다")
    _business_document_scope(found.get("business_id"), user)
    return found


def _business_document_summary(rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    summary: dict[str, dict[str, Any]] = {}
    for row in rows:
        item = summary.setdefault(str(row.get("business_id") or ""), {"count": 0, "has_required": False, "expiring": 0, "expired": 0})
        item["count"] += 1
        if row.get("status") != "current":
            continue
        if row.get("document_type") == BUSINESS_DOCUMENT_REQUIRED_TYPE:
            item["has_required"] = True
        if row.get("expiry_status") == "expiring":
            item["expiring"] += 1
        elif row.get("expiry_status") == "expired":
            item["expired"] += 1
    for item in summary.values():
        item["needs_documents"] = not item["has_required"]
    return summary


def list_business_documents(user: dict[str, Any], business_id: str | None = None) -> dict[str, Any]:
    tenant_id = _tenant_id(user)
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="사업자 서류는 관리자만 조회할 수 있습니다")
    scoped = _business_document_scope(business_id, user)[1] if str(business_id or "").strip() else None
    rows = _business_document_rows(tenant_id, scoped)
    if scoped is None and not _db_available():
        # 파일 모드는 테넌트-사업자 매핑 표가 없으므로 설정에 등록된 사업자만 보여준다.
        known = {str(item.get("id") or "") for item in get_settings(user)["settings"].get("businesses") or []}
        rows = [row for row in rows if str(row.get("business_id") or "") in known]
    today = datetime.now(KST).date()
    documents = sorted(
        (_business_document_view(row, today) for row in rows),
        key=lambda row: (str(row.get("uploaded_at") or ""), str(row.get("id") or "")),
        reverse=True,
    )
    return {
        "documents": documents,
        "document_types": BUSINESS_DOCUMENT_TYPES,
        "required_document_type": BUSINESS_DOCUMENT_REQUIRED_TYPE,
        "summary": _business_document_summary(documents),
    }


def save_business_document(
    *,
    business_id: str,
    document_type: str,
    issue_date: str,
    expires_at: str,
    memo: str,
    filename: str,
    content_type: str,
    data: bytes,
    user: dict[str, Any],
) -> dict[str, Any]:
    tenant_id, normalized_business = _business_document_scope(business_id, user)
    meta = _business_document_meta(document_type)
    issued, expires = _business_document_dates(issue_date, expires_at)
    clean_memo = _business_document_memo(memo)
    if not data:
        raise HTTPException(status_code=400, detail="빈 파일은 등록할 수 없습니다")
    if len(data) > _business_document_upload_limit():
        raise HTTPException(status_code=413, detail="파일은 10MB 이하여야 합니다")
    original = _safe_filename(filename or "document.bin")
    suffix = Path(original).suffix.lower()
    if suffix not in BUSINESS_DOCUMENT_EXTENSIONS:
        raise HTTPException(status_code=400, detail="PDF·이미지·한글/워드 문서만 등록할 수 있습니다")
    mime = _normalize_document_mime(content_type, original)
    document_id = str(uuid4())
    relative = _business_document_relpath(tenant_id, document_id, suffix)
    path = _contract_pdf_root() / relative
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f"{path.name}.{secrets.token_hex(4)}.tmp")
    tmp.write_bytes(data)
    os.chmod(tmp, 0o600)
    tmp.replace(path)
    now = _now()
    record = {
        "id": document_id,
        "tenant_id": tenant_id,
        "business_id": normalized_business,
        "document_type": meta["type"],
        "document_label": meta["label"],
        "original_filename": original,
        "content_type": mime,
        "stored_path": str(relative),
        "sha256": hashlib.sha256(data).hexdigest(),
        "byte_size": len(data),
        "issue_date": issued,
        "expires_at": expires,
        "memo": clean_memo,
        "status": "current",
        "uploaded_by": _email(user),
        "uploaded_at": now,
        "updated_at": now,
        "deleted_at": "",
    }
    if _db_available():
        if not _run_db(_db_business_document_insert(record)):
            # 메타를 못 남긴 원본은 고아가 된다 — 지우고 실패를 그대로 올린다(거짓 성공 금지).
            path.unlink(missing_ok=True)
            raise HTTPException(status_code=503, detail="사업자 서류를 저장하지 못했습니다")
    else:
        rows = _read_file_rows("business_documents")
        rows.insert(0, record)
        _business_document_restatus(rows, tenant_id, normalized_business, {meta["type"]})
        _write_file_rows("business_documents", rows)
    stored = next(
        (row for row in _business_document_rows(tenant_id, normalized_business) if str(row.get("id")) == document_id),
        record,
    )
    return _business_document_view(stored)


def _business_document_upload_limit() -> int:
    # 업로드 상한의 원본은 obys_upload_service.MAX_BYTES(10MB) — API 의 _read_limited_upload 와 같은 값.
    from app.services.obys_upload_service import MAX_BYTES

    return MAX_BYTES


def update_business_document(document_id: str, payload: dict[str, Any], user: dict[str, Any]) -> dict[str, Any]:
    record = _business_document_record(document_id, user)
    tenant_id = str(record["tenant_id"])
    previous_type = str(record.get("document_type") or "")
    updated = dict(record)
    if "document_type" in payload and payload["document_type"] is not None:
        meta = _business_document_meta(payload["document_type"])
        updated["document_type"], updated["document_label"] = meta["type"], meta["label"]
    issued, expires = _business_document_dates(
        payload["issue_date"] if payload.get("issue_date") is not None else record.get("issue_date"),
        payload["expires_at"] if payload.get("expires_at") is not None else record.get("expires_at"),
    )
    updated["issue_date"], updated["expires_at"] = issued, expires
    if payload.get("memo") is not None:
        updated["memo"] = _business_document_memo(payload["memo"])
    updated["updated_at"] = _now()
    if _db_available():
        if not _run_db(_db_business_document_update(updated, previous_type)):
            raise HTTPException(status_code=503, detail="사업자 서류를 수정하지 못했습니다")
    else:
        rows = _read_file_rows("business_documents")
        target = next(
            (row for row in rows if str(row.get("id")) == str(record["id"])
             and str(row.get("tenant_id") or "") == tenant_id
             and str(row.get("business_id") or "") == str(record["business_id"])
             and not row.get("deleted_at")),
            None,
        )
        if target is None:
            raise HTTPException(status_code=404, detail="사업자 서류를 찾을 수 없습니다")
        for key in ("document_type", "document_label", "issue_date", "expires_at", "memo", "updated_at"):
            target[key] = updated[key]
        _business_document_restatus(rows, tenant_id, str(record["business_id"]), {previous_type, updated["document_type"]})
        _write_file_rows("business_documents", rows)
    return _business_document_view(_business_document_record(str(record["id"]), user))


def delete_business_document(document_id: str, user: dict[str, Any]) -> None:
    """soft delete 만 한다. 행과 원본 파일은 남는다."""
    record = _business_document_record(document_id, user)
    tenant_id = str(record["tenant_id"])
    now = _now()
    if _db_available():
        if not _run_db(_db_business_document_soft_delete(record, now)):
            raise HTTPException(status_code=503, detail="사업자 서류를 삭제하지 못했습니다")
        return
    rows = _read_file_rows("business_documents")
    target = next(
        (row for row in rows if str(row.get("id")) == str(record["id"])
         and str(row.get("tenant_id") or "") == tenant_id
         and str(row.get("business_id") or "") == str(record["business_id"])
         and not row.get("deleted_at")),
        None,
    )
    if target is None:
        raise HTTPException(status_code=404, detail="사업자 서류를 찾을 수 없습니다")
    target["deleted_at"] = now
    target["updated_at"] = now
    target["status"] = "deleted"
    _business_document_restatus(rows, tenant_id, str(record["business_id"]), {str(record.get("document_type") or "")})
    _write_file_rows("business_documents", rows)


def get_business_document_file(document_id: str, user: dict[str, Any]) -> tuple[dict[str, Any], Path]:
    record = _business_document_record(document_id, user)
    path = _business_document_file(record)
    if path is None:
        raise HTTPException(status_code=404, detail="사업자 서류 원본 파일이 없습니다")
    if hashlib.sha256(path.read_bytes()).hexdigest() != str(record.get("sha256") or ""):
        logger.error("business document hash mismatch: document=%s", record.get("id"))
        raise HTTPException(status_code=409, detail="사업자 서류 원본이 등록 당시와 다릅니다")
    return record, path


def _contract_business(payload: dict[str, Any], employee: dict[str, Any] | None) -> tuple[str, str]:
    employee = employee or {}
    employee_branch = BRANCH_ALIASES.get(str(employee.get("branch") or "").strip(), str(employee.get("branch") or "").strip())
    employee_business_id = _record_business_id(employee)
    branch = BRANCH_ALIASES.get(
        str(payload.get("branch") or employee_branch or "").strip(),
        str(payload.get("branch") or employee_branch or "").strip(),
    )
    business_id = str(payload.get("business_id") or payload.get("businessId") or BUSINESS_BY_BRANCH.get(branch) or employee_business_id).strip()
    if branch and (business_id not in CANONICAL_BUSINESS_IDS or BUSINESS_BY_BRANCH.get(branch) != business_id):
        raise HTTPException(status_code=400, detail="계약서의 사업자와 지점 연결이 일치하지 않습니다")
    if employee and employee_business_id != business_id:
        raise HTTPException(status_code=400, detail="선택 직원은 해당 사업자 소속이 아닙니다")
    return business_id, branch


def _fill_contract_reference_data(payload: dict[str, Any], user: dict[str, Any]) -> dict[str, Any]:
    result = dict(payload)
    request_id = str(result.get("employee_request_id") or result.get("employeeRequestId") or "").strip()
    employee = None
    if request_id:
        employee = _find(_read_hr("employee_join_requests", user), request_id)
        if not employee or str(employee.get("status") or "").lower() != "approved":
            raise HTTPException(status_code=400, detail="승인된 가입 직원만 계약서에 선택할 수 있습니다")

    business_id, branch = _contract_business(result, employee)
    result["business_id"] = business_id
    result["branch"] = branch
    if employee:
        document_profile = _employee_onboarding_profile(
            _read_hr("onboarding_documents", user),
            employee_email=str(employee.get("email") or ""),
            employee_request_id=request_id,
        )
        employee_defaults = {
            "employee_request_id": request_id,
            "employee_name": employee.get("name") or "",
            "employee_email": employee.get("email") or "",
            "employee_address": employee.get("address") or document_profile.get("address") or "",
            "employee_phone": employee.get("phone") or "",
            "employee_birth_date": employee.get("birth_date") or document_profile.get("birth_date") or "",
            "employee_nationality": employee.get("nationality") or document_profile.get("nationality") or "대한민국",
            "bank_name": document_profile.get("bank_name") or "",
            "bank_account_holder": document_profile.get("bank_account_holder") or "",
            "bank_account_masked": document_profile.get("bank_account_masked") or "",
            "health_certificate_issue_date": document_profile.get("health_certificate_issue_date") or "",
            "health_certificate_valid_until": document_profile.get("health_certificate_valid_until") or "",
            "onboarding_document_summary": document_profile.get("onboarding_document_summary") or "",
        }
        for key, value in employee_defaults.items():
            if not str(result.get(key) or "").strip() and value:
                result[key] = value

    settings = get_settings(user).get("settings") or {}
    businesses = settings.get("businesses") if isinstance(settings.get("businesses"), list) else []
    business = next((item for item in businesses if item.get("id") == business_id), None)
    if not business:
        business = next((item for item in CANONICAL_BUSINESSES if item.get("id") == business_id), {})
    business_defaults = {
        "employer_name": business.get("name") or "",
        "employer_registration_no": business.get("registrationNo") or "",
        "employer_representative": business.get("representative") or "",
        "employer_address": business.get("address") or "",
        "employer_phone": business.get("phone") or "",
        "workplace": branch,
    }
    for key, value in business_defaults.items():
        if not str(result.get(key) or "").strip() and value:
            result[key] = value
    return result


def _contract_payload_value(payload: dict[str, Any], snake_key: str, camel_key: str = "") -> Any:
    value = payload.get(snake_key)
    if value is None and camel_key:
        value = payload.get(camel_key)
    return value


MIN_WAGE_2026 = 10320


def _monthly_paid_hours(weekly_hours: float) -> float:
    """주 소정근로시간에서 월 유급환산시간(주휴 포함)을 구한다."""
    return (weekly_hours + min(8, weekly_hours / 5)) * 365 / 7 / 12


def _probation_wage_rate(payload: dict[str, Any]) -> float | None:
    """수습기간 임금 감액률(%). 감액이 없으면 None 을 돌려준다.

    명시 필드(probation_wage_rate)를 먼저 보고, 없으면 수습기간 자유기재
    문구에서 '90%' 같은 표기를 읽는다. 프런트에 숫자 입력칸이 생기기 전에도
    최저임금 미만 감액이 저장되지 않게 하기 위함이다.
    """
    raw = _contract_payload_value(payload, "probation_wage_rate", "probationWageRate")
    rate: float | None = None
    if raw is not None and str(raw).strip() != "":
        try:
            rate = float(str(raw).strip().rstrip("%"))
        except ValueError:
            raise HTTPException(status_code=400, detail="수습기간 임금 감액률은 숫자로 입력하십시오")
    else:
        text = str(_contract_payload_value(payload, "probation_period", "probationPeriod") or "")
        match = re.search(r"(\d{1,3}(?:\.\d+)?)\s*%", text)
        if match:
            rate = float(match.group(1))
    if rate is None:
        return None
    if rate <= 0 or rate > 100:
        raise HTTPException(
            status_code=400,
            detail="수습기간 임금 감액률은 0 초과 100 이하로 입력하십시오",
        )
    if rate >= 100:
        return None
    return rate


def _missing_contract_value(value: Any) -> bool:
    return str(value or "").strip() in {"", "-", "미등록", "기초등록 필요"}


def _validate_contract_payload(payload: dict[str, Any]) -> dict[str, Any]:
    """Validate legal/operational invariants before a contract is persisted."""
    result = dict(payload)
    contract_type = str(_contract_payload_value(result, "contract_type", "contractType") or "").strip()
    if contract_type not in VALID_CONTRACT_TYPES:
        raise HTTPException(status_code=400, detail="지원하지 않는 계약 유형입니다")

    employee_name = str(_contract_payload_value(result, "employee_name", "employeeName") or "").strip()
    employee_email = str(_contract_payload_value(result, "employee_email", "employeeEmail") or "").strip().lower()
    employee_request_id = str(_contract_payload_value(result, "employee_request_id", "employeeRequestId") or "").strip()
    business_id = str(_contract_payload_value(result, "business_id", "businessId") or "").strip()
    branch = str(result.get("branch") or "").strip()
    contract_date = str(_contract_payload_value(result, "contract_date", "contractDate") or "").strip()
    employee_address = str(_contract_payload_value(result, "employee_address", "employeeAddress") or "").strip()
    employee_phone = str(_contract_payload_value(result, "employee_phone", "employeePhone") or "").strip()
    employee_birth_date = str(_contract_payload_value(result, "employee_birth_date", "employeeBirthDate") or "").strip()

    missing: list[str] = []
    for label, value in (
        ("승인 직원", employee_request_id),
        ("직원명", employee_name),
        ("직원 이메일", employee_email),
        ("근로자 주소", employee_address),
        ("근로자 연락처", employee_phone),
        ("근로자 생년월일", employee_birth_date),
        ("사업자", business_id),
        ("근무 지점", branch),
        ("계약 작성일", contract_date),
        ("사용자 상호", _contract_payload_value(result, "employer_name", "employerName")),
        ("사업자등록번호", _contract_payload_value(result, "employer_registration_no", "employerRegistrationNo")),
        ("대표자", _contract_payload_value(result, "employer_representative", "employerRepresentative")),
        ("사용자 주소", _contract_payload_value(result, "employer_address", "employerAddress")),
    ):
        if _missing_contract_value(value):
            missing.append(label)

    if employee_email and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", employee_email):
        raise HTTPException(status_code=400, detail="직원 이메일 형식이 올바르지 않습니다")
    if employee_phone and len(re.sub(r"\D", "", employee_phone)) < 9:
        raise HTTPException(status_code=400, detail="근로자 연락처 형식이 올바르지 않습니다")
    employer_registration_no = str(
        _contract_payload_value(result, "employer_registration_no", "employerRegistrationNo") or ""
    ).strip()
    if not _missing_contract_value(employer_registration_no) and not re.fullmatch(
        r"\d{3}-?\d{2}-?\d{5}", employer_registration_no
    ):
        raise HTTPException(status_code=400, detail="사업자등록번호 형식이 올바르지 않습니다")

    start_date = str(_contract_payload_value(result, "start_date", "startDate") or "").strip()
    end_date = str(_contract_payload_value(result, "end_date", "endDate") or "").strip()
    if start_date and end_date and end_date < start_date:
        raise HTTPException(status_code=400, detail="계약 종료일은 입사일보다 빠를 수 없습니다")

    birth_date_value: date | None = None
    if employee_birth_date:
        try:
            birth_date_value = date.fromisoformat(employee_birth_date)
        except ValueError as exc:
            raise HTTPException(status_code=400, detail="근로자 생년월일 형식이 올바르지 않습니다") from exc
        reference_text = start_date or contract_date
        try:
            reference_date = date.fromisoformat(reference_text)
        except ValueError:
            reference_date = datetime.now(KST).date()
        if birth_date_value > reference_date:
            raise HTTPException(status_code=400, detail="근로자 생년월일은 계약일보다 늦을 수 없습니다")
        age = reference_date.year - birth_date_value.year - (
            (reference_date.month, reference_date.day) < (birth_date_value.month, birth_date_value.day)
        )
        if age < 18:
            for label, key, camel in (
                ("친권자/후견인 성명", "minor_guardian_name", "minorGuardianName"),
                ("친권자/후견인 연락처", "minor_guardian_phone", "minorGuardianPhone"),
            ):
                if _missing_contract_value(_contract_payload_value(result, key, camel)):
                    missing.append(label)
            consent = str(
                _contract_payload_value(result, "minor_guardian_consent", "minorGuardianConsent") or ""
            ).strip()
            if consent != "confirmed":
                missing.append("친권자/후견인 동의서 확인")

    employment_tax_type = str(
        _contract_payload_value(result, "employment_tax_type", "employmentTaxType") or ""
    ).strip()
    wage_type = str(_contract_payload_value(result, "wage_type", "wageType") or "").strip()
    try:
        wage = float(result.get("wage") or 0)
    except (TypeError, ValueError):
        wage = 0

    if contract_type in EMPLOYMENT_CONTRACT_TYPES:
        if employment_tax_type != "four_insurance":
            raise HTTPException(status_code=400, detail="근로계약서는 4대보험 가입 근로자 구분으로 작성해야 합니다")
        for label, key, camel in (
            ("입사일", "start_date", "startDate"),
            ("근무장소", "workplace", "workplace"),
            ("업무내용", "job_description", "jobDescription"),
            ("근무시간", "work_time", "workTime"),
            ("휴게시간", "rest_time", "restTime"),
            ("주 소정근로시간", "weekly_hours", "weeklyHours"),
            ("근무일/요일", "work_days", "workDays"),
            ("휴일/주휴", "holidays", "holidays"),
            ("급여지급일", "pay_date", "payDate"),
            ("지급방법", "pay_method", "payMethod"),
            ("임금 구성/공제", "wage_composition", "wageComposition"),
            ("연장·야간·휴일근로", "overtime_terms", "overtimeTerms"),
            ("연차/휴가/결근", "leave_terms", "leaveTerms"),
            ("4대보험/세무 처리", "insurance_terms", "insuranceTerms"),
        ):
            if _missing_contract_value(_contract_payload_value(result, key, camel)):
                missing.append(label)
        if contract_type == "part_time" and _missing_contract_value(
            _contract_payload_value(result, "daily_work_schedule", "dailyWorkSchedule")
        ):
            missing.append("근로일별 근로시간")
        if wage_type not in {"hourly", "monthly", "daily"}:
            raise HTTPException(status_code=400, detail="근로계약서의 임금 산정 방식을 확인하십시오")
        if wage <= 0:
            missing.append("확정 임금")
        component_values = []
        for snake, camel in (
            ("base_salary", "baseSalary"),
            ("non_tax_meal_allowance", "nonTaxMealAllowance"),
            ("taxable_allowance", "taxableAllowance"),
        ):
            raw_value = _contract_payload_value(result, snake, camel)
            try:
                component_values.append(float(raw_value or 0))
            except (TypeError, ValueError):
                raise HTTPException(status_code=400, detail="급여 구성 금액을 숫자로 입력하십시오")
        base_salary, non_tax_meal, taxable_allowance = component_values
        component_total = base_salary + non_tax_meal + taxable_allowance
        if component_total and abs(component_total - wage) >= 1:
            raise HTTPException(status_code=400, detail="기본급·비과세 식대·기타 과세수당 합계가 월 총액과 일치해야 합니다")
        if non_tax_meal < 0 or non_tax_meal > 200000:
            raise HTTPException(status_code=400, detail="비과세 식대는 월 200,000원 이내로 입력하십시오")
        meal_provision = str(_contract_payload_value(result, "meal_provision", "mealProvision") or "").strip()
        if non_tax_meal > 0 and meal_provision != "cash_no_meal":
            raise HTTPException(status_code=400, detail="사용자가 식사를 제공하는 경우 현금 식대를 비과세로 분류할 수 없습니다")
        workplace_size = str(
            _contract_payload_value(result, "workplace_size_category", "workplaceSizeCategory") or ""
        ).strip()
        weekly_hours_text = str(_contract_payload_value(result, "weekly_hours", "weeklyHours") or "")
        weekly_match = re.search(r"주\s*(\d+)시간(?:\s*(\d+)분)?", weekly_hours_text)
        contract_date_year = contract_date[:4]
        if (
            contract_type == "regular"
            and workplace_size == "under_5"
            and weekly_match
            and base_salary > 0
            and contract_date_year == "2026"
        ):
            weekly_hours = float(weekly_match.group(1)) + float(weekly_match.group(2) or 0) / 60
            monthly_paid_hours = (weekly_hours + min(8, weekly_hours / 5)) * 365 / 7 / 12
            conservative_hourly = base_salary / monthly_paid_hours
            if conservative_hourly < 10320:
                raise HTTPException(
                    status_code=400,
                    detail=f"과세 기본급 기준 환산시급 {int(conservative_hourly):,}원은 2026년 최저임금 10,320원보다 낮습니다",
                )
        # 수습 감액은 감액 후 환산시급이 최저임금 이상일 때만 허용한다.
        # 최저임금법 제5조②(수습 감액 특례)는 단순노무 직종에 적용되지 않으므로,
        # 음식점에서 쓸 수 있는 것은 '계약임금을 낮추되 최저임금은 넘는' 경우뿐이다.
        probation_rate = _probation_wage_rate(result)
        if probation_rate is not None and weekly_match and base_salary > 0 and contract_date_year == "2026":
            probation_weekly = float(weekly_match.group(1)) + float(weekly_match.group(2) or 0) / 60
            probation_paid_hours = _monthly_paid_hours(probation_weekly)
            probation_hourly = (base_salary * probation_rate / 100) / probation_paid_hours
            if probation_hourly < MIN_WAGE_2026:
                min_base = math.ceil(MIN_WAGE_2026 * probation_paid_hours * 100 / probation_rate)
                raise HTTPException(
                    status_code=400,
                    detail=(
                        f"수습 {probation_rate:g}% 적용 시 환산시급 {int(probation_hourly):,}원으로 "
                        f"2026년 최저임금 {MIN_WAGE_2026:,}원에 미달합니다. "
                        f"이 감액률을 쓰려면 계약 월급이 최소 {min_base:,}원 이상이어야 합니다"
                    ),
                )
        foreign_worker = str(
            _contract_payload_value(result, "foreign_worker", "foreignWorker") or ""
        ).strip().lower() in {"true", "1", "yes"}
        if foreign_worker:
            for label, key, camel in (
                ("국적", "employee_nationality", "employeeNationality"),
                ("체류자격", "visa_status", "visaStatus"),
                ("외국인등록번호(마스킹)", "foreign_registration_no_masked", "foreignRegistrationNoMasked"),
            ):
                if _missing_contract_value(_contract_payload_value(result, key, camel)):
                    missing.append(label)
    elif contract_type == "freelancer":
        if employment_tax_type != "freelancer_33":
            raise HTTPException(status_code=400, detail="프리랜서 용역계약서는 3.3% 원천징수 구분으로 작성해야 합니다")
        if wage_type != "case_fee":
            raise HTTPException(status_code=400, detail="프리랜서 용역계약서는 건별/용역비 방식으로 작성해야 합니다")
        for label, key, camel in (
            ("용역 시작일", "start_date", "startDate"),
            ("용역 업무범위/산출물", "freelancer_scope", "freelancerScope"),
            ("용역비 정산/해지", "freelancer_settlement_terms", "freelancerSettlementTerms"),
        ):
            if _missing_contract_value(_contract_payload_value(result, key, camel)):
                missing.append(label)
        if wage <= 0:
            missing.append("확정 용역비")

    if missing:
        unique_missing = list(dict.fromkeys(missing))
        raise HTTPException(status_code=400, detail=f"계약서 필수 입력값을 확인하십시오: {', '.join(unique_missing)}")
    return result


def _signed_contract_snapshot(contract: dict[str, Any]) -> tuple[dict[str, Any], str]:
    snapshot = {key: value for key, value in contract.items() if key not in CONTRACT_SNAPSHOT_EXCLUDED_FIELDS}
    encoded = json.dumps(snapshot, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return snapshot, hashlib.sha256(encoded).hexdigest()


def _contract_defaults(payload: dict[str, Any]) -> dict[str, Any]:
    now = _now()
    contract_id = str(payload.get("id") or uuid4())
    employee_email = str(payload.get("employee_email") or payload.get("employeeEmail") or "").strip().lower()
    contract_type = str(payload.get("contract_type") or payload.get("contractType") or "part_time")
    meta = CONTRACT_TEMPLATE_META.get(contract_type, CONTRACT_TEMPLATE_META["default"])
    contract = {
        **payload,
        "id": contract_id,
        "contract_type": contract_type,
        "document_kind": str(payload.get("document_kind") or payload.get("documentKind") or meta["document_kind"]),
        "template_version": str(payload.get("template_version") or payload.get("templateVersion") or meta["template_version"]),
        "print_title": str(payload.get("print_title") or payload.get("printTitle") or meta["print_title"]),
        "employee_email": employee_email,
        "employee_email_masked": _mask_email(employee_email),
        "employee_name": str(payload.get("employee_name") or payload.get("employeeName") or "").strip(),
        "status": str(payload.get("status") or "draft"),
        "created_at": payload.get("created_at") or now,
        "updated_at": now,
    }
    return contract


def list_contracts(user: dict[str, Any]) -> list[dict[str, Any]]:
    return sorted(_filter_user(_read_hr("contracts", user), user, "employee_email"), key=lambda row: row.get("updated_at", ""), reverse=True)


def save_contract(payload: dict[str, Any], user: dict[str, Any]) -> dict[str, Any]:
    _tenant_id(user)
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="계약서 작성 권한이 없습니다")
    rows = _read_hr("contracts", user)
    requested_id = str(payload.get("id") or "").strip()
    existing = _find(rows, requested_id) if requested_id else None
    if existing and str(existing.get("status") or "") == "signed":
        raise HTTPException(status_code=409, detail="서명 완료 계약서는 수정할 수 없습니다. 정정 계약서를 새로 작성하십시오")
    contract = _owned_hr_record(_contract_defaults(_validate_contract_payload(_fill_contract_reference_data(payload, user))), user)
    existing = _find(rows, contract["id"])
    if existing:
        if str(existing.get("status") or "") == "requested":
            contract["status"] = "draft"
            contract.pop("sign_token", None)
            contract.pop("requested_at", None)
        existing.update(contract)
        if existing.get("status") == "draft":
            existing.pop("sign_token", None)
            existing.pop("requested_at", None)
        saved = existing
    else:
        contract["status"] = "draft"
        rows.insert(0, contract)
        saved = contract
    _write_hr_record("contracts", saved, user)
    return saved


def request_contract_signature(contract_id: str, user: dict[str, Any]) -> dict[str, Any]:
    _tenant_id(user)
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="서명 요청 권한이 없습니다")
    contract = _require_hr_record(_find(_read_hr("contracts", user), contract_id), user, detail="계약서를 찾을 수 없습니다")
    if str(contract.get("status") or "") == "signed":
        raise HTTPException(status_code=409, detail="서명 완료 계약서는 다시 서명 요청할 수 없습니다")
    _validate_contract_payload(contract)
    contract["status"] = "requested"
    contract["sign_token"] = contract.get("sign_token") or secrets.token_urlsafe(24)
    contract["requested_at"] = _now()
    contract["updated_at"] = contract["requested_at"]
    _write_hr_record("contracts", contract, user)
    return contract


def _contract_signer_email(contract: dict[str, Any], user: dict[str, Any] | None) -> str:
    if not user or not _email(user):
        raise HTTPException(status_code=401, detail="직원 계정 로그인이 필요합니다")
    if _is_admin(user):
        raise HTTPException(status_code=403, detail="관리자는 직원 대신 계약서에 서명할 수 없습니다")
    _require_hr_record(contract, user, detail="서명 요청 계약서를 찾을 수 없습니다")
    signer_email = _email(user)
    employee_email = str(contract.get("employee_email") or "").strip().lower()
    if not employee_email or signer_email != employee_email:
        raise HTTPException(status_code=403, detail="서명 대상 직원 계정이 아닙니다")
    return signer_email


def _validated_signature_image(data_uri: Any) -> tuple[str, str]:
    value = str(data_uri or "").strip()
    prefix = "data:image/png;base64,"
    if not value.startswith(prefix):
        raise HTTPException(status_code=400, detail="자필서명은 PNG 이미지 형식이어야 합니다")
    try:
        raw = base64.b64decode(value[len(prefix) :], validate=True)
    except (binascii.Error, ValueError):
        raise HTTPException(status_code=400, detail="자필서명 이미지가 올바르지 않습니다") from None
    if len(raw) < 100 or len(raw) > MAX_SIGNATURE_BYTES or not raw.startswith(b"\x89PNG\r\n\x1a\n"):
        raise HTTPException(status_code=400, detail="자필서명 이미지 크기 또는 형식을 확인하십시오")
    return value, hashlib.sha256(raw).hexdigest()


def _signing_contract_for_token(rows: list[dict[str, Any]], token: str) -> dict[str, Any]:
    """rows 는 _read_hr 결과(JWT 테넌트 SQL 스코프)다. 다른 테넌트 토큰은 여기서 안 보인다.

    토큰이 없거나 다른 테넌트 것이면 존재 여부를 드러내지 않고 똑같이 403.
    """
    digest = hashlib.sha256(token.encode("utf-8")).hexdigest() if token else ""
    # 서명 뒤에는 sign_token 이 지워지고 해시만 남는다 — 같은 링크 재제출을 "찾을 수 없음"(403)이 아니라 "이미 서명"(409)으로 알리려고 해시로도 찾는다.
    contract = next(
        (row for row in rows if token and (row.get("sign_token") == token or (digest and row.get("sign_token_hash") == digest))),
        None,
    )
    if not contract:
        raise HTTPException(status_code=403, detail="서명 요청 계약서를 찾을 수 없거나 접근 권한이 없습니다")
    return contract


def _require_sign_link_alive(contract: dict[str, Any]) -> None:
    requested_at = _pg_ts(contract.get("requested_at"))
    if requested_at is None:
        return
    if datetime.now(KST) - requested_at > timedelta(days=CONTRACT_SIGN_LINK_TTL_DAYS):
        raise HTTPException(
            status_code=410,
            detail=f"서명 링크가 만료되었습니다({CONTRACT_SIGN_LINK_TTL_DAYS}일). 관리자에게 서명 요청을 다시 받으십시오",
        )


def _require_employee_still_approved(contract: dict[str, Any], user: dict[str, Any] | None) -> None:
    request_id = str(contract.get("employee_request_id") or "").strip()
    if not request_id:
        return
    join = _find(_read_hr("employee_join_requests", user), request_id)
    if not join or str(join.get("status") or "").strip().lower() != "approved":
        raise HTTPException(status_code=403, detail="직원 가입 승인이 확인되지 않아 계약서에 서명할 수 없습니다")


def get_contract_by_token(token: str, user: dict[str, Any] | None = None) -> dict[str, Any]:
    contract = _signing_contract_for_token(_read_hr("contracts", user), token)
    _contract_signer_email(contract, user)
    if str(contract.get("status") or "") == "signed":
        raise HTTPException(status_code=409, detail="이미 서명 완료된 계약서입니다")
    if str(contract.get("status") or "") != "requested":
        raise HTTPException(status_code=409, detail="서명 요청된 계약서가 아닙니다")
    _require_employee_still_approved(contract, user)
    _require_sign_link_alive(contract)
    return contract


def sign_contract(payload: dict[str, Any], user: dict[str, Any] | None = None) -> dict[str, Any]:
    # 더블클릭·두 탭 동시 제출이 둘 다 "requested" 를 읽고 서명본을 두 번 쓰지 못하게 프로세스 안에서 직렬화한다.
    with _CONTRACT_SIGN_LOCK:
        return _sign_contract_locked(payload, user)


def _sign_contract_locked(payload: dict[str, Any], user: dict[str, Any] | None = None) -> dict[str, Any]:
    token = str(payload.get("token") or "")
    contract = _signing_contract_for_token(_read_hr("contracts", user), token)
    signer_email = _contract_signer_email(contract, user)
    if str(contract.get("status") or "") == "signed":
        raise HTTPException(status_code=409, detail="이미 서명 완료된 계약서입니다")
    if str(contract.get("status") or "") != "requested":
        raise HTTPException(status_code=409, detail="서명 요청된 계약서만 서명할 수 있습니다")
    _require_employee_still_approved(contract, user)
    _require_sign_link_alive(contract)
    if payload.get("consent") is not True:
        raise HTTPException(status_code=400, detail="계약 내용 확인 및 전자서명 동의가 필요합니다")
    consent_version = str(payload.get("consent_version") or "").strip()
    if consent_version != CONTRACT_SIGNATURE_CONSENT_VERSION:
        raise HTTPException(status_code=400, detail="전자서명 동의 문구를 새로 확인하십시오")
    signer_name = str(payload.get("signer_name") or "").strip()
    employee_name = str(contract.get("employee_name") or "").strip()
    if not signer_name or re.sub(r"\s+", "", signer_name) != re.sub(r"\s+", "", employee_name):
        raise HTTPException(status_code=400, detail="계약 대상 직원 이름을 정확히 입력하십시오")
    signature_data_uri, signature_sha256 = _validated_signature_image(payload.get("signature_data_uri"))
    _validate_contract_payload(contract)
    contract["status"] = "signed"
    contract["signed_at"] = _now()
    contract["signer_name"] = signer_name
    contract["signer_email"] = signer_email
    contract["signature_data_uri"] = signature_data_uri
    contract["signature_sha256"] = signature_sha256
    contract["signature_consent"] = {
        "accepted": True,
        "version": consent_version,
        "accepted_at": contract["signed_at"],
    }
    contract["signature_audit"] = {
        "authenticated_email": signer_email,
        "client_ip": str(payload.get("audit_ip") or "")[:64],
        "user_agent": str(payload.get("audit_user_agent") or "")[:512],
        "signed_at": contract["signed_at"],
    }
    contract["sign_token_hash"] = hashlib.sha256(token.encode("utf-8")).hexdigest()
    contract.pop("sign_token", None)
    contract["updated_at"] = contract["signed_at"]
    snapshot, snapshot_sha256 = _signed_contract_snapshot(contract)
    contract["signed_snapshot"] = snapshot
    contract["signed_snapshot_sha256"] = snapshot_sha256
    _write_hr_record("contracts", contract, user)
    return contract


def delete_contract(contract_id: str, user: dict[str, Any]) -> None:
    _tenant_id(user)
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="계약서 삭제 권한이 없습니다")
    contract = _require_hr_record(_find(_read_hr("contracts", user), contract_id), user, detail="계약서를 찾을 수 없습니다")
    if str(contract.get("status") or "") == "signed":
        raise HTTPException(status_code=409, detail="서명 완료 계약서는 삭제할 수 없습니다")
    _delete_hr_record("contracts", contract_id, user)


# ---------------------------------------------------------------------------
# 계약서 서명요청 알림 · 서명본 PDF 보관/교부
#
# 알림과 PDF 는 서명요청·서명 트랜잭션이 끝난 **뒤에** 돈다. 어느 쪽이 실패해도
# 서명요청/서명 자체는 성공으로 남는다 — 서명은 이미 봉인된 법적 행위이고,
# 알림은 부가 수단이다. 실패는 이력(sent/failed/skipped)과 signed_pdf_error 로 남긴다.
# ---------------------------------------------------------------------------
CONTRACT_NOTIFICATION_LOG = "contract_notifications"
CONTRACT_NOTIFICATION_TABLE = "yeoljeong_contract_notifications"
CONTRACT_NOTICE_RESEND_COOLDOWN_SECONDS = 300
CONTRACT_DELIVERY_COLUMNS = ("signed_pdf_path", "signed_pdf_sha256", "signed_pdf_bytes", "delivered_at", "delivery_channel")


def _run_coroutine(coro: Any) -> Any:
    """sync 서비스에서 코루틴을 돌린다. ``_run_db`` 와 달리 예외를 삼키지 않는다."""
    try:
        asyncio.get_running_loop()
    except RuntimeError:
        return asyncio.run(coro)
    import concurrent.futures

    with concurrent.futures.ThreadPoolExecutor(max_workers=1) as pool:
        return pool.submit(asyncio.run, coro).result()


async def _db_insert_contract_notifications(rows: list[dict[str, Any]]) -> bool:
    import asyncpg

    conn = await asyncpg.connect(_db_url(), timeout=5)
    try:
        if not await conn.fetchval("SELECT to_regclass($1) IS NOT NULL", f"public.{CONTRACT_NOTIFICATION_TABLE}"):
            return False
        await conn.executemany(
            """
            INSERT INTO yeoljeong_contract_notifications
                (tenant_id, contract_id, event, channel, status, target_masked, error_detail)
            VALUES ($1::uuid, $2, $3, $4, $5, $6, $7)
            """,
            [
                (
                    UUID(row["tenant_id"]),
                    row["contract_id"],
                    row["event"],
                    row["channel"],
                    row["status"],
                    row["target_masked"],
                    row["error_detail"],
                )
                for row in rows
            ],
        )
        return True
    finally:
        await conn.close()


async def _db_fetch_contract_notifications(tenant_id: str, contract_id: str) -> list[dict[str, Any]] | None:
    import asyncpg

    conn = await asyncpg.connect(_db_url(), timeout=5)
    try:
        if not await conn.fetchval("SELECT to_regclass($1) IS NOT NULL", f"public.{CONTRACT_NOTIFICATION_TABLE}"):
            return None
        rows = await conn.fetch(
            """
            SELECT id, tenant_id, contract_id, event, channel, status, target_masked, error_detail, created_at
              FROM yeoljeong_contract_notifications
             WHERE tenant_id = $1::uuid AND contract_id = $2
             ORDER BY created_at DESC, id DESC
             LIMIT 100
            """,
            UUID(tenant_id),
            contract_id,
        )
        return [{**dict(row), "id": str(row["id"]), "tenant_id": str(row["tenant_id"]), "created_at": _iso(row["created_at"])} for row in rows]
    finally:
        await conn.close()


async def _db_update_contract_delivery_columns(contract: dict[str, Any]) -> bool:
    """마이그레이션이 적용된 경우에만 계약서 컬럼에도 보관·교부 메타를 쓴다.

    값의 원본은 contract_payload 이고, 컬럼은 조회 쿼리용 사본이다. 컬럼이 아직
    없을 때 upsert 에 섞으면 계약서 저장 전체가 깨지므로 따로 갱신한다.
    """
    import asyncpg

    conn = await asyncpg.connect(_db_url(), timeout=5)
    try:
        present = await conn.fetchval(
            "SELECT count(*) FROM information_schema.columns "
            "WHERE table_schema = 'public' AND table_name = 'yeoljeong_contracts' AND column_name = ANY($1::text[])",
            list(CONTRACT_DELIVERY_COLUMNS),
        )
        if int(present or 0) < len(CONTRACT_DELIVERY_COLUMNS):
            return False
        result = await conn.execute(
            """
            UPDATE yeoljeong_contracts
               SET signed_pdf_path = $3, signed_pdf_sha256 = $4, signed_pdf_bytes = $5,
                   delivered_at = $6::timestamptz, delivery_channel = $7
             WHERE id = $1 AND tenant_id = $2::uuid AND deleted_at IS NULL
            """,
            str(contract.get("id") or ""),
            UUID(str(contract.get("tenant_id") or "")),
            str(contract.get("signed_pdf_path") or "") or None,
            str(contract.get("signed_pdf_sha256") or "") or None,
            int(contract.get("signed_pdf_bytes") or 0) or None,
            _pg_ts(contract.get("delivered_at")),
            str(contract.get("delivery_channel") or "") or None,
        )
        return result.endswith(" 1")
    finally:
        await conn.close()


def _append_contract_notification_logs(contract: dict[str, Any], event: str, results: list[dict[str, Any]]) -> bool:
    now = _now()
    rows = [
        {
            "id": str(uuid4()),
            "tenant_id": str(contract.get("tenant_id") or ""),
            "contract_id": str(contract.get("id") or ""),
            "event": event,
            "channel": str(item.get("channel") or ""),
            "status": str(item.get("status") or "failed"),
            "target_masked": str(item.get("target_masked") or ""),
            "error_detail": str(item.get("error_detail") or "")[:500],
            "created_at": now,
        }
        for item in results
    ]
    if not rows:
        return True
    if _db_available():
        return bool(_run_db(_db_insert_contract_notifications(rows)))
    _write_file_rows(CONTRACT_NOTIFICATION_LOG, rows + _read_file_rows(CONTRACT_NOTIFICATION_LOG))
    return True


def _contract_notification_history(tenant_id: str, contract_id: str) -> list[dict[str, Any]] | None:
    """최신순 이력. 저장소를 확인할 수 없으면 None."""
    if _db_available():
        rows = _run_db(_db_fetch_contract_notifications(tenant_id, contract_id))
        return rows if isinstance(rows, list) else None
    rows = [
        row
        for row in _read_file_rows(CONTRACT_NOTIFICATION_LOG)
        if str(row.get("tenant_id") or "") == tenant_id and str(row.get("contract_id") or "") == contract_id
    ]
    return sorted(rows, key=lambda row: str(row.get("created_at") or ""), reverse=True)


def _notify_contract_event(contract: dict[str, Any], event: str) -> dict[str, Any]:
    """알림 발송 + 이력 기록. 어떤 실패도 예외로 올리지 않는다."""
    try:
        from app.services import yeoljeong_contract_notify as contract_notify

        results = _run_coroutine(contract_notify.dispatch(contract, event))
        status = contract_notify.summarize(results)
    except Exception as exc:  # noqa: BLE001 — 알림 실패가 서명 흐름을 깨면 안 된다
        logger.warning("contract notify dispatch failed: contract=%s event=%s err=%s", contract.get("id"), event, exc)
        results = [{"channel": "dispatch", "status": "failed", "target_masked": "", "error_detail": f"{type(exc).__name__}: {exc}"[:500]}]
        status = "failed"
    try:
        logged = _append_contract_notification_logs(contract, event, results)
    except Exception as exc:  # noqa: BLE001
        logger.warning("contract notify log failed: contract=%s err=%s", contract.get("id"), exc)
        logged = False
    if not logged:
        logger.warning("contract notify history not stored: contract=%s event=%s", contract.get("id"), event)
    return {"event": event, "status": status, "channels": results, "logged": logged}


def request_contract_signature_with_notice(contract_id: str, user: dict[str, Any]) -> dict[str, Any]:
    contract = request_contract_signature(contract_id, user)
    return {"contract": contract, "notify": _notify_contract_event(contract, "signature_requested")}


def resend_contract_signature_notice(contract_id: str, user: dict[str, Any]) -> dict[str, Any]:
    tenant_id = _tenant_id(user)
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="서명 요청 재발송 권한이 없습니다")
    contract = _require_hr_record(_find(_read_hr("contracts", user), contract_id), user, detail="계약서를 찾을 수 없습니다")
    if str(contract.get("status") or "") != "requested":
        raise HTTPException(status_code=409, detail="서명 요청 상태의 계약서만 알림을 재발송할 수 있습니다")
    history = _contract_notification_history(tenant_id, str(contract.get("id") or ""))
    if history is None:
        # 이력을 못 읽으면 rate limit 도 못 건다 — 열어두지 않고 막는다.
        raise HTTPException(status_code=503, detail="알림 발송 이력을 확인할 수 없어 재발송할 수 없습니다")
    stamps = [
        stamp
        for stamp in (_pg_ts(row.get("created_at")) for row in history if row.get("event") == "signature_requested")
        if stamp
    ]
    if stamps:
        last_sent = max(stamps)
        elapsed = (datetime.now(KST) - last_sent).total_seconds()
        if elapsed < CONTRACT_NOTICE_RESEND_COOLDOWN_SECONDS:
            wait = int(CONTRACT_NOTICE_RESEND_COOLDOWN_SECONDS - elapsed) + 1
            raise HTTPException(status_code=429, detail=f"서명 요청 알림은 5분에 한 번만 보낼 수 있습니다. {wait}초 후 다시 시도하십시오")
    return {"contract_id": str(contract.get("id") or ""), "notify": _notify_contract_event(contract, "signature_requested")}


def _contract_pdf_root() -> Path:
    configured = str(os.getenv("OBYS_UPLOAD_ROOT") or "").strip()
    return Path(configured) if configured else DATA_DIR / "uploads" / "ledgers"


def _contract_pdf_relpath(contract: dict[str, Any]) -> Path:
    tenant = str(UUID(str(contract.get("tenant_id") or "")))
    digest = hashlib.sha256(str(contract.get("id") or "").encode("utf-8")).hexdigest()[:32]
    return Path(tenant) / "contracts" / f"{digest}.signed.pdf"


def _verified_signed_pdf_path(contract: dict[str, Any]) -> Path | None:
    relative = str(contract.get("signed_pdf_path") or "")
    expected = str(contract.get("signed_pdf_sha256") or "")
    if not relative or not expected:
        return None
    root = _contract_pdf_root().resolve()
    path = (root / relative).resolve()
    if root not in path.parents or not path.is_file():
        return None
    if hashlib.sha256(path.read_bytes()).hexdigest() != expected:
        logger.error("signed contract pdf hash mismatch: contract=%s", contract.get("id"))
        return None
    return path


def _persist_contract_delivery(contract: dict[str, Any], user: dict[str, Any] | None) -> None:
    _write_hr_record("contracts", contract, user)
    if _db_available():
        _run_db(_db_update_contract_delivery_columns(contract))


def _store_signed_contract_pdf(contract: dict[str, Any], user: dict[str, Any] | None) -> dict[str, Any]:
    """봉인 스냅샷으로 PDF 를 만들어 보관하고 계약서에 메타를 남긴다. 예외를 올리지 않는다."""
    result: dict[str, Any]
    try:
        from app.services import yeoljeong_contract_pdf as contract_pdf

        data = contract_pdf.render_signed_contract_pdf(contract)
        relative = _contract_pdf_relpath(contract)
        path = _contract_pdf_root() / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        tmp = path.with_name(f"{path.name}.{secrets.token_hex(4)}.tmp")
        tmp.write_bytes(data)
        os.chmod(tmp, 0o600)
        tmp.replace(path)
        digest = hashlib.sha256(data).hexdigest()
        contract["signed_pdf_path"] = str(relative)
        contract["signed_pdf_sha256"] = digest
        contract["signed_pdf_bytes"] = len(data)
        contract["signed_pdf_generated_at"] = _now()
        contract["signed_pdf_error"] = ""
        result = {"status": "stored", "sha256": digest, "bytes": len(data)}
    except Exception as exc:  # noqa: BLE001 — 서명은 유지하고 실패만 기록한다
        logger.error("signed contract pdf failed: contract=%s err=%s", contract.get("id"), exc)
        contract["signed_pdf_error"] = f"{_now()} {type(exc).__name__}: {exc}"[:500]
        result = {"status": "failed", "error": contract["signed_pdf_error"]}
    try:
        _persist_contract_delivery(contract, user)
    except Exception as exc:  # noqa: BLE001
        logger.error("signed contract pdf meta not recorded: contract=%s err=%s", contract.get("id"), exc)
        result = {**result, "recorded": False}
    return result


def sign_contract_and_deliver(payload: dict[str, Any], user: dict[str, Any] | None = None) -> dict[str, Any]:
    contract = sign_contract(payload, user)
    signed_pdf = _store_signed_contract_pdf(contract, user)
    employment = _sync_employment_after_signature(contract, user)
    return {
        "contract": contract,
        "signed_pdf": signed_pdf,
        "employment": employment,
        "notify": _notify_contract_event(contract, "signed"),
    }


def regenerate_signed_contract_pdf(contract_id: str, user: dict[str, Any]) -> dict[str, Any]:
    _tenant_id(user)
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="서명본 PDF 재생성 권한이 없습니다")
    contract = _require_hr_record(_find(_read_hr("contracts", user), contract_id), user, detail="계약서를 찾을 수 없습니다")
    if str(contract.get("status") or "") != "signed":
        raise HTTPException(status_code=409, detail="서명 완료 계약서만 PDF 를 만들 수 있습니다")
    return {"contract_id": str(contract.get("id") or ""), "signed_pdf": _store_signed_contract_pdf(contract, user)}


def signed_contract_pdf_for_download(contract_id: str, user: dict[str, Any]) -> tuple[Path, str]:
    """관리자와 계약 당사자 본인만. 다른 테넌트·제3자는 존재 여부와 무관하게 403."""
    denied = HTTPException(status_code=403, detail="이 계약서 PDF 에 접근할 권한이 없습니다")
    _tenant_id(user)
    try:
        contract = _require_hr_record(_find(_read_hr("contracts", user), contract_id), user, detail="")
    except HTTPException as exc:
        if exc.status_code == 404:
            raise denied from None
        raise
    email = _email(user)
    admin = _is_admin(user)
    is_party = bool(email) and email == str(contract.get("employee_email") or "").strip().lower()
    if not admin and not is_party:
        raise denied
    if str(contract.get("status") or "") != "signed":
        raise HTTPException(status_code=409, detail="서명 완료 계약서만 PDF 를 내려받을 수 있습니다")
    path = _verified_signed_pdf_path(contract)
    if path is None:
        stored = _store_signed_contract_pdf(contract, user)
        path = _verified_signed_pdf_path(contract) if stored.get("status") == "stored" else None
        if path is None:
            raise HTTPException(status_code=503, detail="서명본 PDF 를 준비하지 못했습니다. 관리자에게 재생성을 요청하십시오")
    if is_party and not admin:
        # 교부 기록: 최초 교부 시각은 덮어쓰지 않고, 매 다운로드는 이력에 남긴다.
        if not contract.get("delivered_at"):
            contract["delivered_at"] = _now()
            contract["delivery_channel"] = "download"
            try:
                _persist_contract_delivery(contract, user)
            except Exception as exc:  # noqa: BLE001
                logger.error("contract delivery not recorded: contract=%s err=%s", contract.get("id"), exc)
        try:
            from app.services.yeoljeong_contract_notify import mask_email

            _append_contract_notification_logs(
                contract, "delivered", [{"channel": "download", "status": "sent", "target_masked": mask_email(email)}]
            )
        except Exception as exc:  # noqa: BLE001
            logger.warning("contract delivery log failed: contract=%s err=%s", contract.get("id"), exc)
    title = str(contract.get("print_title") or "계약서")
    filename = f"{_safe_filename(title)}_{str(contract.get('employee_name') or '').strip() or 'signed'}_서명본.pdf"
    return path, _safe_filename(filename)


# ---------------------------------------------------------------------------
# 서명 완료 계약 → 직원 고용조건 스냅샷 → 급여 기준값
#
# 원본은 서명 완료 계약서 하나뿐이다. 직원 레코드의 current_employment 는
# 최신 서명 근로·용역 계약을 가리키는 파생 스냅샷이고, 계약서에서 언제든 다시
# 만들 수 있다(resync). 서명 완료 계약서는 불변이므로 계약서 체인이 이력이고,
# 스냅샷 교체는 yeoljeong_audit_logs 에 남긴다. 급여는 기본값만 채운다 —
# 관리자가 보낸 값이 항상 우선이고, 시급·일급제 월 총액은 추정하지 않는다.
# ---------------------------------------------------------------------------
EMPLOYMENT_SOURCE_CONTRACT_TYPES = EMPLOYMENT_CONTRACT_TYPES | {"freelancer"}
EMPLOYMENT_SNAPSHOT_FIELDS = (
    ("contract_type", "contractType"),
    ("employment_tax_type", "employmentTaxType"),
    ("wage", "wage"),
    ("wage_type", "wageType"),
    ("base_salary", "baseSalary"),
    ("non_tax_meal_allowance", "nonTaxMealAllowance"),
    ("taxable_allowance", "taxableAllowance"),
    ("meal_provision", "mealProvision"),
    ("start_date", "startDate"),
    ("end_date", "endDate"),
    ("work_days", "workDays"),
    ("work_time", "workTime"),
    ("rest_time", "restTime"),
    ("weekly_hours", "weeklyHours"),
    ("pay_date", "payDate"),
    ("pay_method", "payMethod"),
    ("workplace", "workplace"),
    ("business_id", "businessId"),
    ("branch", "branch"),
)
EMPLOYMENT_AMOUNT_FIELDS = {"wage", "base_salary", "non_tax_meal_allowance", "taxable_allowance"}
EMPLOYMENT_AUDIT_RESOURCE = "employee_employment"
PAYROLL_AUDIT_RESOURCE = "payroll_statement"
EMPLOYMENT_AUDIT_LOG = "employment_audit_logs"
PAYROLL_DEFAULT_FIELDS = ("gross_pay", "taxable_pay", "non_tax_meal_allowance", "meal_provision", "employment_tax_type")
PAYROLL_AMOUNT_DEFAULT_FIELDS = {"gross_pay", "taxable_pay", "non_tax_meal_allowance"}


def _employment_amount(value: Any) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        return int(round(float(str(value).replace(",", "").strip())))
    except (TypeError, ValueError):
        return None


def _contract_is_live_signed(contract: dict[str, Any]) -> bool:
    return str(contract.get("status") or "") == "signed" and not contract.get("deleted_at")


def _contract_belongs_to(contract: dict[str, Any], *, employee_email: str, employee_request_id: str) -> bool:
    request_id = str(_contract_payload_value(contract, "employee_request_id", "employeeRequestId") or "").strip()
    email = str(contract.get("employee_email") or "").strip().lower()
    if employee_request_id and request_id:
        return request_id == employee_request_id
    return bool(employee_email) and email == employee_email


def _signed_contract_sort_key(contract: dict[str, Any]) -> tuple[datetime, str]:
    stamp = _pg_ts(contract.get("signed_at")) or datetime.min.replace(tzinfo=KST)
    return stamp, str(contract.get("id") or "")


def _employee_signed_contracts(
    contracts: list[dict[str, Any]], *, employee_email: str, employee_request_id: str, business_id: str = ""
) -> list[dict[str, Any]]:
    """직원의 서명 완료 계약서(삭제 제외)를 서명 시각 오름차순으로. business_id 가 있으면 그 사업자 계약만."""
    email = str(employee_email or "").strip().lower()
    request_id = str(employee_request_id or "").strip()
    rows = [
        row
        for row in contracts
        if _contract_is_live_signed(row)
        and _contract_belongs_to(row, employee_email=email, employee_request_id=request_id)
        and _row_in_business(row, business_id)
    ]
    return sorted(rows, key=_signed_contract_sort_key)


def _employment_terms_from_contract(contract: dict[str, Any]) -> dict[str, Any]:
    # 봉인 스냅샷이 있으면 그것을 읽는다 — 서명 이후 붙은 메타가 섞이지 않는다.
    source = contract.get("signed_snapshot") if isinstance(contract.get("signed_snapshot"), dict) else contract
    terms: dict[str, Any] = {}
    for snake, camel in EMPLOYMENT_SNAPSHOT_FIELDS:
        value = _contract_payload_value(source, snake, camel)
        if snake in EMPLOYMENT_AMOUNT_FIELDS:
            terms[snake] = _employment_amount(value)
        else:
            terms[snake] = str(value).strip() if value is not None else ""
    return terms


def _derive_current_employment(
    contracts: list[dict[str, Any]], *, employee_email: str, employee_request_id: str, business_id: str = ""
) -> dict[str, Any] | None:
    """최신 서명 근로·용역 계약 1건에서 고용조건을 만든다. 비밀유지 서약은 원본이 아니다.

    business_id 가 있으면 그 사업자의 서명계약 중 최신 1건(겸직 직원의 매장별 고용조건), 없으면 전체 최신 1건.
    """
    sources = [
        row
        for row in _employee_signed_contracts(
            contracts,
            employee_email=employee_email,
            employee_request_id=employee_request_id,
            business_id=business_id,
        )
        if str(row.get("contract_type") or "") in EMPLOYMENT_SOURCE_CONTRACT_TYPES
    ]
    if not sources:
        return None
    latest = sources[-1]
    employment = _employment_terms_from_contract(latest)
    return {
        "contract_id": str(latest.get("id") or ""),
        "signed_at": str(latest.get("signed_at") or ""),
        "effective_from": employment.get("start_date") or str(latest.get("signed_at") or ""),
        "employment": employment,
    }


def _employment_snapshot_is_current(employee: dict[str, Any], derived: dict[str, Any] | None) -> bool:
    if derived is None:
        return True
    return (
        str(employee.get("current_employment_contract_id") or "") == derived["contract_id"]
        and employee.get("current_employment") == derived["employment"]
        and not employee.get("current_employment_error")
    )


def _append_employment_audit(
    *,
    tenant_id: str,
    business_id: str,
    actor: str,
    action: str,
    resource_type: str,
    resource_id: str,
    details: dict[str, Any],
) -> bool:
    """감사로그 한 줄. DB 는 ops 서비스의 _insert_audit_log_conn 을 재사용한다. 예외를 올리지 않는다."""
    entry_details = {**details, "tenant_id": tenant_id}
    try:
        if _db_available():
            return bool(_run_db(_db_insert_employment_audit(
                business_id=business_id, actor=actor, action=action,
                resource_type=resource_type, resource_id=resource_id, details=entry_details,
            )))
        row = {
            "id": f"aud-{uuid4().hex[:12]}",
            "business_id": business_id,
            "actor": actor,
            "action": action,
            "resource_type": resource_type,
            "resource_id": resource_id,
            "details": entry_details,
            "created_at": _now(),
        }
        _write_file_rows(EMPLOYMENT_AUDIT_LOG, [row] + _read_file_rows(EMPLOYMENT_AUDIT_LOG))
        return True
    except Exception as exc:  # noqa: BLE001 — 감사로그 실패가 서명·급여 저장을 되돌리면 안 된다
        logger.error("employment audit not recorded: action=%s resource=%s err=%s", action, resource_id, exc)
        return False


async def _db_insert_employment_audit(**kwargs: Any) -> dict[str, Any]:
    import asyncpg

    from app.services.yeoljeong_ops_service import _insert_audit_log_conn

    conn = await asyncpg.connect(_db_url(), timeout=5)
    try:
        return await _insert_audit_log_conn(conn, **kwargs)
    finally:
        await conn.close()


def _employment_audit_rows(tenant_id: str, business_id: str, request_id: str) -> list[dict[str, Any]]:
    if _db_available():
        from app.services import yeoljeong_ops_service as ops

        fetched = _run_db(ops.list_audit_logs(business_id, resource_type=EMPLOYMENT_AUDIT_RESOURCE, limit=1000))
        rows = fetched if isinstance(fetched, list) else []
    else:
        rows = [row for row in _read_file_rows(EMPLOYMENT_AUDIT_LOG) if row.get("business_id") == business_id]
    result = []
    for row in rows:
        details = _payload_dict(row.get("details"))
        if (
            row.get("resource_type") == EMPLOYMENT_AUDIT_RESOURCE
            and str(row.get("resource_id") or "") == request_id
            and str(details.get("tenant_id") or "") == tenant_id
        ):
            result.append({**row, "details": details, "created_at": _iso(row.get("created_at"))})
    return result


def _find_employee_record(
    user: dict[str, Any] | None, *, employee_email: str = "", employee_request_id: str = ""
) -> dict[str, Any] | None:
    email = str(employee_email or "").strip().lower()
    request_id = str(employee_request_id or "").strip()
    rows = [row for row in _read_hr("employee_join_requests", user) if str(row.get("status") or "").lower() == "approved"]
    if request_id:
        found = _find(rows, request_id)
        if found:
            return found
    if email:
        return next((row for row in rows if str(row.get("email") or "").strip().lower() == email), None)
    return None


def _employment_changes(before: dict[str, Any], after: dict[str, Any]) -> dict[str, dict[str, Any]]:
    keys = sorted(set(before) | set(after))
    return {key: {"before": before.get(key), "after": after.get(key)} for key in keys if before.get(key) != after.get(key)}


def _sync_employee_employment(
    employee: dict[str, Any], user: dict[str, Any] | None, *, actor: str, trigger: str
) -> dict[str, Any]:
    """계약서에서 스냅샷을 다시 만든다. 바뀐 게 없으면 쓰지 않는다."""
    tenant_id = _tenant_id(user)
    request_id = str(employee.get("id") or "")
    derived = _derive_current_employment(
        _read_hr("contracts", user),
        employee_email=str(employee.get("email") or ""),
        employee_request_id=request_id,
        business_id=_record_business_id(employee),
    )
    if derived is None:
        # 원본이 없으면 기존 스냅샷도 지우지 않는다(이력 보존).
        return {"status": "no_source", "employee_request_id": request_id}
    if _employment_snapshot_is_current(employee, derived):
        return {"status": "unchanged", "employee_request_id": request_id, "contract_id": derived["contract_id"]}

    previous_contract_id = str(employee.get("current_employment_contract_id") or "")
    previous = employee.get("current_employment") if isinstance(employee.get("current_employment"), dict) else {}
    now = _now()
    record = dict(employee)
    record.update(
        {
            "current_employment": derived["employment"],
            "current_employment_contract_id": derived["contract_id"],
            "current_employment_signed_at": derived["signed_at"],
            "current_employment_effective_from": derived["effective_from"],
            "current_employment_synced_at": now,
            "current_employment_error": "",
            "updated_at": now,
        }
    )
    if previous_contract_id and previous_contract_id != derived["contract_id"]:
        record["previous_employment_contract_id"] = previous_contract_id
    _write_hr_record("employee_join_requests", record, user)
    employee.clear()
    employee.update(record)

    if trigger == "resync":
        action = "employment.snapshot_resynced"
    elif previous_contract_id and previous_contract_id != derived["contract_id"]:
        action = "employment.snapshot_replaced"
    else:
        action = "employment.snapshot_created"
    audited = _append_employment_audit(
        tenant_id=tenant_id,
        business_id=_record_business_id(record) or str(record.get("business_id") or ""),
        actor=actor,
        action=action,
        resource_type=EMPLOYMENT_AUDIT_RESOURCE,
        resource_id=request_id,
        details={
            "trigger": trigger,
            "employee_email_masked": _mask_email(str(record.get("email") or "")),
            "previous_contract_id": previous_contract_id,
            "contract_id": derived["contract_id"],
            "contract_signed_at": derived["signed_at"],
            "effective_from": derived["effective_from"],
            "changes": _employment_changes(previous, derived["employment"]),
        },
    )
    return {
        "status": "updated",
        "action": action,
        "employee_request_id": request_id,
        "contract_id": derived["contract_id"],
        "previous_contract_id": previous_contract_id,
        "audit_logged": audited,
    }


def _record_employment_sync_error(employee_ref: dict[str, str], user: dict[str, Any] | None, error: str) -> bool:
    try:
        employee = _find_employee_record(user, **employee_ref)
        if not employee:
            return False
        employee = dict(employee)
        employee["current_employment_error"] = error
        employee["current_employment_error_at"] = _now()
        _write_hr_record("employee_join_requests", employee, user)
        return True
    except Exception as exc:  # noqa: BLE001
        logger.error("employment sync error not recorded: employee=%s err=%s", employee_ref, exc)
        return False


def _sync_employment_after_signature(contract: dict[str, Any], user: dict[str, Any] | None) -> dict[str, Any]:
    """서명 직후 훅. 서명은 이미 성립했다 — 실패는 current_employment_error 로만 남긴다."""
    employee_ref = {
        "employee_email": str(contract.get("employee_email") or ""),
        "employee_request_id": str(_contract_payload_value(contract, "employee_request_id", "employeeRequestId") or ""),
    }
    if str(contract.get("contract_type") or "") not in EMPLOYMENT_SOURCE_CONTRACT_TYPES:
        return {"status": "skipped", "reason": "고용조건 원본이 아닌 계약 유형", "contract_type": contract.get("contract_type")}
    try:
        employee = _find_employee_record(user, **employee_ref)
        if not employee:
            raise LookupError("승인 직원 레코드를 찾을 수 없습니다")
        return _sync_employee_employment(
            employee, user, actor=str(contract.get("signer_email") or _email(user or {})), trigger="contract_signed"
        )
    except Exception as exc:  # noqa: BLE001
        error = f"{_now()} {type(exc).__name__}: {exc}"[:500]
        logger.error("employment snapshot failed: contract=%s err=%s", contract.get("id"), exc)
        return {"status": "failed", "error": error, "error_recorded": _record_employment_sync_error(employee_ref, user, error)}


def _employee_for_admin_or_403(request_id: str, user: dict[str, Any], *, detail: str) -> dict[str, Any]:
    _tenant_id(user)
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail=detail)
    employee = _find_employee_record(user, employee_request_id=request_id)
    if not employee or str(employee.get("id") or "") != str(request_id):
        # 다른 테넌트의 식별자 존재 여부를 드러내지 않는다.
        raise HTTPException(status_code=403, detail=detail)
    return employee


def resync_employee_employment(request_id: str, user: dict[str, Any]) -> dict[str, Any]:
    employee = _employee_for_admin_or_403(request_id, user, detail="고용조건 재동기화 권한이 없습니다")
    result = _sync_employee_employment(employee, user, actor=_email(user), trigger="resync")
    return {
        "result": result,
        "current_employment": employee.get("current_employment"),
        "current_employment_contract_id": employee.get("current_employment_contract_id") or "",
        "current_employment_effective_from": employee.get("current_employment_effective_from") or "",
        "current_employment_synced_at": employee.get("current_employment_synced_at") or "",
        "current_employment_error": employee.get("current_employment_error") or "",
    }


def employee_employment_history(request_id: str, user: dict[str, Any]) -> dict[str, Any]:
    """관리자와 본인만. 서명 계약서 체인 + 스냅샷 감사로그를 시간순으로."""
    denied = HTTPException(status_code=403, detail="이 직원의 고용 이력에 접근할 권한이 없습니다")
    tenant_id = _tenant_id(user)
    employee = _find_employee_record(user, employee_request_id=request_id)
    if not employee or str(employee.get("id") or "") != str(request_id):
        raise denied
    email = str(employee.get("email") or "").strip().lower()
    if not _is_admin(user) and (not _email(user) or _email(user) != email):
        raise denied
    contracts = []
    for row in _employee_signed_contracts(
        _read_hr("contracts", user),
        employee_email=email,
        employee_request_id=request_id,
        business_id=_record_business_id(employee),
    ):
        terms = _employment_terms_from_contract(row)
        contracts.append(
            {
                "contract_id": str(row.get("id") or ""),
                "contract_type": str(row.get("contract_type") or ""),
                "print_title": str(row.get("print_title") or ""),
                "signed_at": str(row.get("signed_at") or ""),
                "employment_source": str(row.get("contract_type") or "") in EMPLOYMENT_SOURCE_CONTRACT_TYPES,
                "signed_snapshot_sha256": str(row.get("signed_snapshot_sha256") or ""),
                "terms": terms,
            }
        )
    business_id = _record_business_id(employee) or str(employee.get("business_id") or "")
    audit_logs = _employment_audit_rows(tenant_id, business_id, str(employee.get("id") or ""))
    timeline = [
        {"at": item["signed_at"], "kind": "contract_signed", "contract_id": item["contract_id"],
         "contract_type": item["contract_type"], "employment_source": item["employment_source"]}
        for item in contracts
    ] + [
        {"at": row["created_at"], "kind": "audit", "action": row.get("action"), "actor": row.get("actor"),
         "contract_id": row["details"].get("contract_id"), "previous_contract_id": row["details"].get("previous_contract_id")}
        for row in audit_logs
    ]
    timeline.sort(key=lambda item: _pg_ts(item["at"]) or datetime.min.replace(tzinfo=KST))
    return {
        "employee_request_id": str(employee.get("id") or ""),
        "current_employment": employee.get("current_employment"),
        "current_employment_contract_id": employee.get("current_employment_contract_id") or "",
        "current_employment_effective_from": employee.get("current_employment_effective_from") or "",
        "contracts": contracts,
        "audit_logs": sorted(audit_logs, key=lambda row: _pg_ts(row["created_at"]) or datetime.min.replace(tzinfo=KST)),
        "timeline": timeline,
    }


def _payroll_defaults_from_employment(derived: dict[str, Any] | None) -> dict[str, Any]:
    """급여 기본값과 근거. 시급·일급제는 월 총액을 추정하지 않는다."""
    if derived is None:
        return {
            "source_contract_id": "",
            "wage_type": "",
            "contract_type": "",
            "estimated": False,
            "defaults": {},
            "contract_terms": {},
            "basis": "",
            "reason": "서명 완료된 근로·용역 계약서가 없어 기본값이 없습니다",
        }
    terms = derived["employment"]
    contract_type = terms.get("contract_type") or ""
    wage_type = terms.get("wage_type") or ""
    wage = terms.get("wage")
    result: dict[str, Any] = {
        "source_contract_id": derived["contract_id"],
        "wage_type": wage_type,
        "contract_type": contract_type,
        "estimated": False,
        "defaults": {},
        "contract_terms": {
            "wage": wage,
            "weekly_hours": terms.get("weekly_hours") or "",
            "work_days": terms.get("work_days") or "",
            "work_time": terms.get("work_time") or "",
            "start_date": terms.get("start_date") or "",
            "end_date": terms.get("end_date") or "",
            "pay_date": terms.get("pay_date") or "",
        },
        "basis": "",
        "reason": "",
    }
    defaults: dict[str, Any] = result["defaults"]
    if contract_type == "freelancer" or wage_type == "case_fee":
        defaults["employment_tax_type"] = "freelancer_33"
        result["basis"] = f"프리랜서 용역계약({derived['contract_id']}) — 3.3% 원천징수 구분만 채움"
        result["reason"] = "용역비는 건별로 확정되므로 총지급액 기본값이 없습니다"
        return result
    if terms.get("employment_tax_type"):
        defaults["employment_tax_type"] = terms["employment_tax_type"]
    if terms.get("meal_provision"):
        defaults["meal_provision"] = terms["meal_provision"]
    if wage_type == "monthly" and wage:
        non_tax = terms.get("non_tax_meal_allowance") or 0
        defaults["gross_pay"] = wage
        defaults["taxable_pay"] = wage - non_tax
        defaults["non_tax_meal_allowance"] = non_tax
        result["basis"] = (
            f"월급제 계약({derived['contract_id']}) 임금 {wage:,}원 = 과세 {wage - non_tax:,}원 + 비과세 식대 {non_tax:,}원"
        )
        return result
    if wage_type == "hourly":
        result["hourly_wage"] = wage
        result["contract_terms"]["hourly_wage"] = wage
    elif wage_type == "daily":
        result["daily_wage"] = wage
        result["contract_terms"]["daily_wage"] = wage
    result["weekly_hours"] = terms.get("weekly_hours") or ""
    result["reason"] = "시급제는 근태 확정 후 산정" if wage_type == "hourly" else "일급제는 근태 확정 후 산정"
    result["basis"] = f"{wage_type} 계약({derived['contract_id']}) 단가 {wage or 0:,}원 — 월 총액은 근태 확정 후 관리자가 입력"
    return result


def payroll_defaults(employee_email: str, payroll_month: str, user: dict[str, Any]) -> dict[str, Any]:
    denied = HTTPException(status_code=403, detail="급여 기준값 조회 권한이 없습니다")
    _tenant_id(user)
    if not _is_admin(user):
        raise denied
    month = str(payroll_month or "").strip()
    if month and not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", month):
        raise HTTPException(status_code=400, detail="payroll_month 는 YYYY-MM 형식이어야 합니다")
    employee = _find_employee_record(user, employee_email=employee_email)
    if not employee:
        raise denied
    derived = _derive_current_employment(
        _read_hr("contracts", user),
        employee_email=str(employee.get("email") or ""),
        employee_request_id=str(employee.get("id") or ""),
        business_id=_record_business_id(employee),
    )
    result = _payroll_defaults_from_employment(derived)
    result["employee_email"] = str(employee.get("email") or "")
    result["employee_request_id"] = str(employee.get("id") or "")
    result["payroll_month"] = month
    result["employment_snapshot_in_sync"] = _employment_snapshot_is_current(employee, derived)
    if derived and month:
        start = str(derived["employment"].get("start_date") or "")[:7]
        end = str(derived["employment"].get("end_date") or "")[:7]
        if (start and month < start) or (end and month > end):
            result["period_note"] = "급여 월이 계약 기간 밖입니다"
    return result


def _payroll_value_missing(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def _apply_payroll_contract_defaults(
    payload: dict[str, Any], plan: dict[str, Any]
) -> tuple[dict[str, Any], list[str], dict[str, dict[str, Any]]]:
    """비어 있는 칸만 채운다. 관리자가 보낸 값은 그대로 두고, 계약과 다르면 이탈로 기록한다."""
    result = dict(payload)
    defaults = plan.get("defaults") or {}
    applied: list[str] = []
    deviation: dict[str, dict[str, Any]] = {}
    for field in ("meal_provision", "employment_tax_type"):
        if field not in defaults:
            continue
        if _payroll_value_missing(result.get(field)):
            result[field] = defaults[field]
            applied.append(field)
        elif str(result.get(field)).strip() != str(defaults[field]):
            deviation[field] = {"contract": defaults[field], "input": result.get(field)}
    if "gross_pay" in defaults:
        if _payroll_value_missing(result.get("gross_pay")):
            result["gross_pay"] = defaults["gross_pay"]
            applied.append("gross_pay")
        # 과세/비과세 분할은 둘 다 비었고 총지급액이 계약 임금일 때만 채운다 —
        # 관리자가 총액을 바꿨는데 계약 분할을 끼우면 합계 검증이 깨진다.
        gross_matches = _employment_amount(result.get("gross_pay")) == defaults["gross_pay"]
        if gross_matches and all(_payroll_value_missing(result.get(key)) for key in ("taxable_pay", "non_tax_meal_allowance")):
            result["taxable_pay"] = defaults["taxable_pay"]
            result["non_tax_meal_allowance"] = defaults["non_tax_meal_allowance"]
            applied.extend(["taxable_pay", "non_tax_meal_allowance"])
        for field in ("gross_pay", "taxable_pay", "non_tax_meal_allowance"):
            if field in applied or _payroll_value_missing(result.get(field)):
                continue
            if _employment_amount(result.get(field)) != defaults[field]:
                deviation[field] = {"contract": defaults[field], "input": result.get(field)}
    return result, applied, deviation


def list_payroll(user: dict[str, Any]) -> list[dict[str, Any]]:
    return sorted(_filter_user(_read_hr("payroll_statements", user), user, "employee_email"), key=lambda row: row.get("updated_at", ""), reverse=True)


def save_payroll(payload: dict[str, Any], user: dict[str, Any]) -> dict[str, Any]:
    _tenant_id(user)
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="급여내역서 작성 권한이 없습니다")
    rows = _read_hr("payroll_statements", user)
    employee = _find_employee_record(
        user,
        employee_email=str(payload.get("employee_email") or ""),
        employee_request_id=str(payload.get("employee_request_id") or ""),
    )
    derived = (
        _derive_current_employment(
            _read_hr("contracts", user),
            employee_email=str(employee.get("email") or ""),
            employee_request_id=str(employee.get("id") or ""),
            business_id=_record_business_id(employee),
        )
        if employee
        else None
    )
    plan = _payroll_defaults_from_employment(derived)
    payload, defaults_applied, deviation = _apply_payroll_contract_defaults(payload, plan)
    gross = _payroll_integer_for_api(payload.get("gross_pay"), field="gross_pay")
    taxable_pay = _payroll_integer_for_api(payload.get("taxable_pay"), field="taxable_pay")
    non_tax_meal = _payroll_integer_for_api(payload.get("non_tax_meal_allowance"), field="non_tax_meal_allowance")
    if taxable_pay or non_tax_meal:
        if taxable_pay + non_tax_meal != gross:
            raise HTTPException(status_code=400, detail="과세급여와 비과세 식대 합계가 총지급액과 일치해야 합니다")
        if non_tax_meal < 0 or non_tax_meal > 200000:
            raise HTTPException(status_code=400, detail="비과세 식대는 월 200,000원 이내로 입력하십시오")
        if non_tax_meal > 0 and str(payload.get("meal_provision") or "") != "cash_no_meal":
            raise HTTPException(status_code=400, detail="사용자가 식사를 제공하는 경우 현금 식대를 비과세로 분류할 수 없습니다")
    tax_withholding = _payroll_integer_for_api(payload.get("tax_withholding"), field="tax_withholding")
    insurance_deduction = _payroll_integer_for_api(payload.get("insurance_deduction"), field="insurance_deduction")
    other_deduction = _payroll_integer_for_api(payload.get("other_deduction"), field="other_deduction")
    deductions = tax_withholding + insurance_deduction + other_deduction
    now = _now()
    statement_id = str(payload.get("id") or uuid4())
    email = str(payload.get("employee_email") or "").strip().lower()
    statement = {
        **payload,
        "id": statement_id,
        "employee_email": email,
        "employee_email_masked": _mask_email(email),
        "gross_pay": gross,
        "taxable_pay": taxable_pay,
        "non_tax_meal_allowance": non_tax_meal,
        "tax_withholding": tax_withholding,
        "insurance_deduction": insurance_deduction,
        "other_deduction": other_deduction,
        "net_pay": max(0, gross - deductions),
        "source_contract_id": plan["source_contract_id"],
        "contract_defaults_applied": defaults_applied,
        "contract_deviation": deviation,
        "created_at": payload.get("created_at") or now,
        "updated_at": now,
    }
    statement = _owned_hr_record(statement, user)
    existing = _find(rows, statement_id)
    if existing:
        existing.update(statement)
        saved = existing
    else:
        rows.insert(0, statement)
        saved = statement
    _write_hr_record("payroll_statements", saved, user)
    for action, key, value in (
        ("payroll.contract_defaults_applied", "fields", defaults_applied),
        ("payroll.contract_deviation", "deviation", deviation),
    ):
        if value:
            _append_employment_audit(
                tenant_id=str(saved.get("tenant_id") or ""),
                business_id=str(saved.get("business_id") or ""),
                actor=_email(user),
                action=action,
                resource_type=PAYROLL_AUDIT_RESOURCE,
                resource_id=statement_id,
                details={
                    key: value,
                    "source_contract_id": plan["source_contract_id"],
                    "employee_request_id": str((employee or {}).get("id") or ""),
                    "payroll_month": str(saved.get("payroll_month") or ""),
                },
            )
    return saved


def delete_payroll(statement_id: str, user: dict[str, Any]) -> None:
    _tenant_id(user)
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="급여내역서 삭제 권한이 없습니다")
    _require_hr_record(_find(_read_hr("payroll_statements", user), statement_id), user, detail="급여내역서를 찾을 수 없습니다")
    _delete_hr_record("payroll_statements", statement_id, user)


def _decode_csv(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp949", "euc-kr"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise HTTPException(status_code=400, detail="CSV 문자 인코딩을 확인할 수 없습니다")


def _csv_delimiter(text: str) -> str:
    sample_lines = [line for line in text.splitlines()[:30] if line.strip()]
    counts = {delimiter: max((line.count(delimiter) for line in sample_lines), default=0) for delimiter in (",", "\t", ";")}
    return max(counts, key=counts.get)


def _first_present(row: dict[str, str], *keys: str) -> str:
    for key in keys:
        value = row.get(key)
        if value not in (None, ""):
            return value
    return ""


def _amount(value: Any) -> int:
    cleaned = re.sub(r"[^0-9.-]", "", str(value or ""))
    if cleaned in {"", "-", ".", "-."}:
        return 0
    try:
        return int(round(float(cleaned)))
    except ValueError:
        return 0


def _transaction_date(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    normalized = text.replace("/", "-").replace(".", "-")
    normalized = re.sub(r"\s+", " ", normalized)
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            parsed = datetime.strptime(normalized, fmt)
            return parsed.strftime("%Y-%m-%d %H:%M:%S") if "%H" in fmt else parsed.strftime("%Y-%m-%d")
        except ValueError:
            continue
    return text


def _transaction_category(text: str) -> str:
    lowered = str(text or "").lower()
    for category, keywords in DEFAULT_CATEGORY_RULES:
        if any(keyword.lower() in lowered for keyword in keywords):
            return category
    return "미분류"


def list_transactions() -> list[dict[str, Any]]:
    return sorted(_read("transactions"), key=lambda row: str(row.get("transaction_date") or ""), reverse=True)


def list_transactions_for_user(user: dict[str, Any]) -> list[dict[str, Any]]:
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="거래 원장 조회 권한이 없습니다")
    return list_transactions()


def create_transaction(payload: dict[str, Any]) -> dict[str, Any]:
    rows = _read("transactions")
    now = _now()
    record = {
        **payload,
        "id": str(payload.get("id") or uuid4()),
        "transaction_date": _transaction_date(payload.get("transaction_date")),
        "amount": _amount(payload.get("amount")),
        "created_at": payload.get("created_at") or now,
        "updated_at": now,
    }
    rows.insert(0, record)
    _write("transactions", rows)
    return record


def import_file(
    filename: str,
    content: bytes,
    source_type: str,
    *,
    business_id: str = "",
    branch: str = "",
    service: str = "",
    source_account_id: str = "",
) -> dict[str, Any]:
    normalized_source = str(source_type or "other").strip().lower()
    if normalized_source not in {"bank", "card", "other"}:
        raise HTTPException(status_code=400, detail="source_type은 bank, card, other 중 하나여야 합니다")
    normalized_business = str(business_id or "").strip()
    normalized_branch = BRANCH_ALIASES.get(str(branch or "").strip(), str(branch or "").strip())
    normalized_service = str(service or "").strip()
    decoded = _decode_csv(content)
    reader = csv.DictReader(decoded.splitlines(), delimiter=_csv_delimiter(decoded))
    existing = _read("transactions")
    existing_ids = {str(row.get("id") or "") for row in existing}
    imported: list[dict[str, Any]] = []
    duplicate_rows = 0
    now = _now()
    for source_row in reader:
        raw = {str(key or "").strip(): str(value or "").strip() for key, value in source_row.items()}
        if not any(raw.values()):
            continue
        fingerprint = json.dumps(
            {
                "source_type": normalized_source,
                "service": normalized_service,
                "business_id": normalized_business,
                "branch": normalized_branch,
                "row": raw,
            },
            ensure_ascii=False,
            sort_keys=True,
        )
        record_id = hashlib.sha256(fingerprint.encode("utf-8")).hexdigest()
        if record_id in existing_ids:
            duplicate_rows += 1
            continue
        incoming = _amount(_first_present(raw, "입금액", "입금", "입금금액", "맡기신금액", "받으신금액"))
        outgoing = _amount(_first_present(raw, "출금액", "출금", "출금금액", "찾으신금액", "지급금액"))
        amount = incoming or outgoing or _amount(
            raw.get("합계금액") or raw.get("결제금액") or raw.get("거래금액") or raw.get("금액")
        )
        direction = "income" if incoming > 0 else "expense"
        description = (
            raw.get("상품명")
            or raw.get("적요")
            or raw.get("거래내용")
            or raw.get("내용")
            or raw.get("기재내용")
            or raw.get("보낸분/받는분")
            or raw.get("보낸분")
            or raw.get("받는분")
            or raw.get("거래처")
            or raw.get("판매자상호")
            or ""
        )
        transaction_datetime = (
            raw.get("거래일시")
            or " ".join(
                item
                for item in (raw.get("거래일자") or raw.get("거래일") or raw.get("일자") or "", raw.get("거래시간") or "")
                if item
            )
        )
        searchable = " ".join([description, *raw.values()])
        record = {
            "id": record_id,
            "source_type": normalized_source,
            "service": normalized_service,
            "business_id": normalized_business,
            "branch": normalized_branch,
            "source_account_id": str(source_account_id or ""),
            "source_file": Path(filename or "upload.csv").name,
            "transaction_date": _transaction_date(transaction_datetime),
            "description": description,
            "amount": amount,
            "direction": direction,
            "category": _transaction_category(searchable),
            "approval_number": raw.get("승인번호") or "",
            "order_number": raw.get("주문번호") or "",
            "account_name": raw.get("계좌명") or raw.get("계좌번호") or raw.get("계좌") or "",
            "created_at": now,
            "updated_at": now,
        }
        imported.append(record)
        existing_ids.add(record_id)
    if imported:
        _write("transactions", imported + existing)
    return {
        "import": {
            "filename": Path(filename or "upload.csv").name,
            "source_type": normalized_source,
            "imported_rows": len(imported),
            "duplicate_rows": duplicate_rows,
        },
        "rows": imported,
    }


def list_accounts(user: dict[str, Any], business_id: str | None = None) -> list[dict[str, Any]]:
    if not _is_admin(user):
        return []
    rows = _read("platform_accounts")
    changed = _migrate_platform_account_secrets(rows)
    changed = bool(_hydrate_delivery_account_passwords_from_agent_vault(rows)) or changed
    now = _now()
    for row in rows:
        public_status = _public_platform_account_status(row)
        raw_sync_status = str(row.get("last_sync_status") or row.get("portal_status") or "").strip()
        if raw_sync_status == "running" and public_status != "running":
            row["last_sync_status"] = public_status
            row["portal_status"] = public_status
            row["portal_message"] = (
                row.get("portal_message")
                or "이전 연동 실행이 완료 응답 없이 종료되어 연결 상태 확인이 필요합니다."
            )
            row["updated_at"] = now
            changed = True
    if changed:
        _write("platform_accounts", rows)
    result = []
    for row in rows:
        row_business = str(row.get("business_id") or "")
        if business_id and row_business != business_id:
            continue
        item = {k: v for k, v in row.items() if k not in _ACCOUNT_SECRET_FIELDS}
        item["branch"] = BRANCH_ALIASES.get(str(item.get("branch") or ""), str(item.get("branch") or ""))
        item["password_masked"] = "********" if _has_account_secret(row) else ""
        item["status"] = _public_platform_account_status(row)
        item["credential_requirements"] = _missing_connector_requirements(row)
        result.append(item)
    return result


def get_settings(user: dict[str, Any]) -> dict[str, Any]:
    data = _read_json_object("settings")
    ui_settings = data.get("ui_settings")
    if not isinstance(ui_settings, dict):
        ui_settings = {}
    ui_settings = _canonicalize_ui_settings(ui_settings)
    return {
        "settings": ui_settings,
        "meta": {
            "updated_at": data.get("ui_settings_updated_at") or "",
            "updated_by": data.get("ui_settings_updated_by") or "",
            "source": "server-file",
        },
    }


def save_settings(payload: dict[str, Any], user: dict[str, Any]) -> dict[str, Any]:
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="기초설정 저장 권한이 없습니다")
    settings = payload.get("settings") if isinstance(payload, dict) else payload
    if not isinstance(settings, dict):
        raise HTTPException(status_code=400, detail="settings 객체가 필요합니다")
    allowed = {"businesses", "branches", "accounts", "staff", "integrations"}
    raw = {
        key: value
        for key, value in settings.items()
        if key in allowed and isinstance(value, list)
    }
    cleaned = _canonicalize_ui_settings(raw)
    data = _read_json_object("settings")
    now = _now()
    data["ui_settings"] = cleaned
    data["ui_settings_updated_at"] = now
    data["ui_settings_updated_by"] = _email(user)
    _write_json_object("settings", data)
    return {"settings": cleaned, "meta": {"updated_at": now, "updated_by": _email(user), "source": "server-file"}}


async def get_settings_persisted(user: dict[str, Any]) -> dict[str, Any]:
    pool = _get_pool_or_none()
    if pool is None:
        return get_settings(user)
    try:
        async with pool.acquire() as conn:
            ready = await conn.fetchval("SELECT to_regclass('public.yeoljeong_businesses') IS NOT NULL")
            if not ready:
                return get_settings(user)
            business_rows = await conn.fetch(
                """
                SELECT id, entity_type, name, registration_no, representative, tax_type,
                       opened_at, address, memo
                  FROM yeoljeong_businesses
                 WHERE deleted_at IS NULL
                 ORDER BY sort_order, id
                """
            )
            branch_rows = await conn.fetch(
                """
                SELECT id, business_id, name, status, phone, address
                  FROM yeoljeong_branches
                 WHERE deleted_at IS NULL
                 ORDER BY sort_order, id
                """
            )
            extra_row = await conn.fetchrow(
                "SELECT data, updated_at, updated_by FROM yeoljeong_settings WHERE scope = 'ui'"
            )
        extra = _jsonb_object(extra_row["data"]) if extra_row else {}
        settings = {
            "businesses": [
                {
                    "id": row["id"],
                    "entityType": row["entity_type"],
                    "name": row["name"],
                    "registrationNo": row["registration_no"],
                    "representative": row["representative"],
                    "taxType": row["tax_type"],
                    "openedAt": row["opened_at"] or "",
                    "address": row["address"] or "",
                    "memo": row["memo"] or "",
                }
                for row in business_rows
            ],
            "branches": [
                {
                    "id": row["id"],
                    "businessId": row["business_id"],
                    "name": row["name"],
                    "status": row["status"],
                    "phone": row["phone"] or "",
                    "address": row["address"] or "",
                }
                for row in branch_rows
            ],
            "accounts": extra.get("accounts") if isinstance(extra.get("accounts"), list) else [],
            "staff": extra.get("staff") if isinstance(extra.get("staff"), list) else [],
            "integrations": extra.get("integrations") if isinstance(extra.get("integrations"), list) else [],
        }
        return {
            "settings": _canonicalize_ui_settings(settings),
            "meta": {
                "updated_at": extra_row["updated_at"].isoformat(timespec="seconds") if extra_row and extra_row["updated_at"] else "",
                "updated_by": extra_row["updated_by"] if extra_row else "",
                "source": "database",
            },
        }
    except Exception:
        return get_settings(user)


async def save_settings_persisted(payload: dict[str, Any], user: dict[str, Any]) -> dict[str, Any]:
    file_result = save_settings(payload, user)
    pool = _get_pool_or_none()
    if pool is None:
        return file_result
    settings = file_result["settings"]
    now = _now()
    updated_by = _email(user)
    try:
        async with pool.acquire() as conn:
            ready = await conn.fetchval("SELECT to_regclass('public.yeoljeong_businesses') IS NOT NULL")
            if not ready:
                return file_result
            async with conn.transaction():
                for sort_order, item in enumerate(settings["businesses"], start=1):
                    await conn.execute(
                        BUSINESS_SETTINGS_UPSERT_SQL,
                        item["id"],
                        item.get("entityType") or "individual",
                        item["name"],
                        _business_registration_value(item.get("registrationNo")),
                        _business_registration_value(item.get("representative")),
                        item.get("taxType") or "",
                        _business_registration_value(item.get("openedAt")),
                        _business_registration_value(item.get("address")),
                        item.get("memo") or "",
                        sort_order,
                        updated_by,
                    )
                for sort_order, item in enumerate(settings["branches"], start=1):
                    await conn.execute(
                        """
                        INSERT INTO yeoljeong_branches
                            (id, business_id, name, status, phone, address, sort_order, updated_by)
                        VALUES ($1, $2, $3, $4, $5, $6, $7, $8)
                        ON CONFLICT (id) DO UPDATE
                           SET business_id = EXCLUDED.business_id,
                               name = EXCLUDED.name,
                               status = EXCLUDED.status,
                               phone = EXCLUDED.phone,
                               address = EXCLUDED.address,
                               sort_order = EXCLUDED.sort_order,
                               updated_by = EXCLUDED.updated_by,
                               updated_at = NOW(),
                               deleted_at = NULL
                        """,
                        item["id"],
                        item["businessId"],
                        item["name"],
                        item.get("status") or "active",
                        item.get("phone") or "",
                        item.get("address") or "",
                        sort_order,
                        updated_by,
                    )
                business_ids = [str(item["id"]) for item in settings["businesses"]]
                branch_ids = [str(item["id"]) for item in settings["branches"]]
                await conn.execute(
                    """
                    UPDATE yeoljeong_branches
                       SET deleted_at = NOW(),
                           updated_at = NOW(),
                           updated_by = $2
                     WHERE deleted_at IS NULL
                       AND NOT (id = ANY($1::text[]))
                    """,
                    branch_ids,
                    updated_by,
                )
                await conn.execute(
                    """
                    UPDATE yeoljeong_businesses
                       SET deleted_at = NOW(),
                           updated_at = NOW(),
                           updated_by = $2
                     WHERE deleted_at IS NULL
                       AND NOT (id = ANY($1::text[]))
                    """,
                    business_ids,
                    updated_by,
                )
                extra = {
                    "accounts": settings["accounts"],
                    "staff": settings["staff"],
                    "integrations": settings["integrations"],
                }
                await conn.execute(
                    """
                    INSERT INTO yeoljeong_settings (scope, data, updated_by)
                    VALUES ('ui', $1::jsonb, $2)
                    ON CONFLICT (scope) DO UPDATE
                       SET data = EXCLUDED.data,
                           updated_by = EXCLUDED.updated_by,
                           updated_at = NOW()
                    """,
                    json.dumps(extra, ensure_ascii=False),
                    updated_by,
                )
        return {"settings": settings, "meta": {"updated_at": now, "updated_by": updated_by, "source": "database"}}
    except Exception:
        return file_result


def upsert_account(payload: dict[str, Any], user: dict[str, Any]) -> dict[str, Any]:
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="계정 등록 권한이 없습니다")
    rows = _read("platform_accounts")
    account_id = str(payload.get("account_id") or payload.get("server_account_id") or "").strip()
    service = str(payload.get("service") or "").strip()
    username = str(payload.get("username") or "").strip()
    if not service or not username:
        raise HTTPException(status_code=400, detail="연동 서비스와 아이디가 필요합니다")
    if service not in CONNECTOR_LABELS:
        raise HTTPException(status_code=400, detail="지원하지 않는 연동 서비스입니다")
    business_id, branch = _normalize_connector_scope(service, payload.get("business_id"), payload.get("branch"))
    existing = None
    if account_id:
        existing = next((row for row in rows if str(row.get("id") or "") == account_id), None)
        if not existing:
            raise HTTPException(status_code=404, detail="수정할 연동 계정을 찾지 못했습니다")
    if existing is None:
        existing = next(
            (
                row
                for row in rows
                if row.get("service") == service
                and row.get("username") == username
                and BRANCH_ALIASES.get(str(row.get("branch") or ""), str(row.get("branch") or "")) == branch
                and str(row.get("business_id") or "") == business_id
            ),
            None,
        )
    now = _now()
    record = existing or {"id": str(uuid4()), "created_at": now}
    collection_mode = str(payload.get("collection_mode") or "browser-automation").strip()
    if service in BANK_QUICK_SERVICE_CONFIG and collection_mode in {"bank-openbanking", "browser-automation"}:
        collection_mode = "bank-quick-service"
    secret_payload = False
    for plaintext_field, encrypted_field in _ACCOUNT_SECRET_FIELD_MAP.items():
        incoming_secret = str(payload.get(plaintext_field) or "")
        if incoming_secret:
            record[encrypted_field] = _encrypt_secret(incoming_secret)
            record.pop(plaintext_field, None)
            secret_payload = True
    if not secret_payload:
        _migrate_platform_account_secrets([record])
    if service in BANK_QUICK_SERVICE_CONFIG and collection_mode == "bank-quick-service":
        required_secret_fields = [
            ("조회용 계좌번호", "account_no"),
            ("계좌비밀번호", "account_password"),
            ("사업자번호", "business_registration_no"),
        ]
        if service == "shinhan_business":
            required_secret_fields.insert(0, ("로그인 비밀번호", "password"))
        missing = [label for label, key in required_secret_fields if not _has_secret_value(record, key)]
        if missing:
            raise HTTPException(status_code=400, detail=f"은행 간편/빠른조회 필수값을 확인하십시오: {', '.join(missing)}")
    bank_quick_config = BANK_QUICK_SERVICE_CONFIG.get(service) or {}
    account_no_masked = (
        str(payload.get("account_no_masked") or "").strip()
        or _masked_digits(payload.get("account_no"))
        or str(record.get("account_no_masked") or "").strip()
    )
    business_no_masked = (
        str(payload.get("business_registration_no_masked") or "").strip()
        or _masked_digits(payload.get("business_registration_no"))
        or str(record.get("business_registration_no_masked") or "").strip()
    )
    record.update(
        {
            "service": service,
            "label": payload.get("label") or CONNECTOR_LABELS.get(service, service),
            "login_url": _normalize_bank_quick_login_url(
                service,
                payload.get("login_url") or bank_quick_config.get("login_url") or "",
            ),
            "username": username,
            "business_id": business_id,
            "branch": branch,
            "institution_code": str(payload.get("institution_code") or "").strip(),
            "account_no_masked": account_no_masked,
            "business_registration_no_masked": business_no_masked,
            "merchant_no": str(payload.get("merchant_no") or "").strip(),
            "settlement_cycle": str(payload.get("settlement_cycle") or "").strip(),
            "collection_mode": collection_mode,
            "category": str(payload.get("category") or "").strip(),
            "data_scope": str(payload.get("data_scope") or "").strip(),
            "required_proof": str(payload.get("required_proof") or "").strip(),
            "auth_owner": str(payload.get("auth_owner") or "").strip(),
            "mfa_method": str(payload.get("mfa_method") or "").strip(),
            "credential_expires_at": str(payload.get("credential_expires_at") or "").strip(),
            "fallback_auth": str(payload.get("fallback_auth") or "").strip(),
            "sync_scope": str(payload.get("sync_scope") or "").strip(),
            "permission_scope": str(payload.get("permission_scope") or "").strip(),
            "failure_fallback": str(payload.get("failure_fallback") or "").strip(),
            "status": "credential_registered",
            "last_sync_status": record.get("last_sync_status") or "not_started",
            "auto_sync": bool(payload.get("auto_sync")),
            "memo": payload.get("memo") or bank_quick_config.get("enrollment") or "",
            "updated_at": now,
        }
    )
    if not existing:
        rows.insert(0, record)
    _write("platform_accounts", rows)
    public = {k: v for k, v in record.items() if k not in _ACCOUNT_SECRET_FIELDS}
    public["password_masked"] = "********" if _has_account_secret(record) else ""
    return public


def delete_account(account_id: str, user: dict) -> None:
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="외부계정 삭제 권한이 없습니다")
    k = str(account_id or "").strip()
    if not k:
        raise HTTPException(status_code=400, detail="삭제할 계정 ID가 없습니다")
    rows = _read("platform_accounts")
    if not any(str(r.get("id") or "") == k for r in rows):
        raise HTTPException(status_code=404, detail="삭제할 외부계정을 찾지 못했습니다")
    _delete("platform_accounts", k)


def list_settlements(user: dict[str, Any], business_id: str | None = None) -> list[dict[str, Any]]:
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="정산 원장 조회 권한이 없습니다")
    rows = _read("delivery_settlements")
    if business_id:
        rows = [row for row in rows if str(row.get("business_id") or "") == business_id]
    return rows


def list_sales(user: dict[str, Any], business_id: str | None = None) -> list[dict[str, Any]]:
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="매출 원장 조회 권한이 없습니다")
    rows = _read("delivery_sales")
    if business_id:
        rows = [row for row in rows if str(row.get("business_id") or "") == business_id]
    return rows


def list_reviews(user: dict[str, Any], business_id: str | None = None) -> list[dict[str, Any]]:
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="리뷰 원장 조회 권한이 없습니다")
    rows = _read("delivery_reviews")
    if business_id:
        rows = [row for row in rows if str(row.get("business_id") or "") == business_id]
    return rows


def list_ads(user: dict[str, Any], business_id: str | None = None) -> list[dict[str, Any]]:
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="광고 원장 조회 권한이 없습니다")
    rows = _read("delivery_ads")
    if business_id:
        rows = [row for row in rows if str(row.get("business_id") or "") == business_id]
    return rows


def delivery_completion_matrix(user: dict[str, Any], business_id: str | None = None) -> list[dict[str, Any]]:
    """Return completion per account/channel/ledger type."""
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="판매채널 완료 매트릭스 조회 권한이 없습니다")
    accounts = _read("platform_accounts")
    if business_id:
        accounts = [row for row in accounts if str(row.get("business_id") or "") == business_id]
    ledgers = {kind: _read(f"delivery_{kind}") for kind in DELIVERY_RECORD_TYPES}
    statuses = _read("delivery_collection_status")
    latest: dict[tuple[str, str, str], dict[str, Any]] = {}
    for row in sorted(statuses, key=lambda item: str(item.get("updated_at") or item.get("created_at") or ""), reverse=True):
        key = (
            str(row.get("service") or ""),
            str(row.get("business_id") or ""),
            str(row.get("branch") or ""),
        )
        latest.setdefault(key, row)
    matrix: list[dict[str, Any]] = []
    for account in accounts:
        service = str(account.get("service") or "")
        scope = (service, str(account.get("business_id") or ""), str(account.get("branch") or ""))
        status_row = latest.get(scope) or {}
        kinds: dict[str, str] = {}
        for kind, rows in ledgers.items():
            has_rows = any(
                str(row.get("service") or "") == service
                and str(row.get("business_id") or "") == scope[1]
                and str(row.get("branch") or "") == scope[2]
                for row in rows
            )
            if has_rows:
                kinds[kind] = "complete"
            elif str(status_row.get("status") or "") == "action_required":
                kinds[kind] = "action_required"
            else:
                kinds[kind] = "incomplete"
        matrix.append(
            {
                "account_id": str(account.get("id") or ""),
                "service": service,
                "business_id": scope[1],
                "branch": scope[2],
                "status": "complete"
                if all(value == "complete" for value in kinds.values())
                else ("action_required" if "action_required" in kinds.values() else "incomplete"),
                "kinds": kinds,
                "last_collection_status": str(status_row.get("status") or "not_started"),
                "run_id": str(status_row.get("id") or ""),
            }
        )
    return matrix


def list_collection_status(user: dict[str, Any], business_id: str | None = None) -> list[dict[str, Any]]:
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="수집 상태 조회 권한이 없습니다")
    rows = _read("delivery_collection_status")
    if _normalize_stale_delivery_collection_statuses(rows):
        stale_rows = [row for row in rows if str(row.get("error_code") or "") == "BACKGROUND_SYNC_STALE"]
        _write_delivery_collection_statuses(rows)
        for row in stale_rows:
            _run_db(_db_upsert_ledger("delivery_collection_status", row))
    if business_id:
        rows = [row for row in rows if str(row.get("business_id") or "") == business_id]
    return rows


def _normalize_stale_delivery_collection_statuses(rows: list[dict[str, Any]]) -> bool:
    now_dt = datetime.now(KST)
    now_text = now_dt.isoformat(timespec="seconds")
    changed = False
    for row in rows:
        if str(row.get("status") or "").strip() not in {"queued", "running"}:
            continue
        started_at = (
            _pg_ts(row.get("started_at"))
            or _pg_ts(row.get("queued_at"))
            or _pg_ts(row.get("updated_at"))
            or _pg_ts(row.get("created_at"))
        )
        stale_after = _delivery_stale_after_for_status(row)
        if not started_at or now_dt - started_at < stale_after:
            continue
        row["status"] = "failed"
        row["raw_status"] = "stale"
        row["error_code"] = "BACKGROUND_SYNC_STALE"
        row["message"] = (
            f"백그라운드 수집 작업이 {int(stale_after.total_seconds() // 60)}분 이상 "
            "완료 갱신 없이 멈춰 상태를 정리했습니다. 다시 수집 실행이 필요합니다."
        )
        row["finished_at"] = now_text
        row["updated_at"] = now_text
        row.setdefault("counts", _delivery_empty_counts())
        changed = True
    return changed


def _settle_stale_delivery_collection_statuses() -> None:
    statuses = _read("delivery_collection_status")
    if not _normalize_stale_delivery_collection_statuses(statuses):
        return
    _write_delivery_collection_statuses(statuses)
    for row in statuses:
        if str(row.get("error_code") or "") == "BACKGROUND_SYNC_STALE":
            _run_db(_db_upsert_ledger("delivery_collection_status", row))


def _delivery_sync_window(payload: dict[str, Any]) -> tuple[date, date]:
    today = datetime.now(KST).date()
    default_from = today.replace(day=1).isoformat()
    date_from_text = str(payload.get("date_from") or default_from)
    date_to_text = str(payload.get("date_to") or today.isoformat())
    try:
        date_from = date.fromisoformat(date_from_text)
        date_to = date.fromisoformat(date_to_text)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="수집 기간은 YYYY-MM-DD 형식이어야 합니다") from exc
    full_backfill = _delivery_full_backfill_requested(payload)
    if date_from > date_to or (not full_backfill and (date_to - date_from).days > 62):
        raise HTTPException(status_code=400, detail="수집 기간은 시작일 이후 최대 63일입니다")
    return date_from, date_to


def _delivery_requested_services(payload: dict[str, Any]) -> list[str]:
    from app.services.yeoljeong_delivery_collectors import PORTAL_CONFIG

    services = [str(item) for item in (payload.get("services") or []) if str(item).strip()]
    requested_services = services or sorted(PORTAL_CONFIG)
    unsupported = sorted(set(requested_services) - set(PORTAL_CONFIG))
    if unsupported:
        raise HTTPException(status_code=400, detail=f"지원하지 않는 플랫폼: {', '.join(unsupported)}")
    return requested_services


def _delivery_platform_label(service: str) -> str:
    return PLATFORM_LABELS.get(service, service or "배달플랫폼")


def _delivery_all_scope_requested(payload: dict[str, Any]) -> bool:
    markers = {"all", "*", "__all__", "전체"}
    business = str(payload.get("business_id") or payload.get("businessId") or "").strip().lower()
    branch = str(payload.get("branch") or "").strip().lower()
    return bool(payload.get("all_businesses")) or business in markers or branch in markers


def _delivery_full_backfill_requested(payload: dict[str, Any]) -> bool:
    return str(payload.get("mode") or payload.get("collection_mode") or "").strip().lower() == "full_backfill"


def _delivery_backfill_window_days(payload: dict[str, Any]) -> int:
    raw = payload.get("window_days") or payload.get("backfill_window_days") or os.getenv(
        "YEOLJEONG_BAEMIN_BACKFILL_WINDOW_DAYS", "1"
    )
    try:
        return max(1, min(7, int(raw)))
    except (TypeError, ValueError):
        return 1


def _delivery_backfill_batch_limit(payload: dict[str, Any]) -> int:
    raw = payload.get("max_backfill_runs") or payload.get("backfill_batch_limit") or os.getenv(
        "YEOLJEONG_BAEMIN_BACKFILL_BATCH_LIMIT", "1"
    )
    try:
        return max(1, min(4, int(raw)))
    except (TypeError, ValueError):
        return 1


def _delivery_backfill_max_orders(payload: dict[str, Any]) -> int:
    raw = payload.get("max_orders") or payload.get("maxOrders") or 80
    try:
        return max(1, min(300, int(raw)))
    except (TypeError, ValueError):
        return 80


def _delivery_backfill_max_reviews(payload: dict[str, Any]) -> int:
    raw = payload.get("max_reviews") or payload.get("maxReviews") or 80
    try:
        return max(1, min(300, int(raw)))
    except (TypeError, ValueError):
        return 80


def _delivery_backfill_row_payload(
    payload: dict[str, Any],
    *,
    date_from: date,
    date_to: date,
    checkpoint: dict[str, Any] | None = None,
    retry_of: str = "",
    attempt_count: int = 0,
) -> dict[str, Any]:
    return {
        "mode": "full_backfill",
        "window": {
            "date_from": date_from.isoformat(),
            "date_to": date_to.isoformat(),
            "window_days": _delivery_backfill_window_days(payload),
        },
        "checkpoint": checkpoint or {},
        "limits": {
            "concurrency": 1,
            "store_sequence": True,
            "max_orders": _delivery_backfill_max_orders(payload),
            "max_reviews": _delivery_backfill_max_reviews(payload),
            "max_runtime_seconds": int(DELIVERY_BACKFILL_STALE_AFTER.total_seconds()),
        },
        "retry": {
            "retry_of": retry_of,
            "attempt_count": attempt_count,
            "max_attempts": DELIVERY_BACKFILL_MAX_ATTEMPTS,
        },
    }


def _delivery_backfill_status_is_full(row: dict[str, Any]) -> bool:
    payload = row.get("payload") if isinstance(row.get("payload"), dict) else {}
    return str(row.get("service") or "") == "baemin" and str(payload.get("mode") or "") == "full_backfill"


def _delivery_backfill_status_order_key(row: dict[str, Any]) -> tuple[str, str]:
    return (
        str(row.get("date_to") or ""),
        str(row.get("queued_at") or row.get("created_at") or row.get("updated_at") or ""),
    )


def _delivery_backfill_window_from_status(row: dict[str, Any], fallback_from: date, fallback_to: date) -> tuple[date, date]:
    try:
        run_from = date.fromisoformat(str(row.get("date_from") or ""))
        run_to = date.fromisoformat(str(row.get("date_to") or ""))
        return run_from, run_to
    except ValueError:
        return fallback_from, fallback_to


def _delivery_stale_after_for_status(row: dict[str, Any]) -> timedelta:
    return DELIVERY_BACKFILL_STALE_AFTER if _delivery_backfill_status_is_full(row) else DELIVERY_SYNC_STALE_AFTER


def _delivery_clamp_backfill_retry_window(
    run_from: date,
    run_to: date,
    *,
    full_from: date,
    window_days: int,
) -> tuple[date, date]:
    if run_to < full_from:
        return run_from, run_to
    if (run_to - run_from).days + 1 <= window_days:
        return run_from, run_to
    return max(full_from, run_to - timedelta(days=window_days - 1)), run_to


def _delivery_enqueue_baemin_backfill_status(
    statuses: list[dict[str, Any]],
    payload: dict[str, Any],
    *,
    job_id: str,
    business_id: str,
    branch: str,
    date_from: date,
    date_to: date,
    checkpoint: dict[str, Any] | None = None,
    retry_of: str = "",
    attempt_count: int = 0,
) -> dict[str, Any]:
    queued_at = _now()
    existing = [
        row
        for row in statuses
        if _delivery_backfill_status_is_full(row)
        and str(row.get("status") or "") == "queued"
        and str(row.get("business_id") or "") == business_id
        and str(row.get("branch") or "") == branch
        and str(row.get("date_from") or "") == date_from.isoformat()
        and str(row.get("date_to") or "") == date_to.isoformat()
    ]
    if existing:
        return existing[0]
    status_record = {
        "id": str(uuid4()),
        "job_id": job_id,
        "service": "baemin",
        "business_id": business_id,
        "branch": branch,
        "date_from": date_from.isoformat(),
        "date_to": date_to.isoformat(),
        "status": "queued",
        "counts": _delivery_empty_counts(),
        "payload": _delivery_backfill_row_payload(
            payload,
            date_from=date_from,
            date_to=date_to,
            checkpoint=checkpoint,
            retry_of=retry_of,
            attempt_count=attempt_count,
        ),
        "error_code": "",
        "message": f"{branch} 배민 백필 대기 중입니다. window={date_from.isoformat()}~{date_to.isoformat()}",
        "queued_at": queued_at,
        "created_at": queued_at,
        "updated_at": queued_at,
    }
    statuses.insert(0, status_record)
    _run_db(_db_upsert_ledger("delivery_collection_status", status_record))
    return status_record


def _delivery_ensure_baemin_backfill_queue(
    statuses: list[dict[str, Any]],
    payload: dict[str, Any],
    scopes: list[tuple[str, str]],
    full_from: date,
    full_to: date,
    job_id: str,
) -> None:
    if not _delivery_full_backfill_requested(payload):
        return
    window_days = _delivery_backfill_window_days(payload)
    retryable_error_codes = {"BACKGROUND_SYNC_STALE", "ATTEMPT_TIMEOUT", "PC_AGENT_CDP_TIMEOUT"}
    for business_id, branch in scopes:
        scope_rows = [
            row
            for row in statuses
            if _delivery_backfill_status_is_full(row)
            and str(row.get("business_id") or "") == business_id
            and str(row.get("branch") or "") == branch
        ]
        if any(str(row.get("status") or "") in {"queued", "running"} for row in scope_rows):
            continue
        latest = max(scope_rows, key=_delivery_backfill_status_order_key) if scope_rows else {}
        latest_status = str(latest.get("status") or "")
        latest_error = str(latest.get("error_code") or "")
        latest_payload = latest.get("payload") if isinstance(latest.get("payload"), dict) else {}
        latest_retry = latest_payload.get("retry") if isinstance(latest_payload.get("retry"), dict) else {}
        latest_attempt = int(latest.get("attempt_count") or latest_retry.get("attempt_count") or 0)
        latest_from, latest_to = _delivery_backfill_window_from_status(latest, full_from, full_to) if latest else (full_to, full_to)
        retry_of = ""
        checkpoint = payload.get("checkpoint") if isinstance(payload.get("checkpoint"), dict) else None
        if latest_status in {"failed", "action_required"}:
            if latest_error not in retryable_error_codes or latest_attempt >= DELIVERY_BACKFILL_MAX_ATTEMPTS:
                continue
            run_from, run_to = latest_from, latest_to
            run_from, run_to = _delivery_clamp_backfill_retry_window(
                run_from,
                run_to,
                full_from=full_from,
                window_days=window_days,
            )
            checkpoint = latest_payload.get("checkpoint") if isinstance(latest_payload.get("checkpoint"), dict) else None
            retry_of = str(latest.get("id") or "")
            latest_attempt += 1
        elif latest_status in {"succeeded", "partial"}:
            run_to = latest_from - timedelta(days=1)
            if run_to < full_from:
                continue
            run_from = max(full_from, run_to - timedelta(days=window_days - 1))
            checkpoint = {}
            latest_attempt = 0
        else:
            run_to = full_to
            run_from = max(full_from, run_to - timedelta(days=window_days - 1))
            latest_attempt = 0
        _delivery_enqueue_baemin_backfill_status(
            statuses,
            payload,
            job_id=job_id,
            business_id=business_id,
            branch=branch,
            date_from=run_from,
            date_to=run_to,
            checkpoint=checkpoint,
            retry_of=retry_of,
            attempt_count=latest_attempt,
        )


def _delivery_select_baemin_backfill_statuses(
    statuses: list[dict[str, Any]],
    payload: dict[str, Any],
    scopes: list[tuple[str, str]],
) -> list[dict[str, Any]]:
    scope_set = set(scopes)
    queued = [
        row
        for row in statuses
        if _delivery_backfill_status_is_full(row)
        and str(row.get("status") or "") == "queued"
        and (str(row.get("business_id") or ""), str(row.get("branch") or "")) in scope_set
    ]
    queued.sort(key=_delivery_backfill_status_order_key, reverse=True)
    return queued[: _delivery_backfill_batch_limit(payload)]


def _delivery_enqueue_baemin_backfill_continuation(
    statuses: list[dict[str, Any]],
    payload: dict[str, Any],
    *,
    current: dict[str, Any],
    result: dict[str, Any],
    public_status: str,
    public_error_code: str,
    attempt_count: int,
) -> dict[str, Any] | None:
    if not _delivery_backfill_status_is_full(current):
        return None
    full_from, _full_to = _delivery_sync_window(payload)
    run_from, run_to = _delivery_backfill_window_from_status(current, full_from, _full_to)
    diagnostics = result.get("diagnostics") if isinstance(result.get("diagnostics"), dict) else {}
    checkpoint = diagnostics.get("checkpoint") or diagnostics.get("checkpoint_out") or {}
    max_orders = _delivery_backfill_max_orders(payload)
    orders_saved = int(diagnostics.get("order_history_orders_saved") or diagnostics.get("orders_saved") or 0)
    retryable_error_codes = {"BACKGROUND_SYNC_STALE", "ATTEMPT_TIMEOUT", "PC_AGENT_CDP_TIMEOUT"}
    if public_error_code in retryable_error_codes and attempt_count < DELIVERY_BACKFILL_MAX_ATTEMPTS:
        return _delivery_enqueue_baemin_backfill_status(
            statuses,
            payload,
            job_id=str(current.get("job_id") or payload.get("sync_job_id") or ""),
            business_id=str(current.get("business_id") or ""),
            branch=str(current.get("branch") or ""),
            date_from=run_from,
            date_to=run_to,
            checkpoint=checkpoint,
            retry_of=str(current.get("id") or ""),
            attempt_count=attempt_count,
        )
    if public_status in {"action_required", "failed"}:
        return None
    if orders_saved >= max_orders and checkpoint:
        next_from, next_to = run_from, run_to
    else:
        next_to = run_from - timedelta(days=1)
        if next_to < full_from:
            return None
        next_from = max(full_from, next_to - timedelta(days=_delivery_backfill_window_days(payload) - 1))
        checkpoint = {}
    return _delivery_enqueue_baemin_backfill_status(
        statuses,
        payload,
        job_id=str(current.get("job_id") or payload.get("sync_job_id") or ""),
        business_id=str(current.get("business_id") or ""),
        branch=str(current.get("branch") or ""),
        date_from=next_from,
        date_to=next_to,
        checkpoint=checkpoint,
    )


def _delivery_backfill_status_payload(payload: dict[str, Any], result: dict[str, Any] | None = None) -> dict[str, Any]:
    if not _delivery_full_backfill_requested(payload):
        return {}
    diagnostics = result.get("diagnostics") if isinstance(result, dict) and isinstance(result.get("diagnostics"), dict) else {}
    return {
        "mode": "full_backfill",
        "window": {
            "date_from": str(payload.get("date_from") or ""),
            "date_to": str(payload.get("date_to") or ""),
            "window_days": _delivery_backfill_window_days(payload),
        },
        "checkpoint": diagnostics.get("checkpoint") or diagnostics.get("checkpoint_out") or {},
        "metrics": {
            "orders_seen": diagnostics.get("order_history_orders_seen") or diagnostics.get("orders_seen") or 0,
            "orders_saved": diagnostics.get("order_history_orders_saved") or diagnostics.get("orders_saved") or 0,
            "reviews_saved": diagnostics.get("review_backfill_reviews_saved") or 0,
            "ads_saved": diagnostics.get("ads_backfill_ads_saved") or 0,
            "detail_failed": diagnostics.get("order_history_detail_failed") or diagnostics.get("detail_failed") or 0,
            "settlement_pending": diagnostics.get("order_history_settlement_pending") or diagnostics.get("settlement_pending") or 0,
        },
        "limits": {
            "concurrency": 1,
            "store_sequence": True,
            "order_detail_jitter_seconds": [1.0, 1.8],
            "page_jitter_seconds": [2.0, 4.0],
            "max_orders": _delivery_backfill_max_orders(payload),
            "max_reviews": _delivery_backfill_max_reviews(payload),
            "max_runtime_seconds": int(DELIVERY_BACKFILL_STALE_AFTER.total_seconds()),
        },
    }


def _delivery_sync_scopes(
    payload: dict[str, Any],
    requested_services: list[str],
    all_accounts: list[dict[str, Any]],
) -> list[tuple[str, str]]:
    if not _delivery_all_scope_requested(payload):
        return [
            _normalize_delivery_scope(
                str(payload.get("business_id") or MIA_BUSINESS_ID),
                str(payload.get("branch") or MIA_BRANCH_NAME),
            )
        ]

    canonical_scopes = [(str(item["businessId"]), str(item["name"])) for item in CANONICAL_BRANCHES]
    account_scopes: list[tuple[str, str]] = []
    seen: set[tuple[str, str]] = set()
    requested = set(requested_services)
    for row in all_accounts:
        service = str(row.get("service") or "").strip()
        if service not in requested:
            continue
        branch = BRANCH_ALIASES.get(str(row.get("branch") or "").strip(), str(row.get("branch") or "").strip())
        business_id = str(row.get("business_id") or BUSINESS_BY_BRANCH.get(branch) or "").strip()
        if business_id not in CANONICAL_BUSINESS_IDS or BUSINESS_BY_BRANCH.get(branch) != business_id:
            continue
        scope = (business_id, branch)
        if scope not in seen:
            seen.add(scope)
            account_scopes.append(scope)
    if _delivery_full_backfill_requested(payload):
        return account_scopes or canonical_scopes
    return account_scopes or canonical_scopes


def _delivery_run_key(service: str, business_id: str, branch: str) -> str:
    return f"{business_id}|{branch}|{service}"


def _write_delivery_collection_statuses(rows: list[dict[str, Any]], current: dict[str, Any] | None = None) -> None:
    _write_file_rows("delivery_collection_status", rows)
    if current and current.get("id"):
        _run_db(_db_upsert_ledger("delivery_collection_status", current))


def _try_acquire_delivery_sync_lock() -> int | None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    fd = os.open(str(DATA_DIR / ".delivery_sync.lock"), os.O_CREAT | os.O_RDWR, 0o600)
    try:
        fcntl.flock(fd, fcntl.LOCK_EX | fcntl.LOCK_NB)
    except BlockingIOError:
        os.close(fd)
        return None
    except Exception:
        os.close(fd)
        raise
    os.ftruncate(fd, 0)
    os.write(fd, f"{os.getpid()} {_now()}\n".encode("utf-8"))
    return fd


def _release_delivery_sync_lock(fd: int | None) -> None:
    if fd is None:
        return
    try:
        fcntl.flock(fd, fcntl.LOCK_UN)
    finally:
        os.close(fd)


def _delivery_sync_busy_result(payload: dict[str, Any], user: dict[str, Any]) -> dict[str, Any]:
    requested_services = _delivery_requested_services(payload)
    date_from, date_to = _delivery_sync_window(payload)
    all_accounts = _read("platform_accounts")
    scopes = _delivery_sync_scopes(payload, requested_services, all_accounts)
    now_text = _now()
    statuses = _read("delivery_collection_status")
    message = "다른 배달 자동수집 작업이 실행 중이라 중복 실행을 차단했습니다. 현재 작업 완료 후 다시 실행하세요."
    summary: list[dict[str, Any]] = []
    records: list[dict[str, Any]] = []
    for business_id, branch in scopes:
        for service in requested_services:
            status_record = {
                "id": str(uuid4()),
                "job_id": str(payload.get("sync_job_id") or ""),
                "service": service,
                "business_id": business_id,
                "branch": branch,
                "date_from": date_from.isoformat(),
                "date_to": date_to.isoformat(),
                "status": "action_required",
                "raw_status": "busy",
                "counts": _delivery_empty_counts(),
                "payload": _delivery_backfill_status_payload(payload),
                "error_code": "COLLECTION_ALREADY_RUNNING",
                "message": message,
                "started_at": now_text,
                "finished_at": now_text,
                "created_at": now_text,
                "updated_at": now_text,
            }
            statuses.insert(0, status_record)
            records.append(status_record)
            summary.append(
                {
                    "service": service,
                    "status": "action_required",
                    "portal_status": "action_required",
                    "error_code": "COLLECTION_ALREADY_RUNNING",
                    "counts": _delivery_empty_counts(),
                    "run_id": status_record["id"],
                    "account_id": "",
                    "business_id": business_id,
                    "branch": branch,
                    "message": message,
                    "portal_message": message,
                }
            )
    _write_delivery_collection_statuses(statuses)
    for status_record in records:
        _run_db(_db_upsert_ledger("delivery_collection_status", status_record))
    return {
        "queued": False,
        "synced_at": now_text,
        "business_id": scopes[0][0] if len(scopes) == 1 else "all",
        "branch": scopes[0][1] if len(scopes) == 1 else "전체",
        "date_from": date_from.isoformat(),
        "date_to": date_to.isoformat(),
        "totals": _delivery_empty_counts(),
        "summary": summary,
        "sales": [],
        "settlements": [],
        "reviews": [],
        "records": [],
    }


def queue_delivery_sync(payload: dict[str, Any], user: dict[str, Any]) -> dict[str, Any]:
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="자동 수집 실행 권한이 없습니다")

    requested_services = _delivery_requested_services(payload)
    date_from, date_to = _delivery_sync_window(payload)
    all_accounts = _read("platform_accounts")
    scopes = _delivery_sync_scopes(payload, requested_services, all_accounts)
    queued_at = _now()
    job_id = str(payload.get("sync_job_id") or f"delivery-sync-{uuid4().hex[:12]}")
    statuses = _read("delivery_collection_status")
    queued_run_ids: dict[str, str] = {}
    summary: list[dict[str, Any]] = []
    queued_records: list[dict[str, Any]] = []

    for business_id, branch in scopes:
        for service in requested_services:
            run_id = str(uuid4())
            run_key = _delivery_run_key(service, business_id, branch)
            queued_run_ids[run_key] = run_id
            if len(scopes) == 1:
                queued_run_ids[service] = run_id
            run_date_from, run_date_to = date_from, date_to
            status_payload = _delivery_backfill_status_payload(payload)
            if service == "baemin" and _delivery_full_backfill_requested(payload):
                run_date_to = date_to
                run_date_from = max(date_from, run_date_to - timedelta(days=_delivery_backfill_window_days(payload) - 1))
                status_payload = _delivery_backfill_row_payload(
                    payload,
                    date_from=run_date_from,
                    date_to=run_date_to,
                )
            message = (
                f"{branch} {_delivery_platform_label(service)} 백그라운드 수집 대기 중입니다."
                f" window={run_date_from.isoformat()}~{run_date_to.isoformat()}"
            )
            statuses.insert(
                0,
                status_record := {
                    "id": run_id,
                    "job_id": job_id,
                    "service": service,
                    "business_id": business_id,
                    "branch": branch,
                    "date_from": run_date_from.isoformat(),
                    "date_to": run_date_to.isoformat(),
                    "status": "queued",
                    "counts": _delivery_empty_counts(),
                    "payload": status_payload,
                    "error_code": "",
                    "message": message,
                    "queued_at": queued_at,
                    "created_at": queued_at,
                    "updated_at": queued_at,
                },
            )
            queued_records.append(status_record)
            summary.append(
                {
                    "service": service,
                    "status": "queued",
                    "portal_status": "queued",
                    "error_code": "",
                    "counts": _delivery_empty_counts(),
                    "run_id": run_id,
                    "job_id": job_id,
                    "account_id": str(payload.get("account_id") or payload.get("server_account_id") or ""),
                    "business_id": business_id,
                    "branch": branch,
                    "message": message,
                    "portal_message": message,
                }
            )

    _write_delivery_collection_statuses(statuses)
    for status_record in queued_records:
        _run_db(_db_upsert_ledger("delivery_collection_status", status_record))
    return {
        "queued": True,
        "job_id": job_id,
        "queued_run_ids": queued_run_ids,
        "synced_at": queued_at,
        "queued_at": queued_at,
        "business_id": scopes[0][0] if len(scopes) == 1 else "all",
        "branch": scopes[0][1] if len(scopes) == 1 else "전체",
        "date_from": date_from.isoformat(),
        "date_to": date_to.isoformat(),
        "summary": summary,
        "records": [],
        "sales": [],
        "settlements": [],
        "reviews": [],
        "totals": _delivery_empty_counts(),
    }


def _delivery_entry_record(record: dict[str, Any]) -> dict[str, Any]:
    service = str(record.get("service") or "")
    record_type = str(record.get("record_type") or "")
    amount = int(record.get("gross_amount") or record.get("settlement_amount") or 0)
    label = CONNECTOR_LABELS.get(service, service or "배달플랫폼")
    date_value = str(record.get("occurred_on") or datetime.now(KST).date().isoformat())
    if record_type == "settlements":
        entry_type = "bank"
        vendor = f"{label} 정산입금"
        memo = str(record.get("settlement_status") or record.get("settlement_id") or "")
    else:
        entry_type = "sales"
        vendor = f"{label} 매출"
        memo = str(record.get("order_status") or record.get("order_id") or "")
    return {
        "id": f"entry-{record.get('id') or uuid4()}",
        "source_record_id": record.get("id") or "",
        "source_type": "delivery",
        "type": entry_type,
        "service": service,
        "business_id": record.get("business_id") or "",
        "branch": record.get("branch") or "",
        "date": date_value,
        "vendor": vendor,
        "amount": amount,
        "status": "confirmed",
        "memo": memo,
    }


def automation_status(user: dict[str, Any]) -> dict[str, Any]:
    return {
        "enabled": True,
        "mode": "browser-automation",
        "status": "available",
        "message": "계정 기반 포털 수집과 CSV 정산서 가져오기를 사용할 수 있습니다. CAPTCHA·2차 인증은 사용자 조치가 필요합니다.",
        "checked_at": _now(),
    }


def _matching_accounts(
    *,
    services: list[str],
    business_id: str,
    branch: str,
    account_id: str = "",
) -> dict[str, dict[str, Any]]:
    rows = _read("platform_accounts")
    if _migrate_platform_account_secrets(rows):
        _write("platform_accounts", rows)
    requested_account_id = str(account_id or "").strip()
    candidates = [
        row
        for row in rows
        if str(row.get("service") or "") in services
        and (not requested_account_id or str(row.get("id") or "") == requested_account_id)
        and str(row.get("business_id") or "") == business_id
        and (
            not branch
            or not str(row.get("branch") or "").strip()
            or BRANCH_ALIASES.get(str(row.get("branch") or ""), str(row.get("branch") or "")) == branch
        )
    ]
    candidates.sort(key=lambda row: str(row.get("updated_at") or row.get("created_at") or ""), reverse=True)
    accounts_by_service: dict[str, dict[str, Any]] = {}
    for service in services:
        service_rows = [row for row in candidates if str(row.get("service") or "") == service]
        if service_rows:
            accounts_by_service[service] = next((row for row in service_rows if _has_account_secret(row)), service_rows[0])
    return accounts_by_service


def import_transaction_csv(
    payload: dict[str, Any],
    user: dict[str, Any],
) -> dict[str, Any]:
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="거래내역 가져오기 권한이 없습니다")
    service = str(payload.get("service") or "").strip()
    if service not in FINANCIAL_TRANSACTION_SERVICES:
        raise HTTPException(status_code=400, detail="은행/카드 거래 연동 서비스만 가져올 수 있습니다")
    business_id, branch = _normalize_connector_scope(service, payload.get("business_id"), payload.get("branch"))
    csv_text = str(payload.get("csv_text") or "")
    if not csv_text.strip():
        raise HTTPException(status_code=400, detail="거래내역 CSV 내용이 필요합니다")
    result = import_file(
        str(payload.get("filename") or "transactions.csv"),
        csv_text.encode("utf-8-sig"),
        TRANSACTION_SOURCE_BY_SERVICE[service],
        business_id=business_id,
        branch=branch,
        service=service,
        source_account_id=str(payload.get("source_account_id") or ""),
    )
    return {
        **result,
        "business_id": business_id,
        "branch": branch,
        "service": service,
        "transactions": result["rows"],
    }


def sync_financial_transactions(payload: dict[str, Any], user: dict[str, Any]) -> dict[str, Any]:
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="은행/카드 자동연동 실행 권한이 없습니다")
    services = [str(item) for item in (payload.get("services") or []) if str(item).strip()]
    requested_services = services or sorted(FINANCIAL_TRANSACTION_SERVICES)
    unsupported = sorted(set(requested_services) - FINANCIAL_TRANSACTION_SERVICES)
    if unsupported:
        raise HTTPException(status_code=400, detail=f"지원하지 않는 은행/카드 연동: {', '.join(unsupported)}")
    today = datetime.now(KST).date()
    default_from = today.replace(day=1).isoformat()
    date_from_text = str(payload.get("date_from") or default_from)
    date_to_text = str(payload.get("date_to") or today.isoformat())
    try:
        date_from = date.fromisoformat(date_from_text)
        date_to = date.fromisoformat(date_to_text)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="연동 기간은 YYYY-MM-DD 형식이어야 합니다") from exc
    if date_from > date_to or (date_to - date_from).days > 92:
        raise HTTPException(status_code=400, detail="은행/카드 연동 기간은 시작일 이후 최대 93일입니다")
    business_id, branch = _normalize_connector_scope(
        "shinhan_business",
        payload.get("business_id") or MIA_BUSINESS_ID,
        payload.get("branch") or MIA_BRANCH_NAME,
    )
    accounts_by_service = _matching_accounts(
        services=requested_services,
        business_id=business_id,
        branch=branch,
        account_id=str(payload.get("account_id") or payload.get("server_account_id") or "").strip(),
    )
    synced_at = _now()
    transactions = _read("transactions")
    by_id = {str(row.get("id") or ""): row for row in transactions if row.get("id")}
    summary: list[dict[str, Any]] = []
    imported_rows: list[dict[str, Any]] = []
    response_rows: list[dict[str, Any]] = []

    for service in requested_services:
        account = accounts_by_service.get(service)
        if not account:
            summary.append(
                {
                    "service": service,
                    "status": "skipped",
                    "message": "해당 사업자/지점에 등록된 자동수집 대상 계좌가 없어 실행하지 않았습니다.",
                    "error_code": "BANK_ACCOUNT_NOT_PRESENT_FOR_SCOPE",
                    "imported_rows": 0,
                }
            )
            continue
        collection_mode = str(account.get("collection_mode") or "").strip()
        if service in BANK_QUICK_SERVICE_CONFIG and collection_mode == "bank-quick-service":
            quick_missing = [
                label
                for label, key in (
                    ("로그인 비밀번호", "password"),
                    ("조회용 계좌번호", "account_no"),
                    ("계좌비밀번호", "account_password"),
                    ("사업자번호", "business_registration_no"),
                )
                if not _has_secret_value(account, key)
            ]
            if quick_missing:
                summary.append(
                    {
                        "service": service,
                        "status": "skipped",
                        "message": "은행 간편/빠른조회 필수값이 없어 실행하지 않았습니다.",
                        "error_code": "BANK_QUICK_ACCOUNT_NOT_COLLECTABLE",
                        "missing_requirements": quick_missing,
                        "account_id": account.get("id") or "",
                        "collection_mode": collection_mode,
                        "imported_rows": 0,
                    }
                )
                continue
        if collection_mode in {"bank-excel", "card-pg-report", "statement-upload"}:
            summary.append(
                {
                    "service": service,
                    "status": "upload_required",
                    "message": "현재 연동 방식은 파일 업로드입니다. 거래내역 CSV/리포트를 업로드하면 거래원장에 반영됩니다.",
                    "account_id": account.get("id") or "",
                    "imported_rows": 0,
                }
            )
            continue
        if collection_mode in {"api", "bank-openbanking", "browser-automation", "bank-quick-service"}:
            quick_missing = []
            if collection_mode == "bank-quick-service":
                quick_missing = [
                    label
                    for label, key in (
                        ("로그인 비밀번호", "password"),
                        ("조회용 계좌번호", "account_no"),
                        ("계좌비밀번호", "account_password"),
                        ("사업자번호", "business_registration_no"),
                    )
                    if not _has_secret_value(account, key)
                ]
            if not _has_account_secret(account) or quick_missing:
                status = "credential_required"
                message = (
                    f"은행 간편/빠른조회 필수값 누락: {', '.join(quick_missing)}"
                    if quick_missing
                    else "API 키, 인증서 비밀번호 또는 로그인 비밀번호를 설정에서 등록해야 합니다."
                )
                summary.append(
                    {
                        "service": service,
                        "status": status,
                        "message": message,
                        "account_id": account.get("id") or "",
                        "collection_mode": collection_mode,
                        "imported_rows": 0,
                    }
                )
                _mark_platform_account_sync_state(account, status=status, message=message, synced_at=synced_at)
                continue
            if collection_mode == "bank-quick-service":
                branch_id = _branch_id_for_financial_scope(business_id, branch)
                bank_account = _bank_account_for_financial_service(
                    service,
                    business_id=business_id,
                    branch_id=branch_id,
                )
                if bank_account:
                    collect_result = _collect_bank_via_browser(
                        bank_account,
                        {
                            "business_id": business_id,
                            "branch_id": branch_id,
                            "date_from": date_from.isoformat(),
                            "date_to": date_to.isoformat(),
                            "auto_open_browser": True,
                        },
                        business_id=business_id,
                        branch_id=branch_id,
                        date_from=date_from.isoformat(),
                        date_to=date_to.isoformat(),
                        user=user,
                    )
                    collection = collect_result.get("collection") or {}
                    collected_transactions = [
                        row for row in (collect_result.get("transactions") or []) if isinstance(row, dict)
                    ]
                    response_rows.extend(collected_transactions)
                    status = str(collection.get("status") or "failed")
                    message = str(collection.get("message") or "")
                    imported_count = int(collection.get("imported_rows") or 0)
                    summary.append(
                        {
                            "service": service,
                            "status": status,
                            "message": message,
                            "account_id": account.get("id") or "",
                            "bank_account_id": bank_account.get("id") or "",
                            "collection_mode": collection_mode,
                            "imported_rows": imported_count,
                            "duplicate_rows": int(collection.get("duplicate_rows") or 0),
                            "error_code": str(collection.get("error_code") or ""),
                        }
                    )
                    _mark_platform_account_sync_state(account, status=status, message=message, synced_at=synced_at)
                    continue
                status = "connector_not_configured"
                message = "은행계좌 원장에 연결된 브라우저 수집 계좌가 없습니다. 은행 계좌 연결을 먼저 저장하십시오."
            else:
                status = "connector_not_configured"
                message = "자격증명은 저장됐지만 해당 기관 실조회 커넥터가 아직 연결되지 않았습니다. 파일 업로드로 대체 수집할 수 있습니다."
            summary.append(
                {
                    "service": service,
                    "status": status,
                    "message": message,
                    "account_id": account.get("id") or "",
                    "collection_mode": collection_mode,
                    "imported_rows": 0,
                }
            )
            _mark_platform_account_sync_state(account, status=status, message=message, synced_at=synced_at)
            continue
        sample_rows = account.get("last_download_rows") if isinstance(account.get("last_download_rows"), list) else []
        count = 0
        for row in sample_rows:
            record = {
                **row,
                "service": service,
                "source_type": TRANSACTION_SOURCE_BY_SERVICE[service],
                "business_id": business_id,
                "branch": branch,
                "source_account_id": account.get("id") or "",
                "updated_at": synced_at,
            }
            record["id"] = str(record.get("id") or hashlib.sha256(json.dumps(record, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest())
            by_id[record["id"]] = record
            imported_rows.append(record)
            response_rows.append(record)
            count += 1
        summary.append(
            {
                "service": service,
                "status": "completed" if count else "no_records",
                "message": "저장된 다운로드 거래를 반영했습니다." if count else "반영할 신규 거래가 없습니다.",
                "account_id": account.get("id") or "",
                "imported_rows": count,
            }
        )
        status = "completed" if count else "no_records"
        message = "저장된 다운로드 거래를 반영했습니다." if count else "반영할 신규 거래가 없습니다."
        _mark_platform_account_sync_state(account, status=status, message=message, synced_at=synced_at)

    if imported_rows:
        _write("transactions", list(by_id.values()))
    return {
        "synced_at": synced_at,
        "business_id": business_id,
        "branch": branch,
        "date_from": date_from.isoformat(),
        "date_to": date_to.isoformat(),
        "summary": summary,
        "transactions": response_rows,
        "totals": {"transactions": len(response_rows)},
    }


# ---------------------------------------------------------------------------
# 매장비서 은행자동연동 1단계 (AADS/FOOD)
#
# 기존 platform_accounts / transactions 원장을 건드리지 않고, 사업자별 은행계좌
# 등록과 은행거래 원장을 별도 파일 저장소로 분리한다. 계좌번호/비밀번호/OTP/인증서
# 같은 민감정보는 절대 평문 저장하지 않으며, 마스킹·추상 상태 정보만 유지한다.
# 외부 은행 실연동(오픈뱅킹/스크래핑)은 이 단계에서 구현하지 않고, 수동/CSV/목업
# 입력만 원장에 멱등 반영한다.
# ---------------------------------------------------------------------------

BANK_ACCOUNTS_LEDGER = "bank_accounts"
BANK_TRANSACTIONS_LEDGER = "bank_transactions"
BANK_CONNECTION_TYPES = ("open_banking", "csv", "manual", "mock", "browser")
BANK_ACCOUNT_STATUSES = ("active", "paused", "error", "needs_auth")
BANK_TRANSACTION_DIRECTIONS = ("in", "out")
BANK_STATEMENT_MAX_BYTES = 10 * 1024 * 1024
BANK_STATEMENT_MAX_ROWS = 50_000
BANK_STATEMENT_MAX_COLUMNS = 100
BANK_STATEMENT_MAX_EXPANDED_BYTES = 100 * 1024 * 1024
BANK_STATEMENT_EXTENSIONS = {".csv", ".xlsx", ".xlsm"}
BANK_STATEMENT_MIME_TYPES = {
    ".csv": {"", "application/octet-stream", "text/csv", "text/plain", "application/csv", "application/vnd.ms-excel"},
    ".xlsx": {"", "application/octet-stream", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
    ".xlsm": {"", "application/octet-stream", "application/vnd.ms-excel.sheet.macroenabled.12"},
}
BANK_DATE_HEADERS = {"거래일시", "거래일자", "거래일", "일시", "일자", "날짜", "date", "datetime"}
BANK_AMOUNT_HEADERS = {
    "입금액", "입금", "입금금액", "맡기신금액", "받으신금액", "출금액", "출금", "출금금액",
    "찾으신금액", "지급금액", "거래금액", "금액", "amount", "deposit", "credit", "withdrawal", "debit",
}
BANK_CONFIGURED_COLLECTION_TYPES = ("csv", "manual", "mock", "browser")
BANK_SERVICE_CODE_ALIASES: dict[str, tuple[str, str]] = {
    "shinhan_business": ("088", "신한은행"),
    "ibk_business": ("003", "IBK기업은행"),
}
BANK_SERVICE_NAME_KEYWORDS: tuple[tuple[str, str], ...] = (
    ("shinhan_business", "신한"),
    ("ibk_business", "ibk"),
    ("ibk_business", "기업"),
)

# 은행계좌 파일에 저장을 허용하는 필드 화이트리스트. 원본 계좌번호/비밀번호/인증정보는
# 절대 포함하지 않는다(민감정보 제외 원칙).
_BANK_ACCOUNT_PUBLIC_FIELDS = (
    "id",
    "business_id",
    "branch_id",
    "bank_code",
    "bank_name",
    "account_number_masked",
    "account_holder",
    "account_alias",
    "connection_type",
    "connector_type",
    "status",
    "institution_code",
    "auto_sync",
    "memo",
    "last_synced_at",
    "created_at",
    "updated_at",
    # 은행 간편/빠른조회 인증정보 연결 메타 (비밀값은 platform_accounts에 암호화 저장)
    "platform_account_id",
    "credential_service",
    "credential_username",
    "credentials_registered_at",
)
# 어떤 경로로 들어와도 은행계좌 저장소에 남기면 안 되는 민감 키.
_BANK_ACCOUNT_FORBIDDEN_FIELDS = frozenset(
    {
        "account_number",
        "account_no",
        "account_number_raw",
        "password",
        "account_password",
        "login_password",
        "pin",
        "otp",
        "certificate",
        "certificate_password",
        "secret",
        "client_secret",
        "api_key",
        "credential",
        "credentials",
        "access_token",
        "refresh_token",
    }
)

CANONICAL_BRANCH_BY_ID: dict[str, dict[str, Any]] = {item["id"]: item for item in CANONICAL_BRANCHES}


def _write_secure_file_rows(name: str, rows: list[dict[str, Any]]) -> None:
    """Persist a ledger file with 0600 permissions (owner read/write only)."""
    path = _path(name)
    tmp = path.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    try:
        os.chmod(tmp, 0o600)
    except OSError:
        pass
    tmp.replace(path)
    try:
        os.chmod(path, 0o600)
    except OSError:
        pass


def _normalize_bank_scope(business_id: Any, branch_id: Any) -> tuple[str, str]:
    normalized_business = str(business_id or "").strip()
    ui_settings = _canonicalize_ui_settings((_read_json_object("settings").get("ui_settings") or {}))
    business_ids = {str(item.get("id") or "") for item in ui_settings.get("businesses", []) if isinstance(item, dict)}
    branch_businesses = {
        str(item.get("id") or ""): str(item.get("businessId") or "")
        for item in ui_settings.get("branches", [])
        if isinstance(item, dict)
    }
    if normalized_business not in business_ids:
        raise HTTPException(status_code=400, detail="등록되지 않은 사업자입니다")
    normalized_branch = str(branch_id or "").strip()
    if normalized_branch:
        if branch_businesses.get(normalized_branch) != normalized_business:
            raise HTTPException(status_code=400, detail="사업자와 지점 연결이 일치하지 않습니다")
    return normalized_business, normalized_branch


def _bank_connection_type(value: Any, *, default: str = "mock") -> str:
    text = str(value or "").strip().lower()
    if not text:
        return default
    if text not in BANK_CONNECTION_TYPES:
        raise HTTPException(
            status_code=400,
            detail=f"연동 방식은 {', '.join(BANK_CONNECTION_TYPES)} 중 하나여야 합니다",
        )
    return text


def _bank_account_status(value: Any, *, default: str = "needs_auth") -> str:
    text = str(value or "").strip().lower()
    if not text:
        return default
    if text not in BANK_ACCOUNT_STATUSES:
        raise HTTPException(
            status_code=400,
            detail=f"계좌 상태는 {', '.join(BANK_ACCOUNT_STATUSES)} 중 하나여야 합니다",
        )
    return text


def _bank_direction(value: Any) -> str:
    text = str(value or "").strip().lower()
    if text in {"in", "income", "deposit", "credit", "입금", "입금액", "받으신금액", "맡기신금액"}:
        return "in"
    if text in {"out", "expense", "withdrawal", "debit", "출금", "출금액", "찾으신금액", "지급금액"}:
        return "out"
    return ""


def _bank_date_key(value: Any) -> str:
    """Return the YYYY-MM-DD prefix used for range comparisons."""
    normalized = _transaction_date(value)
    return normalized[:10] if normalized else ""


def _bank_within_range(occurred_at: Any, date_from: str, date_to: str) -> bool:
    key = _bank_date_key(occurred_at)
    if not key:
        # Keep undated rows visible unless an explicit range is requested.
        return not (date_from or date_to)
    if date_from and key < date_from:
        return False
    if date_to and key > date_to:
        return False
    return True


def _valid_range_bounds(date_from: Any, date_to: Any) -> tuple[str, str]:
    def _check(label: str, value: Any) -> str:
        text = str(value or "").strip()
        if not text:
            return ""
        try:
            return date.fromisoformat(text).isoformat()
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=f"{label}은 YYYY-MM-DD 형식이어야 합니다") from exc

    start = _check("조회 시작일", date_from)
    end = _check("조회 종료일", date_to)
    if start and end and start > end:
        raise HTTPException(status_code=400, detail="조회 시작일은 종료일보다 이후일 수 없습니다")
    return start, end


def _sanitize_bank_account(record: dict[str, Any]) -> dict[str, Any]:
    """Keep only whitelisted, non-sensitive fields for persistence and output."""
    return {key: record[key] for key in _BANK_ACCOUNT_PUBLIC_FIELDS if key in record}


def _bank_account_number_masked(payload: dict[str, Any], existing: dict[str, Any] | None) -> str:
    provided_mask = str(payload.get("account_number_masked") or "").strip()
    if provided_mask:
        return provided_mask
    raw = str(payload.get("account_number") or "").strip()
    if raw:
        return _masked_digits(raw)
    return str((existing or {}).get("account_number_masked") or "").strip()


def _infer_bank_service_code(*values: Any) -> str:
    for value in values:
        text = str(value or "").strip()
        lowered = text.lower()
        if lowered in BANK_SERVICE_CODE_ALIASES:
            return lowered
        for service_code, keyword in BANK_SERVICE_NAME_KEYWORDS:
            if keyword in lowered or keyword in text:
                return service_code
    return ""


def _same_bank_account_key(account: dict[str, Any], *, business_id: str, branch_id: str, service_code: str, mask: str) -> bool:
    if str(account.get("business_id") or "").strip() != business_id:
        return False
    if str(account.get("branch_id") or "").strip() != branch_id:
        return False
    account_service = _infer_bank_service_code(
        account.get("institution_code"),
        account.get("bank_code"),
        account.get("bank_name"),
    )
    if service_code and account_service != service_code:
        return False
    account_mask = str(account.get("account_number_masked") or "").strip()
    return bool(mask and account_mask and account_mask == mask)


def _find_existing_bank_account_by_key(
    rows: list[dict[str, Any]],
    *,
    business_id: str,
    branch_id: str,
    payload: dict[str, Any],
) -> dict[str, Any] | None:
    connection_type = str(payload.get("connection_type") or "").strip().lower()
    connector_type = str(payload.get("connector_type") or "").strip().lower()
    memo = str(payload.get("memo") or "").strip()
    auto_sync = payload.get("auto_sync") is True
    is_auto_browser_account = connection_type == "browser" and (
        auto_sync
        or connector_type in {"bank-browser", "bank-quick-service"}
        or "platform_accounts" in memo
    )
    if not is_auto_browser_account:
        return None
    service_code = _infer_bank_service_code(
        payload.get("institution_code"),
        payload.get("bank_code"),
        payload.get("bank_name"),
    )
    mask = _bank_account_number_masked(payload, None)
    if not business_id or not branch_id or not service_code or not mask:
        return None
    for row in rows:
        if _same_bank_account_key(
            row,
            business_id=business_id,
            branch_id=branch_id,
            service_code=service_code,
            mask=mask,
        ):
            return row
    return None


def _bank_numeric_code_for_service(service_code: str) -> str:
    return BANK_SERVICE_CODE_ALIASES.get(service_code, ("", ""))[0]


def _apply_bank_account_fields(
    record: dict[str, Any],
    payload: dict[str, Any],
    *,
    creating: bool,
) -> dict[str, Any]:
    inferred_service = _infer_bank_service_code(
        payload.get("institution_code"),
        payload.get("bank_code"),
        payload.get("bank_name"),
        record.get("institution_code"),
        record.get("bank_code"),
        record.get("bank_name"),
    )
    if creating or payload.get("bank_code") is not None:
        bank_code = str(payload.get("bank_code") or record.get("bank_code") or "").strip()
        record["bank_code"] = bank_code or _bank_numeric_code_for_service(inferred_service)
    if creating or payload.get("bank_name") is not None:
        record["bank_name"] = str(payload.get("bank_name") or record.get("bank_name") or "").strip()
    if creating or payload.get("account_holder") is not None:
        record["account_holder"] = str(payload.get("account_holder") or record.get("account_holder") or "").strip()
    if creating or payload.get("account_alias") is not None:
        record["account_alias"] = str(payload.get("account_alias") or record.get("account_alias") or "").strip()
    if creating or payload.get("institution_code") is not None:
        record["institution_code"] = (
            str(payload.get("institution_code") or record.get("institution_code") or "").strip()
            or inferred_service
        )
    if creating or payload.get("memo") is not None:
        record["memo"] = str(payload.get("memo") or record.get("memo") or "").strip()
    if creating or payload.get("connection_type") is not None:
        record["connection_type"] = _bank_connection_type(
            payload.get("connection_type"),
            default=str(record.get("connection_type") or "mock"),
        )
    if creating or payload.get("status") is not None:
        record["status"] = _bank_account_status(
            payload.get("status"),
            default=str(record.get("status") or "needs_auth"),
        )
    if creating or payload.get("auto_sync") is not None:
        record["auto_sync"] = bool(payload.get("auto_sync"))
    if creating or payload.get("connector_type") is not None:
        connector_type = str(payload.get("connector_type") or record.get("connector_type") or "").strip()
        # 허용된 connector_type 값만 저장
        if connector_type not in {"", "bank-browser", "manual", "csv", "mock"}:
            connector_type = ""
        record["connector_type"] = connector_type
    if payload.get("last_synced_at") is not None:
        record["last_synced_at"] = str(payload.get("last_synced_at") or "").strip()
    mask = _bank_account_number_masked(payload, record)
    if creating or mask:
        record["account_number_masked"] = mask
    return record


def list_bank_accounts(
    user: dict[str, Any],
    business_id: str | None = None,
    *,
    branch_id: str | None = None,
    status: str | None = None,
) -> list[dict[str, Any]]:
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="은행계좌 조회 권한이 없습니다")
    rows = _read_file_rows(BANK_ACCOUNTS_LEDGER)
    result: list[dict[str, Any]] = []
    for row in rows:
        if business_id and str(row.get("business_id") or "") != business_id:
            continue
        if branch_id and str(row.get("branch_id") or "") != branch_id:
            continue
        if status and str(row.get("status") or "") != status:
            continue
        result.append(_sanitize_bank_account(row))
    result.sort(key=lambda item: str(item.get("updated_at") or item.get("created_at") or ""), reverse=True)
    return result


def _find_bank_account(rows: list[dict[str, Any]], account_id: str) -> dict[str, Any] | None:
    target = str(account_id or "").strip()
    return next((row for row in rows if str(row.get("id") or "") == target), None)


def _bank_account_matches_scope(account: dict[str, Any], business_id: str, branch_id: str = "") -> bool:
    if str(account.get("business_id") or "") != business_id:
        return False
    account_branch = str(account.get("branch_id") or "")
    return not branch_id or not account_branch or account_branch == branch_id


def _bank_unconfigured_collect_result(
    account: dict[str, Any],
    *,
    business_id: str,
    branch_id: str,
    date_from: str,
    date_to: str,
    reason: str,
    status: str = "needs_auth",
) -> dict[str, Any]:
    return {
        "collection": {
            "bank_account_id": str(account.get("id") or ""),
            "business_id": business_id,
            "branch_id": branch_id,
            "status": status,
            "connector_status": "NOT_CONFIGURED",
            "connection_type": str(account.get("connection_type") or ""),
            "message": reason,
            "collected_rows": 0,
            "imported_rows": 0,
            "duplicate_rows": 0,
            "matched_count": 0,
            "unmatched_count": 0,
            "last_collected_at": str(account.get("last_synced_at") or ""),
            "date_from": date_from,
            "date_to": date_to,
        },
        "transactions": [],
    }


def create_bank_account(payload: dict[str, Any], user: dict[str, Any]) -> dict[str, Any]:
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="은행계좌 등록 권한이 없습니다")
    business_id, branch_id = _normalize_bank_scope(payload.get("business_id"), payload.get("branch_id"))
    rows = _read_file_rows(BANK_ACCOUNTS_LEDGER)
    existing = _find_existing_bank_account_by_key(
        rows,
        business_id=business_id,
        branch_id=branch_id,
        payload=payload,
    )
    if existing is not None:
        return _sanitize_bank_account(existing)
    now = _now()
    record: dict[str, Any] = {
        "id": str(uuid4()),
        "business_id": business_id,
        "branch_id": branch_id,
        "created_at": now,
        "updated_at": now,
    }
    _apply_bank_account_fields(record, payload, creating=True)
    record = _sanitize_bank_account(record)
    rows.insert(0, record)
    _write_secure_file_rows(BANK_ACCOUNTS_LEDGER, rows)
    return _sanitize_bank_account(record)


def update_bank_account(account_id: str, payload: dict[str, Any], user: dict[str, Any]) -> dict[str, Any]:
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="은행계좌 수정 권한이 없습니다")
    rows = _read_file_rows(BANK_ACCOUNTS_LEDGER)
    record = _find_bank_account(rows, account_id)
    if record is None:
        raise HTTPException(status_code=404, detail="수정할 은행계좌를 찾지 못했습니다")
    if payload.get("branch_id") is not None:
        _normalize_bank_scope(record.get("business_id"), payload.get("branch_id"))
        record["branch_id"] = str(payload.get("branch_id") or "").strip()
    _apply_bank_account_fields(record, payload, creating=False)
    record["updated_at"] = _now()
    sanitized = _sanitize_bank_account(record)
    rows = [sanitized if str(row.get("id") or "") == str(account_id) else row for row in rows]
    _write_secure_file_rows(BANK_ACCOUNTS_LEDGER, rows)
    return _sanitize_bank_account(sanitized)


BANK_CREDENTIAL_INPUT_FIELDS: tuple[tuple[str, str], ...] = (
    ("login_password", "password"),
    ("account_no", "account_no"),
    ("account_password", "account_password"),
    ("business_registration_no", "business_registration_no"),
    ("certificate_password", "certificate_password"),
)


def _branch_name_for_bank_account(branch_id: str) -> str:
    """Resolve a bank-account branch_id into the branch name used by platform_accounts."""
    branch_key = str(branch_id or "").strip()
    if not branch_key:
        return ""
    canonical = CANONICAL_BRANCH_BY_ID.get(branch_key)
    if isinstance(canonical, dict):
        name = str(canonical.get("name") or "").strip()
        if name:
            return name
    settings_data = _read_json_object("settings")
    ui_settings = settings_data.get("ui_settings") if isinstance(settings_data.get("ui_settings"), dict) else {}
    settings = _canonicalize_ui_settings(ui_settings)
    branches = settings.get("branches") if isinstance(settings.get("branches"), list) else []
    for item in branches:
        if isinstance(item, dict) and str(item.get("id") or "").strip() == branch_key:
            return str(item.get("name") or "").strip()
    return ""


def save_bank_credentials(account_id: str, payload: dict[str, Any], user: dict[str, Any]) -> dict[str, Any]:
    """Encrypt and persist bank quick-service credentials for a registered bank account.

    Secrets are never written to the bank_accounts ledger.  They are encrypted with
    ``_encrypt_secret`` on a linked ``platform_accounts`` row whose ``collection_mode``
    is ``bank-quick-service`` so PC Agent browser automation can read them through
    ``_bank_quick_credentials_for_account``.
    """
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="은행 인증정보 등록 권한이 없습니다")
    account_key = str(account_id or "").strip()
    rows = _read_file_rows(BANK_ACCOUNTS_LEDGER)
    account = _find_bank_account(rows, account_key)
    if account is None:
        raise HTTPException(status_code=404, detail="인증정보를 등록할 은행계좌를 찾지 못했습니다")
    service_code = _bank_service_code_for_account(account)
    if service_code not in BANK_QUICK_SERVICE_CONFIG:
        raise HTTPException(status_code=400, detail="간편/빠른조회를 지원하지 않는 은행계좌입니다")

    business_id = str(account.get("business_id") or "").strip()
    branch_id = str(account.get("branch_id") or "").strip()
    quick_config = BANK_QUICK_SERVICE_CONFIG[service_code]
    account_mask = str(account.get("account_number_masked") or "").strip()
    login_id = str(payload.get("login_id") or payload.get("username") or "").strip()
    if not login_id:
        login_id = "|".join(
            part
            for part in (
                service_code,
                business_id,
                branch_id,
                account_mask,
            )
            if part
        )
    if not login_id:
        raise HTTPException(status_code=400, detail="은행 인증정보 식별값을 만들 수 없습니다")

    account_payload: dict[str, Any] = {
        "service": service_code,
        "username": login_id,
        "business_id": business_id,
        "branch": _branch_name_for_bank_account(branch_id),
        "institution_code": service_code,
        "collection_mode": "bank-quick-service",
        "label": str(quick_config.get("label") or service_code),
        "login_url": str(quick_config.get("login_url") or ""),
        "category": "bank",
        "auth_owner": str(payload.get("auth_owner") or "").strip(),
        "mfa_method": str(payload.get("mfa_method") or "").strip(),
        "credential_expires_at": str(payload.get("credential_expires_at") or "").strip(),
    }
    registered_fields: list[str] = []
    for input_field, secret_field in BANK_CREDENTIAL_INPUT_FIELDS:
        value = str(payload.get(input_field) or "").strip()
        if value:
            account_payload[secret_field] = value
            registered_fields.append(input_field)

    saved = upsert_account(account_payload, user)

    now = _now()
    rows = _read_file_rows(BANK_ACCOUNTS_LEDGER)
    record = _find_bank_account(rows, account_key)
    if record is None:
        raise HTTPException(status_code=404, detail="인증정보를 등록할 은행계좌를 찾지 못했습니다")
    record["platform_account_id"] = str(saved.get("id") or "")
    record["credential_service"] = service_code
    record["credential_username"] = login_id
    record["credentials_registered_at"] = now
    record["connection_type"] = "browser"
    record["connector_type"] = "bank-quick-service"
    if str(record.get("status") or "").strip() in {"", "needs_auth", "error"}:
        record["status"] = "ready"
    record["updated_at"] = now
    sanitized = _sanitize_bank_account(record)
    rows = [sanitized if str(row.get("id") or "") == account_key else row for row in rows]
    _write_secure_file_rows(BANK_ACCOUNTS_LEDGER, rows)
    return {
        "ok": True,
        "bank_account": sanitized,
        "platform_account_id": str(saved.get("id") or ""),
        "service": service_code,
        "registered_fields": registered_fields,
    }


def _bank_transaction_matches_existing_transaction(bank_row: dict[str, Any], ledger_row: dict[str, Any]) -> bool:
    bank_date = _bank_date_key(bank_row.get("occurred_at"))
    ledger_date = _bank_date_key(ledger_row.get("date") or ledger_row.get("occurred_at"))
    if bank_date and ledger_date and bank_date != ledger_date:
        return False
    bank_amount = int(bank_row.get("amount") or 0)
    ledger_amount = int(abs(_amount(ledger_row.get("amount"))))
    if bank_amount != ledger_amount:
        return False
    bank_text = " ".join(
        str(bank_row.get(key) or "") for key in ("counterparty", "memo", "raw_memo", "category")
    ).lower()
    ledger_text = " ".join(
        str(ledger_row.get(key) or "") for key in ("vendor", "memo", "channel", "category", "source")
    ).lower()
    if not bank_text or not ledger_text:
        return True
    tokens = [token for token in re.split(r"\s+", bank_text) if len(token) >= 2]
    return any(token in ledger_text for token in tokens[:4])


def _annotate_bank_matches(rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], int, int]:
    if not rows:
        return rows, 0, 0
    existing_transactions = _read_file_rows("transactions")
    annotated: list[dict[str, Any]] = []
    matched = 0
    for row in rows:
        record = dict(row)
        if any(_bank_transaction_matches_existing_transaction(record, candidate) for candidate in existing_transactions):
            record["settlement_match"] = record.get("settlement_match") or "matched_existing_transaction"
            matched += 1
        else:
            record["settlement_match"] = record.get("settlement_match") or "unmatched"
        annotated.append(record)
    return annotated, matched, len(annotated) - matched


def _bank_transaction_source_hash(record: dict[str, Any]) -> str:
    fingerprint = json.dumps(
        {
            "business_id": record.get("business_id") or "",
            "bank_account_id": record.get("bank_account_id") or "",
            "occurred_at": record.get("occurred_at") or "",
            "direction": record.get("direction") or "",
            "amount": record.get("amount") or 0,
            "counterparty": record.get("counterparty") or "",
            "raw_memo": record.get("raw_memo") or record.get("memo") or "",
        },
        ensure_ascii=False,
        sort_keys=True,
    )
    return hashlib.sha256(fingerprint.encode("utf-8")).hexdigest()


def _coerce_bank_db_value(value: Any, pg_type: str) -> Any:
    """Coerce a bank-ledger value to the codec expected by the live PostgreSQL column."""
    if value is None:
        return None
    if pg_type in {"text", "varchar", "bpchar"}:
        return str(value)
    if value == "":
        return None if pg_type in {"date", "timestamp", "timestamptz", "uuid"} else value
    if pg_type == "uuid":
        return value if isinstance(value, UUID) else UUID(str(value))
    if pg_type == "date":
        parsed = _pg_ts(value)
        return parsed.date() if parsed else None
    if pg_type in {"timestamp", "timestamptz"}:
        return _pg_ts(value)
    if pg_type in {"int2", "int4", "int8"}:
        return int(value)
    return value


async def _db_insert_bank_transactions(records: list[dict[str, Any]]) -> set[str] | None:
    """Insert canonical rows into the operational ledger using its installed columns."""
    if not records:
        return set()
    import asyncpg

    table = "yeoljeong_bank_transactions"
    allowed_values = {
        "id": lambda row: row.get("id"),
        "business_id": lambda row: row.get("business_id"),
        "branch_id": lambda row: row.get("branch_id"),
        "bank_account_id": lambda row: row.get("bank_account_id"),
        "occurred_at": lambda row: row.get("occurred_at"),
        "occurred_date": lambda row: row.get("occurred_at"),
        "posted_at": lambda row: row.get("posted_at"),
        "direction": lambda row: row.get("direction"),
        "amount": lambda row: row.get("amount"),
        "balance": lambda row: row.get("balance"),
        "counterparty": lambda row: row.get("counterparty"),
        "memo": lambda row: row.get("memo"),
        "raw_memo": lambda row: row.get("raw_memo"),
        "category": lambda row: row.get("category"),
        "platform_match": lambda row: row.get("platform_match"),
        "settlement_match": lambda row: row.get("settlement_match"),
        "source": lambda row: row.get("source"),
        "source_hash": lambda row: row.get("source_hash"),
        "imported_at": lambda row: row.get("imported_at"),
        "created_at": lambda row: row.get("imported_at"),
        "updated_at": lambda row: row.get("imported_at"),
    }
    conn = await asyncpg.connect(_db_url(), timeout=5)
    try:
        column_rows = await conn.fetch(
            """
            SELECT column_name, data_type, udt_name, is_nullable, column_default,
                   identity_generation, is_generated
              FROM information_schema.columns
             WHERE table_schema = 'public' AND table_name = $1
            """,
            table,
        )
        column_types = {str(row["column_name"]): str(row["udt_name"] or row["data_type"] or "") for row in column_rows}
        required_columns = {"id", "business_id", "bank_account_id", "occurred_at", "direction", "amount", "source_hash"}
        if not required_columns.issubset(column_types):
            return None
        unsupported_required = {
            str(row["column_name"])
            for row in column_rows
            if str(row.get("is_nullable") or "").upper() == "NO"
            and row.get("column_default") is None
            and not row.get("identity_generation")
            and str(row.get("is_generated") or "NEVER").upper() == "NEVER"
            and str(row["column_name"]) not in allowed_values
        }
        if unsupported_required:
            return None
        columns = [name for name in allowed_values if name in column_types]
        non_nullable = {
            str(row["column_name"])
            for row in column_rows
            if str(row.get("is_nullable") or "").upper() == "NO"
            and row.get("column_default") is None
            and not row.get("identity_generation")
            and str(row.get("is_generated") or "NEVER").upper() == "NEVER"
        }
        placeholders = ", ".join(f"${index}" for index in range(1, len(columns) + 1))
        query = (
            f"INSERT INTO {table} ({', '.join(columns)}) VALUES ({placeholders}) "
            "ON CONFLICT DO NOTHING RETURNING source_hash"
        )
        prepared_values: list[list[Any]] = []
        for record in records:
            values = [
                _coerce_bank_db_value(allowed_values[column](record), column_types[column])
                for column in columns
            ]
            if any(value is None for column, value in zip(columns, values) if column in non_nullable):
                return None
            prepared_values.append(values)
        inserted_hashes: set[str] = set()
        async with conn.transaction():
            for values in prepared_values:
                inserted = await conn.fetchval(query, *values)
                if inserted:
                    inserted_hashes.add(str(inserted))
        return inserted_hashes
    finally:
        await conn.close()


def _normalize_bank_transaction(
    entry: dict[str, Any],
    *,
    business_id: str,
    branch_id: str,
    bank_account_id: str,
    source: str,
    now: str,
) -> dict[str, Any]:
    direction = _bank_direction(entry.get("direction"))
    if direction not in BANK_TRANSACTION_DIRECTIONS:
        raise HTTPException(status_code=400, detail="거래 방향(direction)은 in 또는 out 이어야 합니다")
    occurred_at = _transaction_date(entry.get("occurred_at"))
    if not occurred_at:
        raise HTTPException(status_code=400, detail="거래 발생일시(occurred_at)가 필요합니다")
    balance_value = entry.get("balance")
    record = {
        "id": str(entry.get("id") or uuid4()),
        "business_id": business_id,
        "branch_id": branch_id,
        "bank_account_id": bank_account_id,
        "occurred_at": occurred_at,
        "posted_at": _transaction_date(entry.get("posted_at")) if entry.get("posted_at") else "",
        "direction": direction,
        "amount": abs(_amount(entry.get("amount"))),
        "balance": _amount(balance_value) if balance_value not in (None, "") else None,
        "counterparty": str(entry.get("counterparty") or "").strip(),
        "memo": str(entry.get("memo") or "").strip(),
        "raw_memo": str(entry.get("raw_memo") or entry.get("memo") or "").strip(),
        "category": str(entry.get("category") or "").strip()
        or _transaction_category(" ".join(str(entry.get(k) or "") for k in ("counterparty", "memo", "raw_memo"))),
        "platform_match": str(entry.get("platform_match") or "").strip(),
        "settlement_match": str(entry.get("settlement_match") or "").strip(),
        "source": str(entry.get("source") or source or "manual").strip(),
        "imported_at": now,
    }
    # Incoming hashes are untrusted and older hashes were not account-scoped.
    # Always derive the canonical key after the server has fixed the scope.
    record["source_hash"] = _bank_transaction_source_hash(record)
    return record


def record_bank_transactions(payload: dict[str, Any], user: dict[str, Any]) -> dict[str, Any]:
    """Idempotently append bank ledger rows (manual/CSV/mock). Dedup by source_hash."""
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="은행거래 원장 기록 권한이 없습니다")
    business_id, branch_id = _normalize_bank_scope(payload.get("business_id"), payload.get("branch_id"))
    bank_account_id = str(payload.get("bank_account_id") or "").strip()
    if not bank_account_id:
        raise HTTPException(status_code=400, detail="bank_account_id가 필요합니다")
    accounts = _read_file_rows(BANK_ACCOUNTS_LEDGER)
    account = _find_bank_account(accounts, bank_account_id)
    if account is None or str(account.get("business_id") or "") != business_id:
        raise HTTPException(status_code=404, detail="등록된 은행계좌를 찾지 못했습니다")
    if branch_id and not _bank_account_matches_scope(account, business_id, branch_id):
        raise HTTPException(status_code=400, detail="선택한 은행계좌와 지점 범위가 일치하지 않습니다")
    entries = payload.get("transactions") or payload.get("rows") or []
    if not isinstance(entries, list):
        raise HTTPException(status_code=400, detail="transactions는 배열이어야 합니다")
    existing = _read_file_rows(BANK_TRANSACTIONS_LEDGER)
    existing_hashes = {str(row.get("source_hash") or "") for row in existing}
    db_enabled = _db_available()
    candidate_hashes: set[str] = set()
    now = _now()
    imported: list[dict[str, Any]] = []
    duplicate_rows = 0
    for entry in entries:
        if not isinstance(entry, dict):
            continue
        record = _normalize_bank_transaction(
            entry,
            business_id=business_id,
            branch_id=branch_id or str(account.get("branch_id") or ""),
            bank_account_id=bank_account_id,
            source=str(payload.get("source") or "manual"),
            now=now,
        )
        if record["source_hash"] in candidate_hashes or (not db_enabled and record["source_hash"] in existing_hashes):
            duplicate_rows += 1
            continue
        candidate_hashes.add(record["source_hash"])
        imported.append(record)
    if imported and db_enabled:
        db_inserted_hashes = _run_db(_db_insert_bank_transactions(imported))
        if db_inserted_hashes is None:
            raise HTTPException(status_code=503, detail="은행거래 DB 원장 반영에 실패했습니다")
        db_duplicate_rows = len(imported) - len(db_inserted_hashes)
        if db_duplicate_rows:
            duplicate_rows += db_duplicate_rows
            imported = [row for row in imported if row["source_hash"] in db_inserted_hashes]
    if imported:
        mirror_rows = [row for row in imported if row["source_hash"] not in existing_hashes]
        if mirror_rows:
            _write_secure_file_rows(BANK_TRANSACTIONS_LEDGER, mirror_rows + existing)
        for row in accounts:
            if str(row.get("id") or "") == bank_account_id:
                row["last_synced_at"] = now
                row["updated_at"] = now
                break
        _write_secure_file_rows(BANK_ACCOUNTS_LEDGER, accounts)
    return {
        "import": {
            "bank_account_id": bank_account_id,
            "business_id": business_id,
            "imported_rows": len(imported),
            "duplicate_rows": duplicate_rows,
        },
        "transactions": imported,
    }


def _bank_transaction_from_csv_row(raw: dict[str, str], *, source: str) -> dict[str, Any]:
    incoming = _amount(
        _first_present(raw, "입금액", "입금", "입금금액", "맡기신금액", "받으신금액", "deposit", "credit")
    )
    outgoing = _amount(
        _first_present(raw, "출금액", "출금", "출금금액", "찾으신금액", "지급금액", "withdrawal", "debit")
    )
    signed_amount = _amount(_first_present(raw, "거래금액", "금액", "amount"))
    if incoming:
        direction = "in"
        amount = incoming
    elif outgoing:
        direction = "out"
        amount = outgoing
    elif signed_amount < 0:
        direction = "out"
        amount = abs(signed_amount)
    else:
        direction = _bank_direction(_first_present(raw, "입출금", "구분", "거래구분", "direction")) or "in"
        amount = abs(signed_amount)
    occurred_at = (
        raw.get("거래일시")
        or raw.get("일시")
        or " ".join(
            item
            for item in (
                raw.get("거래일자") or raw.get("거래일") or raw.get("일자") or raw.get("날짜") or "",
                raw.get("거래시간") or raw.get("시간") or "",
            )
            if item
        )
    )
    memo = (
        raw.get("적요")
        or raw.get("거래내용")
        or raw.get("내용")
        or raw.get("기재내용")
        or raw.get("메모")
        or raw.get("memo")
        or ""
    )
    counterparty = (
        raw.get("보낸분/받는분")
        or raw.get("보낸분")
        or raw.get("받는분")
        or raw.get("거래처")
        or raw.get("상대계좌예금주")
        or raw.get("counterparty")
        or ""
    )
    return {
        "occurred_at": occurred_at,
        "posted_at": raw.get("기산일") or raw.get("처리일") or "",
        "direction": direction,
        "amount": amount,
        "balance": _first_present(raw, "잔액", "balance"),
        "counterparty": counterparty,
        "memo": memo,
        "raw_memo": memo or " / ".join(value for value in raw.values() if value),
        "category": raw.get("분류") or raw.get("카테고리") or "",
        "source": source,
    }


def _bank_header_key(value: Any) -> str:
    return re.sub(r"[\s\n\r\t()\[\]{}_-]+", "", str(value or "").strip()).lower()


_BANK_CANONICAL_HEADER_BY_KEY = {
    _bank_header_key(header): header
    for header in (
        *BANK_DATE_HEADERS,
        *BANK_AMOUNT_HEADERS,
        "거래시간", "시간", "입출금", "구분", "거래구분", "direction",
        "적요", "거래내용", "내용", "기재내용", "메모", "memo",
        "보낸분/받는분", "보낸분", "받는분", "거래처", "상대계좌예금주", "counterparty",
        "기산일", "처리일", "잔액", "balance", "분류", "카테고리",
    )
}


def _bank_canonical_header(value: Any) -> str:
    original = str(value or "").strip()
    return _BANK_CANONICAL_HEADER_BY_KEY.get(_bank_header_key(original), original)


def _bank_header_indexes(row: tuple[Any, ...] | list[Any]) -> tuple[list[str], bool]:
    headers = [_bank_canonical_header(value) for value in row]
    keys = {_bank_header_key(value) for value in headers if value not in (None, "")}
    date_keys = {_bank_header_key(value) for value in BANK_DATE_HEADERS}
    amount_keys = {_bank_header_key(value) for value in BANK_AMOUNT_HEADERS}
    return headers, bool(keys & date_keys) and bool(keys & amount_keys)


def _bank_rows_from_table(rows: Any, *, source: str) -> list[dict[str, Any]]:
    header: list[str] | None = None
    transactions: list[dict[str, Any]] = []
    for row_number, values in enumerate(rows, start=1):
        if row_number > BANK_STATEMENT_MAX_ROWS + 1:
            raise HTTPException(status_code=400, detail=f"은행 파일은 최대 {BANK_STATEMENT_MAX_ROWS:,}행까지 지원합니다")
        cells = list(values)
        if len(cells) > BANK_STATEMENT_MAX_COLUMNS:
            raise HTTPException(status_code=400, detail=f"은행 파일은 최대 {BANK_STATEMENT_MAX_COLUMNS}열까지 지원합니다")
        if header is None:
            candidate, valid = _bank_header_indexes(cells)
            if valid:
                header = candidate
            continue
        raw = {
            key: str(cells[index] if index < len(cells) and cells[index] is not None else "").strip()
            for index, key in enumerate(header)
            if key
        }
        if any(raw.values()):
            transactions.append(_bank_transaction_from_csv_row(raw, source=source))
    if header is None:
        raise HTTPException(status_code=400, detail="은행 파일에 거래일과 금액 필수 헤더가 필요합니다")
    return transactions


def _validate_bank_statement_file(filename: Any, content_type: Any, data: bytes) -> tuple[str, str]:
    safe_name = Path(str(filename or "bank-transactions.csv")).name
    extension = Path(safe_name).suffix.lower()
    if extension == ".xls":
        raise HTTPException(status_code=400, detail=".xls 파일은 .xlsx 또는 CSV로 저장한 뒤 업로드하십시오")
    if extension not in BANK_STATEMENT_EXTENSIONS:
        raise HTTPException(status_code=400, detail="은행 파일은 CSV, XLSX, XLSM 형식만 지원합니다")
    mime = str(content_type or "").split(";", 1)[0].strip().lower()
    if mime not in BANK_STATEMENT_MIME_TYPES[extension]:
        raise HTTPException(status_code=400, detail="파일 확장자와 MIME 형식이 일치하지 않습니다")
    if not data:
        raise HTTPException(status_code=400, detail="빈 은행 파일은 업로드할 수 없습니다")
    if len(data) > BANK_STATEMENT_MAX_BYTES:
        raise HTTPException(status_code=413, detail="은행 파일은 10MB 이하여야 합니다")
    return safe_name, extension


def _validate_excel_archive(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            total_size = 0
            for item in archive.infolist():
                total_size += item.file_size
                if total_size > BANK_STATEMENT_MAX_EXPANDED_BYTES:
                    raise HTTPException(status_code=400, detail="압축 해제 크기가 너무 큰 Excel 파일입니다")
                if item.file_size > 1_000_000 and item.compress_size > 0 and item.file_size / item.compress_size > 200:
                    raise HTTPException(status_code=400, detail="비정상 압축률의 Excel 파일입니다")
    except HTTPException:
        raise
    except (zipfile.BadZipFile, OSError) as exc:
        raise HTTPException(status_code=400, detail="손상되었거나 암호화된 Excel 파일입니다") from exc


def import_bank_transaction_file(
    payload: dict[str, Any], data: bytes, user: dict[str, Any]
) -> dict[str, Any]:
    """Parse a bank statement from bytes and persist through the canonical ledger path."""
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="은행거래 파일 반영 권한이 없습니다")
    filename, extension = _validate_bank_statement_file(payload.get("filename"), payload.get("content_type"), data)
    source = str(payload.get("source") or "file-upload").strip() or "file-upload"
    sheet_count = 1
    if extension == ".csv":
        decoded = _decode_csv(data)
        delimiter = _csv_delimiter(decoded)
        transactions = _bank_rows_from_table(csv.reader(decoded.splitlines(), delimiter=delimiter), source=source)
    else:
        _validate_excel_archive(data)
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_vba=False)
        except (InvalidFileException, OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
            raise HTTPException(status_code=400, detail="손상되었거나 암호화된 Excel 파일입니다") from exc
        try:
            sheet_count = len(workbook.worksheets)
            transactions = []
            header_error: HTTPException | None = None
            for sheet in workbook.worksheets:
                if sheet.max_row > BANK_STATEMENT_MAX_ROWS + 1 or sheet.max_column > BANK_STATEMENT_MAX_COLUMNS:
                    raise HTTPException(status_code=400, detail="Excel 행 또는 열 상한을 초과했습니다")
                try:
                    transactions.extend(_bank_rows_from_table(sheet.iter_rows(values_only=True), source=source))
                except HTTPException as exc:
                    if "필수 헤더" not in str(exc.detail):
                        raise
                    header_error = exc
            if not transactions and header_error:
                raise header_error
        finally:
            workbook.close()
    result = record_bank_transactions(
        {
            "business_id": payload.get("business_id") or MIA_BUSINESS_ID,
            "branch_id": payload.get("branch_id") or "",
            "bank_account_id": payload.get("bank_account_id") or "",
            "source": source,
            "transactions": transactions,
        },
        user,
    )
    metadata = {
        "filename": filename,
        "format": extension.removeprefix("."),
        "sheet_count": sheet_count,
        "parsed_rows": len(transactions),
        "imported_rows": result["import"]["imported_rows"],
        "duplicate_rows": result["import"]["duplicate_rows"],
    }
    return {**result, **metadata, "import": {**result["import"], **metadata, "source": source}}


def import_bank_transaction_csv(payload: dict[str, Any], user: dict[str, Any]) -> dict[str, Any]:
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="은행거래 CSV 반영 권한이 없습니다")
    csv_text = str(payload.get("csv_text") or "")
    if not csv_text.strip():
        raise HTTPException(status_code=400, detail="은행 거래 CSV 내용이 필요합니다")
    filename = Path(str(payload.get("filename") or "bank-transactions.csv")).name
    source = str(payload.get("source") or "csv").strip() or "csv"
    decoded = _decode_csv(csv_text.encode("utf-8-sig"))
    transactions = _bank_rows_from_table(
        csv.reader(decoded.splitlines(), delimiter=_csv_delimiter(decoded)),
        source=source,
    )
    result = record_bank_transactions(
        {
            "business_id": payload.get("business_id") or MIA_BUSINESS_ID,
            "branch_id": payload.get("branch_id") or "",
            "bank_account_id": payload.get("bank_account_id") or "",
            "source": source,
            "transactions": transactions,
        },
        user,
    )
    return {
        **result,
        "import": {
            **result["import"],
            "filename": filename,
            "source": source,
            "parsed_rows": len(transactions),
        },
    }


def collect_bank_account_transactions(account_id: str, payload: dict[str, Any], user: dict[str, Any]) -> dict[str, Any]:
    """Run the configured safe bank connector for one account.

    Real bank login/open-banking connectors intentionally return NOT_CONFIGURED until
    a certified provider or vetted browser connector is attached.
    connection_type="browser" routes to the PC Agent / Browser Bridge connector.
    """
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="은행 자동수집 실행 권한이 없습니다")
    accounts = _read_file_rows(BANK_ACCOUNTS_LEDGER)
    account = _find_bank_account(accounts, account_id)
    if account is None:
        raise HTTPException(status_code=404, detail="수집할 은행계좌를 찾지 못했습니다")
    business_id, requested_branch_id = _normalize_bank_scope(
        payload.get("business_id") or account.get("business_id"),
        payload.get("branch_id") or account.get("branch_id") or "",
    )
    if not _bank_account_matches_scope(account, business_id, requested_branch_id):
        raise HTTPException(status_code=404, detail="사업자/지점 범위에 맞는 은행계좌를 찾지 못했습니다")

    date_from, date_to = _valid_range_bounds(payload.get("date_from"), payload.get("date_to"))
    connection_type = _bank_connection_type(account.get("connection_type"), default="mock")
    branch_id = requested_branch_id or str(account.get("branch_id") or "")

    if str(account.get("status") or "") in {"paused", "error"}:
        return _bank_unconfigured_collect_result(
            account,
            business_id=business_id,
            branch_id=branch_id,
            date_from=date_from,
            date_to=date_to,
            status=str(account.get("status") or "paused"),
            reason="계좌 상태가 active가 아니어서 자동수집을 실행하지 않았습니다.",
        )
    if connection_type not in BANK_CONFIGURED_COLLECTION_TYPES:
        return _bank_unconfigured_collect_result(
            account,
            business_id=business_id,
            branch_id=branch_id,
            date_from=date_from,
            date_to=date_to,
            reason="오픈뱅킹/은행 실시간 조회 커넥터가 아직 연결되지 않았습니다. CSV/수동 대체 수집을 사용하십시오.",
        )

    # ── Browser connector path ───────────────────────────────────────────────
    if connection_type == "browser":
        return _collect_bank_via_browser(
            account,
            payload,
            business_id=business_id,
            branch_id=branch_id,
            date_from=date_from,
            date_to=date_to,
            user=user,
        )

    entries = payload.get("transactions") or payload.get("rows") or []
    if not isinstance(entries, list):
        raise HTTPException(status_code=400, detail="transactions는 배열이어야 합니다")
    # mock 커넥터: 페이로드에 거래 없으면 테스트용 결정적 데이터를 자동 생성.
    if connection_type == "mock" and not entries:
        effective_date = date_from or _bank_date_key(_now())
        entries = [
            {
                "occurred_at": effective_date + " 10:00:00",
                "direction": "in",
                "amount": 1_000_000,
                "counterparty": "mock-deposit",
                "memo": "Mock 입금 테스트",
                "raw_memo": "Mock 입금 테스트",
                "source": "mock",
            },
            {
                "occurred_at": effective_date + " 14:00:00",
                "direction": "out",
                "amount": 300_000,
                "counterparty": "mock-withdrawal",
                "memo": "Mock 출금 테스트",
                "raw_memo": "Mock 출금 테스트",
                "source": "mock",
            },
        ]
    scoped_entries = [
        entry
        for entry in entries
        if isinstance(entry, dict) and _bank_within_range(entry.get("occurred_at"), date_from, date_to)
    ]
    annotated_entries, matched_count, unmatched_count = _annotate_bank_matches(scoped_entries)
    result = record_bank_transactions(
        {
            "business_id": business_id,
            "branch_id": branch_id,
            "bank_account_id": account_id,
            "source": str(payload.get("source") or connection_type),
            "transactions": annotated_entries,
        },
        user,
    )
    imported_rows = int(result.get("import", {}).get("imported_rows") or 0)
    duplicate_rows = int(result.get("import", {}).get("duplicate_rows") or 0)
    account_after = _find_bank_account(_read_file_rows(BANK_ACCOUNTS_LEDGER), account_id) or account
    imported_transactions = result.get("transactions") or []
    total_in = sum(int(row.get("amount") or 0) for row in imported_transactions if row.get("direction") == "in")
    total_out = sum(int(row.get("amount") or 0) for row in imported_transactions if row.get("direction") == "out")
    status = "completed" if imported_rows else "no_records"
    return {
        **result,
        "collection": {
            "bank_account_id": account_id,
            "business_id": business_id,
            "branch_id": branch_id,
            "status": status,
            "connector_status": "CONFIGURED",
            "connection_type": connection_type,
            "message": "은행 거래 수집이 완료되었습니다." if imported_rows else "신규 수집 거래가 없습니다.",
            "collected_rows": len(scoped_entries),
            "imported_rows": imported_rows,
            "duplicate_rows": duplicate_rows,
            "matched_count": matched_count,
            "unmatched_count": unmatched_count,
            "total_in": total_in,
            "total_out": total_out,
            "net_amount": total_in - total_out,
            "last_collected_at": str(account_after.get("last_synced_at") or ""),
            "date_from": date_from,
            "date_to": date_to,
        },
    }


def _run_bank_browser_async(coro: Any) -> Any:
    try:
        asyncio.get_running_loop()
    except RuntimeError:
        return asyncio.run(coro)
    close = getattr(coro, "close", None)
    if callable(close):
        close()
    raise RuntimeError("bank browser automation cannot run inside an active event loop")


async def _with_bank_browser_timeout(coro_factory: Any, timeout_seconds: float) -> Any:
    return await asyncio.wait_for(coro_factory(), timeout=max(1.0, float(timeout_seconds or 1.0)))


def _bank_service_code_for_account(account: dict[str, Any]) -> str:
    institution_code = str(account.get("institution_code") or "").strip()
    if institution_code in BANK_QUICK_SERVICE_CONFIG:
        return institution_code
    bank_code = str(account.get("bank_code") or "").strip()
    bank_name = str(account.get("bank_name") or "").strip()
    if bank_code == "088" or "신한" in bank_name:
        return "shinhan_business"
    if bank_code == "003" or "기업" in bank_name or "IBK" in bank_name.upper():
        return "ibk_business"
    return ""


def _pc_agent_route_execute_json(payload: dict[str, Any], timeout_seconds: float = 25.0) -> dict[str, Any]:
    body = json.dumps(payload).encode("utf-8")
    headers = {"Content-Type": "application/json"}
    urls = [
        str(os.getenv("AADS_PC_AGENT_ROUTE_EXECUTE_URL") or "").strip(),
        "http://127.0.0.1:8080/api/v1/pc-agent/route-execute",
        "http://aads-server:8080/api/v1/pc-agent/route-execute",
        "http://aads-server-green:8080/api/v1/pc-agent/route-execute",
    ]
    seen: set[str] = set()
    for url in [u for u in urls if u and not (u in seen or seen.add(u))]:
        request = urllib.request.Request(url, data=body, headers=headers, method="POST")
        try:
            with urllib.request.urlopen(request, timeout=max(1.0, float(timeout_seconds or 1.0))) as response:
                parsed = json.loads(response.read().decode("utf-8"))
        except urllib.error.HTTPError as exc:
            try:
                parsed = json.loads(exc.read().decode("utf-8", errors="ignore"))
            except Exception:
                continue
            detail = parsed.get("detail") if isinstance(parsed, dict) else None
            if isinstance(detail, dict):
                return detail
            if isinstance(parsed, dict):
                return parsed
            continue
        except (urllib.error.URLError, TimeoutError, OSError, json.JSONDecodeError):
            continue
        if isinstance(parsed, dict):
            return parsed
    return {}


def _bank_browser_tabs_from_route_result(result: dict[str, Any]) -> list[dict[str, Any]]:
    candidates = [
        result.get("tabs"),
        (result.get("result") or {}).get("tabs") if isinstance(result.get("result"), dict) else None,
        ((result.get("result") or {}).get("result") or {}).get("tabs")
        if isinstance(result.get("result"), dict) and isinstance((result.get("result") or {}).get("result"), dict)
        else None,
    ]
    for candidate in candidates:
        if isinstance(candidate, list):
            return [item for item in candidate if isinstance(item, dict)]
    return []


def _reset_shinhan_timeout_probe_to_idpw(
    *,
    browser_work_key: str,
    browser_agent_id: str,
    portal_url: str,
    browser_preferred_port: int | None,
    timeout_seconds: float,
) -> dict[str, str]:
    diagnostics: dict[str, str] = {"attempted": "1"}
    close_payload = {
        "command_type": "browser_close_tab",
        "agent_id": browser_agent_id,
        "job_type": "financial_exclusive",
        "required_capabilities": ["interactive_browser"],
        "queue_if_busy": True,
        "wait_for_turn": True,
        "queue_wait_timeout_seconds": min(45.0, max(10.0, timeout_seconds / 4.0)),
        "lease_ttl_seconds": 90,
        "command_timeout_seconds": 15,
        "params": {
            "work_key": browser_work_key,
            "url_pattern": "fincert|yeskey|cert",
            "keep_last": False,
        },
    }
    close_result = _pc_agent_route_execute_json(close_payload, timeout_seconds=min(75.0, max(25.0, timeout_seconds / 2.0)))
    diagnostics["certificate_tab_close_status"] = str(close_result.get("status") or "")[:40]
    close_tabs = _bank_browser_tabs_from_route_result(close_result)
    if close_tabs:
        diagnostics["certificate_tab_close_remaining_tabs"] = str(len(close_tabs))
    close_data = close_result.get("data") if isinstance(close_result.get("data"), dict) else {}
    close_nested = close_result.get("result") if isinstance(close_result.get("result"), dict) else {}
    if not close_data and isinstance(close_nested.get("result"), dict):
        close_data = close_nested.get("result") or {}
    elif not close_data and isinstance(close_nested.get("data"), dict):
        close_data = close_nested.get("data") or {}
    if isinstance(close_data, dict) and "closed" in close_data:
        diagnostics["certificate_tab_closed"] = str(int(close_data.get("closed") or 0))

    launch_params: dict[str, Any] = {
        "work_key": browser_work_key,
        "url": portal_url,
        "new_window": False,
        "ready_timeout_seconds": 20,
    }
    if browser_preferred_port:
        launch_params["preferred_port"] = browser_preferred_port
    launch_result = _pc_agent_route_execute_json(
        {
            "command_type": "browser_launch",
            "agent_id": browser_agent_id,
            "job_type": "financial_exclusive",
            "required_capabilities": ["interactive_browser"],
            "queue_if_busy": True,
            "wait_for_turn": True,
            "queue_wait_timeout_seconds": min(60.0, max(15.0, timeout_seconds / 3.0)),
            "lease_ttl_seconds": 120,
            "command_timeout_seconds": 35,
            "params": launch_params,
        },
        timeout_seconds=min(110.0, max(55.0, timeout_seconds / 2.0)),
    )
    diagnostics["work_key_relaunch_status"] = str(launch_result.get("status") or "")[:40]
    launch_result_block = launch_result.get("result") if isinstance(launch_result.get("result"), dict) else {}
    launch_data = launch_result.get("data") if isinstance(launch_result.get("data"), dict) else {}
    if not launch_data and isinstance(launch_result_block.get("result"), dict):
        launch_data = launch_result_block.get("result") or {}
    elif not launch_data and isinstance(launch_result_block.get("data"), dict):
        launch_data = launch_result_block.get("data") or {}
    if isinstance(launch_data, dict):
        diagnostics["work_key_relaunch_cdp_ready"] = "1" if launch_data.get("cdp_ready") else "0"
        diagnostics["work_key_relaunch_navigated"] = "1" if launch_data.get("navigated") else "0"
    return diagnostics


def _probe_bank_browser_timeout_state(
    *,
    browser_work_key: str,
    browser_agent_id: str,
    timeout_seconds: float,
    prefer_saved_idpw_login: bool = False,
    portal_url: str = "",
    browser_preferred_port: int | None = None,
) -> dict[str, Any] | None:
    if not browser_work_key or not browser_agent_id:
        return None
    route_result = _pc_agent_route_execute_json(
        {
            "command_type": "browser_tabs",
            "agent_id": browser_agent_id,
            "job_type": "financial_exclusive",
            "required_capabilities": ["interactive_browser"],
            "queue_if_busy": True,
            "wait_for_turn": True,
            "queue_wait_timeout_seconds": min(45.0, max(10.0, float(timeout_seconds or 20.0) / 4.0)),
            "lease_ttl_seconds": 90,
            "command_timeout_seconds": min(20.0, max(5.0, float(timeout_seconds or 20.0))),
            "params": {
                "work_key": browser_work_key,
                "timeout": min(20.0, max(5.0, float(timeout_seconds or 20.0))),
            },
        },
        timeout_seconds=min(75.0, max(25.0, float(timeout_seconds or 25.0) / 2.0)),
    )
    tabs = _bank_browser_tabs_from_route_result(route_result)
    tab_text = " ".join(
        f"{str(tab.get('title') or '')} {str(tab.get('url') or '')}".lower()
        for tab in tabs
    )
    if "4user.yeskey.or.kr/fincert" in tab_text or "fincert" in tab_text:
        reset_diagnostics: dict[str, str] = {}
        if prefer_saved_idpw_login:
            reset_diagnostics = _reset_shinhan_timeout_probe_to_idpw(
                browser_work_key=browser_work_key,
                browser_agent_id=browser_agent_id,
                portal_url=portal_url or BANK_QUICK_SERVICE_CONFIG["shinhan_business"]["login_url"],
                browser_preferred_port=browser_preferred_port,
                timeout_seconds=timeout_seconds,
            )
            return {
                "status": "action_required",
                "error_code": "BANK_BROWSER_IDPW_RETRY_REQUIRED",
                "rows": [],
                "row_count": 0,
                "diagnostics": {
                    "browser_work_key": browser_work_key,
                    "browser_agent_id": browser_agent_id,
                    "browser_timeout_seconds": str(int(timeout_seconds)),
                    "last_observed_stage": "financial certificate iframe",
                    "screen_state": "login_required",
                    "screen_reason_code": "SHINHAN_FINCERT_TIMEOUT_BUT_IDPW_CONFIGURED",
                    "screen_requires_operator": "0",
                    "screen_suggested_action": "retry_saved_idpw_login_same_work_key",
                    "suggested_action": "retry_saved_idpw_login_same_work_key",
                    "shinhan_auth_challenge_policy": "prefer_saved_idpw_login",
                    "shinhan_timeout_idpw_reset": reset_diagnostics,
                    "pc_agent_probe_status": str(route_result.get("status") or ""),
                    "pc_agent_probe_tab_count": str(len(tabs)),
                },
                "message": "신한 금융인증서 탭이 감지됐지만 저장된 ID/PW 로그인 설정이 있으므로 금융인증서 완료 대신 같은 work key에서 ID/PW 로그인으로 재시도하십시오.",
            }
        return {
            "status": "action_required",
            "error_code": "BANK_BROWSER_AUTH_CHALLENGE_DETECTED",
            "rows": [],
            "row_count": 0,
            "diagnostics": {
                "browser_work_key": browser_work_key,
                "browser_agent_id": browser_agent_id,
                "browser_timeout_seconds": str(int(timeout_seconds)),
                "last_observed_stage": "financial certificate iframe",
                "screen_state": "certificate_password_required",
                "screen_reason_code": "SHINHAN_FINCERT_IFRAME_DETECTED_AFTER_TIMEOUT",
                "screen_requires_operator": "1",
                "screen_suggested_action": "complete_financial_certificate_then_retry_same_work_key",
                "suggested_action": "complete_financial_certificate_then_retry_same_work_key",
                "pc_agent_probe_status": str(route_result.get("status") or ""),
                "pc_agent_probe_tab_count": str(len(tabs)),
            },
            "message": "신한 금융인증서 입력 화면이 감지됐습니다. 전용 PC에서 인증을 완료한 뒤 같은 work key로 재시도하십시오.",
        }
    return None


def _branch_names_for_bank_match(branch_id: str) -> set[str]:
    branch_text = str(branch_id or "").strip()
    names = {branch_text} if branch_text else set()
    for item in CANONICAL_BRANCHES:
        if str(item.get("id") or "") == branch_text:
            names.add(str(item.get("name") or ""))
            names.add(str(item.get("label") or ""))
    return {name for name in names if name}


def _branch_id_for_financial_scope(business_id: str, branch: str) -> str:
    business_key = str(business_id or "").strip()
    branch_text = BRANCH_ALIASES.get(str(branch or "").strip(), str(branch or "").strip())
    if not business_key or not branch_text:
        return ""
    settings_data = _read_json_object("settings")
    ui_settings = settings_data.get("ui_settings") if isinstance(settings_data.get("ui_settings"), dict) else {}
    settings = _canonicalize_ui_settings(ui_settings)
    branches = settings.get("branches") if isinstance(settings.get("branches"), list) else []
    for item in [*branches, *CANONICAL_BRANCHES]:
        if not isinstance(item, dict):
            continue
        item_business = str(item.get("businessId") or item.get("business_id") or "").strip()
        item_id = str(item.get("id") or "").strip()
        item_name = BRANCH_ALIASES.get(str(item.get("name") or "").strip(), str(item.get("name") or "").strip())
        if item_business == business_key and branch_text in {item_id, item_name}:
            return item_id
    return ""


def _bank_account_for_financial_service(
    service_code: str,
    *,
    business_id: str,
    branch_id: str,
) -> dict[str, Any] | None:
    rows = _read_file_rows(BANK_ACCOUNTS_LEDGER)
    candidates: list[dict[str, Any]] = []
    for row in rows if isinstance(rows, list) else []:
        if str(row.get("business_id") or "").strip() != business_id:
            continue
        row_branch_id = str(row.get("branch_id") or "").strip()
        if branch_id and row_branch_id and row_branch_id != branch_id:
            continue
        if str(row.get("status") or "active").strip() not in {"active", ""}:
            continue
        if _bank_service_code_for_account(row) != service_code:
            continue
        candidates.append(row)
    if not candidates:
        return None
    candidates.sort(
        key=lambda row: (
            str(row.get("connection_type") or "") == "browser",
            bool(row.get("auto_sync")),
            str(row.get("updated_at") or row.get("created_at") or ""),
        ),
        reverse=True,
    )
    return candidates[0]


def _bank_quick_credentials_for_account(
    account: dict[str, Any],
    *,
    business_id: str,
    branch_id: str,
    tenant_id: str = "",
) -> dict[str, str]:
    """Return decrypted read-only bank quick-service credentials for browser fill.

    Values are used only inside PC Agent browser automation and must not be
    copied into diagnostics, logs, or API responses.
    """
    service_code = _bank_service_code_for_account(account)
    if not service_code:
        return {}
    branch_names = _branch_names_for_bank_match(branch_id or str(account.get("branch_id") or ""))
    rows = _read("platform_accounts")
    candidates: list[dict[str, Any]] = []
    for row in rows if isinstance(rows, list) else []:
        if str(row.get("service") or "").strip() != service_code:
            continue
        if str(row.get("collection_mode") or "").strip() != "bank-quick-service":
            continue
        if business_id and str(row.get("business_id") or "").strip() != business_id:
            continue
        row_branch = str(row.get("branch") or "").strip()
        if branch_names and row_branch and row_branch not in branch_names:
            continue
        candidates.append(row)
    if not candidates:
        return {}
    linked_platform_id = str(account.get("platform_account_id") or "").strip()
    credential_username = str(account.get("credential_username") or "").strip().lower()
    account_mask = str(account.get("account_number_masked") or "").strip()

    def _candidate_rank(row: dict[str, Any]) -> tuple[int, int, int, str]:
        """Prefer the explicitly linked Vault row before scope-only matches.

        Older data can have several bank quick-service rows for one branch.  The
        previous "first row wins" behavior could therefore decrypt another
        account's credentials.  The rank is deterministic and never compares
        plaintext secrets.
        """
        row_id = str(row.get("id") or "").strip()
        row_username = str(row.get("username") or "").strip().lower()
        row_mask = str(row.get("account_no_masked") or "").strip()
        return (
            1 if linked_platform_id and row_id == linked_platform_id else 0,
            1 if credential_username and row_username == credential_username else 0,
            1 if account_mask and row_mask and account_mask == row_mask else 0,
            str(row.get("updated_at") or row.get("created_at") or ""),
        )

    candidates.sort(key=_candidate_rank, reverse=True)
    selected = candidates[0]
    credentials: dict[str, str] = {
        "quick_account_configured": "1",
        "login_username": str(selected.get("username") or "").strip(),
        "portal_url": _normalize_bank_quick_login_url(service_code, selected.get("login_url")),
    }
    for plaintext_field, encrypted_field in (
        ("login_password", "password_enc"),
        ("account_no", "account_no_enc"),
        ("account_password", "account_password_enc"),
        ("business_registration_no", "business_registration_no_enc"),
        ("certificate_password", "certificate_password_enc"),
    ):
        encrypted = str(selected.get(encrypted_field) or "")
        if encrypted:
            credentials[plaintext_field] = _decrypt_secret(encrypted)
    if not credentials.get("login_username") or not credentials.get("login_password"):
        vault_login = _bank_login_from_agent_vault(
            service=service_code,
            tenant_id=tenant_id,
            expected_username=str(credentials.get("login_username") or selected.get("username") or ""),
            business_id=business_id,
            branch_names=branch_names,
        )
        if vault_login:
            credentials["login_username"] = vault_login["login_username"]
            credentials["login_password"] = vault_login["login_password"]
    return credentials


def _public_bank_quick_job(item: dict[str, Any], *, account_id: str = "") -> dict[str, Any]:
    """Return a secret-free quick-inquiry job contract for API clients."""
    payload = item.get("payload") if isinstance(item.get("payload"), dict) else {}
    result = item.get("result") if isinstance(item.get("result"), dict) else {}
    collections = result.get("bank_collections") if isinstance(result.get("bank_collections"), list) else []
    collection = next(
        (
            row
            for row in collections
            if isinstance(row, dict)
            and (not account_id or str(row.get("bank_account_id") or "") == account_id)
        ),
        {},
    )
    retry_after_seconds = 0
    next_run_at = str(item.get("next_run_at") or "").strip()
    if next_run_at and str(item.get("status") or "") in {"queued", "action_required"}:
        try:
            parsed = datetime.fromisoformat(next_run_at.replace("Z", "+00:00"))
            now = datetime.now(parsed.tzinfo or KST)
            retry_after_seconds = max(0, int((parsed - now).total_seconds()))
        except ValueError:
            retry_after_seconds = 0
    status = str(item.get("status") or "queued")
    return {
        "job_id": str(item.get("id") or ""),
        "status": status,
        "action_required": status == "action_required",
        "bank_account_id": account_id or str(payload.get("bank_account_id") or ""),
        "service": str(item.get("service") or payload.get("service") or ""),
        "business_id": str(item.get("business_id") or payload.get("business_id") or ""),
        "branch_id": str(item.get("branch") or payload.get("branch") or ""),
        "date_from": str(payload.get("date_from") or ""),
        "date_to": str(payload.get("date_to") or ""),
        "attempt_count": int(item.get("attempt_count") or 0),
        "retry_after_seconds": retry_after_seconds,
        "next_run_at": next_run_at,
        "error_code": str(item.get("error_code") or collection.get("error_code") or ""),
        "message": str(item.get("message") or collection.get("message") or ""),
        "created_at": str(item.get("created_at") or ""),
        "updated_at": str(item.get("updated_at") or ""),
        "started_at": str(item.get("started_at") or ""),
        "finished_at": str(item.get("finished_at") or ""),
        "result": {
            "collected_rows": int(collection.get("collected_rows") or 0),
            "imported_rows": int(collection.get("imported_rows") or 0),
            "duplicate_rows": int(collection.get("duplicate_rows") or 0),
            "verified_no_records": str(collection.get("status") or "") == "no_records",
            "last_collected_at": str(collection.get("last_collected_at") or ""),
        },
    }


def enqueue_bank_quick_inquiry(
    account_id: str,
    payload: dict[str, Any],
    user: dict[str, Any],
) -> dict[str, Any]:
    """Queue a bank quick/simple inquiry without putting secrets in the job."""
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="은행 빠른조회 실행 권한이 없습니다")
    account_key = str(account_id or "").strip()
    account = _find_bank_account(_read_file_rows(BANK_ACCOUNTS_LEDGER), account_key)
    if account is None:
        raise HTTPException(status_code=404, detail="조회할 은행계좌를 찾지 못했습니다")
    service_code = _bank_service_code_for_account(account)
    if service_code not in BANK_QUICK_SERVICE_CONFIG:
        raise HTTPException(status_code=400, detail="빠른조회/간편조회를 지원하지 않는 은행계좌입니다")
    business_id, requested_branch_id = _normalize_bank_scope(
        payload.get("business_id") or account.get("business_id"),
        payload.get("branch_id") or account.get("branch_id") or "",
    )
    if not _bank_account_matches_scope(account, business_id, requested_branch_id):
        raise HTTPException(status_code=404, detail="사업자/지점 범위에 맞는 은행계좌를 찾지 못했습니다")
    date_from, date_to = _valid_range_bounds(payload.get("date_from"), payload.get("date_to"))
    today = datetime.now(KST).date().isoformat()
    date_from = date_from or today
    date_to = date_to or today
    branch_id = requested_branch_id or str(account.get("branch_id") or "")
    browser_agent_id = str(payload.get("browser_agent_id") or "").strip()
    work_key = str(payload.get("browser_work_key") or "").strip() or (
        f"yeoljeong-bank-{service_code}-{business_id}-{branch_id or 'common'}"
    )
    sync_job_id = f"bank-quick-{uuid4().hex[:16]}"
    queue_payload = {
        "services": [service_code],
        "service": service_code,
        "bank_account_id": account_key,
        "business_id": business_id,
        "branch": branch_id,
        "all_businesses": False,
        "bank_only": True,
        "skip_financial_accounts": False,
        "date_from": date_from,
        "date_to": date_to,
        "source": "bank-quick-api",
        "prefer_pc_agent": True,
        "require_pc_agent": True,
        "auto_open_bank_browser": bool(payload.get("auto_open_browser", True)),
        "force_recreate_bank_browser": bool(payload.get("force_recreate_browser", False)),
        "bank_browser_work_key": work_key,
        "browser_agent_id": browser_agent_id,
        "pc_agent_id": browser_agent_id,
        "required_browser_agent_id": browser_agent_id,
        "sync_job_id": sync_job_id,
    }
    from app.services.pc_agent_collection_queue import enqueue_collection_item

    queued = enqueue_collection_item(
        {
            "tenant_id": str(user.get("tenant_id") or "").strip(),
            "queue_type": "bank",
            "site_key": f"bank:{service_code}",
            "service": service_code,
            "business_id": business_id,
            "branch": branch_id,
            "work_key": work_key,
            "runtime": "pc_agent",
            "priority": 10,
            "min_interval_seconds": 300,
            "latest_only": True,
            "payload": queue_payload,
            "created_by": str(user.get("email") or user.get("id") or "bank-quick-api"),
        }
    )
    return _public_bank_quick_job(queued, account_id=account_key)


def get_bank_quick_inquiry(job_id: str, user: dict[str, Any]) -> dict[str, Any]:
    """Read one quick-inquiry status/result while enforcing tenant scope."""
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="은행 빠른조회 상태 조회 권한이 없습니다")
    from app.services.pc_agent_collection_queue import queue_snapshot

    target = str(job_id or "").strip()
    item = next((row for row in queue_snapshot(200) if str(row.get("id") or "") == target), None)
    if item is None:
        raise HTTPException(status_code=404, detail="은행 빠른조회 작업을 찾지 못했습니다")
    if str(item.get("queue_type") or "") != "bank":
        raise HTTPException(status_code=404, detail="은행 빠른조회 작업을 찾지 못했습니다")
    tenant_id = str(user.get("tenant_id") or "").strip()
    item_tenant_id = str(item.get("tenant_id") or "").strip()
    if tenant_id and item_tenant_id and tenant_id != item_tenant_id:
        raise HTTPException(status_code=404, detail="은행 빠른조회 작업을 찾지 못했습니다")
    payload = item.get("payload") if isinstance(item.get("payload"), dict) else {}
    return _public_bank_quick_job(item, account_id=str(payload.get("bank_account_id") or ""))


def _business_entity_type_for_bank_scope(business_id: str) -> str:
    business_key = str(business_id or "").strip()
    if not business_key:
        return ""
    settings_data = _read_json_object("settings")
    ui_settings = settings_data.get("ui_settings") if isinstance(settings_data.get("ui_settings"), dict) else {}
    settings = _canonicalize_ui_settings(ui_settings)
    businesses = settings.get("businesses") if isinstance(settings.get("businesses"), list) else []
    for item in [*businesses, *CANONICAL_BUSINESSES]:
        if not isinstance(item, dict):
            continue
        if str(item.get("id") or "").strip() != business_key:
            continue
        return str(
            item.get("entityType")
            or item.get("entity_type")
            or item.get("business_type")
            or item.get("businessType")
            or ""
        ).strip()
    return ""


def _collect_bank_via_browser(
    account: dict[str, Any],
    payload: dict[str, Any],
    *,
    business_id: str,
    branch_id: str,
    date_from: str,
    date_to: str,
    user: dict[str, Any],
) -> dict[str, Any]:
    """Orchestrate a bank browser collection via PC Agent / Browser Bridge.

    Security: never attempts headless login.  Credential fields from payload
    are used only as routing keys (session_id / work_key) and are not stored.
    """
    from app.services.yeoljeong_bank_browser_connector import (
        bank_browser_work_key,
        collect_bank_via_browser_session_async,
        ibk_business_browser_work_key,
        shinhan_individual_browser_work_key,
    )

    account_id = str(account.get("id") or "")
    browser_session_id = str(payload.get("browser_session_id") or "").strip()
    browser_work_key_val = str(payload.get("browser_work_key") or "").strip()
    business_entity_type = _business_entity_type_for_bank_scope(business_id)
    service_code = _bank_service_code_for_account(account)
    if not browser_work_key_val:
        if service_code == "shinhan_business" and business_entity_type not in {"corporation", "corporate", "법인"}:
            browser_work_key_val = shinhan_individual_browser_work_key(business_id, branch_id)
        elif service_code == "ibk_business":
            browser_work_key_val = ibk_business_browser_work_key(business_id, branch_id)
        else:
            browser_work_key_val = bank_browser_work_key(account_id, business_id, branch_id)
    bank_credentials = _bank_quick_credentials_for_account(
        account,
        business_id=business_id,
        branch_id=branch_id,
        tenant_id=str(user.get("tenant_id") or ""),
    )
    if service_code in {"shinhan_business", "ibk_business"} and bank_credentials.get("quick_account_configured") == "1":
        required_fields = {
            "account_no": "조회용 계좌번호",
            "account_password": "계좌비밀번호",
            "business_registration_no": "사업자번호",
        }
        if service_code == "shinhan_business":
            required_fields = {"login_password": "로그인 비밀번호", **required_fields}
        missing_labels = [
            label
            for field, label in required_fields.items()
            if not str(bank_credentials.get(field) or "").strip()
        ]
        if missing_labels:
            return {
                "collection": {
                    "bank_account_id": account_id,
                    "business_id": business_id,
                    "branch_id": branch_id,
                    "status": "credential_required",
                    "connector_status": "ACTION_REQUIRED",
                    "connection_type": "browser",
                    "message": "은행 간편/빠른조회 필수값 누락: " + ", ".join(missing_labels),
                    "error_code": "MISSING_CREDENTIALS",
                    "diagnostics": {
                        "auth_mode": "pc_agent_browser",
                        "connector": "bank_browser",
                        "bank_account_id": account_id,
                        "bank_code": str(account.get("bank_code") or ""),
                        "institution_code": service_code,
                        "saved_login_username": "1" if str(bank_credentials.get("login_username") or "").strip() else "0",
                        "saved_login_secret": "1" if str(bank_credentials.get("login_password") or "").strip() else "0",
                        "saved_account_no": "1" if str(bank_credentials.get("account_no") or "").strip() else "0",
                        "saved_account_secret": "1" if str(bank_credentials.get("account_password") or "").strip() else "0",
                        "saved_business_registration_no": "1" if str(bank_credentials.get("business_registration_no") or "").strip() else "0",
                    },
                    "collected_rows": 0,
                    "imported_rows": 0,
                    "duplicate_rows": 0,
                    "matched_count": 0,
                    "unmatched_count": 0,
                    "last_collected_at": str(account.get("last_synced_at") or ""),
                    "date_from": date_from,
                    "date_to": date_to,
                },
                "transactions": [],
            }
    browser_preferred_port_raw = payload.get("browser_preferred_port")
    browser_agent_id_val = str(payload.get("browser_agent_id") or "").strip()
    try:
        browser_preferred_port = int(browser_preferred_port_raw) if browser_preferred_port_raw else None
    except (TypeError, ValueError):
        browser_preferred_port = None

    try:
        timeout_seconds = float(payload.get("browser_timeout_seconds") or 120)
        def _make_browser_coro() -> Any:
            return collect_bank_via_browser_session_async(
                account,
                browser_session_id=browser_session_id,
                browser_work_key=browser_work_key_val,
                date_from=date_from,
                date_to=date_to,
                portal_url=str(payload.get("portal_url") or bank_credentials.get("portal_url") or ""),
                auto_open_browser=bool(payload.get("auto_open_browser")),
                browser_agent_id=browser_agent_id_val,
                browser_preferred_port=browser_preferred_port,
                force_recreate_browser=bool(payload.get("force_recreate_browser")),
                login_username=str(bank_credentials.get("login_username") or ""),
                login_password=str(bank_credentials.get("login_password") or ""),
                account_no=str(bank_credentials.get("account_no") or ""),
                account_password=str(bank_credentials.get("account_password") or ""),
                business_registration_no=str(bank_credentials.get("business_registration_no") or ""),
                business_entity_type=business_entity_type,
                browser_timeout_seconds=timeout_seconds,
            )

        browser_result = _run_bank_browser_async(
            _with_bank_browser_timeout(
                _make_browser_coro,
                timeout_seconds,
            )
        )
    except RuntimeError as exc:
        browser_result = {
            "status": "failed",
            "error_code": "BANK_BROWSER_EVENT_LOOP_ERROR",
            "rows": [],
            "row_count": 0,
            "diagnostics": {"browser_work_key": browser_work_key_val},
            "message": f"은행 브라우저 수집은 이벤트 루프 외부에서만 실행됩니다: {exc!s:.200}",
        }
    except TimeoutError:
        timeout_probe_result = _probe_bank_browser_timeout_state(
            browser_work_key=browser_work_key_val,
            browser_agent_id=browser_agent_id_val,
            timeout_seconds=timeout_seconds,
            portal_url=str(payload.get("portal_url") or bank_credentials.get("portal_url") or ""),
            browser_preferred_port=browser_preferred_port,
            prefer_saved_idpw_login=(
                service_code == "shinhan_business"
                and bool(str(bank_credentials.get("login_username") or "").strip())
                and bool(str(bank_credentials.get("login_password") or "").strip())
            ),
        )
        if (
            timeout_probe_result
            and str(timeout_probe_result.get("error_code") or "") == "BANK_BROWSER_IDPW_RETRY_REQUIRED"
        ):
            retry_timeout_seconds = max(timeout_seconds, 180)
            try:
                browser_result = _run_bank_browser_async(
                    _with_bank_browser_timeout(
                        lambda: collect_bank_via_browser_session_async(
                            account,
                            browser_session_id=browser_session_id,
                            browser_work_key=browser_work_key_val,
                            date_from=date_from,
                            date_to=date_to,
                            portal_url=str(payload.get("portal_url") or bank_credentials.get("portal_url") or ""),
                            auto_open_browser=False,
                            browser_agent_id=browser_agent_id_val,
                            browser_preferred_port=browser_preferred_port,
                            force_recreate_browser=False,
                            login_username=str(bank_credentials.get("login_username") or ""),
                            login_password=str(bank_credentials.get("login_password") or ""),
                            account_no=str(bank_credentials.get("account_no") or ""),
                            account_password=str(bank_credentials.get("account_password") or ""),
                            business_registration_no=str(bank_credentials.get("business_registration_no") or ""),
                            business_entity_type=business_entity_type,
                            browser_timeout_seconds=retry_timeout_seconds,
                        ),
                        retry_timeout_seconds,
                    )
                )
                retry_diag = browser_result.setdefault("diagnostics", {})
                if isinstance(retry_diag, dict):
                    retry_diag["shinhan_idpw_retry_after_timeout"] = "1"
                    retry_diag["previous_timeout_error_code"] = "BANK_BROWSER_IDPW_RETRY_REQUIRED"
                    retry_diag["shinhan_idpw_retry_session_policy"] = "reuse_existing_work_key"
            except TimeoutError:
                browser_result = timeout_probe_result
                diagnostics = browser_result.setdefault("diagnostics", {})
                if isinstance(diagnostics, dict):
                    diagnostics["shinhan_idpw_retry_after_timeout"] = "timeout"
            except Exception as exc:
                browser_result = timeout_probe_result
                diagnostics = browser_result.setdefault("diagnostics", {})
                if isinstance(diagnostics, dict):
                    diagnostics["shinhan_idpw_retry_after_timeout"] = "failed"
                    diagnostics["shinhan_idpw_retry_error_type"] = exc.__class__.__name__[:80]
        else:
            browser_result = timeout_probe_result or {
            "status": "failed",
            "error_code": "BANK_BROWSER_PC_AGENT_TIMEOUT",
            "rows": [],
            "row_count": 0,
            "diagnostics": {
                "browser_work_key": browser_work_key_val,
                "browser_agent_id": browser_agent_id_val,
                "browser_timeout_seconds": str(int(timeout_seconds)),
                "last_observed_stage": "login page",
                "session_recovery_plan": "connect_pc_agent_then_retry_same_work_key",
            },
            "message": f"은행 브라우저 수집이 {int(timeout_seconds)}초 안에 완료되지 않았습니다.",
        }

    br_status = str(browser_result.get("status") or "")
    diagnostics = dict(browser_result.get("diagnostics") or {})

    if br_status != "collected":
        collection_status = (
            br_status
            if br_status in {"action_required", "connector_not_ready", "failed"}
            else "failed"
        )
        connector_status = (
            "ACTION_REQUIRED"
            if collection_status in {"action_required", "connector_not_ready"}
            else "FAILED"
        )
        return {
            "collection": {
                "bank_account_id": account_id,
                "business_id": business_id,
                "branch_id": branch_id,
                "status": collection_status,
                "connector_status": connector_status,
                "connection_type": "browser",
                "message": str(browser_result.get("message") or ""),
                "error_code": str(browser_result.get("error_code") or ""),
                "diagnostics": diagnostics,
                "collected_rows": 0,
                "imported_rows": 0,
                "duplicate_rows": 0,
                "matched_count": 0,
                "unmatched_count": 0,
                "last_collected_at": str(account.get("last_synced_at") or ""),
                "date_from": date_from,
                "date_to": date_to,
            },
            "transactions": [],
        }

    raw_rows = browser_result.get("rows") or []
    entries = [
        {**row, "source": str(row.get("source") or "bank-browser")}
        for row in raw_rows
        if isinstance(row, dict)
    ]
    scoped_entries = [
        entry
        for entry in entries
        if _bank_within_range(entry.get("occurred_at"), date_from, date_to)
    ]
    annotated_entries, matched_count, unmatched_count = _annotate_bank_matches(scoped_entries)
    import_result = record_bank_transactions(
        {
            "business_id": business_id,
            "branch_id": branch_id,
            "bank_account_id": account_id,
            "source": str(diagnostics.get("download_parser") and "bank-browser-download" or "bank-browser"),
            "transactions": annotated_entries,
        },
        user,
    )
    imported_rows = int(import_result.get("import", {}).get("imported_rows") or 0)
    duplicate_rows = int(import_result.get("import", {}).get("duplicate_rows") or 0)
    account_after = _find_bank_account(_read_file_rows(BANK_ACCOUNTS_LEDGER), account_id) or account
    imported_transactions = import_result.get("transactions") or []
    total_in = sum(int(r.get("amount") or 0) for r in imported_transactions if r.get("direction") == "in")
    total_out = sum(int(r.get("amount") or 0) for r in imported_transactions if r.get("direction") == "out")
    final_status = "completed" if imported_rows else "no_records"

    diagnostics["row_count"] = str(len(raw_rows))
    return {
        **import_result,
        "collection": {
            "bank_account_id": account_id,
            "business_id": business_id,
            "branch_id": branch_id,
            "status": final_status,
            "connector_status": "CONFIGURED",
            "connection_type": "browser",
            "message": (
                "은행 브라우저 거래 수집이 완료되었습니다."
                if imported_rows
                else "신규 수집 거래가 없습니다."
            ),
            "collected_rows": len(scoped_entries),
            "imported_rows": imported_rows,
            "duplicate_rows": duplicate_rows,
            "matched_count": matched_count,
            "unmatched_count": unmatched_count,
            "total_in": total_in,
            "total_out": total_out,
            "net_amount": total_in - total_out,
            "last_collected_at": str(account_after.get("last_synced_at") or ""),
            "date_from": date_from,
            "date_to": date_to,
            "browser_work_key": browser_work_key_val,
            "diagnostics": diagnostics,
        },
    }


def match_bank_to_settlements(
    bank_transactions: list[dict[str, Any]],
    settlements: list[dict[str, Any]],
    *,
    tolerance: int = 0,
) -> dict[str, Any]:
    """Match bank 'in' transactions to delivery settlement records by date + amount."""
    matched: list[dict[str, Any]] = []
    unmatched_bank: list[dict[str, Any]] = []
    used_settlement_ids: set[str] = set()
    for txn in bank_transactions:
        txn_date = _bank_date_key(txn.get("occurred_at"))
        txn_amount = int(txn.get("amount") or 0)
        best: dict[str, Any] | None = None
        for settlement in settlements:
            sid = str(settlement.get("id") or "")
            if sid in used_settlement_ids:
                continue
            s_date = _bank_date_key(
                settlement.get("occurred_on") or settlement.get("settled_at") or ""
            )
            s_amount = int(
                settlement.get("settlement_amount") or settlement.get("amount") or 0
            )
            if txn_date == s_date and abs(txn_amount - s_amount) <= tolerance:
                best = settlement
                break
        if best:
            used_settlement_ids.add(str(best.get("id") or ""))
            matched.append(
                {
                    "bank_transaction_id": str(txn.get("id") or ""),
                    "settlement_id": str(best.get("id") or ""),
                    "amount": txn_amount,
                    "matched_on": "date+amount",
                }
            )
        else:
            unmatched_bank.append(txn)
    return {
        "matched": matched,
        "unmatched_bank_transactions": unmatched_bank,
        "unmatched_settlement_count": len(settlements) - len(matched),
        "match_count": len(matched),
    }


def list_bank_transactions(
    user: dict[str, Any],
    *,
    business_id: str | None = None,
    branch_id: str | None = None,
    bank_account_id: str | None = None,
    direction: str | None = None,
    date_from: str | None = None,
    date_to: str | None = None,
) -> list[dict[str, Any]]:
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="은행거래 원장 조회 권한이 없습니다")
    start, end = _valid_range_bounds(date_from, date_to)
    wanted_direction = _bank_direction(direction) if direction else ""
    rows = _read_file_rows(BANK_TRANSACTIONS_LEDGER)
    result: list[dict[str, Any]] = []
    for row in rows:
        if business_id and str(row.get("business_id") or "") != business_id:
            continue
        if branch_id and str(row.get("branch_id") or "") not in {"", branch_id}:
            continue
        if bank_account_id and str(row.get("bank_account_id") or "") != bank_account_id:
            continue
        if wanted_direction and str(row.get("direction") or "") != wanted_direction:
            continue
        if not _bank_within_range(row.get("occurred_at"), start, end):
            continue
        result.append(row)
    result.sort(key=lambda item: str(item.get("occurred_at") or ""), reverse=True)
    return result


def bank_summary(
    user: dict[str, Any],
    *,
    business_id: str | None = None,
    branch_id: str | None = None,
    bank_account_id: str | None = None,
    date_from: str | None = None,
    date_to: str | None = None,
) -> dict[str, Any]:
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="은행 요약 조회 권한이 없습니다")
    start, end = _valid_range_bounds(date_from, date_to)
    accounts = _read_file_rows(BANK_ACCOUNTS_LEDGER)
    transactions = list_bank_transactions(
        user,
        business_id=business_id,
        branch_id=branch_id,
        bank_account_id=bank_account_id,
        date_from=start or None,
        date_to=end or None,
    )

    total_in = sum(int(row.get("amount") or 0) for row in transactions if row.get("direction") == "in")
    total_out = sum(int(row.get("amount") or 0) for row in transactions if row.get("direction") == "out")

    per_account: dict[str, dict[str, Any]] = {}
    for account in accounts:
        acct_id = str(account.get("id") or "")
        if business_id and str(account.get("business_id") or "") != business_id:
            continue
        if branch_id and str(account.get("branch_id") or "") not in {"", branch_id}:
            continue
        if bank_account_id and acct_id != bank_account_id:
            continue
        per_account[acct_id] = {
            "bank_account_id": acct_id,
            "business_id": str(account.get("business_id") or ""),
            "branch_id": str(account.get("branch_id") or ""),
            "account_alias": str(account.get("account_alias") or ""),
            "bank_name": str(account.get("bank_name") or ""),
            "account_number_masked": str(account.get("account_number_masked") or ""),
            "connection_type": str(account.get("connection_type") or ""),
            "status": str(account.get("status") or ""),
            "last_synced_at": str(account.get("last_synced_at") or ""),
            "transaction_count": 0,
            "total_in": 0,
            "total_out": 0,
            "net": 0,
        }

    for row in transactions:
        acct_id = str(row.get("bank_account_id") or "")
        bucket = per_account.get(acct_id)
        if bucket is None:
            # Ledger rows whose account was removed still count toward totals.
            bucket = per_account.setdefault(
                acct_id,
                {
                    "bank_account_id": acct_id,
                    "business_id": str(row.get("business_id") or ""),
                    "branch_id": str(row.get("branch_id") or ""),
                    "account_alias": "",
                    "bank_name": "",
                    "account_number_masked": "",
                    "connection_type": "",
                    "status": "unknown",
                    "last_synced_at": "",
                    "transaction_count": 0,
                    "total_in": 0,
                    "total_out": 0,
                    "net": 0,
                },
            )
        amount = int(row.get("amount") or 0)
        bucket["transaction_count"] += 1
        if row.get("direction") == "in":
            bucket["total_in"] += amount
        else:
            bucket["total_out"] += amount
        bucket["net"] = bucket["total_in"] - bucket["total_out"]

    status_counts: dict[str, int] = {}
    for bucket in per_account.values():
        status_counts[bucket["status"]] = status_counts.get(bucket["status"], 0) + 1

    return {
        "business_id": business_id or "",
        "branch_id": branch_id or "",
        "date_from": start,
        "date_to": end,
        "totals": {
            "total_in": total_in,
            "total_out": total_out,
            "net": total_in - total_out,
            "transaction_count": len(transactions),
            "account_count": len(per_account),
        },
        "account_status_counts": status_counts,
        "accounts": sorted(
            per_account.values(),
            key=lambda item: (str(item.get("account_alias") or ""), str(item.get("bank_account_id") or "")),
        ),
    }


def _delivery_browser_auth_options(payload: dict[str, Any]) -> dict[str, str]:
    storage_state_path = str(payload.get("storage_state_path") or "").strip()
    browser_session_id = str(payload.get("browser_session_id") or "").strip()
    explicit_browser_session_id = browser_session_id
    bridge_mode = ""
    if not storage_state_path:
        try:
            from app.browser_bridge.e2e_adapter import build_e2e_config

            config = build_e2e_config(session_id=browser_session_id or None)
            bridge_mode = str(config.get("mode") or "")
            browser_session_id = browser_session_id or str(config.get("session_id") or "").strip()
            storage_state_path = str(config.get("storage_state_path") or "").strip()
        except Exception:
            storage_state_path = ""
    return {
        "storage_state_path": storage_state_path if storage_state_path and Path(storage_state_path).is_file() else "",
        "browser_session_id": browser_session_id,
        "browser_bridge_mode": bridge_mode,
        "browser_session_id_explicit": "1" if explicit_browser_session_id else "",
    }


def _delivery_browser_auth_for_account(
    payload: dict[str, Any],
    account: dict[str, Any],
    service: str,
    business_id: str,
    branch: str,
) -> dict[str, str]:
    auth = dict(_delivery_browser_auth_options(payload))
    legacy_explicit_session = "browser_session_id_explicit" not in auth and bool(auth.get("browser_session_id"))
    if auth.get("storage_state_path") or auth.get("browser_session_id_explicit") or legacy_explicit_session:
        return auth
    collection_mode = str(account.get("collection_mode") or account.get("collectionMode") or "").strip()
    prefer_pc_agent = bool(
        payload.get("prefer_pc_agent")
        or payload.get("preferPcAgent")
        or payload.get("force_pc_agent")
        or payload.get("forcePcAgent")
        or payload.get("require_pc_agent")
        or payload.get("requirePcAgent")
    )
    force_recreate_session = bool(
        payload.get("force_recreate_portal_sessions")
        or payload.get("forceRecreatePortalSessions")
        or payload.get("force_recreate_sessions")
        or payload.get("refresh_portal_sessions")
        or payload.get("refreshPortalSessions")
    )
    if (
        not _has_secret_value(account, "password")
        and not prefer_pc_agent
    ):
        auth["browser_session_id"] = ""
        auth["browser_bridge_mode"] = ""
        auth["browser_auth_strategy"] = "missing_password_no_pc_agent"
        return auth
    if (
        _has_secret_value(account, "password")
        and not prefer_pc_agent
        and service != "ddangyo"
        and (
            collection_mode != "browser-automation"
            or payload.get("allow_server_headless_fallback")
            or payload.get("allowServerHeadlessFallback")
        )
    ):
        auth["browser_session_id"] = ""
        auth["browser_bridge_mode"] = ""
        auth["browser_auth_strategy"] = "server_headless_password_first"
        return auth
    ambient_browser_session_id = str(auth.get("browser_session_id") or "").strip()
    if ambient_browser_session_id:
        auth["ambient_browser_session_id"] = ambient_browser_session_id
        auth["browser_session_id"] = ""
        auth["ambient_browser_bridge_mode"] = str(auth.get("browser_bridge_mode") or "")
        auth["browser_bridge_mode"] = ""
    if payload.get("disable_pc_agent") or payload.get("disablePcAgent"):
        return auth

    try:
        from app.browser_bridge.service import get_browser_bridge_service
        from app.services.yeoljeong_delivery_collectors import PORTAL_CONFIG

        config = PORTAL_CONFIG.get(service) or {}
        label = _delivery_platform_label(service)
        url = str(
            account.get("portal_home_url")
            or account.get("home_url")
            or account.get("login_url")
            or config.get("login_url")
            or "about:blank"
        )
        work_key = _delivery_browser_work_key(service, business_id, branch)
        auth["browser_service"] = service
        auth["browser_target_url"] = url
        pc_agent_id = str(
            payload.get("pc_agent_id")
            or payload.get("pcAgentId")
            or account.get("pc_agent_id")
            or os.getenv("YEOLJEONG_DELIVERY_PC_AGENT_ID", "")
            or ""
        ).strip()
        bridge_service = get_browser_bridge_service()
        session = None
        errors: list[str] = []
        close_flag = payload.get("close_portal_browser_on_complete")
        if close_flag is None:
            close_flag = payload.get("closePortalBrowserOnComplete")
        close_on_complete = True if close_flag is None else bool(close_flag)
        for attempt in range(3):
            try:
                ensure_kwargs: dict[str, Any] = {
                    "work_key": work_key,
                    "label": f"열정국밥 {branch} {label} 자동수집",
                    "url": url,
                    "force_recreate": force_recreate_session or attempt > 0,
                }
                if pc_agent_id:
                    ensure_kwargs["agent_id"] = pc_agent_id
                session = _run_delivery_browser_async(
                    bridge_service.ensure_work_session(**ensure_kwargs)
                )
                if force_recreate_session:
                    auth["browser_bridge_recovered"] = "force_recreate_requested"
                elif attempt > 0:
                    auth["browser_bridge_recovered"] = f"force_recreate_attempt_{attempt + 1}"
                break
            except Exception as exc:
                errors.append(str(exc)[:300])
                auth["browser_bridge_error"] = errors[-1]
                auth["browser_bridge_errors"] = " | ".join(errors)[-900:]
                if attempt < 2:
                    time.sleep(2 + attempt)
                    continue
                raise
        if session is not None:
            auth["browser_session_id"] = str(getattr(session, "session_id", "") or "")
            auth["browser_bridge_mode"] = "local_agent"
            auth["browser_work_key"] = work_key
            auth["browser_close_on_complete"] = "1" if close_on_complete else ""
            if pc_agent_id:
                auth["browser_agent_id"] = pc_agent_id
            if force_recreate_session:
                auth["browser_session_recreated"] = "1"
            if service == "baemin":
                auth["browser_session_policy"] = "shared_reuse"
            _append_delivery_browser_session_event(
                "session_ensured",
                {
                    "service": service,
                    "business_id": business_id,
                    "branch": branch,
                    "session_id": auth["browser_session_id"],
                    "work_key": work_key,
                    "agent_id": pc_agent_id,
                    "close_on_complete": auth["browser_close_on_complete"],
                    "recovered": auth.get("browser_bridge_recovered") or "",
                    "recreated": auth.get("browser_session_recreated") or "",
                    "session_policy": auth.get("browser_session_policy") or "",
                },
            )
    except Exception as exc:
        auth["browser_bridge_error"] = str(exc)[:300]
    return auth


def _delivery_browser_work_key(service: str, business_id: str, branch: str) -> str:
    normalized_service = re.sub(r"[^a-z0-9._:-]+", "-", str(service or "").strip().lower()).strip("-")
    if normalized_service == "baemin":
        shared_key = os.getenv("YEOLJEONG_BAEMIN_BROWSER_WORK_KEY", "yeoljeong-delivery-baemin-shared")
        return re.sub(r"[^a-z0-9._:-]+", "-", shared_key.strip().lower()).strip("-") or "yeoljeong-delivery-baemin-shared"
    normalized_business = re.sub(r"[^a-z0-9._:-]+", "-", str(business_id or "").strip().lower()).strip("-")
    branch_hash = hashlib.sha256(str(branch or "").encode("utf-8")).hexdigest()[:10]
    return f"yeoljeong-delivery-{normalized_service or 'portal'}-{normalized_business or 'business'}-{branch_hash}"


def _append_delivery_browser_session_event(event: str, payload: dict[str, Any]) -> None:
    try:
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        record = {
            "event": event,
            "recorded_at": _now(),
            **{
                key: value
                for key, value in payload.items()
                if value not in (None, "")
            },
        }
        with (DATA_DIR / "delivery_browser_session_events.jsonl").open("a", encoding="utf-8") as handle:
            handle.write(json.dumps(record, ensure_ascii=False, sort_keys=True) + "\n")
    except Exception:
        return


def _append_delivery_stage_log(
    stage_logs: list[dict[str, str]],
    *,
    service: str,
    stage: str,
    status: str,
    started_at: float,
    error_code: str = "",
    reason: str = "",
    **fields: Any,
) -> None:
    service_key = re.sub(r"[^a-z0-9_:-]+", "_", str(service or "portal").strip().lower()).strip("_") or "portal"
    append_site_stage_log(
        stage_logs,
        stage=f"{service_key}_{stage}",
        status=status,
        started_at=started_at,
        event_name="delivery_browser_collection_stage",
        error_code=error_code,
        reason=reason,
        **fields,
    )


def _delivery_browser_diagnostics(
    browser_auth: dict[str, Any],
    session_id: str,
    *,
    url: str = "",
    stage_logs: list[dict[str, str]] | None = None,
    extra: dict[str, Any] | None = None,
) -> dict[str, Any]:
    diagnostics: dict[str, Any] = {
        "auth_mode": "pc_agent_browser",
        "browser_session_id": session_id,
        "browser_work_key": str(browser_auth.get("browser_work_key") or ""),
        "browser_bridge_mode": str(browser_auth.get("browser_bridge_mode") or ""),
    }
    if url:
        diagnostics["url"] = url
    if stage_logs is not None:
        diagnostics["browser_stage_logs"] = stage_logs
        diagnostics["browser_stage_log_schema"] = SITE_STAGE_LOG_SCHEMA
    if extra:
        diagnostics.update(extra)
    return diagnostics


def _run_delivery_browser_async(coro: Any) -> Any:
    try:
        asyncio.get_running_loop()
    except RuntimeError:
        return asyncio.run(coro)
    close = getattr(coro, "close", None)
    if callable(close):
        close()
    raise RuntimeError("delivery browser automation cannot run inside an active event loop")


async def _close_delivery_browser_work_session_async(
    browser_auth: dict[str, Any],
    *,
    reason: str,
) -> None:
    if str(browser_auth.get("browser_close_on_complete") or "") != "1":
        return
    session_id = str(browser_auth.get("browser_session_id") or "").strip()
    work_key = str(browser_auth.get("browser_work_key") or "").strip()
    if not session_id and not work_key:
        return
    try:
        from app.browser_bridge.service import get_browser_bridge_service
        from app.services.pc_agent_manager import pc_agent_manager

        bridge = get_browser_bridge_service()
        session = bridge.sessions.get(session_id) if session_id else None
        metadata = dict(getattr(getattr(session, "endpoint", None), "metadata", None) or {})
        agent_id = str(browser_auth.get("browser_agent_id") or metadata.get("agent_id") or "").strip()
        close_work_key = work_key or str(metadata.get("work_key") or "").strip()
        browser_service = str(browser_auth.get("browser_service") or "").strip()
        if agent_id and close_work_key:
            close_params = {
                "work_key": close_work_key,
                "close_browser": True,
                "close_tabs": True,
                "reason": reason,
                "command_timeout_seconds": 10,
            }
            cleanup_job_type = f"browser_bridge_cleanup_{session_id or close_work_key}"
            route_first = bool(
                getattr(bridge, "_route_pc_agent_via_active_api_first", lambda: False)()
            )
            close_result: dict[str, Any] | None = None
            if route_first:
                close_result = await bridge._execute_pc_agent_route_via_active_api(
                    command_type="browser_close_session",
                    params=close_params,
                    agent_id=agent_id,
                    job_type=cleanup_job_type,
                    required_capabilities=["interactive_browser"],
                    queue_wait_timeout_seconds=10,
                    lease_ttl_seconds=30,
                    command_timeout_seconds=10,
                )
            if not route_first or close_result is None:
                close_result = await pc_agent_manager.execute_routed_command(
                    command_type="browser_close_session",
                    params=close_params,
                    agent_id=agent_id,
                    job_type=cleanup_job_type,
                    required_capabilities=["interactive_browser"],
                    queue_if_busy=True,
                    wait_for_turn=True,
                    queue_wait_timeout_seconds=10,
                    lease_ttl_seconds=30,
                    command_timeout_seconds=10,
                )
                if (
                    isinstance(close_result, dict)
                    and close_result.get("status") != "success"
                    and str(close_result.get("error_code") or "") in {"PC_AGENT_OFFLINE", "NO_CAPABLE_AGENT"}
                ):
                    await bridge._execute_pc_agent_route_via_active_api(
                        command_type="browser_close_session",
                        params=close_params,
                        agent_id=agent_id,
                        job_type=cleanup_job_type,
                        required_capabilities=["interactive_browser"],
                        queue_wait_timeout_seconds=10,
                        lease_ttl_seconds=30,
                        command_timeout_seconds=10,
                    )
            close_payload: dict[str, Any] = {}
            if isinstance(close_result, dict):
                raw_command_result = close_result.get("result") if isinstance(close_result.get("result"), dict) else {}
                raw_payload = raw_command_result.get("result") if isinstance(raw_command_result, dict) else None
                if isinstance(raw_payload, dict):
                    close_payload = raw_payload
            close_process = close_payload.get("process") if isinstance(close_payload, dict) else {}
            close_process = close_process if isinstance(close_process, dict) else {}
            if (
                browser_service == "baemin"
                and close_payload.get("session_released") is False
                and str(close_process.get("reason") or "") == "session_not_found"
            ):
                orphan_params = {
                    "url_pattern": "self.baemin.com",
                    "keep_last": False,
                    "command_timeout_seconds": 15,
                }
                orphan_job_type = f"browser_bridge_orphan_cleanup_{session_id or close_work_key}"
                orphan_result: dict[str, Any] | None = None
                if route_first:
                    orphan_result = await bridge._execute_pc_agent_route_via_active_api(
                        command_type="browser_close_tab",
                        params=orphan_params,
                        agent_id=agent_id,
                        job_type=orphan_job_type,
                        required_capabilities=["interactive_browser"],
                        queue_wait_timeout_seconds=10,
                        lease_ttl_seconds=30,
                        command_timeout_seconds=15,
                    )
                if not route_first or orphan_result is None:
                    orphan_result = await pc_agent_manager.execute_routed_command(
                        command_type="browser_close_tab",
                        params=orphan_params,
                        agent_id=agent_id,
                        job_type=orphan_job_type,
                        required_capabilities=["interactive_browser"],
                        queue_if_busy=True,
                        wait_for_turn=True,
                        queue_wait_timeout_seconds=10,
                        lease_ttl_seconds=30,
                        command_timeout_seconds=15,
                    )
                _append_delivery_browser_session_event(
                    "orphan_tabs_close_requested",
                    {
                        "reason": reason,
                        "agent_id": agent_id,
                        "session_id": session_id,
                        "work_key": close_work_key,
                        "status": orphan_result.get("status") if isinstance(orphan_result, dict) else "",
                        "error_code": orphan_result.get("error_code") if isinstance(orphan_result, dict) else "",
                        "message": orphan_result.get("message") if isinstance(orphan_result, dict) else "",
                    },
                )
            _append_delivery_browser_session_event(
                "close_requested",
                {
                    "reason": reason,
                    "agent_id": agent_id,
                    "session_id": session_id,
                    "work_key": close_work_key,
                    "status": close_result.get("status") if isinstance(close_result, dict) else "",
                    "error_code": close_result.get("error_code") if isinstance(close_result, dict) else "",
                    "message": close_result.get("message") if isinstance(close_result, dict) else "",
                },
            )
        if session_id:
            bridge.sessions.retire_session(
                session_id,
                stale_reason=reason,
                clear_work_key=True,
                clear_active=False,
                clear_lease=True,
            )
    except Exception as exc:
        _append_delivery_browser_session_event(
            "close_failed",
            {
                "reason": reason,
                "session_id": session_id,
                "work_key": work_key,
                "error": str(exc)[:300],
            },
        )
        return


_DELIVERY_SERVICE_URL_MARKERS = {
    "baemin": ("baemin.com",),
    "coupangeats": ("coupangeats.com",),
    "yogiyo": ("yogiyo.co.kr",),
    "ddangyo": ("ddangyo.com",),
}


_DELIVERY_LOGIN_SELECTORS = {
    "baemin": {
        "username": (
            "input[autocomplete='username']",
            "input[name='id']",
            "input[name*='id' i]",
            "input[name*='user' i]",
            "input[type='email']",
            "input[type='text']",
        ),
        "password": (
            "input[autocomplete='current-password']",
            "input[name='password']",
            "input[name*='pw' i]",
            "input[type='password']",
        ),
        "submit": (
            "button[type='submit']",
            "input[type='submit']",
            "button:has-text('로그인')",
            "text=로그인",
        ),
    },
    "coupangeats": {
        "username": (
            "#loginId",
            "input[autocomplete='username']",
            "input[name*='email' i]",
            "input[name*='id' i]",
            "input[name*='user' i]",
            "input[placeholder*='아이디']",
            "input[placeholder*='이메일']",
            "input[type='email']",
            "input[type='text']",
        ),
        "password": (
            "#password",
            "input[autocomplete='current-password']",
            "input[name*='password' i]",
            "input[name*='pw' i]",
            "input[placeholder*='비밀번호']",
            "input[type='password']",
        ),
        "submit": (
            "button[type='submit']",
            "input[type='submit']",
            "button:has-text('로그인')",
            "text=로그인",
        ),
    },
    "yogiyo": {
        "username": (
            "input[autocomplete='username']",
            "input[name*='id' i]",
            "input[name*='email' i]",
            "input[name*='user' i]",
            "input[placeholder*='아이디']",
            "input[placeholder*='이메일']",
            "input[type='email']",
            "input[type='text']",
        ),
        "password": (
            "input[autocomplete='current-password']",
            "input[name*='password' i]",
            "input[name*='pw' i]",
            "input[placeholder*='비밀번호']",
            "input[type='password']",
        ),
        "submit": (
            "button[type='submit']",
            "input[type='submit']",
            "button:has-text('로그인')",
            "text=로그인",
        ),
    },
    "ddangyo": {
        "username": (
            "#mf_wfm_login_id",
            "#mf_ipt_usrId",
            "#userId",
            "input[name*='user' i]",
            "input[name*='id' i]",
            "input[placeholder*='아이디']",
            "input[type='text']",
        ),
        "password": (
            "#mf_wfm_login_pw",
            "#mf_ipt_pw",
            "#password",
            "input[name*='password' i]",
            "input[name*='pw' i]",
            "input[placeholder*='비밀번호']",
            "input[type='password']",
        ),
        "submit": (
            "#mf_btn_webLogin",
            "input[type='button'][value*='로그인']",
            "button[type='submit']",
            "input[type='submit']",
            "button:has-text('로그인')",
            "text=로그인",
        ),
    },
}

_DELIVERY_CAPTCHA_SELECTORS = {
    "ddangyo": (
        "#mf_wfm_login_captcha",
        "#mf_ipt_captcha",
        "#captcha",
        "input[name*='captcha' i]",
        "input[id*='captcha' i]",
        "input[placeholder*='보안문자']",
        "input[placeholder*='자동입력방지']",
        "input[placeholder*='숫자']",
        "input[type='text']",
    ),
}


def _delivery_login_selectors(service: str) -> dict[str, tuple[str, ...]]:
    fallback = {
        "username": (
            "input[autocomplete='username']",
            "input[name*='id' i]",
            "input[name*='user' i]",
            "input[type='email']",
            "input[type='text']",
        ),
        "password": ("input[autocomplete='current-password']", "input[type='password']"),
        "submit": (
            "button[type='submit']",
            "input[type='submit']",
            "input[type='button'][value*='로그인']",
            "button:has-text('로그인')",
            "text=로그인",
        ),
    }
    configured = _DELIVERY_LOGIN_SELECTORS.get(str(service or "").strip().lower(), {})
    return {
        "username": tuple(configured.get("username") or fallback["username"]),
        "password": tuple(configured.get("password") or fallback["password"]),
        "submit": tuple(configured.get("submit") or fallback["submit"]),
    }


def _delivery_result_is_wrong_portal(service: str, result: dict[str, Any]) -> bool:
    diagnostics = result.get("diagnostics") if isinstance(result.get("diagnostics"), dict) else {}
    url = str(diagnostics.get("url") or "").lower()
    if not url:
        return False
    markers = _DELIVERY_SERVICE_URL_MARKERS.get(service, ())
    return bool(markers and not any(marker in url for marker in markers))


def _delivery_result_has_no_visible_source(result: dict[str, Any]) -> bool:
    error_code = str(result.get("error_code") or "").upper()
    if error_code not in {"AUTHENTICATED_NO_ROWS", "EMPTY_SOURCE", "NO_PARSEABLE_ROWS"}:
        return False
    diagnostics = result.get("diagnostics") if isinstance(result.get("diagnostics"), dict) else {}
    has_section_diagnostics = any(kind in diagnostics for kind in DELIVERY_RECORD_TYPES)
    section_values = [str(diagnostics.get(kind) or "").upper() for kind in DELIVERY_RECORD_TYPES]
    if has_section_diagnostics and all(value in {"SECTION_NOT_FOUND", "NO_EXPORT_OR_TABLE", "NO_PARSEABLE_ROWS", ""} for value in section_values):
        return True
    return error_code in {"EMPTY_SOURCE", "NO_PARSEABLE_ROWS"} and str(diagnostics.get("auth_mode") or "") == "pc_agent_browser"


async def _delivery_bridge_page_for_service(context: Any, service: str) -> Any:
    pages = list(getattr(context, "pages", None) or [])
    markers = _DELIVERY_SERVICE_URL_MARKERS.get(str(service or "").strip(), ())
    for page in pages:
        try:
            url = str(await page.evaluate("window.location.href") or getattr(page, "url", "") or "").lower()
        except Exception:
            url = str(getattr(page, "url", "") or "").lower()
        if markers and any(marker in url for marker in markers):
            return page
    if markers:
        return await context.new_page()
    return pages[0] if pages else await context.new_page()


def _baemin_bridge_page_kind(text: str) -> str:
    lowered = text.lower()
    if "리뷰" in text or "review" in lowered:
        return "reviews"
    if "정산" in text or "입금" in text or "settlement" in lowered:
        return "settlements"
    return "sales"


def _delivery_bridge_page_kind(text: str) -> str:
    lowered = str(text or "").lower()
    if "리뷰" in text or "review" in lowered:
        return "reviews"
    if any(term in text for term in ("정산", "입금", "지급")) or any(
        term in lowered for term in ("settlement", "deposit", "payout")
    ):
        return "settlements"
    return "sales"


def _money_from_text(value: str) -> int:
    digits = re.sub(r"[^0-9]", "", str(value or ""))
    return int(digits or 0)


def _baemin_dashboard_records(text: str, business_id: str, branch: str) -> dict[str, list[dict[str, Any]]]:
    """Capture summary data visible on the authenticated Baemin home dashboard."""
    compact = re.sub(r"\r\n?", "\n", str(text or ""))
    compact = "\n".join(line.strip() for line in compact.splitlines() if line.strip())
    today = datetime.now(KST).date()
    yesterday = today - timedelta(days=1)
    collected_at = _now()
    records = _delivery_empty_record_lists()

    sales_match = re.search(r"어제\s*주문금액\s*([0-9,]+)\s*원.*?어제\s*주문수\s*([0-9,]+)\s*건", compact, re.S)
    if sales_match:
        amount = _money_from_text(sales_match.group(1))
        order_count = _money_from_text(sales_match.group(2))
        source_id = f"baemin-dashboard-sales-{yesterday.isoformat()}"
        records["sales"].append(
            {
                "id": hashlib.sha256(f"{business_id}|{branch}|{source_id}".encode("utf-8")).hexdigest(),
                "source_id": source_id,
                "business_id": business_id,
                "branch": branch,
                "service": "baemin",
                "platform": "baemin",
                "record_type": "sales",
                "occurred_on": yesterday.isoformat(),
                "gross_amount": amount,
                "order_count": order_count,
                "order_status": "dashboard_summary",
                "collected_at": collected_at,
            }
        )

    settlement_match = re.search(r"입금\s*예정\s*금액\s*([0-9,]+)\s*원", compact)
    if settlement_match:
        amount = _money_from_text(settlement_match.group(1))
        source_id = f"baemin-dashboard-settlement-{today.isoformat()}"
        records["settlements"].append(
            {
                "id": hashlib.sha256(f"{business_id}|{branch}|{source_id}".encode("utf-8")).hexdigest(),
                "source_id": source_id,
                "settlement_id": source_id,
                "business_id": business_id,
                "branch": branch,
                "service": "baemin",
                "platform": "baemin",
                "record_type": "settlements",
                "occurred_on": today.isoformat(),
                "settlement_amount": amount,
                "settlement_status": "입금예정",
                "collected_at": collected_at,
            }
        )

    review_matches = re.finditer(r"(오늘|어제)\n(.{8,800}?)\n열정국밥\s+중랑구중화점", compact, re.S)
    for index, match in enumerate(review_matches, start=1):
        review_text = re.sub(r"\s+", " ", match.group(2)).strip()
        if not review_text or review_text == "열정국밥 중랑구중화점":
            continue
        occurred_on = today if match.group(1) == "오늘" else yesterday
        source_material = f"baemin-dashboard-review-{occurred_on.isoformat()}-{index}-{review_text[:80]}"
        source_id = hashlib.sha256(source_material.encode("utf-8")).hexdigest()[:32]
        records["reviews"].append(
            {
                "id": hashlib.sha256(f"{business_id}|{branch}|baemin|reviews|{source_id}".encode("utf-8")).hexdigest(),
                "source_id": source_id,
                "review_id": source_id,
                "business_id": business_id,
                "branch": branch,
                "service": "baemin",
                "platform": "baemin",
                "record_type": "reviews",
                "occurred_on": occurred_on.isoformat(),
                "rating": 0,
                "review_text": review_text[:4000],
                "reply_status": "",
                "collected_at": collected_at,
            }
        )
    return records


def _baemin_bridge_login_state(url: str, text: str) -> str:
    lowered_url = str(url or "").lower()
    lowered_text = str(text or "").lower()
    if any(
        term in lowered_text
        for term in (
            "보안 위배 접근 제한",
            "올바르지 않은 요청",
            "잠시 이용이 제한",
            "비정상 동작",
            "비정상적인 동작",
            "잠시 후 다시 시도",
            "access denied",
            "forbidden",
        )
    ):
        return "blocked"
    if any(term in lowered_text for term in ("captcha", "캡차", "보안문자", "2차 인증", "추가 인증", "본인인증", "휴대폰 인증", "기기 인증", "인증번호")):
        return "challenge"
    if "login" in lowered_url or all(marker in text for marker in ("로그인", "회원가입")):
        return "login"
    return "authenticated"


def _delivery_bridge_login_state(url: str, text: str) -> str:
    decision = classify_portal_state(url, text)
    if decision.state == "portal_error":
        return "blocked"
    if decision.state in {"captcha_required", "otp_required"}:
        return "challenge"
    if decision.state == "login_required":
        return "login"
    return "authenticated" if decision.state == "collectable_page" else "challenge"


def _delivery_bridge_challenge_code(service: str, text: str) -> str:
    lowered_text = str(text or "").lower()
    if service == "ddangyo" and any(
        term in lowered_text for term in ("captcha", "캡차", "보안문자", "자동입력방지", "숫자를 입력")
    ):
        return "DDANGYO_NUMERIC_CAPTCHA_REQUIRED"
    return "PORTAL_AUTH_CHALLENGE"


def _delivery_captcha_value_for_account(
    payload: dict[str, Any],
    account: dict[str, Any],
    service: str,
    business_id: str,
    branch: str,
) -> str:
    if service != "ddangyo" or payload.get("operator_approved") is not True:
        return ""
    values = payload.get("captcha_values") if isinstance(payload.get("captcha_values"), dict) else {}
    run_key = _delivery_run_key(service, business_id, branch)
    candidates = (
        approved_operator_input(payload),
        payload.get("captcha_value"),
        payload.get("captcha"),
        values.get(run_key),
        values.get(str(account.get("id") or "")),
        values.get(service),
    )
    for candidate in candidates:
        digits = re.sub(r"[^0-9]", "", str(candidate or ""))
        if 3 <= len(digits) <= 8:
            return digits
    return ""


def _url_origin(value: str) -> str:
    parsed = urlparse(str(value or ""))
    if parsed.scheme and parsed.netloc:
        return f"{parsed.scheme.lower()}://{parsed.netloc.lower()}"
    return str(value or "").strip().lower()


def _delivery_captcha_automation_approval(
    payload: dict[str, Any],
    user: dict[str, Any],
    service: str,
    business_id: str,
    branch: str,
) -> dict[str, Any]:
    if service != "ddangyo":
        return {}
    approved = payload.get("captcha_auto_approved") is True
    manual_value = _delivery_captcha_value_for_account(payload, {}, service, business_id, branch)
    if not approved and payload.get("operator_approved") is True and not manual_value:
        approved = True
    if not approved:
        return {}
    scope = payload.get("captcha_approval_scope") if isinstance(payload.get("captcha_approval_scope"), dict) else {}
    allowed_origins = scope.get("origins") if isinstance(scope.get("origins"), list) else []
    return {
        "approved": True,
        "approved_by": str(scope.get("approved_by") or user.get("email") or user.get("id") or "unknown"),
        "approved_at": _now(),
        "approval_source": str(scope.get("approval_source") or "operator_request"),
        "service": service,
        "business_id": business_id,
        "branch": branch,
        "origin": _url_origin(str(scope.get("origin") or scope.get("page_url") or "https://boss.ddangyo.com/")),
        "origins": [_url_origin(str(item)) for item in allowed_origins if str(item).strip()],
        "challenge_kind": "captcha",
        "automation": "llm_vision_read_and_fill",
        "max_executions": max(1, min(int(scope.get("max_executions") or 1), 10)),
    }


def _delivery_captcha_automation_allowed(approval: dict[str, Any], page_url: str) -> tuple[bool, str]:
    if not approval or approval.get("approved") is not True:
        return False, "captcha_auto_approval_missing"
    actual_origin = _url_origin(page_url)
    approved_origin = _url_origin(str(approval.get("origin") or ""))
    allowed_origins = {
        _url_origin(str(item))
        for item in approval.get("origins", [])
        if str(item).strip()
    }
    if approved_origin and actual_origin != approved_origin:
        return False, "captcha_auto_origin_out_of_scope"
    if allowed_origins and actual_origin not in allowed_origins:
        return False, "captcha_auto_origin_out_of_scope"
    if str(approval.get("challenge_kind") or "") != "captcha":
        return False, "captcha_auto_challenge_kind_out_of_scope"
    return True, "captcha_auto_approval_scope_match"


def _record_delivery_captcha_automation_event(
    event: str,
    approval: dict[str, Any],
    *,
    url: str,
    session_id: str,
    work_key: str,
    run_id: str = "",
    reason: str = "",
) -> None:
    _append_delivery_browser_session_event(
        event,
        {
            "service": approval.get("service") or "ddangyo",
            "business_id": approval.get("business_id") or "",
            "branch": approval.get("branch") or "",
            "run_id": run_id,
            "session_id": session_id,
            "work_key": work_key,
            "approved_by": approval.get("approved_by") or "",
            "approved_at": approval.get("approved_at") or "",
            "approval_source": approval.get("approval_source") or "",
            "approved_origin": approval.get("origin") or "",
            "actual_origin": _url_origin(url),
            "challenge_kind": "captcha",
            "automation": approval.get("automation") or "llm_vision_read_and_fill",
            "reason": reason,
        },
    )


def _delivery_challenge_message(service: str, service_label: str) -> str:
    if service == "ddangyo":
        return (
            "땡겨요 ID/PW 자동입력은 완료됐고 숫자 캡챠 승인이 필요합니다. "
            "승인 범위가 있으면 같은 세션에서 모델 판독과 입력을 자동 수행하고, 승인 범위가 없으면 대기합니다."
        )
    return f"{service_label} 포털이 추가 인증을 요구합니다. PC 브라우저에서 인증을 완료한 뒤 다시 수집해야 합니다."


async def _delivery_bridge_fill_ddangyo_numeric_captcha(page: Any, captcha_value: str) -> bool:
    digits = re.sub(r"[^0-9]", "", str(captcha_value or ""))
    if not (3 <= len(digits) <= 8):
        return False
    selectors = _DELIVERY_CAPTCHA_SELECTORS.get("ddangyo", ())
    for selector in selectors:
        locator = page.locator(selector).first
        try:
            if hasattr(locator, "count") and hasattr(locator, "is_visible"):
                if not await locator.count() or not await locator.is_visible(timeout=700):
                    continue
            await locator.fill(digits)
            clicked = await _delivery_bridge_click_login(page, "ddangyo")
            if not clicked and hasattr(locator, "press"):
                await locator.press("Enter")
            await page.wait_for_timeout(3000)
            try:
                await page.wait_for_load_state("networkidle", timeout=5000)
            except Exception:
                pass
            return True
        except Exception:
            continue
    try:
        result = await page.evaluate(
            r"""
            ({digits, selectors}) => {
              const visible = element => {
                if (!element) return false;
                const style = window.getComputedStyle(element);
                const rect = element.getBoundingClientRect();
                return style.visibility !== 'hidden'
                  && style.display !== 'none'
                  && rect.width > 0
                  && rect.height > 0
                  && element.type !== 'hidden'
                  && !element.disabled;
              };
              const setNativeValue = (element, value) => {
                const descriptor = Object.getOwnPropertyDescriptor(Object.getPrototypeOf(element), 'value');
                if (descriptor && descriptor.set) descriptor.set.call(element, value);
                else element.value = value;
                element.dispatchEvent(new Event('input', {bubbles: true}));
                element.dispatchEvent(new Event('change', {bubbles: true}));
                element.dispatchEvent(new KeyboardEvent('keyup', {bubbles: true, key: 'Unidentified'}));
              };
              let input = null;
              for (const selector of selectors) {
                try {
                  input = [...document.querySelectorAll(selector)].find(visible);
                  if (input) break;
                } catch (_) {}
              }
              if (!input) return {filled: false, reason: 'CAPTCHA_INPUT_NOT_FOUND'};
              input.focus();
              setNativeValue(input, digits);
              const submit = [...document.querySelectorAll('button,a,[role="button"],input[type="button"],input[type="submit"]')]
                .find(element => {
                  if (!visible(element)) return false;
                  const text = String(element.innerText || element.textContent || element.value || '').trim().toLowerCase();
                  return ['로그인', '확인', 'login', 'sign in'].some(label => text.includes(label));
                });
              if (submit) submit.click();
              else input.dispatchEvent(new KeyboardEvent('keydown', {bubbles: true, key: 'Enter', code: 'Enter'}));
              return {filled: true, clicked: Boolean(submit), reason: ''};
            }
            """,
            {"digits": digits, "selectors": list(selectors)},
        )
        if not (isinstance(result, dict) and result.get("filled")):
            return False
        await page.wait_for_timeout(3000)
        try:
            await page.wait_for_load_state("networkidle", timeout=5000)
        except Exception:
            pass
        return True
    except Exception:
        return False


def _delivery_challenge_screenshot_path(service: str, business_id: str, branch: str, session_id: str) -> Path:
    safe_service = re.sub(r"[^A-Za-z0-9_.-]+", "-", str(service or "portal")).strip("-") or "portal"
    safe_business = re.sub(r"[^A-Za-z0-9_.-]+", "-", str(business_id or "business")).strip("-") or "business"
    branch_hash = hashlib.sha256(str(branch or "").encode("utf-8")).hexdigest()[:10]
    session_hash = hashlib.sha256(str(session_id or "").encode("utf-8")).hexdigest()[:10]
    timestamp = datetime.now(KST).strftime("%Y%m%d-%H%M%S")
    return DATA_DIR / "delivery_auth_challenges" / f"{timestamp}-{safe_service}-{safe_business}-{branch_hash}-{session_hash}.png"


async def _capture_delivery_challenge_screenshot(
    page: Any,
    *,
    service: str,
    business_id: str,
    branch: str,
    session_id: str,
) -> str:
    try:
        image = await page.screenshot(full_page=True)
    except Exception:
        return ""
    if not image:
        return ""
    path = _delivery_challenge_screenshot_path(service, business_id, branch, session_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(image)
    return str(path)


async def _baemin_bridge_first_visible(page: Any, selectors: tuple[str, ...]) -> Any | None:
    for selector in selectors:
        locator = page.locator(selector).first
        try:
            if not hasattr(locator, "count") or not hasattr(locator, "is_visible"):
                return locator
            if await locator.count() and await locator.is_visible(timeout=700):
                return locator
        except Exception:
            continue
    return None


async def _baemin_bridge_fill_first(page: Any, selectors: tuple[str, ...], value: str) -> bool:
    for selector in selectors:
        locator = page.locator(selector).first
        try:
            if hasattr(locator, "count") and hasattr(locator, "is_visible"):
                if not await locator.count() or not await locator.is_visible(timeout=700):
                    continue
            await locator.fill(value)
            return True
        except Exception:
            continue
    return False


async def _delivery_bridge_fill_login_dom(page: Any, service: str, username: str, password: str) -> dict[str, Any]:
    selectors = _delivery_login_selectors(service)
    try:
        result = await page.evaluate(
            r"""
            ({username, password, usernameSelectors, passwordSelectors, submitSelectors}) => {
              const visible = element => {
                if (!element) return false;
                const style = window.getComputedStyle(element);
                const rect = element.getBoundingClientRect();
                return style.visibility !== 'hidden'
                  && style.display !== 'none'
                  && rect.width > 0
                  && rect.height > 0
                  && element.type !== 'hidden'
                  && !element.disabled;
              };
              const firstVisible = selectors => {
                for (const selector of selectors) {
                  try {
                    const match = [...document.querySelectorAll(selector)].find(visible);
                    if (match) return match;
                  } catch (_) {}
                }
                return null;
              };
              const setNativeValue = (element, value) => {
                const prototype = Object.getPrototypeOf(element);
                const descriptor = Object.getOwnPropertyDescriptor(prototype, 'value');
                if (descriptor && descriptor.set) descriptor.set.call(element, value);
                else element.value = value;
                element.dispatchEvent(new Event('input', {bubbles: true}));
                element.dispatchEvent(new Event('change', {bubbles: true}));
                element.dispatchEvent(new KeyboardEvent('keyup', {bubbles: true, key: 'Unidentified'}));
              };
              const userInput = firstVisible(usernameSelectors);
              const passwordInput = firstVisible(passwordSelectors);
              if (!userInput || !passwordInput) {
                return {filled: false, clicked: false, reason: 'LOGIN_FORM_NOT_FOUND'};
              }
              userInput.focus();
              setNativeValue(userInput, username);
              passwordInput.focus();
              setNativeValue(passwordInput, password);

              let submit = firstVisible(submitSelectors);
              if (!submit) {
                const labels = ['로그인', 'login', 'sign in', '확인'];
                submit = [...document.querySelectorAll('button,a,[role="button"],input[type="button"],input[type="submit"]')]
                  .find(element => {
                    if (!visible(element)) return false;
                    const text = String(element.innerText || element.textContent || element.value || '').trim().toLowerCase();
                    return labels.some(label => text.includes(label));
                  });
              }
              if (submit) {
                submit.click();
                return {filled: true, clicked: true, reason: ''};
              }
              const form = passwordInput.closest('form') || userInput.closest('form');
              if (form) {
                if (typeof form.requestSubmit === 'function') form.requestSubmit();
                else form.submit();
                return {filled: true, clicked: true, reason: 'FORM_SUBMIT'};
              }
              passwordInput.dispatchEvent(new KeyboardEvent('keydown', {bubbles: true, key: 'Enter', code: 'Enter'}));
              passwordInput.dispatchEvent(new KeyboardEvent('keyup', {bubbles: true, key: 'Enter', code: 'Enter'}));
              return {filled: true, clicked: false, reason: 'ENTER_DISPATCHED'};
            }
            """,
            {
                "username": username,
                "password": password,
                "usernameSelectors": list(selectors["username"]),
                "passwordSelectors": list(selectors["password"]),
                "submitSelectors": list(selectors["submit"]),
            },
        )
        return result if isinstance(result, dict) else {"filled": bool(result), "clicked": False, "reason": ""}
    except Exception as exc:
        return {"filled": False, "clicked": False, "reason": exc.__class__.__name__}


async def _delivery_bridge_click_login(page: Any, service: str = "") -> bool:
    for selector in _delivery_login_selectors(service)["submit"]:
        locator = page.locator(selector).first
        try:
            if hasattr(locator, "count") and hasattr(locator, "is_visible"):
                if not await locator.count() or not await locator.is_visible(timeout=700):
                    continue
            await locator.click(timeout=4000)
            return True
        except Exception:
            continue
    try:
        result = await page.evaluate(
            r"""
            () => {
              const visible = element => {
                const style = window.getComputedStyle(element);
                const rect = element.getBoundingClientRect();
                return style.visibility !== 'hidden'
                  && style.display !== 'none'
                  && rect.width > 0
                  && rect.height > 0;
              };
              const buttons = [
                ...document.querySelectorAll('button,input[type="submit"],input[type="button"],a,[role="button"]')
              ];
              const target = buttons.find(element => {
                if (!visible(element) || element.disabled) return false;
                const text = String(element.innerText || element.textContent || element.value || '').trim().toLowerCase();
                return element.type === 'submit'
                  || text.includes('로그인')
                  || text.includes('login')
                  || text.includes('sign in');
              });
              if (target) {
                target.click();
                return true;
              }
              const password = [...document.querySelectorAll('input[type="password"]')].find(visible);
              if (!password) return false;
              password.dispatchEvent(new KeyboardEvent('keydown', {bubbles: true, key: 'Enter', code: 'Enter'}));
              password.dispatchEvent(new KeyboardEvent('keyup', {bubbles: true, key: 'Enter', code: 'Enter'}));
              return true;
            }
            """
        )
        return result is True
    except Exception:
        return False
    return False


async def _baemin_bridge_click_login(page: Any) -> bool:
    return await _delivery_bridge_click_login(page, "baemin")


async def _baemin_bridge_login_with_saved_secret(page: Any, account: dict[str, Any]) -> dict[str, Any] | None:
    username = str(account.get("username") or "").strip()
    password = _decrypt_secret(str(account.get("password_enc") or "")) if _has_secret_value(account, "password") else ""
    if not username or not password:
        return {
            "status": "credential_required",
            "error_code": "PC_AGENT_LOGIN_REQUIRED",
            "records": {},
            "message": "PC Agent 브라우저가 배민 로그인 화면입니다. 저장된 배민 계정 비밀번호가 필요합니다.",
        }
    try:
        try:
            await page.wait_for_load_state("networkidle", timeout=8000)
        except Exception:
            pass
        try:
            await page.wait_for_selector("input[type='password']", state="visible", timeout=8000)
        except Exception:
            pass
        selectors = _delivery_login_selectors("baemin")
        username_filled = await _baemin_bridge_fill_first(
            page,
            selectors["username"],
            username,
        )
        password_filled = await _baemin_bridge_fill_first(
            page,
            selectors["password"],
            password,
        )
        if not username_filled or not password_filled:
            dom_result = await _delivery_bridge_fill_login_dom(page, "baemin", username, password)
            if not dom_result.get("filled"):
                return {
                    "status": "portal_action_required",
                    "error_code": "LOGIN_FORM_NOT_FOUND",
                    "records": {},
                    "diagnostics": {"login_automation": "dom_fallback_failed", "login_reason": str(dom_result.get("reason") or "")},
                }
            clicked = bool(dom_result.get("clicked"))
        else:
            clicked = await _baemin_bridge_click_login(page)
        if not clicked:
            password_input = await _baemin_bridge_first_visible(
                page,
                selectors["password"],
            )
            if password_input is not None and hasattr(password_input, "press"):
                await password_input.press("Enter")
            elif hasattr(page, "press_key"):
                await page.press_key("Enter")
        await page.wait_for_timeout(5000)
        try:
            await page.wait_for_load_state("networkidle", timeout=5000)
        except Exception:
            pass
    finally:
        password = ""
    return None


async def _delivery_bridge_login_with_saved_secret(
    page: Any,
    account: dict[str, Any],
    service_label: str,
) -> dict[str, Any] | None:
    service = str(account.get("service") or "").strip()
    username = str(account.get("username") or "").strip()
    password = _decrypt_secret(str(account.get("password_enc") or "")) if _has_secret_value(account, "password") else ""
    if not username or not password:
        return {
            "status": "credential_required",
            "error_code": "PC_AGENT_LOGIN_REQUIRED",
            "records": {},
            "message": f"PC Agent 브라우저가 {service_label} 로그인 화면입니다. 저장된 계정 ID/PW가 필요합니다.",
        }
    try:
        try:
            await page.wait_for_load_state("networkidle", timeout=8000)
        except Exception:
            pass
        try:
            await page.wait_for_selector("input[type='password']", state="visible", timeout=8000)
        except Exception:
            pass
        selectors = _delivery_login_selectors(service)
        username_filled = await _baemin_bridge_fill_first(
            page,
            selectors["username"],
            username,
        )
        password_filled = await _baemin_bridge_fill_first(
            page,
            selectors["password"],
            password,
        )
        if not username_filled or not password_filled:
            dom_result = await _delivery_bridge_fill_login_dom(page, service, username, password)
            if not dom_result.get("filled"):
                return {
                    "status": "portal_action_required",
                    "error_code": "LOGIN_FORM_NOT_FOUND",
                    "records": {},
                    "diagnostics": {"login_automation": "dom_fallback_failed", "login_reason": str(dom_result.get("reason") or "")},
                }
            clicked = bool(dom_result.get("clicked"))
        else:
            clicked = await _delivery_bridge_click_login(page, service)
        if not clicked:
            password_input = await _baemin_bridge_first_visible(
                page,
                selectors["password"],
            )
            if password_input is not None and hasattr(password_input, "press"):
                await password_input.press("Enter")
            elif hasattr(page, "press_key"):
                await page.press_key("Enter")
        await page.wait_for_timeout(5000)
        try:
            await page.wait_for_load_state("networkidle", timeout=5000)
        except Exception:
            pass
    finally:
        password = ""
    return None


async def _delivery_bridge_click_first(page: Any, labels: tuple[str, ...], timeout: int = 2500) -> bool:
    if not hasattr(page, "get_by_role"):
        try:
            clicked = await page.evaluate(
                r"""
                labels => {
                  const normalizedLabels = labels.map(value => String(value || '').toLowerCase());
                  const candidates = [
                    ...document.querySelectorAll('button,a,[role="button"],input[type="button"],input[type="submit"],li,span,div')
                  ];
                  const visible = element => {
                    const style = window.getComputedStyle(element);
                    const rect = element.getBoundingClientRect();
                    return style.visibility !== 'hidden' && style.display !== 'none' && rect.width > 0 && rect.height > 0;
                  };
                  const textOf = element => String(element.innerText || element.textContent || element.value || '').trim().toLowerCase();
                  const target = candidates.find(element => {
                    if (!visible(element)) return false;
                    const text = textOf(element);
                    return text && normalizedLabels.some(label => text.includes(label));
                  });
                  if (!target) return false;
                  target.click();
                  return true;
                }
                """,
                list(labels),
            )
            if clicked:
                await page.wait_for_timeout(800)
                return True
        except Exception:
            return False
    for label in labels:
        pattern = re.compile(re.escape(label), re.I)
        for role in ("button", "link", None):
            try:
                matches = page.get_by_role(role, name=pattern) if role else page.get_by_text(pattern)
                count = await matches.count() if hasattr(matches, "count") else 0
            except Exception:
                continue
            for index in range(min(count, 20)):
                locator = matches.nth(index)
                try:
                    if await locator.is_visible(timeout=500):
                        await locator.click(timeout=timeout)
                        await page.wait_for_timeout(800)
                        return True
                except Exception:
                    continue
    return False


async def _delivery_bridge_set_period(page: Any, date_from: str, date_to: str) -> None:
    try:
        if not hasattr(page.locator("body"), "count"):
            await page.evaluate(
                r"""
                ({dateFrom, dateTo}) => {
                  const dateInputs = [...document.querySelectorAll('input[type="date"]')];
                  if (dateInputs[0]) dateInputs[0].value = dateFrom;
                  if (dateInputs[1]) dateInputs[1].value = dateTo;
                  const startInputs = [...document.querySelectorAll('input[title*="시작 날짜"],input[placeholder*="시작"]')];
                  const endInputs = [...document.querySelectorAll('input[title*="종료 날짜"],input[placeholder*="종료"]')];
                  if (!dateInputs[0] && startInputs[0]) startInputs[0].value = dateFrom;
                  if (!dateInputs[1] && endInputs[0]) endInputs[0].value = dateTo;
                  [...dateInputs, ...startInputs, ...endInputs].forEach(input => {
                    input.dispatchEvent(new Event('input', {bubbles: true}));
                    input.dispatchEvent(new Event('change', {bubbles: true}));
                  });
                }
                """,
                {"dateFrom": date_from, "dateTo": date_to},
            )
            await _delivery_bridge_click_first(page, ("조회", "검색", "적용"), timeout=2500)
            return
        date_inputs = page.locator("input[type='date']")
        count = await date_inputs.count()
        if count >= 1:
            await date_inputs.nth(0).fill(date_from)
        if count >= 2:
            await date_inputs.nth(1).fill(date_to)
        if count < 2:
            for selector, value in (("input[title*='시작 날짜']", date_from), ("input[title*='종료 날짜']", date_to)):
                locator = page.locator(selector).first
                try:
                    if await locator.count() and await locator.is_visible(timeout=400):
                        await locator.fill(value)
                except Exception:
                    continue
        await _delivery_bridge_click_first(page, ("조회", "검색", "적용"), timeout=2500)
    except Exception:
        return


async def _collect_delivery_from_browser_bridge_session_async(
    account: dict[str, Any],
    browser_auth: dict[str, str],
    date_from: str,
    date_to: str,
) -> dict[str, Any]:
    service = str(account.get("service") or "").strip()
    session_id = str(browser_auth.get("browser_session_id") or "").strip()
    stage_logs: list[dict[str, str]] = []
    if not session_id:
        started_at = time.monotonic()
        _append_delivery_stage_log(
            stage_logs,
            service=service,
            stage="browser_session",
            status="failed",
            started_at=started_at,
            error_code="PC_AGENT_SESSION_REQUIRED",
            reason="session_id_missing",
        )
        return {
            "status": "credential_required",
            "error_code": "PC_AGENT_SESSION_REQUIRED",
            "records": {},
            "diagnostics": _delivery_browser_diagnostics(browser_auth, session_id, stage_logs=stage_logs),
        }
    try:
        from app.browser_bridge.service import get_browser_bridge_service
        from app.services.yeoljeong_delivery_collectors import PORTAL_CONFIG, parse_portal_export

        config = PORTAL_CONFIG.get(service)
        if not config:
            return {"status": "failed", "error_code": "UNSUPPORTED_PLATFORM", "records": {}}
        service_label = _delivery_platform_label(service)
        bridge = get_browser_bridge_service()
        session = bridge.sessions.get(session_id)
        if not session:
            started_at = time.monotonic()
            _append_delivery_stage_log(
                stage_logs,
                service=service,
                stage="browser_session",
                status="failed",
                started_at=started_at,
                error_code="PC_AGENT_SESSION_NOT_FOUND",
                reason="session_not_found",
            )
            return {
                "status": "credential_required",
                "error_code": "PC_AGENT_SESSION_NOT_FOUND",
                "records": {},
                "diagnostics": _delivery_browser_diagnostics(browser_auth, session_id, stage_logs=stage_logs),
            }
        session_started_at = time.monotonic()
        _append_delivery_stage_log(
            stage_logs,
            service=service,
            stage="browser_session",
            status="success",
            started_at=session_started_at,
            reason="session_ready",
        )
        context = await bridge._context_for_session(session)
        page = await _delivery_bridge_page_for_service(context, service)

        if service == "baemin":
            home_url = "https://self.baemin.com/"
        else:
            home_url = str(account.get("portal_home_url") or account.get("home_url") or account.get("login_url") or config["login_url"])
        site_started_at = time.monotonic()
        try:
            await page.goto(home_url, wait_until="domcontentloaded", timeout=45000)
            _append_delivery_stage_log(
                stage_logs,
                service=service,
                stage="site_access",
                status="success",
                started_at=site_started_at,
                reason="goto_domcontentloaded",
                timeout_ms=45000,
            )
        except Exception:
            _append_delivery_stage_log(
                stage_logs,
                service=service,
                stage="site_access",
                status="failed",
                started_at=site_started_at,
                error_code="PORTAL_NAVIGATION_FAILED",
                reason="page_goto_failed",
                timeout_ms=45000,
            )
        try:
            await page.wait_for_load_state("networkidle", timeout=8000)
        except Exception:
            pass

        url = ""
        text = ""
        html = ""
        try:
            url = str(await page.evaluate("window.location.href") or "")
            text = str(await page.evaluate("document.body ? document.body.innerText : ''") or "")
            html = str(await page.evaluate("document.body ? document.body.innerHTML : ''") or "")
        except Exception:
            pass

        auth_diagnostics: dict[str, str] = {}
        login_state = _baemin_bridge_login_state(url, text) if service == "baemin" else _delivery_bridge_login_state(url, text)
        _append_delivery_stage_log(
            stage_logs,
            service=service,
            stage="auth_state",
            status="success" if login_state == "authenticated" else "pending",
            started_at=time.monotonic(),
            reason=login_state,
            current_url=url[:120],
        )
        if login_state == "login":
            login_result = (
                await _baemin_bridge_login_with_saved_secret(page, account)
                if service == "baemin"
                else await _delivery_bridge_login_with_saved_secret(page, account, service_label)
            )
            if login_result is not None:
                login_result.setdefault("diagnostics", {}).update(
                    {
                        "auth_mode": "pc_agent_browser",
                        "browser_session_id": session_id,
                        "browser_work_key": str(browser_auth.get("browser_work_key") or ""),
                        "url": url,
                        "browser_stage_logs": stage_logs,
                        "browser_stage_log_schema": SITE_STAGE_LOG_SCHEMA,
                    }
                )
                return login_result
            try:
                url = str(await page.evaluate("window.location.href") or url)
                text = str(await page.evaluate("document.body ? document.body.innerText : ''") or "")
                html = str(await page.evaluate("document.body ? document.body.innerHTML : ''") or "")
            except Exception:
                pass
            login_state = _baemin_bridge_login_state(url, text) if service == "baemin" else _delivery_bridge_login_state(url, text)
        if login_state == "blocked":
            return {
                "status": "portal_action_required",
                "error_code": f"{service.upper()}_SECURITY_BLOCKED",
                "records": {},
                "diagnostics": _delivery_browser_diagnostics(browser_auth, session_id, url=url, stage_logs=stage_logs),
                "message": f"{service_label} 포털이 접속을 보안 정책으로 차단했습니다. PC 브라우저에서 인증 또는 정산 CSV 업로드가 필요합니다.",
            }
        if login_state == "challenge":
            challenge_code = _delivery_bridge_challenge_code(service, text)
            challenge_screenshot = await _capture_delivery_challenge_screenshot(
                page,
                service=service,
                business_id=str(account.get("business_id") or ""),
                branch=str(account.get("branch") or ""),
                session_id=session_id,
            )
            diagnostics = {
                "auth_mode": "pc_agent_browser",
                "browser_session_id": session_id,
                "browser_work_key": str(browser_auth.get("browser_work_key") or ""),
                "url": url,
                "browser_stage_logs": stage_logs,
                "browser_stage_log_schema": SITE_STAGE_LOG_SCHEMA,
            }
            if challenge_screenshot:
                diagnostics["challenge_screenshot_path"] = challenge_screenshot
            captcha_value = str(account.get("captcha_value") or "")
            captcha_approval = account.get("_captcha_approval") if isinstance(account.get("_captcha_approval"), dict) else {}
            captcha_auto_allowed, captcha_auto_reason = _delivery_captcha_automation_allowed(captcha_approval, url)
            if captcha_approval:
                diagnostics["captcha_auto_approval"] = captcha_auto_reason
                diagnostics["captcha_approved_by"] = str(captcha_approval.get("approved_by") or "")
                diagnostics["captcha_approved_at"] = str(captcha_approval.get("approved_at") or "")
                diagnostics["captcha_approved_origin"] = str(captcha_approval.get("origin") or "")
            if challenge_code == "DDANGYO_NUMERIC_CAPTCHA_REQUIRED" and not captcha_value:
                if captcha_auto_allowed:
                    _record_delivery_captcha_automation_event(
                        "captcha_auto_approval_used",
                        captcha_approval,
                        url=url,
                        session_id=session_id,
                        work_key=str(browser_auth.get("browser_work_key") or ""),
                        run_id=str(account.get("_run_id") or ""),
                        reason=captcha_auto_reason,
                    )
                    try:
                        from app.services.captcha_vision_solver import solve_captcha_with_vision

                        captcha_value = await solve_captcha_with_vision(
                            page,
                            screenshot_path=str(challenge_screenshot or ""),
                            approval_context={**captcha_approval, "page_url": url},
                        )
                        if captcha_value:
                            diagnostics["captcha_mode"] = "approved_vision_auto_solved"
                            diagnostics["captcha_source"] = "approved_vision_model"
                        else:
                            diagnostics["captcha_vision"] = "no_digits"
                    except Exception as _captcha_vision_exc:
                        diagnostics["captcha_vision_error"] = str(_captcha_vision_exc)[:150]
                else:
                    diagnostics["captcha_vision"] = "approval_required"
            if challenge_code == "DDANGYO_NUMERIC_CAPTCHA_REQUIRED":
                if not captcha_value:
                    diagnostics["captcha_input"] = "operator_input_required"
                    return {
                        "status": "portal_action_required",
                        "error_code": challenge_code,
                        "message": "땡겨요 숫자 캡챠 자동입력 승인이 필요합니다. 승인 범위를 전달하면 같은 PC Agent 세션에서 모델 판독과 입력을 자동 수행합니다.",
                        "records": {},
                        "diagnostics": diagnostics,
                    }
                captcha_accepted = False
                _cred_username = str(account.get("username") or "").strip()
                _cred_password = _decrypt_secret(str(account.get("password_enc") or "")) if _has_secret_value(account, "password") else ""
                for _captcha_attempt in range(3):
                    if not captcha_value:
                        break
                    if _cred_username and _cred_password:
                        _login_sels = _delivery_login_selectors(service)
                        await _baemin_bridge_fill_first(page, _login_sels["username"], _cred_username)
                        await _baemin_bridge_fill_first(page, _login_sels["password"], _cred_password)
                    if not await _delivery_bridge_fill_ddangyo_numeric_captcha(page, captcha_value):
                        diagnostics["captcha_input"] = "input_failed"
                        break
                    try:
                        url = str(await page.evaluate("window.location.href") or url)
                        text = str(await page.evaluate("document.body ? document.body.innerText : ''") or "")
                        html = str(await page.evaluate("document.body ? document.body.innerHTML : ''") or "")
                    except Exception:
                        pass
                    login_state = _delivery_bridge_login_state(url, text)
                    diagnostics.update({"captcha_input": "submitted", "url": url})
                    if not diagnostics.get("captcha_mode"):
                        diagnostics["captcha_mode"] = "operator_confirmed_input"
                    if login_state == "challenge":
                        diagnostics["captcha_input"] = f"rejected_attempt_{_captcha_attempt + 1}"
                        if captcha_auto_allowed:
                            _record_delivery_captcha_automation_event(
                                "captcha_auto_retry_approval_used",
                                captcha_approval,
                                url=url,
                                session_id=session_id,
                                work_key=str(browser_auth.get("browser_work_key") or ""),
                                run_id=str(account.get("_run_id") or ""),
                                reason=f"{captcha_auto_reason}:retry_{_captcha_attempt + 1}",
                            )
                            try:
                                from app.services.captcha_vision_solver import solve_captcha_with_vision

                                captcha_value = await solve_captcha_with_vision(
                                    page,
                                    max_retries=1,
                                    approval_context={**captcha_approval, "page_url": url},
                                )
                                diagnostics["captcha_source"] = "approved_vision_model_retry"
                            except Exception as _captcha_retry_exc:
                                diagnostics["captcha_vision_error"] = str(_captcha_retry_exc)[:150]
                                captcha_value = ""
                        else:
                            captcha_value = ""
                        continue
                    if login_state != "authenticated":
                        diagnostics["captcha_input"] = f"submitted_{login_state}"
                        return {
                            "status": "portal_action_required",
                            "error_code": _delivery_bridge_challenge_code(service, text),
                            "records": {},
                            "diagnostics": diagnostics,
                            "message": _delivery_challenge_message(service, service_label),
                        }
                    diagnostics["captcha_input"] = "accepted"
                    auth_diagnostics.update({key: str(value) for key, value in diagnostics.items()})
                    captcha_accepted = True
                    break
                if not captcha_accepted:
                    return {
                        "status": "portal_action_required",
                        "error_code": challenge_code,
                        "records": {},
                        "diagnostics": diagnostics,
                        "message": _delivery_challenge_message(service, service_label),
                    }
            else:
                return {
                    "status": "portal_action_required",
                    "error_code": challenge_code,
                    "records": {},
                    "diagnostics": diagnostics,
                    "message": _delivery_challenge_message(service, service_label),
                }
        if login_state == "login":
            login_diagnostics: dict[str, str] = {
                "auth_mode": "pc_agent_browser",
                "browser_session_id": session_id,
                "browser_work_key": str(browser_auth.get("browser_work_key") or ""),
                "url": url,
                "browser_stage_log_schema": SITE_STAGE_LOG_SCHEMA,
            }
            login_diagnostics["browser_stage_logs"] = stage_logs
            login_screenshot = await _capture_delivery_challenge_screenshot(
                page,
                service=service,
                business_id=str(account.get("business_id") or ""),
                branch=str(account.get("branch") or ""),
                session_id=session_id,
            )
            if login_screenshot:
                login_diagnostics["login_failure_screenshot_path"] = login_screenshot
            return {
                "status": "credential_required",
                "error_code": "PC_AGENT_LOGIN_REQUIRED",
                "records": {},
                "diagnostics": login_diagnostics,
                "message": f"PC Agent 브라우저가 {service_label} 로그인 화면입니다. 먼저 해당 포털 로그인이 필요합니다.",
            }

        records = _delivery_empty_record_lists()
        diagnostics = _delivery_browser_diagnostics(browser_auth, session_id, url=url, stage_logs=stage_logs)
        diagnostics.update(auth_diagnostics)
        if service == "baemin":
            dashboard_records = _baemin_dashboard_records(
                text,
                str(account.get("business_id") or ""),
                str(account.get("branch") or ""),
            )
            for name, rows in dashboard_records.items():
                if rows:
                    records[name] = rows
            diagnostics["dashboard_sales"] = str(len(dashboard_records["sales"]))
            diagnostics["dashboard_settlements"] = str(len(dashboard_records["settlements"]))
            diagnostics["dashboard_reviews"] = str(len(dashboard_records["reviews"]))

        for kind, labels in config["sections"].items():
            section_started_at = time.monotonic()
            clicked = await _delivery_bridge_click_first(page, tuple(labels), timeout=3500)
            if clicked:
                await _delivery_bridge_set_period(page, date_from, date_to)
                try:
                    await page.wait_for_load_state("networkidle", timeout=5000)
                except Exception:
                    pass
            try:
                section_text = str(await page.evaluate("document.body ? document.body.innerText : ''") or "")
                section_html = str(await page.evaluate("document.body ? document.body.innerHTML : ''") or "")
            except Exception:
                section_text = text
                section_html = html
            parsed = parse_portal_export(
                service,
                kind,
                section_html or section_text,
                str(account.get("business_id") or ""),
                str(account.get("branch") or ""),
            )
            incoming = parsed.get("records", {}).get(kind) or []
            if incoming:
                records[kind] = incoming
            diagnostics[kind] = str(parsed.get("diagnostics", {}).get("source") or ("clicked" if clicked else "section_not_found"))
            _append_delivery_stage_log(
                stage_logs,
                service=service,
                stage=f"{kind}_data_collection",
                status="success" if incoming else "failed",
                started_at=section_started_at,
                error_code="" if incoming else str(parsed.get("error_code") or "NO_PARSEABLE_ROWS"),
                reason="rows_parsed" if incoming else diagnostics[kind],
                row_count=str(len(incoming)),
                timeout_ms=8500,
            )

        total_records = sum(len(rows) for rows in records.values())
        _append_delivery_stage_log(
            stage_logs,
            service=service,
            stage="data_collection",
            status="success" if total_records else "failed",
            started_at=time.monotonic(),
            error_code="" if total_records else "AUTHENTICATED_NO_ROWS",
            reason="records_found" if total_records else "authenticated_no_rows",
            row_count=str(total_records),
        )
        return {
            "status": "succeeded" if total_records else "partial",
            "error_code": "" if total_records else "AUTHENTICATED_NO_ROWS",
            "records": records,
            "diagnostics": diagnostics,
            "message": "" if total_records else f"{service_label} 로그인은 확인됐지만 조회 구간에서 표 데이터를 찾지 못했습니다.",
        }
    except Exception as exc:
        _append_delivery_stage_log(
            stage_logs,
            service=service,
            stage="collector_exception",
            status="failed",
            started_at=time.monotonic(),
            error_code=f"PC_AGENT_COLLECTOR_{exc.__class__.__name__.upper()}",
            reason=str(exc)[:160],
        )
        return {
            "status": "failed",
            "error_code": f"PC_AGENT_COLLECTOR_{exc.__class__.__name__.upper()}",
            "records": {},
            "diagnostics": _delivery_browser_diagnostics(browser_auth, session_id, stage_logs=stage_logs),
            "message": str(exc)[:300],
        }
    finally:
        await _close_delivery_browser_work_session_async(
            browser_auth,
            reason=f"delivery_collect_complete_{service or 'portal'}",
        )


async def _collect_baemin_from_browser_bridge_session_async(
    account: dict[str, Any],
    browser_auth: dict[str, str],
    backfill: dict[str, Any] | None = None,
) -> dict[str, Any]:
    session_id = str(browser_auth.get("browser_session_id") or "").strip()
    stage_logs: list[dict[str, str]] = []
    if not session_id:
        _append_delivery_stage_log(
            stage_logs,
            service="baemin",
            stage="browser_session",
            status="failed",
            started_at=time.monotonic(),
            error_code="PC_AGENT_SESSION_REQUIRED",
            reason="session_id_missing",
        )
        return {
            "status": "credential_required",
            "error_code": "PC_AGENT_SESSION_REQUIRED",
            "records": {},
            "diagnostics": _delivery_browser_diagnostics(browser_auth, session_id, stage_logs=stage_logs),
        }
    try:
        from app.browser_bridge.service import get_browser_bridge_service
        from app.services.baemin_order_history_collector import BackfillLimits, collect_baemin_order_history
        from app.services.baemin_review_collector import collect_reviews
        from app.services.baemin_ads_collector import collect_ads
        from app.services.yeoljeong_delivery_collectors import parse_portal_export

        bridge = get_browser_bridge_service()
        session = bridge.sessions.get(session_id)
        if not session:
            _append_delivery_stage_log(
                stage_logs,
                service="baemin",
                stage="browser_session",
                status="failed",
                started_at=time.monotonic(),
                error_code="PC_AGENT_SESSION_NOT_FOUND",
                reason="session_not_found",
            )
            return {
                "status": "credential_required",
                "error_code": "PC_AGENT_SESSION_NOT_FOUND",
                "records": {},
                "diagnostics": _delivery_browser_diagnostics(browser_auth, session_id, stage_logs=stage_logs),
            }
        _append_delivery_stage_log(
            stage_logs,
            service="baemin",
            stage="browser_session",
            status="success",
            started_at=time.monotonic(),
            reason="session_ready",
        )
        context = await bridge._context_for_session(session)
        page = await _delivery_bridge_page_for_service(context, "baemin")
        url = str(getattr(page, "url", "") or "")
        try:
            url = str(await page.evaluate("window.location.href") or url)
        except Exception:
            pass
        if "baemin.com" not in url.lower():
            site_started_at = time.monotonic()
            try:
                await page.goto("https://self.baemin.com/", wait_until="domcontentloaded", timeout=45000)
                try:
                    await page.wait_for_load_state("networkidle", timeout=8000)
                except Exception:
                    pass
                url = str(await page.evaluate("window.location.href") or "")
                _append_delivery_stage_log(
                    stage_logs,
                    service="baemin",
                    stage="site_access",
                    status="success",
                    started_at=site_started_at,
                    reason="goto_domcontentloaded",
                    timeout_ms=45000,
                    current_url=url[:120],
                )
            except Exception:
                _append_delivery_stage_log(
                    stage_logs,
                    service="baemin",
                    stage="site_access",
                    status="failed",
                    started_at=site_started_at,
                    error_code="PORTAL_NAVIGATION_FAILED",
                    reason="page_goto_failed",
                    timeout_ms=45000,
                )
        else:
            _append_delivery_stage_log(
                stage_logs,
                service="baemin",
                stage="site_access",
                status="success",
                started_at=time.monotonic(),
                reason="reused_baemin_tab",
                current_url=url[:120],
            )
        text = ""
        html = ""
        try:
            text = str(await page.evaluate("document.body ? document.body.innerText : ''") or "")
        except Exception:
            text = ""
        try:
            html = str(await page.evaluate("document.body ? document.body.innerHTML : ''") or "")
        except Exception:
            html = text
        login_state = _baemin_bridge_login_state(url, text)
        _append_delivery_stage_log(
            stage_logs,
            service="baemin",
            stage="auth_state",
            status="success" if login_state == "authenticated" else "pending",
            started_at=time.monotonic(),
            reason=login_state,
            current_url=url[:120],
        )
        if login_state == "blocked":
            return {
                "status": "portal_action_required",
                "error_code": "BAEMIN_SECURITY_BLOCKED",
                "records": {},
                "diagnostics": _delivery_browser_diagnostics(browser_auth, session_id, url=url, stage_logs=stage_logs),
                "message": "배민 포털이 접속을 보안 정책으로 차단했습니다. PC 브라우저에서 직접 인증 또는 정산 CSV 업로드가 필요합니다.",
            }
        if login_state == "challenge":
            challenge_screenshot = await _capture_delivery_challenge_screenshot(
                page,
                service="baemin",
                business_id=str(account.get("business_id") or ""),
                branch=str(account.get("branch") or ""),
                session_id=session_id,
            )
            diagnostics = {
                "auth_mode": "pc_agent_browser",
                "browser_session_id": session_id,
                "browser_work_key": str(browser_auth.get("browser_work_key") or ""),
                "browser_bridge_mode": str(browser_auth.get("browser_bridge_mode") or ""),
                "url": url,
                "browser_stage_logs": stage_logs,
                "browser_stage_log_schema": SITE_STAGE_LOG_SCHEMA,
            }
            if challenge_screenshot:
                diagnostics["challenge_screenshot_path"] = challenge_screenshot
            return {
                "status": "portal_action_required",
                "error_code": "PORTAL_AUTH_CHALLENGE",
                "records": {},
                "diagnostics": diagnostics,
                "message": "배민 포털이 추가 인증을 요구합니다. PC 브라우저에서 인증을 완료한 뒤 다시 수집해야 합니다.",
            }
        if login_state == "login":
            login_result = await _baemin_bridge_login_with_saved_secret(page, account)
            if login_result is not None:
                login_result.setdefault("diagnostics", {}).update(
                    {
                        "auth_mode": "pc_agent_browser",
                        "browser_session_id": session_id,
                        "browser_work_key": str(browser_auth.get("browser_work_key") or ""),
                        "browser_bridge_mode": str(browser_auth.get("browser_bridge_mode") or ""),
                        "url": url,
                        "browser_stage_logs": stage_logs,
                        "browser_stage_log_schema": SITE_STAGE_LOG_SCHEMA,
                    }
                )
                return login_result
            try:
                url = str(await page.evaluate("window.location.href") or url)
            except Exception:
                pass
            try:
                text = str(await page.evaluate("document.body ? document.body.innerText : ''") or "")
            except Exception:
                text = ""
            try:
                html = str(await page.evaluate("document.body ? document.body.innerHTML : ''") or "")
            except Exception:
                html = text
            login_state = _baemin_bridge_login_state(url, text)
        if login_state == "login":
            return {
                "status": "credential_required",
                "error_code": "PC_AGENT_LOGIN_REQUIRED",
                "records": {},
                "diagnostics": _delivery_browser_diagnostics(browser_auth, session_id, url=url, stage_logs=stage_logs),
                "message": "PC Agent 브라우저가 배민 로그인 화면입니다. 먼저 해당 브라우저에서 배민 관리자 로그인이 필요합니다.",
            }
        if login_state == "challenge":
            challenge_screenshot = await _capture_delivery_challenge_screenshot(
                page,
                service="baemin",
                business_id=str(account.get("business_id") or ""),
                branch=str(account.get("branch") or ""),
                session_id=session_id,
            )
            diagnostics = {
                "auth_mode": "pc_agent_browser",
                "browser_session_id": session_id,
                "browser_work_key": str(browser_auth.get("browser_work_key") or ""),
                "browser_bridge_mode": str(browser_auth.get("browser_bridge_mode") or ""),
                "url": url,
                "browser_stage_logs": stage_logs,
                "browser_stage_log_schema": SITE_STAGE_LOG_SCHEMA,
            }
            if challenge_screenshot:
                diagnostics["challenge_screenshot_path"] = challenge_screenshot
            return {
                "status": "portal_action_required",
                "error_code": "PORTAL_AUTH_CHALLENGE",
                "records": {},
                "diagnostics": diagnostics,
                "message": "배민 포털이 추가 인증을 요구합니다. PC 브라우저에서 인증을 완료한 뒤 다시 수집해야 합니다.",
            }
        if login_state == "blocked":
            return {
                "status": "portal_action_required",
                "error_code": "BAEMIN_SECURITY_BLOCKED",
                "records": {},
                "diagnostics": _delivery_browser_diagnostics(browser_auth, session_id, url=url, stage_logs=stage_logs),
                "message": "배민 포털이 접속을 보안 정책으로 차단했습니다. PC 브라우저에서 직접 인증 또는 정산 CSV 업로드가 필요합니다.",
            }
        full_backfill = bool(backfill and backfill.get("mode") == "full_backfill")
        max_orders = min(300, max(1, int((backfill or {}).get("max_orders") or account.get("_max_orders") or 300)))
        history_result = await collect_baemin_order_history(
            page,
            account,
            str(account.get("_date_from") or ""),
            str(account.get("_date_to") or ""),
            {
                "max_orders": max_orders,
                "checkpoint": (backfill or {}).get("checkpoint") or {},
                "limits": BackfillLimits(max_records=max_orders, max_runtime_seconds=12 * 60),
            },
        )
        history_records = history_result.get("records") if isinstance(history_result, dict) else {}
        history_total = sum(
            len(rows)
            for rows in (history_records or {}).values()
            if isinstance(rows, list)
        )
        _append_delivery_stage_log(
            stage_logs,
            service="baemin",
            stage="order_history_data_collection",
            status="success" if history_total else "failed",
            started_at=time.monotonic(),
            error_code="" if history_total else str((history_result or {}).get("error_code") or "NO_ORDER_HISTORY_ROWS"),
            reason="rows_parsed" if history_total else str((history_result or {}).get("status") or "empty"),
            row_count=str(history_total),
        )

        source = html or text
        dashboard_records = _baemin_dashboard_records(
            text,
            str(account.get("business_id") or ""),
            str(account.get("branch") or ""),
        )
        kind = _baemin_bridge_page_kind(text)
        parsed = parse_portal_export(
            "baemin",
            kind,
            source,
            str(account.get("business_id") or ""),
            str(account.get("branch") or ""),
        )
        records = _delivery_empty_record_lists()
        if history_total:
            for name, rows in (history_records or {}).items():
                if rows:
                    records[name] = list(rows)
        review_result: dict[str, Any] = {}
        ads_result: dict[str, Any] = {}
        if full_backfill:
            max_reviews = min(300, max(1, int((backfill or {}).get("max_reviews") or 300)))
            review_result = await collect_reviews(
                page,
                business_id=str(account.get("business_id") or ""),
                branch=str(account.get("branch") or ""),
                max_records=max_reviews,
            )
            ads_result = await collect_ads(
                page,
                business_id=str(account.get("business_id") or ""),
                branch=str(account.get("branch") or ""),
            )
            _append_delivery_stage_log(
                stage_logs,
                service="baemin",
                stage="review_data_collection",
                status="success" if ((review_result or {}).get("records") or {}).get("reviews") else "failed",
                started_at=time.monotonic(),
                error_code="" if ((review_result or {}).get("records") or {}).get("reviews") else str((review_result or {}).get("error_code") or "NO_REVIEW_ROWS"),
                reason=str((review_result or {}).get("status") or ""),
                row_count=str(len(((review_result or {}).get("records") or {}).get("reviews") or [])),
            )
            _append_delivery_stage_log(
                stage_logs,
                service="baemin",
                stage="ads_data_collection",
                status="success" if ((ads_result or {}).get("records") or {}).get("ads") else "failed",
                started_at=time.monotonic(),
                error_code="" if ((ads_result or {}).get("records") or {}).get("ads") else str((ads_result or {}).get("error_code") or "NO_ADS_ROWS"),
                reason=str((ads_result or {}).get("status") or ""),
                row_count=str(len(((ads_result or {}).get("records") or {}).get("ads") or [])),
            )
            for name, rows in (review_result.get("records") or {}).items():
                if rows:
                    records[name] = list(rows)
            for name, rows in (ads_result.get("records") or {}).items():
                if rows:
                    records[name] = list(rows)
        for name, rows in dashboard_records.items():
            if rows and not records.get(name):
                records[name] = rows
        for name, rows in (parsed.get("records") or {}).items():
            if rows and not records.get(name):
                records[name] = rows
        total_records = sum(len(rows) for rows in records.values())
        _append_delivery_stage_log(
            stage_logs,
            service="baemin",
            stage="data_collection",
            status="success" if total_records else "failed",
            started_at=time.monotonic(),
            error_code="" if total_records else str(parsed.get("error_code") or "NO_PARSEABLE_ROWS"),
            reason="records_found" if total_records else str(parsed.get("status") or "empty"),
            row_count=str(total_records),
        )
        status = "succeeded" if total_records else str(parsed.get("status") or "partial")
        error_code = "" if total_records else str(parsed.get("error_code") or "")
        diagnostics = dict(parsed.get("diagnostics") or {})
        diagnostics.update(
            {
                "auth_mode": "pc_agent_browser",
                "browser_session_id": session_id,
                "browser_work_key": str(browser_auth.get("browser_work_key") or ""),
                "browser_bridge_mode": str(browser_auth.get("browser_bridge_mode") or ""),
                "url": url,
                "parsed_page_kind": kind,
                "order_history_status": str((history_result or {}).get("status") or ""),
                "order_history_error_code": str((history_result or {}).get("error_code") or ""),
                "order_history_orders_seen": str(
                    ((history_result or {}).get("diagnostics") or {}).get("orders_seen") or 0
                ),
                "order_history_orders_saved": str(
                    ((history_result or {}).get("diagnostics") or {}).get("orders_saved") or 0
                ),
                "order_history_detail_failed": str(
                    ((history_result or {}).get("diagnostics") or {}).get("detail_failed") or 0
                ),
                "order_history_settlement_pending": str(
                    ((history_result or {}).get("diagnostics") or {}).get("settlement_pending") or 0
                ),
                "checkpoint": ((history_result or {}).get("diagnostics") or {}).get("checkpoint_out") or {},
                "review_backfill_status": str((review_result or {}).get("status") or ""),
                "review_backfill_reviews_saved": len(((review_result or {}).get("records") or {}).get("reviews") or []),
                "ads_backfill_status": str((ads_result or {}).get("status") or ""),
                "ads_backfill_ads_saved": len(((ads_result or {}).get("records") or {}).get("ads") or []),
                "dashboard_sales": str(len(dashboard_records["sales"])),
                "dashboard_settlements": str(len(dashboard_records["settlements"])),
                "dashboard_reviews": str(len(dashboard_records["reviews"])),
                "browser_stage_logs": stage_logs,
                "browser_stage_log_schema": SITE_STAGE_LOG_SCHEMA,
            }
        )
        return {
            "status": status,
            "error_code": error_code,
            "records": records,
            "diagnostics": diagnostics,
            "message": parsed.get("message") or "",
        }
    except Exception as exc:
        _append_delivery_stage_log(
            stage_logs,
            service="baemin",
            stage="collector_exception",
            status="failed",
            started_at=time.monotonic(),
            error_code=f"PC_AGENT_COLLECTOR_{exc.__class__.__name__.upper()}",
            reason=str(exc)[:160],
        )
        return {
            "status": "failed",
            "error_code": f"PC_AGENT_COLLECTOR_{exc.__class__.__name__.upper()}",
            "records": {},
            "diagnostics": _delivery_browser_diagnostics(browser_auth, session_id, stage_logs=stage_logs),
            "message": str(exc)[:300],
        }
    finally:
        await _close_delivery_browser_work_session_async(
            browser_auth,
            reason="delivery_collect_complete_baemin",
        )


def _collect_baemin_from_browser_bridge_session(
    account: dict[str, Any],
    browser_auth: dict[str, str],
    backfill: dict[str, Any] | None = None,
) -> dict[str, Any] | None:
    if str(browser_auth.get("browser_bridge_mode") or "") != "local_agent":
        return None
    if not str(browser_auth.get("browser_session_id") or "").strip():
        return None
    return _run_async(_collect_baemin_from_browser_bridge_session_async(account, browser_auth, backfill))


def _collect_delivery_from_browser_bridge_session(
    account: dict[str, Any],
    browser_auth: dict[str, str],
    date_from: str,
    date_to: str,
) -> dict[str, Any] | None:
    if str(browser_auth.get("browser_bridge_mode") or "") != "local_agent":
        return None
    if not str(browser_auth.get("browser_session_id") or "").strip():
        return None
    return _run_async(_collect_delivery_from_browser_bridge_session_async(account, browser_auth, date_from, date_to))


def _normalize_delivery_collection_result(service: str, result: dict[str, Any]) -> dict[str, Any]:
    if service == "ddangyo" and str(result.get("error_code") or "").upper() == "DDANGYO_NUMERIC_CAPTCHA_REQUIRED":
        diagnostics = result.get("diagnostics") if isinstance(result.get("diagnostics"), dict) else {}
        return {
            **result,
            "status": "portal_action_required",
            "error_code": "DDANGYO_NUMERIC_CAPTCHA_REQUIRED",
            "message": _delivery_challenge_message("ddangyo", _delivery_platform_label("ddangyo")),
            "diagnostics": diagnostics,
        }
    if _delivery_result_is_wrong_portal(service, result):
        label = _delivery_platform_label(service)
        diagnostics = result.get("diagnostics") if isinstance(result.get("diagnostics"), dict) else {}
        rejected_counts = {
            kind: len(result.get("records", {}).get(kind) or [])
            for kind in DELIVERY_RECORD_TYPES
        }
        diagnostics = {**diagnostics, "wrong_portal_rejected_counts": rejected_counts}
        return {
            **result,
            "status": "portal_action_required",
            "error_code": "PC_AGENT_WRONG_PORTAL_SESSION",
            "records": _delivery_empty_record_lists(),
            "message": (
                f"{label} 자동수집 세션이 다른 포털 화면에 연결됐습니다. "
                "플랫폼별 PC Agent 작업 세션을 다시 생성한 뒤 재수집해야 합니다."
            ),
            "diagnostics": diagnostics,
        }
    if _delivery_result_has_no_visible_source(result):
        label = _delivery_platform_label(service)
        diagnostics = result.get("diagnostics") if isinstance(result.get("diagnostics"), dict) else {}
        return {
            **result,
            "status": "portal_action_required",
            "error_code": "PORTAL_TABLE_NOT_FOUND",
            "message": (
                f"{label} 로그인은 확인됐지만 매출/정산/리뷰 표를 찾지 못했습니다. "
                "포털 메뉴 구조 또는 조회 조건 확인 후 PC Agent 세션에서 다시 수집해야 합니다."
            ),
            "diagnostics": diagnostics,
        }
    if service != "baemin" and str(result.get("error_code") or "") == "BAEMIN_SECURITY_BLOCKED":
        label = _delivery_platform_label(service)
        result = {**result}
        result["error_code"] = f"{service.upper()}_SECURITY_BLOCKED"
        result["message"] = f"{label} 포털이 서버 자동접속을 보안 정책으로 차단했습니다. PC 인증 세션 또는 정산 CSV 업로드가 필요합니다."
    filtered_records, rejected_counts = _delivery_filter_meaningful_records(result.get("records"))
    rejected_total = sum(rejected_counts.values())
    if rejected_total:
        diagnostics = result.get("diagnostics") if isinstance(result.get("diagnostics"), dict) else {}
        result = {**result, "records": filtered_records}
        result["diagnostics"] = {
            **diagnostics,
            "invalid_record_rejected_counts": rejected_counts,
        }
        if not any(filtered_records.values()):
            label = _delivery_platform_label(service)
            result["status"] = "portal_action_required"
            result["error_code"] = "PORTAL_EMPTY_RECORDS_REJECTED"
            result["message"] = (
                f"{label} 포털에서 금액·주문번호 없는 빈 데이터만 감지되어 저장을 차단했습니다. "
                "조회 기간, 포털 메뉴, 엑셀/CSV 원본을 확인한 뒤 다시 수집해야 합니다."
            )
    return result


def sync_delivery(payload: dict[str, Any], user: dict[str, Any]) -> dict[str, Any]:
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="자동 수집 실행 권한이 없습니다")
    _settle_stale_delivery_collection_statuses()
    from app.services.bank_collection_lock import bank_lock_is_active, default_bank_lock_path
    bank_lock_path = os.getenv("YEOLJEONG_BANK_AUTO_COLLECT_LOCK_PATH", default_bank_lock_path())
    if bank_lock_is_active(bank_lock_path):
        return {
            "queued": False,
            "status": "deferred",
            "error_code": "DELIVERY_DEFERRED_DUE_TO_BANK_LOCK",
            "summary": [],
            "totals": _delivery_empty_counts(),
            "diagnostics": {"delivery_deferred_due_to_bank_lock": "1"},
            "message": "은행 전용 PC Agent 수집 중이어서 배달 브라우저 수집을 보류했습니다.",
        }
    lock_fd = _try_acquire_delivery_sync_lock()
    if lock_fd is None:
        return _delivery_sync_busy_result(payload, user)
    try:
        return _sync_delivery_unlocked(payload, user)
    finally:
        _release_delivery_sync_lock(lock_fd)


def _sync_delivery_unlocked(payload: dict[str, Any], user: dict[str, Any]) -> dict[str, Any]:
    from app.services.yeoljeong_delivery_collectors import collect_account

    requested_services = _delivery_requested_services(payload)
    date_from, date_to = _delivery_sync_window(payload)
    requested_account_id = str(payload.get("account_id") or payload.get("server_account_id") or "").strip()
    queued_run_ids = payload.get("queued_run_ids") if isinstance(payload.get("queued_run_ids"), dict) else {}
    sync_job_id = str(payload.get("sync_job_id") or "").strip()

    all_accounts = _read("platform_accounts")
    accounts_changed = _migrate_platform_account_secrets(all_accounts)
    accounts_changed = bool(_hydrate_delivery_account_passwords_from_agent_vault(all_accounts)) or accounts_changed
    if accounts_changed:
        _write("platform_accounts", all_accounts)
    def _delivery_account_score(row: dict[str, Any], service: str) -> tuple[int, int, int, int, int, str]:
        mode = str(row.get("collection_mode") or row.get("collectionMode") or "").strip()
        upload_mode = mode in DELIVERY_UPLOAD_COLLECTION_MODES
        has_password = _has_secret_value(row, "password")
        has_any_secret = _has_account_secret(row)
        is_canonical_account = str(row.get("id") or "") == f"acct-{service}"
        is_canonical_upload = str(row.get("id") or "") == f"acct-{service}" and upload_mode
        return (
            1 if has_password and not upload_mode else 0,
            1 if has_any_secret and not upload_mode else 0,
            1 if not upload_mode else 0,
            1 if is_canonical_account else 0,
            0 if is_canonical_upload else 1,
            str(row.get("updated_at") or row.get("created_at") or ""),
        )

    scopes = _delivery_sync_scopes(payload, requested_services, all_accounts)
    synced_at = _now()
    summary = []
    ledger_names = {kind: f"delivery_{kind}" for kind in DELIVERY_RECORD_TYPES}
    ledgers = {name: _read(name) for name in ledger_names.values()}
    statuses = _read("delivery_collection_status")
    full_baemin_backfill = _delivery_full_backfill_requested(payload) and requested_services == ["baemin"]
    selected_backfill_statuses: dict[tuple[str, str], dict[str, Any]] = {}
    if full_baemin_backfill:
        effective_job_id = sync_job_id or f"delivery-backfill-{uuid4().hex[:12]}"
        if not sync_job_id:
            sync_job_id = effective_job_id
        _delivery_ensure_baemin_backfill_queue(statuses, payload, scopes, date_from, date_to, effective_job_id)
        queued_statuses = _delivery_select_baemin_backfill_statuses(statuses, payload, scopes)
        selected_backfill_statuses = {
            (str(row.get("business_id") or ""), str(row.get("branch") or "")): row
            for row in queued_statuses
        }
        if selected_backfill_statuses:
            scopes = list(selected_backfill_statuses)
    response_ledgers = _delivery_empty_record_lists()
    response_records: list[dict[str, Any]] = []
    baemin_security_blocked = False

    for business_id, branch in scopes:
        candidates = [
            row
            for row in all_accounts
            if (not requested_services or row.get("service") in requested_services)
            and (not requested_account_id or str(row.get("id") or "") == requested_account_id)
            and str(row.get("business_id") or BUSINESS_BY_BRANCH.get(str(row.get("branch") or "")) or "") == business_id
            and BRANCH_ALIASES.get(str(row.get("branch") or ""), str(row.get("branch") or "")) == branch
        ]
        candidates.sort(key=lambda row: str(row.get("updated_at") or row.get("created_at") or ""), reverse=True)
        accounts_by_service: dict[str, dict[str, Any]] = {}
        for requested_service in requested_services:
            service_rows = [row for row in candidates if str(row.get("service") or "") == requested_service]
            if requested_account_id and service_rows:
                accounts_by_service[requested_service] = service_rows[0]
            elif service_rows:
                accounts_by_service[requested_service] = max(
                    service_rows,
                    key=lambda row: _delivery_account_score(row, requested_service),
                )

        for service in requested_services:
            if service == "baemin" and baemin_security_blocked:
                continue
            account = accounts_by_service.get(service)
            selected_backfill_status = selected_backfill_statuses.get((business_id, branch))
            run_date_from, run_date_to = date_from, date_to
            if selected_backfill_status:
                run_date_from, run_date_to = _delivery_backfill_window_from_status(
                    selected_backfill_status,
                    date_from,
                    date_to,
                )
            run_id = str(
                queued_run_ids.get(_delivery_run_key(service, business_id, branch))
                or (queued_run_ids.get(service) if len(scopes) == 1 else "")
                or (selected_backfill_status.get("id") if selected_backfill_status else "")
                or uuid4()
            )
            queued_status = next((row for row in statuses if str(row.get("id") or "") == run_id), None)
            status_record = {
                "id": run_id,
                "job_id": sync_job_id,
                "service": service,
                "business_id": business_id,
                "branch": branch,
                "date_from": run_date_from.isoformat(),
                "date_to": run_date_to.isoformat(),
                "status": "running",
                "counts": _delivery_empty_counts(),
                "payload": _delivery_backfill_row_payload(
                    payload,
                    date_from=run_date_from,
                    date_to=run_date_to,
                    checkpoint=(
                        (selected_backfill_status.get("payload") or {}).get("checkpoint")
                        if selected_backfill_status and isinstance(selected_backfill_status.get("payload"), dict)
                        else None
                    ),
                    attempt_count=int((selected_backfill_status or {}).get("attempt_count") or 0),
                )
                if full_baemin_backfill
                else _delivery_backfill_status_payload(payload),
                "error_code": "",
                "started_at": synced_at,
                "created_at": synced_at,
                "updated_at": synced_at,
            }
            if queued_status:
                queued_status.update(status_record)
                status_record = queued_status
            else:
                statuses.insert(0, status_record)
            _write_delivery_collection_statuses(statuses, status_record)
            if not account:
                result = {"status": "credential_required", "error_code": "ACCOUNT_NOT_REGISTERED", "records": {}}
            else:
                collection_mode = str(account.get("collection_mode") or account.get("collectionMode") or "").strip()
                collection_account = dict(account)
                collection_account["_date_from"] = run_date_from.isoformat()
                collection_account["_date_to"] = run_date_to.isoformat()
                collection_account["_max_orders"] = _delivery_backfill_max_orders(payload)
                collection_account["_max_reviews"] = _delivery_backfill_max_reviews(payload)
                collection_account["_mode"] = str(payload.get("mode") or "")
                collection_account["_run_id"] = run_id
                captcha_value = _delivery_captcha_value_for_account(payload, collection_account, service, business_id, branch)
                if captcha_value:
                    collection_account["captcha_value"] = captcha_value
                captcha_approval = _delivery_captcha_automation_approval(payload, user, service, business_id, branch)
                if captcha_approval:
                    collection_account["_captcha_approval"] = captcha_approval
                browser_auth = _delivery_browser_auth_for_account(payload, collection_account, service, business_id, branch)
                if browser_auth["storage_state_path"]:
                    collection_account["storage_state_path"] = browser_auth["storage_state_path"]
                can_use_browser_auth = bool(browser_auth["storage_state_path"] or browser_auth["browser_session_id"])
                bridge_result = None
                if can_use_browser_auth:
                    backfill_context = (
                        {
                            "mode": "full_backfill",
                            "date_from": run_date_from.isoformat(),
                            "date_to": run_date_to.isoformat(),
                            "max_orders": _delivery_backfill_max_orders(payload),
                            "max_reviews": _delivery_backfill_max_reviews(payload),
                            "checkpoint": (
                                (selected_backfill_status.get("payload") or {}).get("checkpoint")
                                if selected_backfill_status and isinstance(selected_backfill_status.get("payload"), dict)
                                else (payload.get("checkpoint") if isinstance(payload.get("checkpoint"), dict) else {})
                            ),
                        }
                        if service == "baemin" and _delivery_full_backfill_requested(payload)
                        else None
                    )
                    bridge_result = (
                        (
                            _collect_baemin_from_browser_bridge_session(collection_account, browser_auth, backfill_context)
                            if backfill_context
                            else _collect_baemin_from_browser_bridge_session(collection_account, browser_auth)
                        )
                        if service == "baemin"
                        else _collect_delivery_from_browser_bridge_session(
                            collection_account,
                            browser_auth,
                            date_from.isoformat(),
                            date_to.isoformat(),
                        )
                    )
                if bridge_result is not None:
                    result = bridge_result
                elif collection_mode in DELIVERY_UPLOAD_COLLECTION_MODES and not can_use_browser_auth:
                    label = _delivery_platform_label(service)
                    result = {
                        "status": "upload_required",
                        "error_code": "CSV_UPLOAD_REQUIRED",
                        "records": {},
                        "diagnostics": {"collection_mode": collection_mode},
                        "message": f"{label} 포털 CSV/엑셀 정산서 업로드가 필요한 계정입니다.",
                    }
                elif (
                    collection_mode == "browser-automation"
                    and not can_use_browser_auth
                    and not _has_secret_value(account, "password")
                    and payload.get("require_pc_agent")
                ):
                    label = _delivery_platform_label(service)
                    result = {
                        "status": "credential_required",
                        "error_code": "PC_AGENT_SESSION_REQUIRED",
                        "records": {},
                        "diagnostics": {
                            "collection_mode": collection_mode,
                            "browser_bridge_error": browser_auth.get("browser_bridge_error") or "",
                            "ambient_browser_session_id": browser_auth.get("ambient_browser_session_id") or "",
                            "ambient_browser_bridge_mode": browser_auth.get("ambient_browser_bridge_mode") or "",
                        },
                        "message": f"{label} 자동수집은 저장된 비밀번호 또는 PC Agent 전용 세션이 필요합니다.",
                    }
                elif (
                    collection_mode == "browser-automation"
                    and not can_use_browser_auth
                    and _has_secret_value(account, "password")
                    and not payload.get("allow_server_headless_fallback")
                    and not payload.get("allowServerHeadlessFallback")
                ):
                    label = _delivery_platform_label(service)
                    if service == "ddangyo":
                        message = "땡겨요는 숫자 캡챠 입력이 필요하므로 PC Agent 브라우저 세션이 연결되어야 자동로그인과 수집을 계속할 수 있습니다."
                    else:
                        message = f"{label} 자동수집은 PC Agent 전용 브라우저 세션 생성 후 실행해야 합니다."
                    result = {
                        "status": "credential_required",
                        "error_code": "PC_AGENT_SESSION_REQUIRED",
                        "records": {},
                        "diagnostics": {
                            "collection_mode": collection_mode,
                            "browser_bridge_error": browser_auth.get("browser_bridge_error") or "",
                            "browser_auth_strategy": browser_auth.get("browser_auth_strategy") or "",
                        },
                        "message": message,
                    }
                elif not _has_secret_value(account, "password") and not can_use_browser_auth:
                    label = _delivery_platform_label(service)
                    result = {
                        "status": "credential_required",
                        "error_code": "CREDENTIAL_REQUIRED",
                        "records": {},
                        "message": f"{label} 계정 비밀번호가 등록되지 않았습니다.",
                    }
                elif service == "ddangyo" and not can_use_browser_auth:
                    result = {
                        "status": "portal_action_required",
                        "error_code": "PC_AGENT_SESSION_REQUIRED",
                        "records": {},
                        "diagnostics": {
                            "collection_mode": collection_mode,
                            "browser_bridge_error": browser_auth.get("browser_bridge_error") or "",
                            "browser_auth_strategy": browser_auth.get("browser_auth_strategy") or "",
                        },
                        "message": "땡겨요는 숫자 캡챠 입력이 필요하므로 PC Agent 브라우저 세션이 연결되어야 자동로그인과 수집을 계속할 수 있습니다.",
                    }
                else:
                    secret = _decrypt_secret(str(account.get("password_enc") or "")) if _has_secret_value(account, "password") else ""
                    result = collect_account(collection_account, secret, date_from.isoformat(), date_to.isoformat())
                    if browser_auth["browser_session_id"]:
                        result.setdefault("diagnostics", {})["browser_session_id"] = browser_auth["browser_session_id"]
                    if browser_auth.get("browser_work_key"):
                        result.setdefault("diagnostics", {})["browser_work_key"] = browser_auth["browser_work_key"]
                    if browser_auth["browser_bridge_mode"]:
                        result.setdefault("diagnostics", {})["browser_bridge_mode"] = browser_auth["browser_bridge_mode"]
                    if browser_auth.get("browser_bridge_error"):
                        result.setdefault("diagnostics", {})["browser_bridge_error"] = browser_auth["browser_bridge_error"]
                result = _normalize_delivery_collection_result(service, result)
                if (
                    browser_auth.get("browser_session_id")
                    and str(browser_auth.get("browser_close_on_complete") or "") == "1"
                ):
                    _run_delivery_browser_async(
                        _close_delivery_browser_work_session_async(
                            browser_auth,
                            reason=f"delivery_sync_result_{service}",
                        )
                    )

            counts = _delivery_empty_counts()
            for kind, ledger_name in ledger_names.items():
                incoming = result.get("records", {}).get(kind) or []
                by_id = {str(row.get("id") or ""): row for row in ledgers[ledger_name] if row.get("id")}
                for record in incoming:
                    by_id[str(record["id"])] = record
                ledgers[ledger_name] = list(by_id.values())
                counts[kind] = len(incoming)
                response_ledgers[kind].extend(incoming)
                if kind in {"sales", "settlements"}:
                    response_records.extend(_delivery_entry_record(record) for record in incoming)
            finished_at = _now()
            public_status = _delivery_public_collection_status(result.get("status"))
            public_error_code = _delivery_public_error_code(public_status, result.get("error_code"))
            baemin_security_blocked_result = (
                service == "baemin"
                and str(result.get("error_code") or "").strip().upper()
                in BAEMIN_SECURITY_BLOCK_CODES
            )
            result_diagnostics = dict(result.get("diagnostics") or {})
            if baemin_security_blocked_result:
                result_diagnostics["hard_stop_remaining_baemin_scopes"] = True
                result_diagnostics["cooldown_minutes"] = BAEMIN_SECURITY_BLOCK_COOLDOWN_MINUTES
                result_diagnostics["cooldown_until"] = _delivery_baemin_security_cooldown_until(finished_at and datetime.fromisoformat(finished_at))
            elif _delivery_result_requires_operator_cooldown(public_status, public_error_code):
                result_diagnostics["cooldown_minutes"] = DELIVERY_OPERATOR_ACTION_COOLDOWN_MINUTES
                result_diagnostics["cooldown_until"] = _delivery_operator_action_cooldown_until(
                    finished_at and datetime.fromisoformat(finished_at)
                )
            browser_session_id = str(browser_auth.get("browser_session_id") or "").strip()
            browser_work_key = str(browser_auth.get("browser_work_key") or "").strip()
            if browser_session_id:
                result_diagnostics["browser_session_id"] = browser_session_id
            if browser_work_key:
                result_diagnostics["browser_work_key"] = browser_work_key
            if browser_session_id or browser_work_key:
                result_diagnostics["resume_token"] = make_resume_token(browser_work_key, browser_session_id, run_id)
                _append_delivery_browser_session_event(
                    "collection_result",
                    {
                        "service": service,
                        "business_id": business_id,
                        "branch": branch,
                        "run_id": run_id,
                        "session_id": browser_session_id,
                        "work_key": browser_work_key,
                        "status": public_status,
                        "raw_status": result.get("status") or "",
                        "error_code": public_error_code,
                        "message": result.get("message") or "",
                        "counts": counts,
                        "hard_stop_remaining_baemin_scopes": result_diagnostics.get("hard_stop_remaining_baemin_scopes") or False,
                        "cooldown_until": result_diagnostics.get("cooldown_until") or "",
                    },
                )
            previous_attempts = int(status_record.get("attempt_count") or 0)
            attempt_count = previous_attempts + 1
            if public_status == "action_required" and attempt_count >= DELIVERY_CHALLENGE_MAX_ATTEMPTS:
                public_status = "failed"
                public_error_code = "CHALLENGE_MAX_ATTEMPTS_EXCEEDED"
                result_diagnostics["challenge_terminal"] = "max_attempts"
            continuation = _delivery_enqueue_baemin_backfill_continuation(
                statuses,
                {
                    **payload,
                    "date_from": date_from.isoformat(),
                    "date_to": date_to.isoformat(),
                },
                current={
                    **status_record,
                    "date_from": run_date_from.isoformat(),
                    "date_to": run_date_to.isoformat(),
                },
                result={**result, "diagnostics": result_diagnostics},
                public_status=public_status,
                public_error_code=public_error_code,
                attempt_count=attempt_count,
            )
            status_record.update(
                {
                    "status": public_status,
                    "raw_status": result.get("status") or "",
                    "counts": counts,
                    "payload": (
                        {
                            **_delivery_backfill_status_payload(
                                {
                                    **payload,
                                    "date_from": run_date_from.isoformat(),
                                    "date_to": run_date_to.isoformat(),
                                },
                                result,
                            ),
                            "continuation_run_id": continuation.get("id") if continuation else "",
                        }
                        if full_baemin_backfill
                        else _delivery_backfill_status_payload(payload, result)
                    ),
                    "error_code": public_error_code,
                    "diagnostics": result_diagnostics,
                    "cooldown_until": result_diagnostics.get("cooldown_until") or "",
                    "attempt_count": attempt_count,
                    "max_attempts": DELIVERY_CHALLENGE_MAX_ATTEMPTS,
                    "challenge_timeout_seconds": int(DELIVERY_CHALLENGE_TIMEOUT.total_seconds()),
                    "resume_token": result_diagnostics.get("resume_token") or status_record.get("resume_token") or "",
                    "message": result.get("message") or "",
                    "finished_at": finished_at,
                    "updated_at": finished_at,
                }
            )
            _write_delivery_collection_statuses(statuses, status_record)
            if account:
                account["last_sync_status"] = status_record["status"]
                account["portal_status"] = status_record["status"]
                account["portal_message"] = status_record["message"] or status_record["error_code"]
                account["last_sync_at"] = finished_at
                account["updated_at"] = finished_at
            summary.append(
                {
                    "service": service,
                    "status": status_record["status"],
                    "portal_status": status_record["status"],
                    "error_code": status_record["error_code"],
                    "counts": counts,
                    "run_id": run_id,
                    "account_id": account.get("id") if account else "",
                    "collection_mode": str(account.get("collection_mode") or account.get("collectionMode") or "") if account else "",
                    "business_id": business_id,
                    "branch": branch,
                    "message": status_record["message"],
                    "portal_message": status_record["message"],
                    "cooldown_until": status_record.get("cooldown_until") or "",
                }
            )
            if baemin_security_blocked_result:
                baemin_security_blocked = True

    for kind, ledger_name in ledger_names.items():
        _write_file_rows(ledger_name, ledgers[ledger_name])
        if ledger_name in DB_LEDGER_TABLE_BY_NAME:
            for record in response_ledgers[kind]:
                _run_db(_db_upsert_ledger(ledger_name, record))
    _write("platform_accounts", all_accounts)
    _write_delivery_collection_statuses(statuses)
    return {
        "synced_at": synced_at,
        "business_id": scopes[0][0] if len(scopes) == 1 else "all",
        "branch": scopes[0][1] if len(scopes) == 1 else "전체",
        "date_from": date_from.isoformat(),
        "date_to": date_to.isoformat(),
        "summary": summary,
        "records": response_records,
        "sales": response_ledgers["sales"],
        "settlements": response_ledgers["settlements"],
        "reviews": response_ledgers["reviews"],
        "ads": response_ledgers["ads"],
        "totals": {kind: sum(item["counts"].get(kind, 0) for item in summary) for kind in DELIVERY_RECORD_TYPES},
    }


def import_delivery_portal_text(payload: dict[str, Any], user: dict[str, Any]) -> dict[str, Any]:
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="배달 포털 파싱 반영 권한이 없습니다")
    from app.services.yeoljeong_delivery_collectors import PORTAL_CONFIG, parse_portal_export

    service = str(payload.get("service") or "").strip()
    record_type = str(payload.get("record_type") or payload.get("recordType") or "").strip()
    if service not in PORTAL_CONFIG:
        raise HTTPException(status_code=400, detail="지원하지 않는 배달 플랫폼입니다")
    if record_type not in DELIVERY_RECORD_TYPES:
        raise HTTPException(status_code=400, detail="반영 구분은 sales, settlements, reviews, ads 중 하나여야 합니다")

    business_id, branch = _normalize_delivery_scope(payload.get("business_id"), payload.get("branch"))
    source_text = str(payload.get("source_text") or payload.get("sourceText") or "")
    parsed = parse_portal_export(service, record_type, source_text, business_id, branch)
    parsed_records, rejected_counts = _delivery_filter_meaningful_records(parsed.get("records"))
    ledger_names = {kind: f"delivery_{kind}" for kind in DELIVERY_RECORD_TYPES}
    ledger_name = ledger_names[record_type]
    existing_rows = _read(ledger_name)
    by_id = {str(row.get("id") or ""): row for row in existing_rows if row.get("id")}
    imported: list[dict[str, Any]] = []
    duplicate_rows = 0
    now = _now()
    for record in parsed_records.get(record_type) or []:
        record["source_file"] = Path(str(payload.get("filename") or "pc-browser-copy.html")).name
        record["collection_mode"] = "pc-browser-parse"
        record["created_at"] = record.get("created_at") or now
        record["updated_at"] = now
        record_id = str(record.get("id") or "")
        if record_id in by_id:
            duplicate_rows += 1
        by_id[record_id] = record
        imported.append(record)
    if imported or parsed.get("status") != "succeeded":
        _write(ledger_name, list(by_id.values()))

    statuses = _read("delivery_collection_status")
    counts = _delivery_empty_counts()
    counts[record_type] = len(imported)
    run_id = str(uuid4())
    status = parsed.get("status") or "failed"
    error_code = parsed.get("error_code") or ""
    message = parsed.get("message") or "PC에서 로그인 후 복사/저장한 배민 화면 데이터를 파싱해 반영했습니다."
    rejected_total = sum(rejected_counts.values())
    if rejected_total and not imported:
        status = "portal_action_required"
        error_code = "PORTAL_EMPTY_RECORDS_REJECTED"
        message = (
            f"{_delivery_platform_label(service)} 포털에서 금액·식별자 없는 빈 데이터만 감지되어 저장을 차단했습니다. "
            "원본 파일 또는 복사 화면을 확인한 뒤 다시 반영해야 합니다."
        )
    statuses.insert(
        0,
        {
            "id": run_id,
            "service": service,
            "business_id": business_id,
            "branch": branch,
            "date_from": str(payload.get("date_from") or ""),
            "date_to": str(payload.get("date_to") or ""),
            "status": status,
            "counts": counts,
            "error_code": error_code,
            "diagnostics": {
                **(parsed.get("diagnostics") or {}),
                "collection_mode": "pc-browser-parse",
                "record_type": record_type,
                "duplicate_rows": duplicate_rows,
                "invalid_record_rejected_counts": rejected_counts,
            },
            "message": message,
            "started_at": now,
            "finished_at": now,
            "created_at": now,
            "updated_at": now,
        },
    )
    _write_delivery_collection_statuses(statuses, statuses[0] if statuses else None)

    response_ledgers = _delivery_empty_record_lists()
    response_ledgers[record_type] = imported
    records = [_delivery_entry_record(record) for record in imported if record_type in {"sales", "settlements"}]
    return {
        "synced_at": now,
        "business_id": business_id,
        "branch": branch,
        "summary": [
            {
                "service": service,
                "status": status,
                "portal_status": status,
                "error_code": error_code,
                "counts": counts,
                "run_id": run_id,
                "collection_mode": "pc-browser-parse",
                "message": message,
                "portal_message": message,
            }
        ],
        "records": records,
        "sales": response_ledgers["sales"],
        "settlements": response_ledgers["settlements"],
        "reviews": response_ledgers["reviews"],
        "ads": response_ledgers["ads"],
        "totals": counts,
        "import": {"imported": len(imported), "duplicate_rows": duplicate_rows},
    }


def import_settlement_csv(
    text: str,
    user: dict[str, Any],
    *,
    service: str,
    business_id: str,
    branch: str,
    filename: str = "settlement.csv",
) -> dict[str, Any]:
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="정산서 가져오기 권한이 없습니다")
    from app.services.yeoljeong_delivery_collectors import PORTAL_CONFIG, normalize_record

    normalized_service = str(service or "").strip()
    if normalized_service not in PORTAL_CONFIG:
        raise HTTPException(status_code=400, detail="지원하지 않는 배달 플랫폼입니다")
    normalized_business, normalized_branch = _normalize_delivery_scope(business_id, branch)
    rows = _read("delivery_settlements")
    reader = csv.DictReader(text.splitlines())
    by_id = {str(row.get("id") or ""): row for row in rows if row.get("id")}
    imported: list[dict[str, Any]] = []
    duplicate_rows = 0
    now = _now()
    for raw in reader:
        source_row = {str(key or "").strip(): value for key, value in raw.items()}
        if not any(str(value or "").strip() for value in source_row.values()):
            continue
        record = normalize_record(
            normalized_service,
            "settlements",
            source_row,
            normalized_business,
            normalized_branch,
        )
        record["source_file"] = Path(filename or "settlement.csv").name
        record["created_at"] = now
        record["updated_at"] = now
        if str(record["id"]) in by_id:
            duplicate_rows += 1
            continue
        by_id[str(record["id"])] = record
        imported.append(record)
    if imported:
        _write("delivery_settlements", list(by_id.values()))
    return {
        "imported": len(imported),
        "duplicate_rows": duplicate_rows,
        "business_id": normalized_business,
        "branch": normalized_branch,
        "service": normalized_service,
        "records": imported,
        "settlements": imported,
    }


def reset_data_for_tests() -> None:
    if DATA_DIR.exists():
        shutil.rmtree(DATA_DIR)
